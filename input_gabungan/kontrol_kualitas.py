#!/usr/bin/env python3
"""
kontrol_kualitas.py — KONTROL KUALITAS SUMBER DATA input usaha (format standar
`input_usaha.xlsx` & format tahap 2), TANPA browser/VPN, TANPA mengubah apa pun.

ATURANNYA TIDAK DITULIS ULANG DI SINI. Yang dipakai adalah hasil pemeriksaan offline
yang SAMA PERSIS dgn `main_gabungan.py --cek` (`muat_sumber()` -> loader ->
`periksa_semua` / `periksa_semua_tahap2`), termasuk koreksi & nilai pengganti yang
dipasang loader. Setiap kali aturan pemeriksaan berubah, laporan ini ikut berubah
dgn sendirinya — tidak ada daftar aturan kedua yang bisa tertinggal.

Yang ditambahkan program ini: hasil pemeriksaan itu diterjemahkan jadi daftar yang
bisa langsung dikerjakan pengelola data —
  - SEL mana (huruf kolom Excel + nomor baris) yang perlu dibetulkan;
  - kategori tiap temuan:
      DITOLAK    baris TIDAK akan diinput (status SKIP_DATA_*) — wajib dibetulkan
      DRAFT      koordinat belum ada/tidak terbaca — diinput tapi ditahan sbg draft
      DIGANTI    isian kosong/tidak valid -> skrip memakai NILAI PENGGANTI (asumsi)
      DIKOREKSI  isian diubah/dilengkapi skrip dgn aturan dari data baris itu sendiri
      TINJAU     mungkin tidak konsisten / pernah memicu GALAT — perlu dicek manusia
      INFO       tafsiran & default biasa format ini — tidak perlu tindakan
  - saran perbaikan, rekap per jenis / kolom / PPL;
  - salinan sheet dgn sel bermasalah diwarnai + komentar (nomor baris & huruf kolom
    SAMA dgn sheet asli; kolom hasil QC ditambahkan di paling kanan).

Pesan pemeriksaan (teks bebas) dipetakan ke kategori & kolom lewat POLA_TANDA /
KEYS_MASALAH. Pesan yang belum dikenali TIDAK hilang: masuk kategori TINJAU tanpa
sel & dicetak di akhir run — itu tanda POLA_TANDA perlu ditambah.

Keluaran berisi data responden -> hanya lokal (sudah diabaikan .gitignore: *.xlsx, *.csv).

Contoh (jalankan dari root proyek):
    python input_gabungan/kontrol_kualitas.py --sumber input_usaha.xlsx
    python input_gabungan/kontrol_kualitas.py --sumber bahan/input_tahap2.xlsx --format tahap2
    python input_gabungan/kontrol_kualitas.py --sumber bahan/input_tahap2.xlsx --format tahap2 \\
        --dari 2 --sampai 500 --per-ppl
Panduan: docs/PANDUAN_KONTROL_KUALITAS.md
"""

from __future__ import annotations

import argparse
import csv
import datetime
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from functools import cached_property
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

# -- jalankan dari root proyek; sisipkan root ke sys.path utk paket inti/ dll --
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
import input_gabungan.main_gabungan as mg  # noqa: E402
from inti.gabungan_loader import (  # noqa: E402
    KEY_16B, KEY_26, KEY_27, KEY_29, KEY_PEKERJA, KOLOM, NAMA_SHEET_DITERIMA, OPSI_FORM, GabunganRow,
    Pemeriksaan, _baca_mentah, _cari_indeks, _sel, parse_pilihan_baris,
)
from inti.tahap2_loader import (  # noqa: E402
    CEK_TOTAL, KOLOM_TAHAP2, KOLOM_TAHAP2_TAMBAHAN, NAMA_SHEET_TAHAP2, Tahap2Row, _baca_mentah_tahap2,
    _indeks_tahap2,
)

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

KELUARAN_PATH = Path("./kontrol_kualitas.xlsx")
FOLDER_PER_PPL = "kontrol_kualitas_per_ppl"

# ---------------------------------------------------------------------------
# Kategori temuan — urutan = tingkat keparahan (paling atas paling parah)
# ---------------------------------------------------------------------------
KATEGORI = ("DITOLAK", "DRAFT", "DIGANTI", "DIKOREKSI", "TINJAU", "INFO")
PERINGKAT = {k: i for i, k in enumerate(KATEGORI)}
KATEGORI_TINDAKAN = KATEGORI[:-1]   # semua selain INFO
ARTI_KATEGORI = {
    "DITOLAK": "Baris TIDAK akan diinput (dilewati pemeriksaan sebelum dokumen dibuat). Wajib dibetulkan di sheet.",
    "DRAFT": "Koordinat belum ada/tidak terbaca: dokumen dibuat & diisi, tapi DITAHAN sbg draft (tidak dikirim) "
             "sampai koordinat diisi di sheet.",
    "DIGANTI": "Isian sheet kosong/tidak valid -> skrip memakai NILAI PENGGANTI (asumsi/ketetapan). Dokumen tetap "
               "dikirim dgn nilai itu; betulkan sheet dari kuesioner kalau datanya ada.",
    "DIKOREKSI": "Isian sheet diubah/dilengkapi skrip dgn aturan dari data baris itu sendiri (format rusak, pos "
                 "dipindah, teks dilengkapi). Cek hasilnya; sebaiknya sheet ikut dibetulkan.",
    "TINJAU": "Mungkin tidak konsisten atau pernah memicu GALAT form. Tidak menghentikan input; cek kebenarannya.",
    "INFO": "Tafsiran/default biasa format ini (ketetapan yang berlaku utk semua baris). Tidak perlu tindakan.",
}
WARNA = {"DITOLAK": "FFC7CE", "DRAFT": "D9D2E9", "DIGANTI": "F8CBAD", "DIKOREKSI": "FFEB9C",
         "TINJAU": "DDEBF7", "INFO": "EDEDED", "BERSIH": "E2EFDA"}
ARTI_STATUS = {
    "SIAP": "akan diinput; dikirim kalau batch dijalankan dgn --submit",
    mg.STATUS_SIAP_TANPA_KOORDINAT: "akan diinput tapi disimpan DRAFT (tidak dikirim) sampai koordinat diisi",
}


# ---------------------------------------------------------------------------
# Nama rincian <-> key internal (diturunkan dari KOLOM loader, bukan daftar kedua)
# ---------------------------------------------------------------------------
def _label_dari_awalan(awalan: str) -> str:
    """Awalan judul KOLOM -> nomor rincian: "12.c." -> "12c", "24.a1." -> "24a1", "18." -> "18"."""
    m = re.fullmatch(r"(\d+)\.(?:([a-z]\d?)\.)?", awalan)
    return f"{m.group(1)}{m.group(2) or ''}" if m else ""


LABEL_KEY: dict[str, str] = {k: lbl for k, a in KOLOM.items() if (lbl := _label_dari_awalan(a))}
LABEL_KEY.update({
    "kbli": "13g KBLI", "internet_semua": "16b1-b6", "kodepos": "kodepos", "idsubsls": "idsubsls",
    "nama": "nama usaha", "akun_ppl": "akun PPL", "hp": "no HP/WA", "jalan_domisili": "Nama Jalan",
    "latitude": "latitude", "longitude": "longitude",
})

_dari_label: dict[str, tuple[str, ...]] = defaultdict(tuple)
for _k, _l in LABEL_KEY.items():
    if re.fullmatch(r"\d+[a-z]?\d?", _l):
        _dari_label[_l] += (_k,)
_dari_label.update({
    "13g": ("kbli",),
    # varian bulanan (30-33) diisi dari kolom yang sama dgn 26-29
    "30a": ("gaji",), "30b": ("biaya_produksi",), "30c": ("biaya_pembelian",), "30d": ("operasional",),
    "30e": ("non_operasional",), "31a": ("nilai_pendapatan",), "31b": ("pendapatan_lain",),
    "31d": ("pendapatan_online",),
    # total yang dihitung form -> semua rinciannya
    "26f": KEY_26, "30f": KEY_26, "27c": KEY_27, "31c": KEY_27, "28c": ("aset_usaha_thn", "aset_lain_thn"),
    "24c1": ("tk_laki", "tk_pr"),
})
KEY_DARI_LABEL: dict[str, tuple[str, ...]] = dict(_dari_label)

# Key yang kolomnya di sheet tertentu bernama lain.
ALIAS_KOLOM: dict[str, tuple[str, ...]] = {
    "nama": ("nama_komersial",),          # tahap 2: nama usaha = kolom 8b
    "akun_ppl": ("info_nama_ppl",),       # tahap 2 tanpa kolom email PPL
    "judul_kbli": ("info_judul_kbli",),
    **{k: ("internet_semua",) for k in KEY_16B},   # tahap 2: satu kolom "16b1-b6"
}
_KEY_DIKENAL = set(KOLOM) | set(KOLOM_TAHAP2) | set(KOLOM_TAHAP2_TAMBAHAN)
_PILIH_WILAYAH = ("pilih_prov", "pilih_kab", "pilih_kec", "pilih_desa", "pilih_sls", "pilih_subsls")


def _keys_label(*labels: str) -> tuple[str, ...]:
    return tuple(k for lbl in labels for k in KEY_DARI_LABEL.get((lbl or "").lower(), ()))


def _keys_daftar(teks: str) -> tuple[str, ...]:
    """"gaji, biaya_produksi, produk (13f), latitude/longitude (…)" -> key-keynya."""
    out: list[str] = []
    for bagian in (teks or "").split(","):
        m = re.match(r"\s*([a-z][a-z0-9_]*(?:/[a-z][a-z0-9_]*)*)", bagian)
        if m:
            out.extend(m.group(1).split("/"))
    return tuple(out)


def _keys_total(nama: str) -> tuple[str, ...]:
    """"26a+26b+26c+26d+26e" (nama total CEK_TOTAL tahap 2) -> (kolom total, rincian...)."""
    for key_cek, rincian, n in CEK_TOTAL:
        if n == (nama or "").strip():
            return (key_cek, *rincian)
    return ()


def _keys_umum(teks: str) -> tuple[str, ...]:
    """Cadangan utk pesan yang belum dikenali: hanya bentuk yang jelas menunjuk rincian —
    `key=nilai` / `key='nilai'` dan nomor rincian spt "12c"/"26a"."""
    keys = [k for k in re.findall(r"\b([a-z][a-z0-9_]{1,})(?==)", teks or "") if k in _KEY_DIKENAL]
    keys += [k for lbl in re.findall(r"\b(\d{1,2}[a-f]\d?)\b", teks or "") for k in KEY_DARI_LABEL.get(lbl, ())]
    return tuple(dict.fromkeys(keys))


# ---------------------------------------------------------------------------
# Pola pesan TANDA (review/koreksi) -> kategori, kolom, saran.
# Teksnya dihasilkan inti/gabungan_loader.py & inti/tahap2_loader.py. URUTAN PENTING:
# pola yang lebih khusus di atas. Pesan baru yang belum ada di sini -> TINJAU
# (TANDA_LAIN) & dilaporkan, jadi tidak pernah hilang diam-diam.
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Pola:
    jenis: str
    regex: str
    kategori: str
    keys: tuple[str, ...] = ()
    saran: str = ""
    keys_dari: Callable | None = None   # (re.Match, baris) -> key tambahan


def _13f(_m, row) -> tuple[str, ...]:
    """13f diambil dari kolom 13f kalau terisi, kalau tidak dari 13a (GABUNGAN_13F_DARI_13A)."""
    return ("produk",) if row is not None and row["produk"] else ("keg_utama",)


POLA_TANDA: tuple[Pola, ...] = (
    # --- DRAFT ---------------------------------------------------------------
    Pola("KOORDINAT_BELUM_ADA", r"KOORDINAT BELUM ADA\b", "DRAFT", ("latitude", "longitude"),
         "Isi latitude & longitude (desimal bertitik, mis. -8.1234567 & 115.1234567; format sel Teks). "
         "Sampai diisi, dokumennya ditahan sbg draft."),
    # --- DIGANTI (nilai pengganti / asumsi) ----------------------------------
    Pola("UMUR_PENGGANTI", r"12c umur kosong -> .*\(nilai pengganti", "DIGANTI", ("umur",),
         "Isi umur pengusaha (10-99) dari kuesioner; sekarang dikirim nilai pengganti."),
    Pola("TAHUN_OPERASI_PENGGANTI", r"25 tahun operasi kosong -> .*\(nilai pengganti", "DIGANTI", ("tahun_operasi",),
         "Isi tahun mulai beroperasi dari kuesioner; sekarang dikirim nilai pengganti."),
    Pola("UMUR_DIKOREKSI_PER_BARIS", r"12c umur '.*' -> '.*' \(koreksi per baris", "DIGANTI", ("umur",),
         "Umur pengusaha diganti ketetapan per baris (TAHAP2_KOREKSI_BARIS); betulkan dari kuesioner."),
    Pola("JALAN_KOSONG_JADI_WILAYAH", r"Nama Jalan kosong -> nama wilayah", "DIGANTI", ("jalan_domisili",),
         "Isi Nama Jalan/alamat dari kuesioner; sekarang diisi nama wilayah (desa/kecamatan)."),
    Pola("13A_KOSONG_JADI_JUDUL_KBLI", r"13a kosong -> judul KBLI", "DIGANTI", ("keg_utama",),
         "Isi kegiatan utama (13a) dari kuesioner; sekarang diisi judul KBLI."),
    Pola("HP_TIDAK_VALID_JADI_9999", r"no WA '.*' tidak valid -> ", "DIGANTI", ("hp",),
         "Tulis ulang nomor HP sbg TEKS (08…, 10-13 digit). Notasi ilmiah Excel (8,13E+10) menghilangkan "
         "digitnya. Tidak ada nomor -> 9999."),
    Pola("NIK_TIDAK_VALID_JADI_9999", r"12d NIK '.*' tidak valid", "DIGANTI", ("nik_pengusaha",),
         "Tulis NIK 16 digit sbg TEKS; >16 digit 7777, belum punya 8888, lainnya 9999."),
    Pola("PENGUSAHA_PENGGANTI", r"12a kosong/'-' -> '.*' \(tidak ada nama dalam kurung", "DIGANTI", ("pengusaha",),
         "Isi nama pengusaha/penanggung jawab (12a) dari kuesioner."),
    Pola("PEKERJA_KOSONG_JADI_1", r"24 kosong semua -> ", "DIGANTI", KEY_PEKERJA,
         "Isi jumlah pekerja (24) dari kuesioner; pemilik yang ikut bekerja ikut dihitung."),
    Pola("GAJI_DIISI_OTOMATIS", r"26a diisi [\d.,]+ \(", "DIGANTI", ("gaji", "tk_dibayar"),
         "Isi 26a upah/gaji dari kuesioner, atau betulkan 24a2 kalau pekerjanya sebenarnya tidak dibayar."),
    Pola("TOTAL_NOL_JADI_MINIMAL", r"(\d\d[a-f]) 0 \(semua pos kosong/nol\) -> (\d\d[a-f]) diisi minimal", "DIGANTI",
         saran="Semua pos kosong/nol -> skrip mengisi nilai minimal form (100.000; usaha baru tahun ini 10.000). "
               "Isi angka sebenarnya dari kuesioner.",
         keys_dari=lambda m, r: _keys_label(m.group(1), m.group(2))),
    Pola("TOTAL_DINAIKKAN_KE_MINIMAL", r"(\d\d[a-f]) [\d.,]+ < minimal [\d.,]+ -> (\w+) ditambah", "DIGANTI",
         saran="Total di bawah minimal form -> pos terbesar dinaikkan skrip. Periksa angkanya di kuesioner.",
         keys_dari=lambda m, r: (*_keys_label(m.group(1)), m.group(2))),
    Pola("16B_SATU_KODE_YA", r"16b1-b6 '.*' \(satu kolom\) -> Ya hanya", "DIGANTI", ("internet_semua",),
         "Satu kode Ya utk 16b1-b6 -> rincian yang Ya diasumsikan. Tulis per rincian (6 nilai 1/2, mis. "
         "1,2,2,1,1,2) kalau kuesionernya merinci."),
    Pola("16B_LIMA_NILAI", r"16b1-b6 '.*' berisi 5 nilai", "DIGANTI", ("internet_semua",),
         "Tulis 6 nilai (b1..b6); b6 Lainnya sekarang diasumsikan."),
    Pola("16B_TANPA_YA_JADI_B6", r"16a Ya tapi 16b '.*' tanpa Ya -> 16b6", "DIGANTI", ("internet", "internet_semua"),
         "16a Ya berarti minimal satu 16b Ya: betulkan 16b (atau 16a kalau memang tidak memakai internet)."),
    Pola("27D_DIISI_KRN_PESANAN", r"27d (?:0|kosong) -> [\d.,]+% krn 16b1", "DIGANTI",
         ("pendapatan_online", "internet_semua"),
         "Isi 27d (persen pendapatan online) dari kuesioner — 16b1 menerima pesanan online = Ya."),
    Pola("KBLI_DIGANTI_GENAI", r"KBLI \d+ kategori [A-Z] ditolak form -> 13g diisi rekomendasi GenAI", "DIGANTI",
         ("kbli",), "Ganti kode KBLI dgn KBLI usaha yang sesuai (kategori P/U ditolak form; sekarang diisi "
                    "rekomendasi GenAI pertama)."),
    # --- DIKOREKSI (aturan dari data baris itu sendiri) ----------------------
    Pola("UMUR_DARI_USAHA_LAIN", r"12c umur kosong -> .*\(disalin", "DIKOREKSI", ("umur",),
         "Isi umur di baris ini juga (disalin dari usaha lain pemilik yang sama)."),
    Pola("TAHUN_DARI_USAHA_LAIN", r"25 tahun operasi kosong -> .*\(disalin", "DIKOREKSI", ("tahun_operasi",),
         "Isi tahun mulai beroperasi di baris ini juga (disalin dari usaha lain pemilik yang sama)."),
    Pola("PENGUSAHA_DARI_NAMA_USAHA", r"12a kosong/'-' -> '.*' \(nama dalam kurung", "DIKOREKSI", ("pengusaha",),
         "Isi nama pengusaha (12a); sekarang diambil dari nama dalam kurung di nama usaha."),
    Pola("KOORDINAT_DIPULIHKAN", r"koordinat sheet '.*' \(format rusak Excel\)", "DIKOREKSI",
         ("latitude", "longitude"),
         "Excel merusak koordinat (titik ribuan/minus/satu sel); skrip memulihkannya. Tulis ulang sbg TEKS "
         "desimal bertitik & cek titiknya."),
    Pola("IDSUBSLS_AWALAN_DIBETULKAN", r"idsubsls '.*' -> '.*' \(awalan kabupaten", "DIKOREKSI", ("idsubsls",),
         "Betulkan 4 digit pertama idsubsls (kode provinsi+kabupaten)."),
    Pola("PEKERJA_DARI_TOTAL", r"(tk_\w+) kosong -> \d+ \(total", "DIKOREKSI",
         saran="Isi rincian pekerja yang kosong (sekarang dihitung dari kolom 24.Total).",
         keys_dari=lambda m, r: (m.group(1),
                                 "cek_tk_gender" if m.group(1) in ("tk_laki", "tk_pr") else "cek_tk_bayar")),
    Pola("PEKERJA_IKUT_JK_PEMILIK", r"24 laki/perempuan .* \(ikut jenis kelamin pemilik", "DIKOREKSI",
         ("tk_laki", "tk_pr"),
         "24 laki+perempuan tidak cocok dgn total pekerja: betulkan dari kuesioner (sekarang mengikuti jenis "
         "kelamin pemilik)."),
    Pola("PEKERJA_JADI_DIBAYAR", r"26a [\d.,]+ terisi tapi 24a2 = 0", "DIKOREKSI",
         ("tk_dibayar", "tk_tdk_dibayar", "gaji"),
         "26a terisi berarti ada pekerja dibayar: betulkan 24a2/24b2 (atau 26a kalau memang tidak ada upah)."),
    Pola("26C_DIPINDAH_KE_26B", r"KBLI \d+ tanpa 26c di form", "DIKOREKSI", ("biaya_pembelian", "biaya_produksi"),
         "KBLI bukan perdagangan tidak punya 26c di form: tulis nilainya di 26b."),
    Pola("POS_PENGELUARAN_DIPINDAH", r"(\d\d[a-f]) 0 \(.*\) -> (\d\d[a-f]) [\d.,]+ dipindah ke (\d\d[a-f])",
         "DIKOREKSI",
         saran="Pos pengeluaran yang wajib > 0 utk jenis usaha ini kosong -> diisi dari pos lain. Betulkan "
               "rincian 26 dari kuesioner.",
         keys_dari=lambda m, r: _keys_label(m.group(3), m.group(2))),
    Pola("27D_DIBULATKAN", r"27d '.*' dibulatkan", "DIKOREKSI", ("pendapatan_online",),
         "Tulis 27d sbg bilangan bulat (persen)."),
    Pola("WILAYAH_BENTROK_ALAMAT", r"Nama Jalan dilengkapi nama wilayah: .*\(desa idsubsls dipakai", "TINJAU",
         ("jalan_domisili", "idsubsls", "alamat_usaha_view"),
         "idsubsls & alamat 8c menunjuk desa berbeda (desa menurut idsubsls yang dipakai): pastikan idsubsls benar."),
    Pola("JALAN_DILENGKAPI", r"Nama Jalan dilengkapi nama wilayah", "DIKOREKSI", ("jalan_domisili",),
         "Tulis Nama Jalan/alamat lebih lengkap (minimal 10 huruf); sekarang dilengkapi nama wilayah."),
    Pola("13A_DILENGKAPI", r"13a (?:dilengkapi judul KBLI|'.*' < \d+ karakter -> dilengkapi)", "DIKOREKSI",
         ("keg_utama",), "Tulis kegiatan utama (13a) lebih lengkap (minimal 15 karakter); sekarang dilengkapi "
                         "judul KBLI."),
    Pola("13F_DILENGKAPI", r"13f dilengkapi", "DIKOREKSI",
         saran="Tulis produk utama (13f) minimal 4 karakter; sekarang dilengkapi.", keys_dari=_13f),
    Pola("PEKERJA_POLA_DIKOREKSI", r"24 \(laki, perempuan, dibayar, tidak dibayar\)", "DIKOREKSI", KEY_PEKERJA,
         "Betulkan rincian 24: laki+perempuan harus = dibayar+tidak dibayar."),
    Pola("BADAN_USAHA_DARI_AWALAN", r"11a '.*' -> '.*' dari awalan nama", "DIKOREKSI", ("badan_usaha",),
         "Pilih opsi 11a yang ada di form (PT -> 1.a, CV -> 7, UD -> 13)."),
    Pola("BUMDES_DIKOREKSI", r"BUMDES: ", "DIKOREKSI", ("badan_usaha", "lap_keuangan", *KEY_29),
         "BUM Desa: 11a '6. BUM Desa', 11d Ya, modal pemerintah (29e) dominan."),
    Pola("NAMA_USAHA_DIGANTI", r"nama usaha diganti -> ", "DIKOREKSI", ("nama",),
         "Nama usaha diganti sesuai ketetapan (pembeda nama); samakan di sheet kalau perlu."),
    # --- TINJAU --------------------------------------------------------------
    Pola("NAMA_KEMBAR_DINOMORI", r"nama kembar persis .* -> nama dibedakan penomoran", "TINJAU",
         ("nama_komersial",),
         "Nama, 13a & 13f sama persis di wilayah yang sama, hanya angkanya beda -> pastikan BUKAN entri ganda "
         "(hapus duplikatnya kalau ya)."),
    Pola("TOTAL_BEDA_DGN_RINCIAN", r"kolom total (\S+) di sheet = [\d.,]+, jumlah rincian", "TINJAU",
         saran="Kolom total tidak sama dgn jumlah rinciannya -> salah satu angka mungkin salah ketik "
               "(skrip memakai rinciannya).",
         keys_dari=lambda m, r: _keys_total(m.group(1))),
    Pola("KBLI_TIDAK_NYAMBUNG", r"KBLI \d+ '.*' tidak berbagi satu kata", "TINJAU", ("kbli", "keg_utama", "produk"),
         "Judul KBLI tidak nyambung dgn 13a/13f: periksa kode KBLI-nya."),
    Pola("JK_VS_PEKERJA", r"12b '.*' tapi 24[ab]1", "TINJAU", ("jk", "tk_laki", "tk_pr"),
         "Pemilik ikut dihitung di 24: pekerja berjenis kelamin sama dgn pemilik (12b) tidak boleh 0."),
    Pola("KODEPOS_MINORITAS", r"kodepos \d+ beda dgn mayoritas desa", "TINJAU", ("kodepos",),
         "Kodepos beda dgn baris lain di desa yang sama: samakan kalau salah."),
    Pola("WILAYAH_PILIH_BEDA", r"wilayah tujuan ubah alokasi ambigu", "TINJAU", ("idsubsls", *_PILIH_WILAYAH),
         "Kolom Pilih PROVINSI..SUBSLS beda dgn idsubsls: samakan (menentukan wilayah tujuan pindah wilayah)."),
    Pola("VARIAN_BULANAN_DARI_KOLOM", r"varian bulanan \(mulai beroperasi", "TINJAU", ("tahun_operasi",),
         "Usaha mulai beroperasi tahun ini: pastikan angka 26-29 di sheet adalah nilai SEBULAN (rincian 30-33)."),
    Pola("NAMA_TERMUAT_NAMA_LAIN", r"nama '.*' terkandung di nama dokumen baris", "TINJAU", ("nama",),
         "Nama usaha termuat di nama dokumen baris lain: beri pembeda supaya pencarian dokumen tidak salah buka."),
    Pola("JALAN_KOSONG", r"Nama Jalan KOSONG", "TINJAU", ("jalan_domisili",), "Isi Nama Jalan."),
    # --- INFO ----------------------------------------------------------------
    Pola("HP_NOL_DEPAN_DIKEMBALIKAN", r"no WA '.*' -> '.*' \(nol di depan", "INFO", ("hp",),
         "Format sel no HP sbg Teks supaya nol di depan tidak hilang."),
    Pola("HP_KOSONG_JADI_9999", r"no WA '.*' kosong -> ", "INFO", ("hp",),
         "Kosong dianggap tidak ada/tidak bersedia (9999)."),
    Pola("UANG_KOSONG_JADI_0", r"indikator ekonomi kosong dianggap 0: (.+)$", "INFO",
         saran="Sel kosong dianggap 0 (ketetapan). Isi kalau sebenarnya ada nilainya.",
         keys_dari=lambda m, r: _keys_daftar(m.group(1))),
    Pola("TOTAL_TIDAK_DIISI", r"kolom total (\S+) di sheet = 0 padahal", "INFO",
         saran="Kolom total tidak diisi; form menghitung totalnya sendiri.",
         keys_dari=lambda m, r: _keys_total(m.group(1))[:1]),
    Pola("TOTAL_TIDAK_TERBACA_DIABAIKAN", r"kolom total '.*' \((\S+)\) tidak terbaca", "INFO",
         saran="Kolom total tidak terbaca (mis. '####' krn sel sempit); rinciannya yang dipakai.",
         keys_dari=lambda m, r: _keys_total(m.group(1))[:1]),
    Pola("16B_SATU_KODE", r"16b1-b6 diisi '.*' dari satu kolom", "INFO", ("internet_semua",)),
    Pola("16B_DIURAI", r"16b1-b6 (?:'.*' -> Ya utk|dari daftar)", "INFO", ("internet_semua",)),
    Pola("13B_DARI_KBLI", r"13b1/b2/b3 diturunkan dari golongan KBLI", "INFO",
         ("produk_sendiri", "layanan_mamin", "keg_penjualan")),
    Pola("13DE_DARI_JUDUL_KBLI", r"13d/13e diisi dari judul KBLI", "INFO", ("input_produksi", "proses_produksi")),
    Pola("13C_DEFAULT_MAKAN_MINUM", r"13c default '.*' \(usaha makan-minum", "INFO", ("lokasi_usaha",)),
    Pola("DEFAULT_TAHAP2", r"(\w+) default '.*' \(tidak ada di kuesioner tahap 2\)", "INFO",
         keys_dari=lambda m, r: (m.group(1),)),
    Pola("KODEPOS_DARI_DAFTAR", r"kodepos \d+ dari daftar wilayah", "INFO", ("kodepos",)),
    Pola("13F_DISALIN_DARI_13A", r"13f disalin dari 13a", "INFO", ("produk",)),
    Pola("NAMA_TANPA_12A", r"(nama dokumen|8b) tanpa \(12a\)", "INFO",
         saran="'<nama> (<12a>)' lebih dari 50 karakter -> nama dokumen/8b memakai nama usaha saja.",
         keys_dari=lambda m, r: ("nama",) if m.group(1) == "nama dokumen" else ("nama_komersial",)),
    Pola("NAMA_DIBEDAKAN", r"(?:usaha pecahan bernama sama|nama kembar persis) .* -> nama dibedakan", "INFO",
         ("nama_komersial",), "Usaha bernama sama dibedakan otomatis (13f/13a/wilayah)."),
)
_POLA_RE = tuple((p, re.compile(p.regex)) for p in POLA_TANDA)
JENIS_TANDA_LAIN = "TANDA_LAIN"
SARAN_TANDA_LAIN = "Pesan pemeriksaan yang belum dikenali program kontrol kualitas — baca keterangannya."
_cache_tanda: dict[str, tuple[Pola | None, re.Match | None]] = {}


def klasifikasi_tanda(teks: str, row=None) -> tuple[Pola | None, tuple[str, ...]]:
    """Pesan TANDA -> (pola yang cocok atau None, key rincian yang terlibat)."""
    if teks not in _cache_tanda:
        _cache_tanda[teks] = next(((p, m) for p, rx in _POLA_RE if (m := rx.match(teks))), (None, None))
    pola, m = _cache_tanda[teks]
    if pola is None:
        return None, _keys_umum(teks)
    return pola, tuple(dict.fromkeys((*pola.keys, *(pola.keys_dari(m, row) if pola.keys_dari else ()))))


# ---------------------------------------------------------------------------
# Kode MASALAH (baris ditolak) -> kolom & saran
# ---------------------------------------------------------------------------
KEYS_MASALAH: dict[str, tuple[str, ...]] = {
    "JALAN_KURANG_10_HURUF": ("jalan_domisili",),
    "13A_KURANG_15_KARAKTER": ("keg_utama",),
    "8B_TERLALU_PANJANG": ("nama_komersial", "pengusaha"),
    "BUMDES_BUKAN_KODE_6": ("badan_usaha",),
    "8B_DIAWALI_CV": ("nama_komersial",),
    "13C_MAMIN_BUKAN_5_11": ("lokasi_usaha",),
    "KBLI_KATEGORI_DITOLAK": ("kbli",),
    "WILAYAH_TIDAK_KONSISTEN": ("idsubsls", *_PILIH_WILAYAH),
    "PEKERJA_24_TIDAK_KONSISTEN": KEY_PEKERJA,
    "26A_HARUS_0_TANPA_PEKERJA_DIBAYAR": ("gaji", "tk_dibayar"),
    "26A_PER_PEKERJA_DI_BAWAH_MINIMAL": ("gaji", "tk_dibayar"),
    "MODAL_29_BUKAN_100": KEY_29,
    "HP_TIDAK_VALID": ("hp",),
    "NIK_TIDAK_VALID": ("nik_pengusaha",),
    "UMUR_DI_LUAR_10_99": ("umur",),
    "VARIAN_BULANAN": ("tahun_operasi",),
    "16B_TANPA_YA": ("internet", *KEY_16B),
    "16B_TIDAK_JELAS": ("internet_semua",),
    "26C_KATEGORI_TANPA_26C": ("biaya_pembelian", "kbli"),
    "26B_HARUS_LEBIH_0": ("biaya_produksi", "kbli"),
    "30C_HARUS_LEBIH_0": ("biaya_pembelian",),
    "NAMA_TUMPANG_TINDIH": ("nama", "nama_komersial", "pengusaha"),
    "13DE_TIDAK_ADA_DI_TAHAP2": ("input_produksi", "proses_produksi", "kbli"),
    "KODEPOS_TIDAK_DIKETAHUI": ("kodepos", "idsubsls"),
}
SARAN_UMUM = "Baca keterangan pemeriksaan, betulkan di sheet, lalu jalankan ulang kontrol kualitas."
SARAN_MASALAH: dict[str, str] = {
    "WAJIB_KOSONG": "Isi kolom yang kosong dari kuesioner.",
    "JALAN_KURANG_10_HURUF": "Tulis Nama Jalan/alamat lebih lengkap (minimal 10 huruf), atau '-' kalau memang tidak ada.",
    "13A_KURANG_15_KARAKTER": "Tulis kegiatan utama (13a) lebih lengkap, minimal 15 karakter.",
    "13F_KURANG_4_KARAKTER": "Tulis produk utama (13f) minimal 4 karakter.",
    "8B_TERLALU_PANJANG": "Singkat nama usaha (8b) atau nama pengusaha (12a): '<nama usaha> (<12a>)' maksimal "
                          "50 karakter.",
    "BUMDES_BUKAN_KODE_6": "Nama memuat BUMDes: isi 11a '6. BUM Desa', 11d Ya, dan modal pemerintah (29e) dominan.",
    "8B_DIAWALI_CV": "Tulis CV di belakang nama, mis. 'NAMA USAHA, CV'.",
    "NILAI_TIDAK_DIDUKUNG": "Nilai ini butuh pengisian manual (skrip hanya mendukung nilai standar); betulkan kalau "
                            "salah ketik, atau isi dokumennya manual.",
    "OPSI_TIDAK_ADA_DI_FORM": "Isi persis salah satu opsi form.",
    "13C_MAMIN_BUKAN_5_11": "Usaha makan-minum: 13c harus kode 5-11 (mis. '5. Kedai, stan, tenda' atau "
                            "'9. Restoran, warung makan, dan sejenisnya').",
    "KBLI_KATEGORI_DITOLAK": "Ganti KBLI dgn kode usaha yang sesuai; kategori P & U ditolak form.",
    "ANGKA_TIDAK_VALID": "Tulis angka bulat tanpa titik/koma/Rp (idsubsls 16 digit, kodepos 5, KBLI 5, tahun 4, "
                         "koordinat desimal di Indonesia). Sesudah dibetulkan jalankan ulang: pemeriksaan "
                         "konsistensi angka (24, 26-29, dst.) baru berjalan kalau semua angka valid.",
    "WILAYAH_TIDAK_KONSISTEN": "Samakan kolom Pilih PROVINSI..SUBSLS dgn idsubsls (salah satunya salah).",
    "PEKERJA_24_TIDAK_KONSISTEN": "24a1+24b1 (laki+perempuan) harus sama dgn 24a2+24b2 (dibayar+tidak dibayar).",
    "26A_HARUS_0_TANPA_PEKERJA_DIBAYAR": "Tidak ada pekerja dibayar (24a2 = 0) -> 26a harus 0; atau betulkan 24a2.",
    "26A_PER_PEKERJA_DI_BAWAH_MINIMAL": "26a dibagi 24a2 harus > Rp 50.000 per pekerja dibayar; betulkan 26a atau 24a2.",
    "MODAL_29_BUKAN_100": "29a-29f (persen modal) harus berjumlah 100.",
    "DI_BAWAH_MINIMAL": "Total pengeluaran (26f) & pendapatan (27c) minimal 100.000 (usaha baru tahun ini: 10.000). "
                        "Periksa angkanya di kuesioner.",
    "HP_TIDAK_VALID": "No HP: 08 + 8-11 digit (total 10-13 digit), sel format Teks; tidak ada -> 9999.",
    "NIK_TIDAK_VALID": "NIK 16 digit (sel format Teks), atau 7777 (>16 digit) / 8888 (belum punya) / 9999 (lainnya).",
    "UMUR_DI_LUAR_10_99": "Umur pengusaha 10-99; betulkan dari kuesioner (tidak ditebak skrip).",
    "VARIAN_BULANAN": "Usaha mulai beroperasi tahun ini memakai rincian 30-33 (angka SEBULAN): isi dokumennya "
                      "manual, atau betulkan tahun mulai beroperasi kalau salah.",
    "16B_TANPA_YA": "16a Ya -> minimal satu 16b1-16b6 Ya; kalau tidak ada, 16a seharusnya Tidak.",
    "16B_TIDAK_JELAS": "Tulis 16b1-b6 sbg 6 nilai 1/2 (mis. 2,1,2,2,2,2), satu kode 1/2, atau B1,B3 utk rincian "
                       "yang Ya.",
    "26C_KATEGORI_TANPA_26C": "KBLI bukan perdagangan tidak punya 26c: pindahkan nilainya ke 26b.",
    "26B_HARUS_LEBIH_0": "KBLI kategori B-F / golongan 56 wajib biaya produksi (26b) > 0.",
    "30C_HARUS_LEBIH_0": "Usaha dagang yang mulai tahun ini wajib 30c (pembelian barang dagangan) > 0: isi 26c.",
    "BARIS_GANDA": "Baris-baris ini identik (akun + wilayah + nama sama): hapus duplikatnya, atau bedakan namanya "
                   "kalau memang usaha berbeda.",
    "NAMA_TUMPANG_TINDIH": "Nama usaha satu baris termuat di nama baris lain -> pencarian dokumen bisa salah buka; "
                           "beri pembeda nama.",
    "TOTAL_TIDAK_COCOK": "Kolom total tidak sama dgn jumlah rinciannya: cari angka yang salah ketik.",
    "TOTAL_TIDAK_TERBACA": "Kolom total tidak terbaca sbg angka (mis. '####' krn sel terlalu sempit).",
    "13DE_TIDAK_ADA_DI_TAHAP2": "Usaha industri: form mewajibkan 13d (bahan baku) & 13e (proses). Tambahkan kolom "
                                "'13d' & '13e' di sheet, atau isi dokumennya manual.",
    "KODEPOS_TIDAK_DIKETAHUI": "Kodepos desa ini belum ada di daftar: tambahkan kolom 'kodepos' di sheet atau "
                               "lengkapi daftar kodepos di konfigurasi lokal.",
}


def keys_masalah(kode: str, pesan: str, row=None) -> tuple[str, ...]:
    """Kode masalah + pesannya -> key rincian yang harus dibetulkan."""
    if kode == "WAJIB_KOSONG" and ":" in pesan:
        return _keys_daftar(pesan.split(":", 1)[1])
    if kode == "ANGKA_TIDAK_VALID":
        if pesan.startswith("kolom:"):
            return _keys_daftar(pesan[len("kolom:"):])
        m = re.match(r"(\w+)=", pesan)          # "27d=150 > 100 persen", "tahun_operasi=2030 (...)"
        return (KEY_DARI_LABEL.get(m.group(1)) or (m.group(1),)) if m else ()
    if kode in ("OPSI_TIDAK_ADA_DI_FORM", "NILAI_TIDAK_DIDUKUNG"):
        m = re.match(r"(\w+)='", pesan)
        return (m.group(1),) if m else ()
    if kode == "DI_BAWAH_MINIMAL":
        m = re.match(r"(\w+)=", pesan)          # "26f=50000 < 100000"
        return KEY_DARI_LABEL.get(m.group(1), ()) if m else ()
    if kode in ("TOTAL_TIDAK_COCOK", "TOTAL_TIDAK_TERBACA"):
        m = re.search(r"kolom total (?:'.*' \((\S+)\)|(\S+) di sheet)", pesan)
        return _keys_total(m.group(1) or m.group(2)) if m else ()
    if kode == "13F_KURANG_4_KARAKTER":
        return _13f(None, row)
    if kode == "BARIS_GANDA":
        return ("nama", "idsubsls", "pengusaha" if isinstance(row, Tahap2Row) else "akun_ppl")
    return KEYS_MASALAH.get(kode) or _keys_umum(pesan)


def saran_masalah(kode: str, pesan: str, row=None) -> str:
    if kode == "OPSI_TIDAK_ADA_DI_FORM":
        m = re.match(r"(\w+)='", pesan)
        opsi = OPSI_FORM.get(m.group(1), ()) if m else ()
        if opsi:
            kode_saja = " (sheet tahap 2 boleh berisi nomor opsinya saja)" if isinstance(row, Tahap2Row) else ""
            return "Isi persis salah satu opsi form: " + " | ".join(opsi) + kode_saja
    return SARAN_MASALAH.get(kode, SARAN_UMUM)


# ---------------------------------------------------------------------------
# Sheet sumber (posisi kolom & nilai mentah)
# ---------------------------------------------------------------------------
@dataclass
class Sheet:
    format: str
    tab: str
    mentah: list[list]             # semua baris sheet apa adanya (indeks 0 = baris judul)
    idx: dict[str, int]            # key -> indeks kolom (0 = kolom A)

    @cached_property
    def judul(self) -> list[str]:
        return [_sel(j) for j in (self.mentah[0] if self.mentah else [])]

    @cached_property
    def lebar(self) -> int:
        return max((len(r) for r in self.mentah), default=0)

    def kolom(self, key: str) -> int | None:
        if key in self.idx:
            return self.idx[key]
        return next((self.idx[a] for a in ALIAS_KOLOM.get(key, ()) if a in self.idx), None)

    def nilai(self, baris: int, i: int) -> str:
        sel = self.mentah[baris - 1] if 0 < baris <= len(self.mentah) else []
        return _sel(sel[i]) if i < len(sel) else ""

    def nama_kolom(self, i: int) -> str:
        judul = self.judul
        return judul[i] if i < len(judul) else ""


def huruf_kolom(i: int) -> str:
    """Indeks kolom 0-based -> huruf Excel (0 -> A, 26 -> AA)."""
    n, huruf = i + 1, ""
    while n:
        n, sisa = divmod(n - 1, 26)
        huruf = chr(65 + sisa) + huruf
    return huruf


def nama_tab(path: Path, format_sumber: str) -> str:
    """Tab yang dibaca loader (aturan pilihnya SAMA dgn _baca_mentah / _baca_mentah_tahap2)."""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return path.name
    wb = load_workbook(path, read_only=True)
    try:
        per_nama = {n.strip().lower(): n for n in wb.sheetnames}
        calon = NAMA_SHEET_TAHAP2 if format_sumber == "tahap2" else NAMA_SHEET_DITERIMA
        cocok = next((per_nama[n] for n in calon if n in per_nama), "")
        return cocok or (wb.sheetnames[0] if format_sumber == "tahap2" and wb.sheetnames else "")
    finally:
        wb.close()


def baca_sheet(sumber: str, format_sumber: str) -> Sheet:
    """Baris mentah & posisi kolom dgn fungsi baca yang SAMA dgn loader."""
    path = Path(sumber)
    if format_sumber == "tahap2":
        mentah = _baca_mentah_tahap2(path)
        idx = _indeks_tahap2([_sel(j) for j in mentah[0]]) if mentah else {}
    else:
        mentah = _baca_mentah(path)
        idx = _cari_indeks([_sel(j) for j in mentah[0]]) if mentah else {}
    return Sheet(format_sumber, nama_tab(path, format_sumber), mentah, idx)


# ---------------------------------------------------------------------------
# Temuan
# ---------------------------------------------------------------------------
@dataclass
class Temuan:
    baris: int
    kategori: str
    jenis: str
    pesan: str                       # teks ASLI dari pemeriksaan
    keys: tuple[str, ...] = ()
    kolom: tuple[int, ...] = ()      # indeks kolom sheet (0-based) yang terlibat
    saran: str = ""

    @property
    def seragam(self) -> bool:
        """INFO utk rincian yang tidak punya kolom di sheet (default/tafsiran yang sama utk
        banyak baris) -> cukup diringkas, tidak didaftar per baris."""
        return self.kategori == "INFO" and not self.kolom


def temuan_baris(row: GabunganRow, cek: Pemeriksaan, sheet: Sheet) -> list[Temuan]:
    out: list[Temuan] = []

    def tambah(kategori: str, jenis: str, pesan: str, keys: tuple, saran: str) -> None:
        keys = tuple(dict.fromkeys(k for k in keys if k))
        kolom = tuple(dict.fromkeys(i for i in (sheet.kolom(k) for k in keys) if i is not None))
        out.append(Temuan(row.baris, kategori, jenis, pesan, keys, kolom, saran))

    for kode, pesan in cek.masalah:
        tambah("DITOLAK", kode, pesan, keys_masalah(kode, pesan, row), saran_masalah(kode, pesan, row))
    for teks in cek.tanda:
        pola, keys = klasifikasi_tanda(teks, row)
        if pola is None:
            tambah("TINJAU", JENIS_TANDA_LAIN, teks, keys, SARAN_TANDA_LAIN)
        else:
            tambah(pola.kategori, pola.jenis, teks, keys, pola.saran)
    return out


def status_qc(temuan: list[Temuan]) -> str:
    """Kategori terparah selain INFO; "BERSIH" kalau tidak ada."""
    kat = [t.kategori for t in temuan if t.kategori != "INFO"]
    return min(kat, key=PERINGKAT.__getitem__) if kat else "BERSIH"


def ppl_baris(row: GabunganRow) -> str:
    """Nama PPL utk rekap: kolom "Nama PPL" (tahap 2) kalau ada, kalau tidak akun PPL."""
    nama = row.info.get("nama_ppl", "") if isinstance(row, Tahap2Row) else ""
    return " ".join(str(nama or "").split()) or row.akun_ppl or "(tanpa PPL)"


def _kunci_ppl(row: GabunganRow) -> str:
    """Pengelompokan PPL tanpa membedakan huruf besar/kecil & spasi ganda."""
    return " ".join(ppl_baris(row).lower().split())


def nilai_skrip(row: GabunganRow, keys: tuple[str, ...]) -> str:
    """Nilai yang akan DIKETIK skrip utk rincian yang terlibat (sesudah koreksi loader)."""
    bagian, sudah_16b = [], False
    for k in keys:
        if k in KEY_16B or k == "internet_semua":
            if not sudah_16b:
                sudah_16b = True
                bagian.append("16b1-b6: " + ",".join(
                    "Ya" if row[x].startswith("1") else "Tidak" if row[x].startswith("2") else "-" for x in KEY_16B))
            continue
        if k in row.v:
            bagian.append(f"{LABEL_KEY.get(k, k)}: {row[k] or '(kosong)'}")
    return "; ".join(bagian)


def nomor_ringkas(baris) -> str:
    """[2,3,4,9] -> '2-4,9' (bisa langsung dipakai sbg --baris)."""
    urut, bagian = sorted(set(baris)), []
    for n in urut:
        if bagian and n == bagian[-1][1] + 1:
            bagian[-1][1] = n
        else:
            bagian.append([n, n])
    return ",".join(f"{a}-{b}" if b > a else f"{a}" for a, b in bagian)


@dataclass
class Laporan:
    sumber: str
    format: str
    mode: str
    sheet: Sheet
    rows: list[GabunganRow]              # SEMUA baris sheet (pemeriksaan lintas baris selalu atas semuanya)
    dipilih: list[GabunganRow]           # baris yang dilaporkan
    hasil: dict[int, Pemeriksaan]
    temuan: dict[int, list[Temuan]]      # per baris dipilih
    status_audit: dict[str, str] = field(default_factory=dict)
    dibuat: str = field(default_factory=lambda: datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))

    def semua_temuan(self, dgn_seragam: bool = False) -> list[Temuan]:
        return [t for r in self.dipilih for t in self.temuan[r.baris] if dgn_seragam or not t.seragam]

    def baris(self, nomor: int) -> GabunganRow:
        return self._per_baris[nomor]

    def __post_init__(self):
        self._per_baris = {r.baris: r for r in self.rows}


def keterangan_mode(format_sumber: str, satu_subsls: bool, koordinat: str | None, rows) -> str:
    otomatis = mg.koordinat_otomatis(koordinat, format_sumber)
    bagian = ["satu subsls (semua dokumen di satu list)" if satu_subsls else "per baris (alur lama)",
              ("koordinat kirim (tanpa koordinat tetap dikirim)" if mg.koordinat_dikirim(koordinat)
               else "koordinat otomatis (tanpa koordinat -> DRAFT)" if otomatis else "koordinat wajib")]
    if format_sumber == "tahap2":
        bagian.append("format tahap 2 (rincian di luar kuesioner kertas = default config)")
    else:
        murni = rows[0].murni if rows else mg.GABUNGAN_MODE_MURNI
        bagian.append("mode murni (isian apa adanya)" if murni else "mode normal (aturan & koreksi aktif)")
    return " | ".join(bagian)


def periksa_sumber(sumber: str, format_sumber: str = "standar", *, mode_satu_subsls: bool = True,
                   koordinat: str | None = None, kodepos: str = "", cek_total: bool = True,
                   baris: str | None = None, dari: int | None = None, sampai: int | None = None,
                   hanya_belum_terkirim: bool = False, audit: list[dict] | None = None) -> Laporan:
    """Pemeriksaan offline (SAMA dgn main_gabungan --cek) + pemetaan ke sel/kategori.
    `audit` None = baca audit_log_gabungan.csv di folder kerja (kalau ada)."""
    rows, hasil = mg.muat_sumber(sumber, format_sumber, mode_satu_subsls, kodepos, cek_total=cek_total,
                                 izinkan_tanpa_koordinat=mg.koordinat_otomatis(koordinat, format_sumber))
    sheet = baca_sheet(sumber, format_sumber)
    status_audit = mg.status_terakhir_dari(mg._baca_audit() if audit is None else audit)
    dipilih = list(rows)
    if baris:
        ingin = parse_pilihan_baris(baris)
        dipilih = [r for r in dipilih if r.baris in ingin]
    if dari is not None or sampai is not None:
        dipilih = mg.saring_rentang(dipilih, dari, sampai)
    if hanya_belum_terkirim:
        dipilih = [r for r in dipilih if status_audit.get(r.kunci, "") not in mg.STATUS_TERKIRIM]
    temuan = {r.baris: temuan_baris(r, hasil[r.baris], sheet) for r in dipilih}
    return Laporan(sumber, format_sumber, keterangan_mode(format_sumber, mode_satu_subsls, koordinat, rows),
                   sheet, rows, dipilih, hasil, temuan, status_audit)


# ---------------------------------------------------------------------------
# Rekap
# ---------------------------------------------------------------------------
def rekap_jenis(lap: Laporan) -> list[dict]:
    """Per (kategori, jenis): baris terkena, contoh pesan, saran. Urut keparahan lalu jumlah."""
    grup: dict[tuple, dict] = {}
    for t in lap.semua_temuan():
        g = grup.setdefault((t.kategori, t.jenis), {"kategori": t.kategori, "jenis": t.jenis, "baris": set(),
                                                    "temuan": 0, "contoh": t.pesan, "saran": t.saran})
        g["baris"].add(t.baris)
        g["temuan"] += 1
    return sorted(grup.values(), key=lambda g: (PERINGKAT[g["kategori"]], -len(g["baris"]), g["jenis"]))


def rekap_seragam(lap: Laporan) -> list[dict]:
    """INFO utk rincian tanpa kolom di sheet, dikelompokkan per (jenis, rincian)."""
    grup: dict[tuple, dict] = {}
    for r in lap.dipilih:
        for t in lap.temuan[r.baris]:
            if t.seragam:
                rincian = ", ".join(LABEL_KEY.get(k, k) for k in t.keys)
                g = grup.setdefault((t.jenis, rincian), {"jenis": t.jenis, "rincian": rincian, "baris": set(),
                                                         "contoh": t.pesan})
                g["baris"].add(t.baris)
    return sorted(grup.values(), key=lambda g: (-len(g["baris"]), g["jenis"], g["rincian"]))


def rekap_ppl(lap: Laporan) -> list[dict]:
    grup: dict[str, dict] = {}
    for r in lap.dipilih:
        g = grup.setdefault(_kunci_ppl(r), {"ppl": ppl_baris(r), "baris": 0, "jenis": Counter(),
                                           **{k: 0 for k in (*KATEGORI_TINDAKAN, "BERSIH")}})
        g["baris"] += 1
        temuan = lap.temuan[r.baris]
        for kat in {t.kategori for t in temuan if t.kategori != "INFO"}:
            g[kat] += 1
        g["BERSIH"] += status_qc(temuan) == "BERSIH"
        g["jenis"].update({t.jenis for t in temuan if t.kategori != "INFO"})
    return sorted(grup.values(), key=lambda g: (-g["DITOLAK"], -g["DIGANTI"], -g["baris"], g["ppl"]))


def rekap_kolom(lap: Laporan) -> list[dict]:
    grup: dict[int, Counter] = defaultdict(Counter)
    for t in lap.semua_temuan():
        for i in t.kolom:
            grup[i][t.kategori] += 1
    keluar = [{"kolom": i, **{k: c[k] for k in KATEGORI}, "total": sum(c[k] for k in KATEGORI_TINDAKAN)}
              for i, c in grup.items()]
    return sorted(keluar, key=lambda d: (-d["total"], d["kolom"]))


def tanda_tak_dikenal(lap: Laporan) -> Counter:
    return Counter(t.pesan for t in lap.semua_temuan() if t.jenis == JENIS_TANDA_LAIN)


# ---------------------------------------------------------------------------
# Tulis Excel
# ---------------------------------------------------------------------------
KOLOM_TEMUAN = ("No", "Baris", "Sel", "Kolom", "Kategori", "Jenis", "Nilai di sheet", "Nilai dipakai skrip",
                "Keterangan pemeriksaan", "Saran perbaikan", "PPL", "Nama usaha (nama dokumen)", "Status input",
                "Status audit")


def _aman(v):
    """Teks tanpa karakter kontrol yang ditolak openpyxl."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


_ISI: dict[str, PatternFill] = {}


def _isi(warna: str) -> PatternFill:
    if warna not in _ISI:
        _ISI[warna] = PatternFill("solid", start_color=warna, end_color=warna)
    return _ISI[warna]


def _kepala(ws, baris: int, kepala, warna: str = "D9E1F2") -> None:
    for j, h in enumerate(kepala, 1):
        c = ws.cell(baris, j, h)
        c.font = Font(bold=True)
        c.fill = _isi(warna)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _lebar(ws, lebar: dict[str, float]) -> None:
    for huruf, w in lebar.items():
        ws.column_dimensions[huruf].width = w


def baris_temuan(lap: Laporan, temuan: list[Temuan]) -> list[list]:
    """Isi lembar 'Temuan' (urut keparahan, lalu nomor baris)."""
    urut = sorted(temuan, key=lambda t: (PERINGKAT[t.kategori], t.baris, t.jenis))
    keluar = []
    for no, t in enumerate(urut, 1):
        row = lap.baris(t.baris)
        sel = ", ".join(f"{huruf_kolom(i)}{t.baris}" for i in t.kolom)
        judul = " | ".join(lap.sheet.nama_kolom(i) for i in t.kolom)
        if len(t.kolom) == 1:
            nilai = lap.sheet.nilai(t.baris, t.kolom[0]) or "(kosong)"
        else:
            nilai = "; ".join(f"{lap.sheet.nama_kolom(i)[:24]}: {lap.sheet.nilai(t.baris, i) or '(kosong)'}"
                              for i in t.kolom)
        dipakai = "" if t.kategori == "DITOLAK" else nilai_skrip(row, t.keys)
        keluar.append([no, t.baris, sel, judul, t.kategori, t.jenis, nilai, dipakai, t.pesan, t.saran,
                       ppl_baris(row), row.nama_dokumen, lap.hasil[t.baris].status,
                       lap.status_audit.get(row.kunci, "")])
    return keluar


def _lembar_temuan(ws, lap: Laporan, temuan: list[Temuan], tautan: bool) -> None:
    ws.title = "Temuan"
    _kepala(ws, 1, KOLOM_TEMUAN)
    rapat = Alignment(wrap_text=True, vertical="top")
    isi_semua = baris_temuan(lap, temuan)
    # Nomor baris dihitung sendiri: ws.max_row menyisir SEMUA sel tiap dipanggil (O(n²)).
    for r, isi in enumerate(isi_semua, start=2):
        ws.append([_aman(v) for v in isi])
        ws.cell(r, 5).fill = _isi(WARNA[isi[4]])
        for j in (9, 10):          # keterangan & saran: teks panjang
            ws.cell(r, j).alignment = rapat
        if tautan and isi[2]:
            c = ws.cell(r, 3)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'Data bertanda'!{isi[2].split(',')[0].strip()}")
            c.style = "Hyperlink"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{huruf_kolom(len(KOLOM_TEMUAN) - 1)}{len(isi_semua) + 1}"
    _lebar(ws, {"A": 6, "B": 7, "C": 12, "D": 24, "E": 11, "F": 26, "G": 26, "H": 30, "I": 60, "J": 50,
                "K": 22, "L": 34, "M": 26, "N": 24})


def _lembar_ringkasan(ws, lap: Laporan) -> None:
    ws.title = "Ringkasan"
    rapat = Alignment(wrap_text=True, vertical="top")
    nomor = nomor_ringkas(r.baris for r in lap.dipilih)
    rentang = f" (baris {nomor if len(nomor) <= 80 else nomor[:77] + '...'})" if lap.dipilih else ""
    ws["A1"] = "Kontrol kualitas sumber data input usaha"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = _aman(f"Sumber: {lap.sumber} — tab '{lap.sheet.tab}' — format {lap.format} — dibuat {lap.dibuat}")
    ws["A3"] = f"{len(lap.dipilih)} dari {len(lap.rows)} baris diperiksa{rentang}. Mode: {lap.mode}"
    ws["A4"] = ("Aturan = pemeriksaan offline main_gabungan --cek saat ini (tidak ada aturan tambahan). "
                "Sel bermasalah diwarnai di lembar 'Data bertanda'; rinciannya di lembar 'Temuan'.")
    r = 6

    def tabel(judul: str, kepala, isi, warna_kolom: int | None = None) -> None:
        nonlocal r
        ws.cell(r, 1, judul).font = Font(bold=True, size=12)
        _kepala(ws, r + 1, kepala)
        r += 2
        for baris in isi:
            for j, v in enumerate(baris, 1):
                c = ws.cell(r, j, _aman(v))
                c.alignment = rapat
            if warna_kolom and baris[warna_kolom - 1] in WARNA:
                ws.cell(r, warna_kolom).fill = _isi(WARNA[baris[warna_kolom - 1]])
            r += 1
        r += 1

    status = Counter(lap.hasil[b.baris].status for b in lap.dipilih)
    tabel("Status input baris (keputusan main_gabungan)", ("Status", "Baris", "Arti"),
          [(s, n, ARTI_STATUS.get(s, "TIDAK diinput — betulkan sheet (lihat kategori DITOLAK)"))
           for s, n in sorted(status.items(), key=lambda t: (t[0].startswith("SKIP"), -t[1]))])

    per_kat = {k: [set(), 0] for k in KATEGORI}
    for t in lap.semua_temuan(dgn_seragam=True):
        per_kat[t.kategori][0].add(t.baris)
        per_kat[t.kategori][1] += 1
    bersih = sum(status_qc(lap.temuan[b.baris]) == "BERSIH" for b in lap.dipilih)
    tabel("Kategori temuan", ("Kategori", "Baris", "Temuan", "Arti / tindakan"),
          [*[(k, len(per_kat[k][0]), per_kat[k][1], ARTI_KATEGORI[k]) for k in KATEGORI],
           ("BERSIH", bersih, "", "Baris tanpa temuan selain INFO.")], warna_kolom=1)

    tabel("Jenis temuan (urut keparahan, lalu jumlah baris)",
          ("Kategori", "Jenis", "Baris", "Temuan", "Nomor baris (format --baris)", "Saran perbaikan",
           "Contoh keterangan"),
          [(g["kategori"], g["jenis"], len(g["baris"]), g["temuan"], nomor_ringkas(g["baris"]), g["saran"],
            g["contoh"]) for g in rekap_jenis(lap)], warna_kolom=1)

    seragam = rekap_seragam(lap)
    if seragam:
        tabel("Asumsi utk rincian yang TIDAK ada kolomnya di sheet (sama utk banyak baris; tidak perlu tindakan)",
              ("Jenis", "Rincian", "Baris", "Contoh keterangan"),
              [(g["jenis"], g["rincian"], len(g["baris"]), g["contoh"]) for g in seragam])
    tak_dikenal = tanda_tak_dikenal(lap)
    if tak_dikenal:
        tabel("Pesan pemeriksaan yang BELUM dikenali program ini (dimasukkan TINJAU)", ("Pesan", "Temuan"),
              tak_dikenal.most_common())
    _lebar(ws, {"A": 26, "B": 30, "C": 10, "D": 10, "E": 34, "F": 60, "G": 70})


def _lembar_ppl(ws, lap: Laporan) -> None:
    ws.title = "Per PPL"
    kepala = ("PPL", "Baris", *KATEGORI_TINDAKAN, "BERSIH", "% bersih", "Jenis terbanyak")
    _kepala(ws, 1, kepala)
    for g in rekap_ppl(lap):
        ws.append([_aman(g["ppl"]), g["baris"], *(g[k] for k in KATEGORI_TINDAKAN), g["BERSIH"],
                   round(100 * g["BERSIH"] / g["baris"], 1) if g["baris"] else 0,
                   ", ".join(f"{j} ({n})" for j, n in g["jenis"].most_common(3))])
    for j, k in enumerate(KATEGORI_TINDAKAN, 3):
        ws.cell(1, j).fill = _isi(WARNA[k])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{huruf_kolom(len(kepala) - 1)}{max(ws.max_row, 1)}"
    _lebar(ws, {"A": 32, "J": 70})


def _lembar_kolom(ws, lap: Laporan) -> None:
    ws.title = "Per kolom"
    kepala = ("Kolom", "Judul kolom", *KATEGORI, "Total (selain INFO)")
    _kepala(ws, 1, kepala)
    for d in rekap_kolom(lap):
        ws.append([huruf_kolom(d["kolom"]), _aman(lap.sheet.nama_kolom(d["kolom"])),
                   *(d[k] for k in KATEGORI), d["total"]])
    for j, k in enumerate(KATEGORI, 3):
        ws.cell(1, j).fill = _isi(WARNA[k])
    ws.freeze_panes = "C2"
    _lebar(ws, {"A": 8, "B": 44})


def _lembar_data(ws, lap: Laporan) -> None:
    """Salinan sheet (nomor baris & huruf kolom SAMA dgn aslinya) + sel bermasalah diwarnai."""
    ws.title = "Data bertanda"
    for sel in lap.sheet.mentah:
        ws.append([_aman(v) for v in sel])
    lebar = lap.sheet.lebar
    k_status, k_ringkas = lebar + 1, lebar + 2
    for j in range(1, lebar + 1):
        ws.cell(1, j).font = Font(bold=True)
        ws.cell(1, j).fill = _isi("D9E1F2")
    for j, h in ((k_status, "QC_STATUS"), (k_ringkas, "QC_TEMUAN")):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = _isi("BDD7EE")

    per_sel: dict[tuple[int, int], list[Temuan]] = defaultdict(list)
    for r in lap.dipilih:
        temuan = [t for t in lap.temuan[r.baris] if t.kategori != "INFO"]
        for t in temuan:
            for i in t.kolom:
                per_sel[(r.baris, i)].append(t)
        st = status_qc(lap.temuan[r.baris])
        c = ws.cell(r.baris, k_status, st)
        c.fill = _isi(WARNA[st])
        ringkas = " | ".join(dict.fromkeys(f"{t.kategori}: {t.jenis}" for t in
                                           sorted(temuan, key=lambda t: PERINGKAT[t.kategori])))
        ws.cell(r.baris, k_ringkas, _aman(ringkas))
    for (b, i), daftar in per_sel.items():
        daftar.sort(key=lambda t: PERINGKAT[t.kategori])
        c = ws.cell(b, i + 1)
        c.fill = _isi(WARNA[daftar[0].kategori])
        teks = "\n".join(f"[{t.kategori}] {t.pesan}" + (f"\n-> {t.saran}" if t.saran else "") for t in daftar)
        c.comment = Comment(_aman(teks[:1500]), "kontrol_kualitas", width=360, height=180)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{huruf_kolom(k_ringkas - 1)}{max(len(lap.sheet.mentah), 1)}"
    ws.column_dimensions[huruf_kolom(k_ringkas - 1)].width = 60


def tulis_excel(lap: Laporan, path: Path, salinan: bool = True) -> None:
    wb = Workbook()
    _lembar_ringkasan(wb.active, lap)
    _lembar_temuan(wb.create_sheet(), lap, lap.semua_temuan(), tautan=salinan)
    _lembar_ppl(wb.create_sheet(), lap)
    _lembar_kolom(wb.create_sheet(), lap)
    if salinan:
        _lembar_data(wb.create_sheet(), lap)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _nama_berkas(teks: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", teks.replace("@", "_at_")).strip("_") or "tanpa_ppl"


def tulis_per_ppl(lap: Laporan, folder: Path) -> list[Path]:
    """Satu berkas per PPL: temuan yang perlu tindakan (tanpa INFO) di baris milik PPL itu."""
    per_ppl: dict[str, list[Temuan]] = defaultdict(list)
    nama_ppl: dict[str, str] = {}
    for r in lap.dipilih:
        nama_ppl.setdefault(_kunci_ppl(r), ppl_baris(r))
        per_ppl[_kunci_ppl(r)].extend(t for t in lap.temuan[r.baris] if t.kategori != "INFO")
    folder.mkdir(parents=True, exist_ok=True)
    dipakai: Counter = Counter()
    keluar = []
    for kunci, temuan in sorted(per_ppl.items()):
        ppl = nama_ppl[kunci]
        if not temuan:
            continue
        nama = _nama_berkas(ppl)
        dipakai[nama] += 1
        path = folder / (f"{nama}.xlsx" if dipakai[nama] == 1 else f"{nama}_{dipakai[nama]}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Ringkasan"
        ws.append([f"Kontrol kualitas — {ppl}"])
        ws.append([f"Sumber: {lap.sumber} — dibuat {lap.dibuat}"])
        ws.append([])
        ws.append(["Kategori", "Baris", "Arti / tindakan"])
        for k in KATEGORI_TINDAKAN:
            n = len({t.baris for t in temuan if t.kategori == k})
            if n:
                ws.append([k, n, ARTI_KATEGORI[k]])
        _lebar(ws, {"A": 14, "B": 8, "C": 110})
        _lembar_temuan(wb.create_sheet(), lap, temuan, tautan=False)
        wb.save(path)
        keluar.append(path)
    return keluar


def tulis_csv(lap: Laporan, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(KOLOM_TEMUAN)
        w.writerows(baris_temuan(lap, lap.semua_temuan()))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def cetak_ringkasan(lap: Laporan, keluaran: Path) -> None:
    print(f"=== KONTROL KUALITAS: {lap.sumber} (format {lap.format}, tab '{lap.sheet.tab}') ===")
    print(f"{len(lap.dipilih)} dari {len(lap.rows)} baris diperiksa. Mode: {lap.mode}\n")
    status = Counter(lap.hasil[b.baris].status for b in lap.dipilih)
    bisa = sum(n for s, n in status.items() if mg.hasil_ok(s))
    print(f"Status input: {bisa} bisa diinput "
          f"({status.get('SIAP', 0)} SIAP, {status.get(mg.STATUS_SIAP_TANPA_KOORDINAT, 0)} tanpa koordinat -> "
          f"DRAFT), {len(lap.dipilih) - bisa} DITOLAK")

    per_kat: dict[str, set] = defaultdict(set)
    for t in lap.semua_temuan(dgn_seragam=True):
        per_kat[t.kategori].add(t.baris)
    print("\nKategori (jumlah baris):")
    for k in KATEGORI:
        print(f"  {k:<10}{len(per_kat[k]):6d}  {ARTI_KATEGORI[k][:92]}")
    bersih = sum(status_qc(lap.temuan[b.baris]) == "BERSIH" for b in lap.dipilih)
    print(f"  {'BERSIH':<10}{bersih:6d}  baris tanpa temuan selain INFO")

    jenis = [g for g in rekap_jenis(lap) if g["kategori"] != "INFO"]
    if jenis:
        print("\nJenis temuan yang perlu tindakan (baris, kategori, jenis, contoh nomor baris):")
        for g in jenis[:20]:
            print(f"  {len(g['baris']):6d}  {g['kategori']:<10} {g['jenis']:<34} {nomor_ringkas(g['baris'])[:60]}")
        if len(jenis) > 20:
            print(f"  ... {len(jenis) - 20} jenis lain — lihat lembar Ringkasan")
    tak_dikenal = tanda_tak_dikenal(lap)
    if tak_dikenal:
        print(f"\n⚠️ {sum(tak_dikenal.values())} pesan pemeriksaan belum dikenali pola kolomnya "
              "(dimasukkan TINJAU tanpa sel) — tambahkan polanya di POLA_TANDA:")
        for pesan, n in tak_dikenal.most_common(5):
            print(f"  {n:6d}  {pesan[:120]}")
    print(f"\nLaporan: {keluaran} (lembar Ringkasan, Temuan, Per PPL, Per kolom, Data bertanda)")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Kontrol kualitas sumber data input usaha (format standar / tahap 2) — offline, "
                    "memakai pemeriksaan yang sama dgn main_gabungan --cek.")
    ap.add_argument("--sumber", required=True, help="File .xlsx/.csv sumber (sama dgn --sumber main_gabungan)")
    ap.add_argument("--format", choices=("standar", "tahap2"), default="standar",
                    help="standar = input_usaha.xlsx / Agenda*.xlsx; tahap2 = hasil pendataan kertas tahap 2")
    ap.add_argument("--baris", default=None, help="Hanya laporkan nomor baris sheet ini, mis. 2,5,10-20")
    ap.add_argument("--dari", type=int, default=None, help="Laporkan mulai baris sheet ke-N")
    ap.add_argument("--sampai", type=int, default=None, help="Laporkan sampai baris sheet ke-N")
    ap.add_argument("--koordinat", choices=("otomatis", "wajib", "kirim"), default=None,
                    help="Sama dgn main_gabungan (bawaan: otomatis utk tahap2, wajib utk standar)")
    ap.add_argument("--per-baris", action="store_true",
                    help="Periksa spt alur lama (dokumen per akun & subsls baris); bawaan mode satu subsls")
    ap.add_argument("--abaikan-cek-total", action="store_true", help="Tahap 2: jangan bandingkan kolom TOTAL")
    ap.add_argument("--kodepos", default="", help="Tahap 2: kodepos cadangan (sama dgn main_gabungan)")
    ap.add_argument("--hanya-belum-terkirim", action="store_true",
                    help="Lewati baris yang menurut audit_log_gabungan.csv sudah terkirim")
    ap.add_argument("--keluaran", default=str(KELUARAN_PATH), help=f"Berkas Excel hasil (bawaan {KELUARAN_PATH})")
    ap.add_argument("--csv", default=None, help="Tulis juga daftar temuan sbg CSV ke path ini")
    ap.add_argument("--tanpa-salinan", action="store_true",
                    help="Jangan tulis lembar 'Data bertanda' (salinan sheet berwarna) — lebih cepat & kecil")
    ap.add_argument("--per-ppl", nargs="?", const=FOLDER_PER_PPL, default=None, metavar="FOLDER",
                    help=f"Tulis juga satu berkas per PPL (bawaan folder {FOLDER_PER_PPL}/) utk dibagikan")
    mg.opsi_audit(ap)
    args = ap.parse_args(argv)
    mg.pakai_audit(args.audit)

    if args.dari is not None and args.sampai is not None and args.dari > args.sampai:
        print(f"❌ --dari {args.dari} lebih besar dari --sampai {args.sampai}.", file=sys.stderr)
        return 2
    try:
        lap = periksa_sumber(args.sumber, args.format, mode_satu_subsls=not args.per_baris,
                             koordinat=args.koordinat, kodepos=args.kodepos, cek_total=not args.abaikan_cek_total,
                             baris=args.baris, dari=args.dari, sampai=args.sampai,
                             hanya_belum_terkirim=args.hanya_belum_terkirim)
    except (FileNotFoundError, ValueError) as e:
        print(f"❌ {e}", file=sys.stderr)
        if args.format == "standar" and ("tidak sesuai" in str(e) or "tidak ada di" in str(e)):
            print("   Kalau ini sheet hasil pendataan kertas tahap 2, tambahkan --format tahap2.", file=sys.stderr)
        return 2
    if not lap.dipilih:
        print("Tidak ada baris yang cocok dgn pilihan --baris/--dari/--sampai.", file=sys.stderr)
        return 2

    keluaran = Path(args.keluaran)
    try:
        tulis_excel(lap, keluaran, salinan=not args.tanpa_salinan)
        if args.csv:
            tulis_csv(lap, Path(args.csv))
    except PermissionError as e:
        print(f"❌ Tidak bisa menulis {e.filename}: tutup berkasnya di Excel dulu, lalu jalankan ulang.",
              file=sys.stderr)
        return 2
    cetak_ringkasan(lap, keluaran)
    if args.csv:
        print(f"Daftar temuan (CSV): {args.csv}")
    if args.per_ppl:
        berkas = tulis_per_ppl(lap, Path(args.per_ppl))
        print(f"Berkas per PPL: {len(berkas)} di {args.per_ppl}/ (hanya temuan yang perlu tindakan)")
    return 1 if any(not lap.hasil[b.baris].bisa_diproses for b in lap.dipilih) else 0


if __name__ == "__main__":
    sys.exit(main())
