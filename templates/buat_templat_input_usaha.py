#!/usr/bin/env python3
"""
buat_templat_input_usaha.py — bangkitkan templates/input_usaha.xlsx: templat SATU-SATUNYA format
input usaha (format tahap 2 = kuesioner kertas SE2026), berisi judul kolom + SATU baris contoh
fiktif + tab petunjuk.

    python templates/buat_templat_input_usaha.py

Baris contoh ditandai "CONTOH" di kolom "Uraian:" -> loader SELALU menolaknya
(SKIP_DATA_BARIS_CONTOH), jadi aman walau lupa dihapus. Skrip ini MEMBUKTIKAN baris contoh
benar: dibaca inti/tahap2_loader.py, ditolak HANYA karena penanda contoh, dan berstatus SIAP
begitu penandanya dihapus. Kalau judul kolom / kode opsi di loader berubah, jalankan ulang.
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ["FASIH_ABAIKAN_CONFIG_LOKAL"] = "1"   # verifikasi memakai data contoh inti/config.py
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from inti.config import TAHAP2_DEFAULT  # noqa: E402
from inti.id_dokumen import JUDUL_ID  # noqa: E402
from inti.tahap2_loader import (  # noqa: E402
    KOLOM_TAHAP2_OPSIONAL, KOLOM_TAHAP2_TAMBAHAN, PENANDA_CONTOH, load_tahap2, periksa_semua_tahap2,
)

KELUARAN = Path(__file__).resolve().parent / "input_usaha.xlsx"
NAMA_TAB = "input_usaha"

# (judul kolom persis, key loader, rincian kuesioner, cara mengisi, nilai contoh fiktif)
KOLOM: list[tuple[str, str, str, str, str]] = [
    ("Sumber/Kec.", "info_sumber_kec", "info", "bebas (tidak dikirim)", "010"),
    ("Periode", "info_periode", "info", "bebas (tidak dikirim)", "21Sep"),
    ("Uraian:", "info_uraian", "info", f"bebas; baris yang diawali {PENANDA_CONTOH} SELALU ditolak",
     f"{PENANDA_CONTOH} - hapus baris ini"),
    ("Nama PPL", "info_nama_ppl", "petugas", "nama pencacah (rekap kontrol kualitas per PPL)", "I KETUT CONTOH"),
    ("3", "info_kec", "3. kecamatan", "info: nama + kode", "GEROKGAK 510801"),
    ("4", "info_desa", "4. desa", "info: nama + kode", "PATAS 0010"),
    ("5", "idsubsls", "5-7. wilayah", "idsubsls 16 digit wilayah ASLI usaha (sel Teks)", "5108010008000101"),
    ("8b.", "nama_komersial", "8b. nama usaha", "nama dokumen = '<8b> (<12a>)', maks 50 karakter",
     "WARUNG CONTOH SEJAHTERA"),
    ("8c.", "jalan_domisili", "8c. alamat", "-> Nama Jalan SE2026-P; pendek dilengkapi nama wilayah",
     "JALAN RAYA CONTOH NOMOR 1"),
    ("no WA", "hp", "12e. no HP", "08 + 8-11 digit (sel Teks); tidak ada -> 9999", "081234567890"),
    ("12a", "pengusaha", "12a. nama pengusaha", "nama orang", "NI LUH CONTOH"),
    ("12b", "jk", "12b. jenis kelamin", "1 Laki-laki | 2 Perempuan", "2"),
    ("12c", "umur", "12c. umur", "10-99", "34"),
    ("12d", "nik_pengusaha", "12d. NIK", "16 digit (sel Teks), atau 7777 / 8888 / 9999", "9999"),
    ("13a", "keg_utama", "13a. kegiatan utama", "min 15 karakter (pendek dilengkapi judul KBLI)",
     "MENJUAL BERAS DAN KEBUTUHAN POKOK SECARA ECERAN"),
    ("13f", "produk", "13f. produk utama", "min 4 karakter", "BERAS ECERAN"),
    ("14a", "jaringan", "14a. jaringan usaha", "1 Tunggal | 2 Kantor pusat | 3 Cabang | 4 Perwakilan | "
     "5 Pabrik | 6 Unit pembantu", "1"),
    ("16a", "internet", "16a. pakai internet", "1 Ya | 2 Tidak", "2"),
    ("16b1-b6", "internet_semua", "16b1-b6. tujuan internet", "kalau 16a = 1: enam nilai 1/2 dipisah koma "
     "(b1..b6), atau B1,B3 (yang Ya); kosong kalau 16a = 2", ""),
    ("17b", "perlindungan_lingkungan", "17b. perlindungan lingkungan", "1 Ya | 2 Tidak", "2"),
    ("21", "mitra_kdkmp", "21. mitra KDKMP", "1 Ya | 2 Tidak", "2"),
    ("22", "peran_mbg", "22. peran MBG", "1 SPPG | 2 supplier | 3 penerima manfaat | 4 lainnya | 5 Tidak", "5"),
    ("24.L", "tk_laki", "24a1. pekerja laki-laki", "angka", "0"),
    ("24.P", "tk_pr", "24b1. pekerja perempuan", "angka", "1"),
    ("24.Total", "cek_tk_gender", "pemeriksa", "= 24.L + 24.P (tidak dikirim)", "1"),
    ("24.Dibayar", "tk_dibayar", "24a2. dibayar", "angka", "0"),
    ("24.Tidak dibayar", "tk_tdk_dibayar", "24b2. tidak dibayar", "angka", "1"),
    ("24.Total", "cek_tk_bayar", "pemeriksa", "= dibayar + tidak dibayar (tidak dikirim)", "1"),
    ("25", "tahun_operasi", "25. tahun mulai beroperasi", "4 digit; tahun berjalan -> varian bulanan 30-33",
     "2019"),
    ("Rp26", "cek_26f", "26f. total pengeluaran", "pemeriksa = 26a..26e (tidak dikirim)", "Rp10.700.000"),
    ("26a", "gaji", "26a. upah/gaji", "rupiah setahun", "Rp0"),
    ("26b", "biaya_produksi", "26b. bahan baku/produksi", "rupiah setahun", "Rp0"),
    ("26c", "biaya_pembelian", "26c. barang dagangan", "rupiah setahun (hanya perdagangan)", "Rp10.500.000"),
    ("26d", "operasional", "26d. operasional lain", "rupiah setahun", "Rp200.000"),
    ("26e", "non_operasional", "26e. non operasional", "rupiah setahun", "Rp0"),
    ("27a", "nilai_pendapatan", "27a. pendapatan usaha", "rupiah setahun", "Rp11.800.000"),
    ("27b", "pendapatan_lain", "27b. pendapatan lain", "rupiah setahun", "Rp0"),
    ("27c", "cek_27c", "27c. total pendapatan", "pemeriksa = 27a + 27b (tidak dikirim)", "Rp11.800.000"),
    ("27d", "pendapatan_online", "27d. % penjualan online", "0-100 (dipakai kalau 16a = 1)", "0"),
    ("28a", "aset_usaha_thn", "28a. aset tanah & bangunan", "rupiah", "Rp0"),
    ("28b", "aset_lain_thn", "28b. aset selain tanah & bangunan", "rupiah", "Rp2.000.000"),
    ("28c", "cek_28c", "28c. total aset", "pemeriksa = 28a + 28b (tidak dikirim)", "Rp2.000.000"),
    ("28c1", "info_28c1", "info", "tidak dikirim", "-"),
    ("28d", "luas_tanah_thn", "28d. luas tanah (m2)", "angka", "0"),
    ("Latitude", "latitude", "geotag", "desimal, mis. -8,2004731; kosong -> dokumen DRAFT tanpa geotag",
     "-8,2004731"),
    ("Longitude", "longitude", "geotag", "desimal, mis. 114,7987732", "114,7987732"),
    ("Kode KBLI", "kbli", "13g. KBLI", "5 digit KBLI 2025 (sel Teks)", "47241"),
    ("Judul KBLI", "info_judul_kbli", "info", "judul KBLI (melengkapi 13a yang pendek)", "Perdagangan Eceran Beras"),
    (JUDUL_ID, "", "diisi skrip", "BIARKAN KOSONG — ID dokumen fasih-web ditulis otomatis sesudah dibuat", ""),
]

ARTI_TAMBAHAN = {
    "akun_ppl": "email PPL (hanya utk --per-baris)", "kodepos": "kodepos 5 digit (menang atas daftar config)",
    "nomor_domisili": "SE2026-P Blok/Nomor", "jenis_kawasan": "8d", "punya_nib": "10a NIB (1 Ya | 2 Tidak)",
    "nib_nomor": "10b nomor NIB", "tidak_nib": "10c alasan tanpa NIB", "badan_usaha": "11a badan usaha",
    "lap_keuangan": "11d catatan keuangan", "produk_sendiri": "13b1", "layanan_mamin": "13b2",
    "keg_penjualan": "13b3", "keg_jasa": "13b4", "lokasi_usaha": "13c tempat usaha", "input_produksi": "13d",
    "proses_produksi": "13e", "digital": "16c", "produksi_lingkungan": "17a", "produk_seni": "18",
    "halal": "19a", "sudah_halal": "19b", "belum_halal": "19c", "izin_edar": "20a", "sudah_bpom": "20b",
    "belum_bpom": "20c", "barang_non_pddk": "23a", "jasa_non_pddk": "23b", "beli_jasa_non_pddk": "23c",
    "pribadi": "29a % modal pribadi", "non_profit": "29b", "publik": "29c", "non_publik": "29d",
    "pemerintah": "29e", "asing": "29f",
}

BIRU = PatternFill("solid", fgColor="DDEBF7")
KUNING = PatternFill("solid", fgColor="FFF2CC")
ABU = PatternFill("solid", fgColor="EDEDED")
TEBAL = Font(bold=True)


def tulis(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = NAMA_TAB
    for i, (judul, key, *_ , contoh) in enumerate(KOLOM, start=1):
        h = ws.cell(row=1, column=i, value=judul)
        h.font, h.fill = TEBAL, (ABU if key.startswith(("info_", "cek_")) or not key else BIRU)
        c = ws.cell(row=2, column=i, value=contoh)
        c.fill = KUNING
        huruf = get_column_letter(i)
        ws.column_dimensions[huruf].width = max(10, min(32, len(str(contoh)) + 2, len(judul) + 12))
        for r in range(1, 1001):   # SEMUA sel Teks: NIK/idsubsls/KBLI tidak dipotong Excel
            ws.cell(row=r, column=i).number_format = "@"
    ws.freeze_panes = "A2"

    p = wb.create_sheet("petunjuk")
    baris = [
        ["TEMPLAT INPUT USAHA SE2026 (format tahap 2 = kuesioner kertas)"],
        ["1. Satu baris = satu usaha = satu dokumen fasih-web. Isi tab 'input_usaha' mulai baris 2."],
        [f"2. Baris 2 (kuning) = CONTOH fiktif. Timpa dgn data Anda; selama kolom 'Uraian:' diawali "
         f"{PENANDA_CONTOH}, baris itu SELALU ditolak skrip."],
        ["3. Semua sel berformat Teks — jangan diubah ke Angka (NIK, idsubsls & kode KBLI akan rusak)."],
        ["4. Kolom berkode (12b, 14a, 16a, 17b, 21, 22): isi ANGKA kode opsi, lihat kolom 'Cara mengisi'."],
        ["5. Kolom abu-abu = info/pemeriksa, tidak dikirim ke form. Kolom 'ID Dokumen FASIH' diisi skrip."],
        ["6. Periksa tanpa browser: python input_usaha/jalankan.py --sumber bahan/input_usaha.xlsx --cek"],
        [],
        ["Kolom", "Rincian kuesioner", "Cara mengisi", "Contoh", "Wajib ada?"],
    ]
    for judul, key, rincian, cara, contoh in KOLOM:
        wajib = "tidak" if (key in KOLOM_TAHAP2_OPSIONAL or key.startswith("info_") or not key) else "ya"
        if key == "info_nama_ppl":
            wajib = "ya"
        baris.append([judul, rincian, cara, contoh, wajib])
    baris += [[], ["KOLOM TAMBAHAN (opsional) — tambahkan di kanan kalau nilai bawaan tidak sesuai utk baris tertentu"],
              ["Judul kolom", "Rincian", "Nilai bawaan kalau kolomnya tidak ada / kosong"]]
    for key, judul in KOLOM_TAHAP2_TAMBAHAN.items():
        baris.append([judul, ARTI_TAMBAHAN.get(key, key), TAHAP2_DEFAULT.get(key, "(dari KBLI / tidak diisi)")])
    for r in baris:
        p.append(r)
    for sel in (p["A1"], *p[9], *p[len(baris) - len(KOLOM_TAHAP2_TAMBAHAN) - 1],
                *p[len(baris) - len(KOLOM_TAHAP2_TAMBAHAN)]):
        sel.font = TEBAL
    for huruf, lebar in zip("ABCDE", (18, 32, 70, 30, 10)):
        p.column_dimensions[huruf].width = lebar
    for r in p.iter_rows():
        for sel in r:
            sel.alignment = Alignment(vertical="top", wrap_text=sel.column == 3)
    wb.save(path)


def verifikasi(path: Path) -> list[str]:
    """Baris contoh: dibaca loader, ditolak HANYA krn penanda contoh, SIAP tanpa penandanya."""
    masalah = []
    rows = load_tahap2(path)
    if len(rows) != 1:
        return [f"loader membaca {len(rows)} baris, harusnya 1"]
    h = periksa_semua_tahap2(rows, mode_satu_subsls=True)[rows[0].baris]
    if [k for k, _ in h.masalah] != ["BARIS_CONTOH"]:
        masalah.append(f"baris contoh: masalah {h.masalah} (harusnya hanya BARIS_CONTOH)")
    with tempfile.TemporaryDirectory() as d:
        salinan = Path(d) / "uji.xlsx"
        wb = load_workbook(path)
        ws = wb[NAMA_TAB]
        ws.cell(row=2, column=[k for _, k, *_ in KOLOM].index("info_uraian") + 1, value="Responden 1")
        wb.save(salinan)
        r = load_tahap2(salinan)
        h2 = periksa_semua_tahap2(r, mode_satu_subsls=True)[r[0].baris]
        if h2.status != "SIAP":
            masalah.append(f"tanpa penanda contoh: {h2.status} {h2.pesan}")
        else:
            print(f"  baris contoh tanpa penanda: SIAP — nama dokumen '{r[0].nama_dokumen}'")
    return masalah


def main() -> int:
    tulis(KELUARAN)
    masalah = verifikasi(KELUARAN)
    for m in masalah:
        print("❌", m)
    if masalah:
        return 1
    print(f"✅ {KELUARAN} ({len(KOLOM)} kolom, 1 baris contoh) — terverifikasi loader.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
