# -*- coding: utf-8 -*-
"""Uji koordinat pengganti tahap 2 — offline, poligon & titik FIKTIF.
Jalankan: python tests/test_koordinat_pengganti.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import os  # noqa: E402
os.environ.setdefault("FASIH_ABAIKAN_CONFIG_LOKAL", "1")  # hasil uji tidak bergantung inti/config_lokal.py

from koordinat.peta import peta_dari_fitur  # noqa: E402
from koordinat.koordinat_pengganti import (  # noqa: E402
    Jalan, TitikListing, baca_koordinat, dms_tanpa_simbol, kunci_kelompok, meter, nilai_wilayah, pilih_subsls,
    rencana, tingkat_beda, tulis_tempel,
)

ok_all = True


def check(label, got, want):
    global ok_all
    ok = got == want
    ok_all &= ok
    print(f"{'PASS' if ok else 'FAIL'} | {label}: got={got!r} want={want!r}")


# --- membaca teks koordinat (pola nyata sheet 2026-09-24) ---
check("desimal koma", baca_koordinat("-8,14435", "lat"), -8.14435)
check("desimal titik", baca_koordinat("115.0794770", "lon"), 115.079477)
check("akhiran S", baca_koordinat("-8.12736124S", "lat"), -8.12736124)
check("akhiran E", baca_koordinat("115.070484E", "lon"), 115.070484)
check("derajat-menit-detik", round(baca_koordinat("8°7'35,424\"S", "lat"), 6), -8.126507)
check("derajat-menit-detik bujur", round(baca_koordinat("115°4'48,129\"E", "lon"), 6), 115.080036)
check("° sbg titik desimal", baca_koordinat("8°12887404S", "lat"), -8.12887404)
check("titik desimal hilang (lintang)", baca_koordinat("-812399973S", "lat"), -8.12399973)
check("titik desimal hilang (bujur)", baca_koordinat("115093422", "lon"), 115.093422)
check("koma dobel", baca_koordinat("-8,10356,1", "lat"), -8.103561)
check("spasi", baca_koordinat("114, 96437916", "lon"), 114.96437916)
check("lintang tanpa minus", baca_koordinat("8.130898", "lat"), -8.130898)
check("kosong", baca_koordinat("", "lat"), None)
check("angka dari sel numerik", baca_koordinat(-8.2, "lat"), -8.2)

check("derajat tanpa simbol lintang", round(dms_tanpa_simbol("-8,747", "lat"), 6), -8.129722)
check("derajat tanpa simbol bujur (155 -> 115)", round(dms_tanpa_simbol("155,2058", "lon"), 6), 115.349444)
check("menit >= 60 ditolak", dms_tanpa_simbol("115,7512", "lon"), None)

# --- kelompok pemilik + alamat ---
check("pemilik & alamat sama walau beda tulisan", kunci_kelompok(2, "I Made  Contoh.", "Jl. Contoh, No 1"),
      kunci_kelompok(3, "I MADE CONTOH", "JL CONTOH NO 1"))
check("alamat kosong -> baris sendiri", kunci_kelompok(5, "I MADE CONTOH", "-"), "baris:5")
anggota = [{"baris": 1, "idsubsls": "A"}, {"baris": 2, "idsubsls": "B"}, {"baris": 3, "idsubsls": "B"}]
check("subsls baris terbanyak", pilih_subsls(anggota, {}), "B")
seri = [{"baris": 1, "idsubsls": "B"}, {"baris": 2, "idsubsls": "A"}]
check("seri -> subsls tempat titik jatuh", pilih_subsls(seri, {1: "", 2: "A"}), "A")
check("seri tanpa titik -> kode terkecil", pilih_subsls(seri, {}), "A")

# --- rencana dgn peta fiktif: dua subsls persegi ±1,1 km di Buleleng ---
def persegi(ids, lat0, lon0, d=0.01):
    ring = [[lon0, lat0], [lon0 + d, lat0], [lon0 + d, lat0 + d], [lon0, lat0 + d], [lon0, lat0]]
    return {"properties": {"idsubsls": ids}, "geometry": {"type": "Polygon", "coordinates": [ring]}}


A, B = "5108990001000101", "5108990001000102"
peta = peta_dari_fitur([persegi(A, -8.20, 115.00), persegi(B, -8.20, 115.02)])
listing = TitikListing([(-8.195, 115.005), (-8.193, 115.007), (-8.195, 115.025)])
baris = [
    # 2: di dalam A -> tetap
    {"baris": 2, "idsubsls": A, "pemilik": "I MADE CONTOH", "alamat": "BANJAR SATU", "lat_mentah": "-8,1950",
     "lon_mentah": "115,0050"},
    # 3: pemilik & alamat sama, titik 3 km -> ikut koordinat baris 2
    {"baris": 3, "idsubsls": A, "pemilik": "I Made Contoh", "alamat": "Banjar Satu", "lat_mentah": "-8,1950",
     "lon_mentah": "115,0350"},
    # 4: sendirian, bujur 145 (salah ketik 115) -> diperbaiki
    {"baris": 4, "idsubsls": A, "pemilik": "NI LUH CONTOH", "alamat": "BANJAR DUA", "lat_mentah": "-8,1960",
     "lon_mentah": "145,0060"},
    # 5: sendirian, 2 km di luar B (di kabupaten) -> acak listing di B
    {"baris": 5, "idsubsls": B, "pemilik": "I KETUT CONTOH", "alamat": "BANJAR TIGA", "lat_mentah": "-8,1950",
     "lon_mentah": "115,0500"},
    # 6: kosong -> acak
    {"baris": 6, "idsubsls": B, "pemilik": "I WAYAN CONTOH", "alamat": "BANJAR EMPAT", "lat_mentah": "",
     "lon_mentah": ""},
    # 7-9: satu kelompok, 2 baris di B & 1 di A -> subsls B
    {"baris": 7, "idsubsls": A, "pemilik": "KADEK CONTOH", "alamat": "PASAR", "lat_mentah": "-8,1950",
     "lon_mentah": "115,0050"},
    {"baris": 8, "idsubsls": B, "pemilik": "KADEK CONTOH", "alamat": "PASAR", "lat_mentah": "", "lon_mentah": ""},
    {"baris": 9, "idsubsls": B, "pemilik": "KADEK CONTOH", "alamat": "PASAR", "lat_mentah": "", "lon_mentah": ""},
]
h = {x["baris"]: x for x in rencana(baris, peta, listing)}
check("2: titik benar tetap", (h[2]["sumber"], h[2]["lat"], h[2]["lon"]), ("ASLI", -8.195, 115.005))
check("3: ikut koordinat kelompok", (h[3]["sumber"], h[3]["lat"], h[3]["lon"]),
      ("KELOMPOK (dari baris 2)", -8.195, 115.005))
check("4: salah ketik 145 -> 115", (h[4]["sumber"], h[4]["lon"]), ("SALAH_KETIK_DIPERBAIKI", 115.006))
check("5: diganti titik listing di B (digeser 5–15 m)",
      (h[5]["sumber"], 4.9 <= meter(h[5]["lat"], h[5]["lon"], -8.195, 115.025) <= 15.1), ("ACAK_LISTING", True))
check("6: kosong -> titik listing di B", h[6]["sumber"], "ACAK_LISTING")
check("7-9: satu koordinat di subsls mayoritas B",
      ({(h[i]["lat"], h[i]["lon"]) for i in (7, 8, 9)}.__len__(), h[7]["subsls_koordinat"],
       peta.jarak_m(B, h[7]["lon"], h[7]["lat"])), (1, B, 0.0))
h2 = {x["baris"]: x for x in rencana(baris, peta, listing)}
check("acak tapi tetap (dijalankan ulang = sama)", [(h2[i]["lat"], h2[i]["lon"]) for i in (5, 6, 8)],
      [(h[i]["lat"], h[i]["lon"]) for i in (5, 6, 8)])
check("titik jauh tidak 'diperbaiki' dgn tebakan satu digit (di kabupaten)", h[5]["sumber"] != "SALAH_KETIK_DIPERBAIKI",
      True)

# --- dekat jalan: titik listing dekat ruas didahulukan; tanpa listing -> tepi ruas ---
jalan = Jalan([(-8.1931, 115.0010, -8.1931, 115.0090)])   # ruas di lintang -8.1931 dalam A
baris_j = [{"baris": 20, "idsubsls": A, "pemilik": "X CONTOH", "alamat": "Y", "lat_mentah": "", "lon_mentah": ""}]
hj = rencana(baris_j, peta, listing, jalan, jarak_jalan_m=50)[0]
check("listing dekat jalan dipilih (-8.193 bukan -8.195)",
      (hj["sumber"], jalan.jarak_m(hj["lat"], hj["lon"]) <= 50), ("ACAK_LISTING_DEKAT_JALAN", True))
hk = rencana(baris_j, peta, None, jalan, jarak_jalan_m=50)[0]
check("tanpa listing -> dekat ruas jalan di poligon",
      (hk["sumber"], jalan.jarak_m(hk["lat"], hk["lon"]) <= 20, peta.jarak_m(A, hk["lon"], hk["lat"])),
      ("ACAK_POLIGON_DEKAT_JALAN", True, 0.0))

# --- --subsls-dari-koordinat (2026-09-26): subsls = poligon tempat titik jatuh ---
check("tingkat beda", [tingkat_beda(A, A), tingkat_beda(A, B), tingkat_beda(A, "5108990001000201"),
                       tingkat_beda(A, "5108990002000101"), tingkat_beda(A, "5108980001000101")],
      ["", "SUB_SLS", "SLS", "DESA", "KECAMATAN"])
C = "5108990002000101"                                          # desa LAIN, di sebelah timur B
D = "5108990001000103"                                          # menempel di timur A (celah A-B)
fitur = [persegi(A, -8.20, 115.00), persegi(B, -8.20, 115.02), persegi(C, -8.20, 115.04), persegi(D, -8.20, 115.01)]
fitur[2]["properties"].update(nmkec="KECAMATAN CONTOH", nmdesa="DESA BARU CONTOH")
peta3 = peta_dari_fitur(fitur)


def b_(n, ids, lat, lon, pemilik=None, alamat="", **x):
    return {"baris": n, "idsubsls": ids, "pemilik": pemilik or f"PEMILIK {n}", "alamat": alamat,
            "lat_mentah": lat, "lon_mentah": lon, "kec_mentah": "Kec Asli", "desa_mentah": "Desa Asli",
            "subsls_mentah": int(ids) if ids else None, **x}


baris_k = [
    b_(30, A, "-8,1950", "115,0250"),                           # titik di B -> subsls B, koordinat tetap
    b_(31, A, "-8,1950", "115,0250", berdokumen=True),          # sama, tapi SUDAH berdokumen -> aturan lama
    b_(32, A, "", ""),                                          # tanpa titik -> kolom 5, titik acak di A
    b_(33, B, "-8,1950", "115,0320"),                           # di luar semua poligon, 220 m dari B -> tetap B
    b_(34, A, "-8,1950", "115,0450"),                           # titik di C (desa lain) -> C
    # 35-37: satu kelompok, kolom 5 semuanya A; dua titik di B, satu di A -> B, titik dari dalam B
    b_(35, A, "-8,1950", "115,0050", "KOMANG CONTOH", "PASAR"),
    b_(36, A, "-8,1940", "115,0260", "KOMANG CONTOH", "PASAR"),
    b_(37, A, "-8,1940", "115,0260", "KOMANG CONTOH", "PASAR"),
    b_(38, A, "-8,1950", "115,0101"),                           # ±11 m di seberang batas A, di D -> tetap A
]
hk = {x["baris"]: x for x in rencana(baris_k, peta3, listing, subsls_dari_koordinat=True)}
check("30: titik di B -> subsls B, koordinat asli",
      (hk[30]["idsubsls_baru"], hk[30]["asal_subsls"], hk[30]["sumber"], hk[30]["lat"], hk[30]["lon"]),
      (B, "KOORDINAT", "ASLI", -8.195, 115.025))
check("31: sudah berdokumen -> kolom 5 tetap A, titik diganti spt aturan lama",
      (hk[31]["idsubsls_baru"], hk[31]["asal_subsls"], hk[31]["subsls_titik"], hk[31]["sumber"],
       peta3.jarak_m(A, hk[31]["lon"], hk[31]["lat"])),
      (A, "TETAP_SUDAH_ADA_DOKUMEN", B, "ACAK_LISTING", 0.0))
check("32: tanpa titik -> kolom 5 & titik acak di dalamnya",
      (hk[32]["idsubsls_baru"], hk[32]["asal_subsls"], peta3.jarak_m(A, hk[32]["lon"], hk[32]["lat"])), (A, "", 0.0))
check("33: titik di luar semua poligon -> kolom 5 (aturan 500 m)",
      (hk[33]["idsubsls_baru"], hk[33]["asal_subsls"], hk[33]["sumber"]), (B, "", "ASLI"))
check("34: titik di desa lain -> subsls desa itu", (hk[34]["idsubsls_baru"], hk[34]["asal_subsls"]), (C, "KOORDINAT"))
check("35-37: kelompok ikut subsls titik terbanyak (B), satu koordinat dari DALAM B",
      ({hk[i]["idsubsls_baru"] for i in (35, 36, 37)}, hk[35]["asal_subsls"], hk[36]["asal_subsls"],
       {(hk[i]["lat"], hk[i]["lon"]) for i in (35, 36, 37)}),
      ({B}, "KOORDINAT_KELOMPOK", "KOORDINAT", {(-8.194, 115.026)}))
check("38: titik <= toleransi (20 m) di luar subsls kolom 5 -> tidak pindah, koordinat asli",
      (hk[38]["idsubsls_baru"], hk[38]["asal_subsls"], hk[38]["subsls_titik"], hk[38]["sumber"]), (A, "", D, "ASLI"))
h0 = {x["baris"]: x for x in rencana(baris_k, peta3, listing, subsls_dari_koordinat=True, toleransi_m=0)}
check("38: toleransi 0 (murni koordinat) -> pindah ke D", (h0[38]["idsubsls_baru"], h0[38]["asal_subsls"]),
      (D, "KOORDINAT"))
hl = {x["baris"]: x for x in rencana(baris_k, peta3, listing)}
check("tanpa opsi: kolom 5 tidak pernah berubah", [i for i in hl if hl[i]["idsubsls_baru"] != hl[i]["idsubsls"]], [])
check("tanpa opsi: 30 tetap diganti titik di A (perilaku lama)", peta3.jarak_m(A, hl[30]["lon"], hl[30]["lat"]), 0.0)

check("kolom 3/4/5: pindah sub-SLS -> nama asli, kode baru", nilai_wilayah({**hk[30], "dinilai": True}, peta3),
      ("Kec Asli", "Desa Asli", B, "SUB_SLS"))
check("kolom 3/4/5: pindah desa -> nama desa dari peta", nilai_wilayah({**hk[34], "dinilai": True}, peta3),
      ("Kec Asli", "Desa Baru Contoh", C, "DESA"))
check("kolom 3/4/5: tidak pindah -> teks asli persis (angka jadi teks)", nilai_wilayah({**hk[32], "dinilai": True}, peta3),
      ("Kec Asli", "Desa Asli", A, ""))

import tempfile  # noqa: E402
import openpyxl  # noqa: E402
with tempfile.TemporaryDirectory() as tmp:
    f = Path(tmp) / "t.xlsx"
    n = tulis_tempel(f, list(hk.values()), 37, peta3, subsls=True)
    ws = openpyxl.load_workbook(f).active
    check("xlsx: C:E = 3/4/5, idsubsls teks '@' (tidak dipotong Excel)",
          (ws["C1"].value, ws["E1"].value, ws["E30"].value, ws["E30"].number_format, ws["D34"].value),
          ("3", "5", B, "@", "Desa Baru Contoh"))
    check("xlsx: baris berubah dihitung (30, 31, 32, 34, 35-37 + 33 tidak)", n, 7)

print("\nSEMUA PASS" if ok_all else "\nADA YANG FAIL")
sys.exit(0 if ok_all else 1)
