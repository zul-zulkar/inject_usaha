# -*- coding: utf-8 -*-
"""Uji antar_pc/gabung_id_sumber.py (satukan kolom ID dokumen sheet sumber
dari beberapa PC) — offline, data fiktif.
Jalankan: python tests/test_gabung_id_sumber.py
"""
import io
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import os  # noqa: E402
os.environ.setdefault("FASIH_ABAIKAN_CONFIG_LOKAL", "1")  # hasil uji tidak bergantung inti/config_lokal.py

import openpyxl  # noqa: E402

from antar_pc import gabung_id_sumber as gis  # noqa: E402
from inti import id_dokumen as idd  # noqa: E402

ok_all = True


def check(label, got, want):
    global ok_all
    ok = got == want
    ok_all &= ok
    print(f"{'PASS' if ok else 'FAIL'} | {label}: got={got!r} want={want!r}")


I1 = "11111111-1111-4111-8111-111111111111"
I2 = "22222222-2222-4222-8222-222222222222"
I3 = "33333333-3333-4333-8333-333333333333"
I4 = "44444444-4444-4444-8444-444444444444"


def salinan(nama, baris: dict, kunci=None):
    """baris = {nomor: (sidik, teks ID)}."""
    return gis.Salinan(nama=nama, mentah={b: v[1] for b, v in baris.items() if v[1]},
                       sidik={b: v[0] for b, v in baris.items()}, kunci=kunci or {})


def kat(keputusan):
    return {k.baris: k.kategori for k in keputusan}


# --- pencocokan & kategori ---
utama = salinan("UTAMA", {2: ("a", ""), 3: ("b", I2), 4: ("c", ""), 5: ("d", ""), 6: ("e", "catatan"),
                          7: ("f", ""), 8: ("g", "")},
                kunci={2: "ka", 3: "kb", 4: "kc", 5: "kd", 6: "ke", 7: "kf", 8: "kg"})
pc1 = salinan("pc1", {2: ("a", I1), 3: ("b", I2), 4: ("c", I3), 6: ("e", I4)})
pc2 = salinan("pc2", {4: ("c", I4), 5: ("d-diedit", f"https://x/survey/s/p/{I1.upper()}/entry")},
              kunci={5: "kd"})
kep, buang = gis.rencana_gabung(utama, [pc1, pc2])
check("kategori per baris utama", kat(kep),
      {2: "ID_GANDA", 3: "SUDAH", 4: "KONFLIK", 5: "ID_GANDA", 6: "SEL_TIDAK_VALID"})
check("ID_GANDA: tidak ada yang ditulis", [k.id_tulis for k in kep if k.kategori == "ID_GANDA"], ["", ""])
check("baris diedit dicocokkan lewat KUNCI", [c for k in kep if k.baris == 5 for v in k.usulan.values()
                                                for _, _, c in v], ["KUNCI"])

kep, _ = gis.rencana_gabung(utama, [salinan("pc1", {2: ("a", I1), 7: ("f", I3)})])
check("sepakat & sel kosong -> TULIS", [(k.baris, k.kategori, k.id_tulis) for k in kep],
      [(2, "TULIS", I1), (7, "TULIS", I3)])

kep, _ = gis.rencana_gabung(utama, [salinan("pc1", {2: ("a", I1)}), salinan("pc2", {2: ("a", I1)})])
check("dua PC membawa ID sama -> tetap TULIS", kat(kep), {2: "TULIS"})

kep, _ = gis.rencana_gabung(utama, [salinan("pc1", {3: ("b", I3)})])
check("utama berisi ID lain -> BEDA, tidak ditimpa", (kat(kep), kep[0].id_tulis), ({3: "BEDA"}, ""))

check("ID yang sudah ada di baris utama lain -> ID_GANDA",
      gis.rencana_gabung(utama, [salinan("pc1", {2: ("a", I2)})])[0][0].kategori, "ID_GANDA")

_, buang = gis.rencana_gabung(utama, [salinan("pc1", {9: ("x", I3), 2: ("a", "salah ketik")})])
check("baris tak dikenal & sel bukan ID -> terbuang",
      sorted(a.split(":")[0] for *_, a in buang), ["SUMBER_TIDAK_VALID", "TIDAK_KETEMU"])

# Baris bergeser (sheet utama diurutkan): cocok lewat isi, bukan nomor baris.
geser = salinan("UTAMA", {2: ("b", ""), 3: ("a", "")})
kep, _ = gis.rencana_gabung(geser, [salinan("pc1", {2: ("a", I1)})])
check("baris bergeser -> ikut isi", [(k.baris, k.id_tulis) for k in kep], [(3, I1)])

# Dua baris identik di utama: nomor baris sama didahulukan, lainnya tidak ditebak.
kembar = salinan("UTAMA", {2: ("z", ""), 3: ("z", "")})
kep, buang = gis.rencana_gabung(kembar, [salinan("pc1", {3: ("z", I1)}), salinan("pc2", {9: ("z", I2)})])
check("baris identik: nomor sama didahulukan", [(k.baris, k.id_tulis) for k in kep], [(3, I1)])
check("baris identik di nomor lain -> tidak ditebak", len(buang), 1)


# --- ujung ke ujung: xlsx utama + folder salinan ---
def buat_xlsx(path, isi, judul=("nama", "pemilik", "kbli", idd.JUDUL_ID)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "input_usaha"
    ws.append(list(judul))
    for b in isi:
        ws.append(list(b))
    wb.save(path)


def kolom_id(path):
    ws = openpyxl.load_workbook(path)["input_usaha"]
    return [r[3].value if len(r) > 3 else None for r in ws.iter_rows(min_row=2)]


with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    data = [("APOTEK A", "I MADE CONTOH", "47721"), ("APOTEK B", "I KETUT CONTOH", "47721"),
            ("APOTEK C", "NI LUH CONTOH", "47721")]
    utama_p = d / "input_usaha.xlsx"
    buat_xlsx(utama_p, [(*data[0], I1), (*data[1], None), (*data[2], None)])
    folder = d / "bahan_pc"
    folder.mkdir()
    buat_xlsx(folder / "pc1.xlsx", [(*data[0], I1), (*data[1], I2), (*data[2], None)])
    buat_xlsx(folder / "pc2.xlsx", [(*data[0], None), (*data[1], None), (*data[2], I3)])
    (folder / "~$pc2.xlsx").write_text("kunci excel")
    lap = d / "laporan.csv"
    arg = ["--utama", str(utama_p), "--sumber", str(folder)]
    hasil = folder / "hasil"

    with redirect_stdout(io.StringIO()):
        kode = gis.main(arg)
    check("laporan saja: sheet utama tidak berubah", (kode, kolom_id(utama_p)), (0, [I1, None, None]))
    check("laporan CSV di <sumber>/hasil", (hasil / "laporan_gabung_id.csv").exists(), True)
    check("laporan saja: sheet hasil belum dibuat", (hasil / "input_usaha.xlsx").exists(), False)

    with redirect_stdout(io.StringIO()):
        kode = gis.main(arg + ["--tulis"])
    check("--tulis: sheet hasil di <sumber>/hasil berisi ID kedua PC",
          (kode, kolom_id(hasil / "input_usaha.xlsx")), (0, [I1, I2, I3]))
    check("--tulis: sheet utama TIDAK diubah", kolom_id(utama_p), [I1, None, None])

    buf = io.StringIO()
    with redirect_stdout(buf):
        gis.main(arg + ["--tulis"])
    check("dijalankan ulang: sheet hasil lama tidak dibaca sbg salinan & dicadangkan",
          (kolom_id(hasil / "input_usaha.xlsx"), len(list(hasil.glob("input_usaha.xlsx.bak-*")))),
          ([I1, I2, I3], 1))
    check("dijalankan ulang: tidak ada salinan dari folder hasil", "hasil/" in buf.getvalue(), False)

    with redirect_stdout(io.StringIO()):
        kode = gis.main(arg + ["--tulis", "--keluaran", str(utama_p), "--laporan", str(lap)])
    check("--keluaran = --utama: sheet utama diisi langsung", (kode, kolom_id(utama_p)), (0, [I1, I2, I3]))
    check("--keluaran = --utama: cadangan utama dibuat", len(list(d.glob("input_usaha.xlsx.bak-*"))), 1)
    buf = io.StringIO()
    with redirect_stdout(buf):
        gis.main(arg + ["--keluaran", str(utama_p), "--laporan", str(lap)])
    check("dijalankan ulang: tidak ada yang ditulis", "Tidak ada ID baru" in buf.getvalue(), True)
    check("berkas ~$ & utama sendiri tidak dibaca sbg salinan",
          [p.name for p in gis.daftar_berkas([str(folder), str(utama_p)], utama_p)], ["pc1.xlsx", "pc2.xlsx"])

# --- SATU folder utk audit & sheet bahan, satu subfolder per PC (berkas disalin apa adanya) ---
import csv  # noqa: E402

import input_usaha.mesin as mg  # noqa: E402
from antar_pc import gabung_audit as ga  # noqa: E402

with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    data = [("APOTEK A", "I MADE CONTOH", "47721"), ("APOTEK B", "I KETUT CONTOH", "47721")]
    utama_p = d / "input_usaha.xlsx"
    buat_xlsx(utama_p, [(*data[0], None), (*data[1], None)])
    folder = d / "audit_pc"
    for pc, isi in (("pc2", [(*data[0], I1), (*data[1], None)]), ("pc3", [(*data[0], None), (*data[1], I2)])):
        (folder / pc).mkdir(parents=True)
        buat_xlsx(folder / pc / "input_usaha.xlsx", isi)
        with (folder / pc / "audit_log_gabungan.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=mg.AUDIT_FIELDS)
            w.writeheader()
            w.writerow({"timestamp": "2026-09-25T10:00:00", "kunci": "0123456789", "status": "DOKUMEN_DIBUAT",
                        "akun_login": f"ppl.{pc}@mail.com", "idsubsls": "5108010002000501"})
    (folder / "laporan_lain.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    buat_xlsx(folder / "sheet_lama.xlsx", [data[0]], judul=("nama", "pemilik", "kbli"))

    with redirect_stdout(io.StringIO()):
        salin = gis.daftar_berkas([str(folder)], utama_p)
        audit = ga.kumpulkan_sumber([str(folder)])
    check("folder bersama: gabung_id_sumber hanya mengambil sheet ber-kolom ID",
          [gis.label_berkas(p, [str(folder)]) for p in salin], ["pc2/input_usaha.xlsx", "pc3/input_usaha.xlsx"])
    check("folder bersama: gabung_audit hanya mengambil audit",
          [ga.label_berkas(p, [str(folder)]) for p in audit],
          ["pc2/audit_log_gabungan.csv", "pc3/audit_log_gabungan.csv"])
    with redirect_stdout(io.StringIO()):
        kode = gis.main(["--utama", str(utama_p), "--sumber", str(folder), "--tulis"])
        kode_a = ga.main(["--sumber", str(folder), "--tulis", "--list-json", str(d / "tidak_ada.json")])
    check("folder bersama: ID kedua PC masuk ke audit_pc/hasil",
          (kode, kolom_id(folder / "hasil" / "input_usaha.xlsx")), (0, [I1, I2]))
    check("folder bersama: gabung_audit menulis semua keluaran ke audit_pc/hasil",
          (kode_a, sorted(p.name for p in (folder / "hasil").iterdir())),
          (0, ["audit_log_gabungan.csv", "daftar_ganda.csv", "input_usaha.xlsx", "laporan_gabung.csv",
               "laporan_gabung_agregat.csv", "laporan_gabung_id.csv"]))
    with redirect_stdout(io.StringIO()):
        audit2 = ga.kumpulkan_sumber([str(folder)])
    check("folder hasil tidak dibaca sbg audit sumber", len(audit2), 2)

print("\nSEMUA LULUS" if ok_all else "\nADA YANG GAGAL")
sys.exit(0 if ok_all else 1)
