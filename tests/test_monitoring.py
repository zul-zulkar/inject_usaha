# -*- coding: utf-8 -*-
"""Uji monitoring/monitoring.py — offline, data FIKTIF.
Jalankan: python tests/test_monitoring.py
"""
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import os  # noqa: E402
os.environ.setdefault("FASIH_ABAIKAN_CONFIG_LOKAL", "1")  # hasil uji tidak bergantung inti/config_lokal.py

import input_usaha.mesin as mg  # noqa: E402
from inti.gabungan_loader import Pemeriksaan  # noqa: E402
from monitoring import monitoring as mon  # noqa: E402

ok_all = True


def check(label, got, want):
    global ok_all
    ok = got == want
    ok_all &= ok
    print(f"{'PASS' if ok else 'FAIL'} | {label}: got={got!r} want={want!r}")


SURVEI = "https://fasih-web.bps.go.id/survey/s/p"


def url(did):
    return f"{SURVEI}/{did}/entry"


def ab(kunci, status, waktu, did="", subsls="5108060006000224", akun="ppl.contoh@mail.com"):
    b = {k: "" for k in mg.AUDIT_FIELDS}
    b.update({"kunci": kunci, "status": status, "timestamp": waktu, "dokumen_url": url(did) if did else "",
              "idsubsls_input": subsls, "akun_login": akun})
    return b


@dataclass
class Baris:
    baris: int
    kunci: str
    nama_dokumen: str
    idsubsls: str
    id_dokumen: str = ""
    v: dict = field(default_factory=dict)
    info: dict = field(default_factory=dict)


def periksa(ditolak=False, tanpa_koordinat=False):
    p = Pemeriksaan(tanpa_koordinat=tanpa_koordinat)
    if ditolak:
        p.masalah.append(("UMUR_DI_LUAR_10_99", "umur 5"))
    return p


# --- fungsi kecil ---
check("kategori server", [mon.kategori_server(a) for a in
                          ("APPROVED BY Pengawas", "SUBMITTED BY Pencacah", "REJECTED BY Pengawas", "DRAFT", "", "COMPLETED")],
      ["APPROVED", "SUBMITTED", "REJECTED", "DRAFT", "DRAFT", "LAIN"])
check("kategori audit", [mon.kategori_audit(s) for s in ("TERKIRIM_TERVERIFIKASI", mg.STATUS_DIPINDAH,
                                                        "DOKUMEN_DIBUAT", "DRAFT_TANPA_KOORDINAT")],
      ["SUBMITTED", "SUBMITTED", "DRAFT", "DRAFT"])
check("id dari url", mon.id_dari_url(url("abc")), "abc")
check("nama dari kode identitas", mon.nama_dari_kode("5108060006000224 - Warung  Sembako (Contoh)"), "WARUNG SEMBAKO (CONTOH)")
check("prelist dikenali", [bool(mon.POLA_PRELIST.match(k)) for k in
                           ("5108060006000224 - UMK - 19", "5108060006000224 - WARUNG (A)")], [True, False])

# --- ringkas_audit: DOKUMEN_DIHAPUS menggugurkan, DIBUAT sesudah terkirim tidak menurunkan status ---
info = mon.ringkas_audit(Path("a.csv"), [
    ab("K1", "DOKUMEN_DIBUAT", "2026-09-20 10:00:00", "d1"),
    ab("K1", "TERKIRIM_TERVERIFIKASI", "2026-09-20 10:01:00", "d1"),
    ab("K1", "DOKUMEN_DIBUAT", "2026-09-21 10:00:00", "d1"),            # catatan sinkron
    ab("K2", "DOKUMEN_DIBUAT", "2026-09-20 11:00:00", "d2"),
    ab("K2", mg.STATUS_DIHAPUS, "2026-09-20 12:00:00"),
    ab("K3", mg.STATUS_TANPA_URL, "2026-09-20 13:00:00"),
])
check("status dokumen tetap terkirim", info.dokumen["d1"]["status"], "TERKIRIM_TERVERIFIKASI")
check("dokumen dihapus", (info.dihapus, "K2" in info.ids_kunci), ({"d2"}, False))
check("tanpa url", info.tanpa_url, {"K3"})
check("dibuat = timestamp pertama", info.dokumen["d1"]["dibuat"], "2026-09-20 10:00:00")

# --- audit_milik: audit batch lain yang kebetulan kenal 1 kunci tidak ikut ---
A = mon.InfoAudit(Path("A"), ids_kunci={f"K{i}": {f"a{i}"} for i in range(100)})
B = mon.InfoAudit(Path("B"), ids_kunci={"K1": {"b1"}})
C = mon.InfoAudit(Path("C"), ids_kunci={f"K{i}": {f"c{i}"} for i in range(40)})
check("audit milik sheet", [a.path.name for a in mon.audit_milik({f"K{i}" for i in range(100)}, [A, B, C])], ["A", "C"])
check("tidak ada audit", mon.audit_milik({"X"}, [A]), [])

# --- snapshot: baris CSV console & item datatable ---
item = {"id": "s1", "codeIdentity": "5108060006000224 - WARUNG (A)", "data1": "warung (a)",
        "assignmentStatusAlias": "APPROVED BY Pengawas", "mode": ["PAPI"], "currentUserUsername": "PML@X",
        "sumError": 0, "dateCreated": "2026-09-20T02:00:00.000+00:00",
        "region": {"level1": {"fullCode": "51", "level2": {"fullCode": "5108", "level3": {"fullCode": "5108060",
                   "level4": {"fullCode": "5108060002", "level5": {"fullCode": "51080600020002",
                   "level6": {"fullCode": "5108060002000203"}}}}}}}}
r = mon.ringkas_snapshot(item)
check("snapshot dari item datatable", (r["nama"], r["mode"], r["subsls"], r["petugas"], r["galat"]),
      ("WARUNG (A)", "PAPI", "5108060002000203", "pml@x", "0"))
r2 = mon.ringkas_snapshot({"id": "s2", "kode_identitas": "x", "nama": "B", "status": "DRAFT", "subsls": "5108",
                           "galat": "2", "diambil": "2026-09-29T02:00:00.000Z", "total_server": "10"})
check("snapshot dari CSV console", (r2["nama"], r2["galat"], r2["total_server"]), ("B", "2", "10"))
check("waktu lokal", len(mon.waktu_lokal("2026-09-29T02:00:00.000Z")), 19)
check("waktu rusak", mon.waktu_lokal("bukan waktu"), "")

# --- bangun: semua kategori + ganda ---
W = "5108060006000224"          # subsls wadah
T1, T2 = "5108060002000203", "5108060002000204"
rows21 = [
    Baris(2, "K1", "WARUNG A (SATU)", T1, id_dokumen="d1"),     # approved (audit approve), sudah dipindah
    Baris(3, "K2", "WARUNG B (DUA)", T1),                        # terkirim, GANDA (d2 + d2b)
    Baris(4, "K3", "WARUNG C (TIGA)", T2),                       # draft tanpa koordinat
    Baris(5, "K4", "WARUNG D (EMPAT)", T2),                      # belum input
    Baris(6, "K5", "WARUNG E (LIMA)", T2),                       # ditolak data
    Baris(7, "K6", "WARUNG F (ENAM)", T2),                       # tanpa URL -> perlu cek
]
periksa21 = {2: periksa(), 3: periksa(), 4: periksa(tanpa_koordinat=True), 5: periksa(), 6: periksa(ditolak=True),
             7: periksa()}
aud = mon.ringkas_audit(Path("audit/batch21/audit_log_gabungan.csv"), [
    ab("K1", "TERKIRIM_TERVERIFIKASI", "2026-09-20 10:00:00", "d1"),
    ab("K1", mg.STATUS_DIPINDAH, "2026-09-27 10:00:00", "d1", subsls=T1),
    ab("K2", "TERKIRIM_TERVERIFIKASI", "2026-09-20 11:00:00", "d2"),
    ab("K2", "TERKIRIM_BELUM_TERVERIFIKASI", "2026-09-20 12:00:00", "d2b"),
    ab("K3", mg.STATUS_DRAFT_TANPA_KOORDINAT, "2026-09-20 13:00:00", "d3"),
    ab("K6", mg.STATUS_TANPA_URL, "2026-09-20 14:00:00"),
])
rows22 = [Baris(2, "L1", "WARUNG B (DUA)", T1, id_dokumen="e1")]   # nama sama, batch lain
aud22 = mon.ringkas_audit(Path("audit/audit_log_gabungan.csv"), [ab("L1", "TERKIRIM_TERVERIFIKASI", "2026-09-25 10:00:00", "e1")])
sheets = [mon.Sheet("batch21", rows21, periksa21, [aud]), mon.Sheet("batch22", rows22, {2: periksa()}, [aud22])]
glob, dihapus = mon.dokumen_global([aud, aud22])
nama_wil = {"5108060": "BULELENG", "5108060002": "ANTURAN", "51080600020002": "BANJAR CONTOH"}
detail, ganda, luar = mon.bangun(sheets, glob, dihapus, {"d1": "APPROVED_TERVERIFIKASI"}, None, None, nama_wil)
kat = {(d["batch"], d["baris"]): d["kategori"] for d in detail}
check("kategori tanpa snapshot", kat, {
    ("batch21", 2): "APPROVED", ("batch21", 3): "SUBMITTED", ("batch21", 4): "DRAFT", ("batch21", 5): "BELUM_INPUT",
    ("batch21", 6): "DITOLAK_DATA", ("batch21", 7): "PERLU_CEK", ("batch22", 2): "SUBMITTED"})
d = {(x["batch"], x["baris"]): x for x in detail}
check("dipindah = di wilayah", (d[("batch21", 2)]["di_wilayah"], d[("batch21", 2)]["subsls_dokumen"]), ("ya", T1))
check("masih di wadah", d[("batch21", 3)]["di_wilayah"], "tidak")
check("ganda dihitung", (d[("batch21", 3)]["jumlah_dokumen"], d[("batch21", 3)]["ganda"]), (2, "ya"))
check("draft tanpa koordinat diberi keterangan", "tanpa koordinat" in d[("batch21", 4)]["keterangan"], True)
check("nama wilayah dari peta", (d[("batch21", 2)]["kecamatan"], d[("batch21", 2)]["desa"], d[("batch21", 2)]["sls"]),
      ("BULELENG", "ANTURAN", "BANJAR CONTOH"))
check("jenis ganda", sorted({g["jenis"] for g in ganda}), ["NAMA_SAMA_BEDA_BATCH", "USAHA_DOKUMEN_GANDA"])
check("anggota grup ganda", sorted(g["id_dokumen"] for g in ganda if g["jenis"] == "USAHA_DOKUMEN_GANDA"), ["d2", "d2b"])

# --- rekap ---
rk = mon.rekap(detail, ("desa_kode",), ("desa",))
check("rekap per desa", [(r["desa_kode"], r["total"], r["APPROVED"], r["usaha_ganda"], r["dokumen_berlebih"]) for r in rk],
      [("5108060002", 7, 1, 1, 1)])
check("persen dari usaha lolos pemeriksaan", (rk[0]["pct_terkirim"], rk[0]["pct_approved"]), (round(3 / 6, 4), round(1 / 6, 4)))

# --- dengan snapshot: server menang, ganda di luar audit, dokumen hilang, luar sheet ---
snap = {
    "d1": mon.ringkas_snapshot({"id": "d1", "nama": "WARUNG A (SATU)", "status": "APPROVED BY Pengawas", "subsls": T1}),
    "d2": mon.ringkas_snapshot({"id": "d2", "nama": "WARUNG B (DUA)", "status": "SUBMITTED BY Pencacah", "subsls": W}),
    # d2b tidak ada di server (dihapus admin); dokumen tercatat itu dibuat SEBELUM snapshot diambil
    "d3": mon.ringkas_snapshot({"id": "d3", "nama": "WARUNG C (TIGA)", "status": "DRAFT", "subsls": W, "galat": "1"}),
    "x9": mon.ringkas_snapshot({"id": "x9", "nama": "WARUNG D (EMPAT)", "status": "DRAFT", "subsls": W}),   # di luar audit
    "e1": mon.ringkas_snapshot({"id": "e1", "nama": "WARUNG B (DUA)", "status": "APPROVED BY Pengawas", "subsls": T1}),
    "z1": mon.ringkas_snapshot({"id": "z1", "nama": "TOKO LAIN (X)", "status": "DRAFT", "subsls": W}),
    "z2": mon.ringkas_snapshot({"id": "z2", "nama": "TOKO LAIN (X)", "status": "DRAFT", "subsls": W}),
    "p1": mon.ringkas_snapshot({"id": "p1", "kode_identitas": f"{W} - UMK - 19", "status": "DRAFT", "subsls": W}),
}
info_snap = {"diambil": "2030-01-01T00:00:00Z", "total_server": len(snap)}
detail2, ganda2, luar2 = mon.bangun(sheets, glob, dihapus, {}, snap, info_snap, nama_wil)
d2 = {(x["batch"], x["baris"]): x for x in detail2}
check("server menang (approved tanpa audit approve)", (d2[("batch21", 2)]["kategori"], d2[("batch21", 2)]["sumber_status"]),
      ("APPROVED", "server"))
check("ganda gugur kalau satu dokumen tidak di server", (d2[("batch21", 3)]["jumlah_dokumen"], d2[("batch21", 3)]["ganda"]), (1, ""))
check("galat server di keterangan", "galat server 1" in d2[("batch21", 4)]["keterangan"], True)
check("dokumen di luar audit dikenali lewat nama", (d2[("batch21", 5)]["kategori"], d2[("batch21", 5)]["id_dokumen"]),
      ("DRAFT", "x9"))
check("luar sheet (prelist dilewati)", sorted(x["id_dokumen"] for x in luar2), ["z1", "z2"])
check("ganda luar sheet", sorted({g["jenis"] for g in ganda2}), ["NAMA_SAMA_BEDA_BATCH", "NAMA_SAMA_LUAR_SHEET"])

# snapshot TIDAK lengkap: dokumen tak terbaca tidak dianggap hilang
detail3, _, _ = mon.bangun(sheets, glob, dihapus, {}, snap, {"diambil": "2030-01-01T00:00:00Z", "total_server": 999}, nama_wil)
d3 = {(x["batch"], x["baris"]): x for x in detail3}
check("snapshot tidak lengkap -> d2b tetap hidup", d3[("batch21", 3)]["jumlah_dokumen"], 2)
# dokumen dibuat SESUDAH snapshot diambil -> status audit
detail4, _, _ = mon.bangun(sheets, glob, dihapus, {}, snap, {"diambil": "2026-09-20T00:00:00Z", "total_server": len(snap)}, nama_wil)
d4 = {(x["batch"], x["baris"]): x for x in detail4}
check("dokumen lebih baru dari snapshot tidak dianggap hilang", d4[("batch21", 3)]["jumlah_dokumen"], 2)

# snapshot tanpa jumlah server (list_api_*.json) -> tidak pernah menganggap dokumen hilang
detail5, _, _ = mon.bangun(sheets, glob, dihapus, {}, snap, {"diambil": "", "total_server": 0}, nama_wil)
check("snapshot tanpa total -> tidak ada TIDAK_DI_SERVER",
      [x["jumlah_dokumen"] for x in detail5 if x["baris"] == 3 and x["batch"] == "batch21"], [2])

# --- dokumen dipakai >1 baris (baris identik) ---
rows_id = [Baris(2, "M1", "SAMA (A)", T1), Baris(3, "M1", "SAMA (A)", T1)]
aud_id = mon.ringkas_audit(Path("x"), [ab("M1", "TERKIRIM_TERVERIFIKASI", "2026-09-20 10:00:00", "m1")])
det, gd, _ = mon.bangun([mon.Sheet("b", rows_id, {}, [aud_id])], *mon.dokumen_global([aud_id]), {}, None, None, {})
check("dokumen dipakai banyak baris", sorted({g["jenis"] for g in gd}), ["DOKUMEN_DIPAKAI_BANYAK_BARIS"])

# --- keluaran: Excel & HTML terbentuk, data HTML padat bisa dibentangkan ---
with tempfile.TemporaryDirectory() as tmp:
    lembar = [("Per batch", mon.rekap(detail2, ("batch",)), ["batch", *mon.KOLOM_REKAP_ANGKA]),
              ("Per desa", mon.rekap(detail2, ("desa_kode",), ("desa",)), ["desa_kode", "desa", *mon.KOLOM_REKAP_ANGKA])]
    x = mon.tulis_excel(Path(tmp) / "m.xlsx", [("Dibuat", "uji")], lembar, ganda2, luar2, detail2)
    from openpyxl import load_workbook
    wb = load_workbook(x, read_only=True)
    check("lembar Excel", wb.sheetnames, ["Ringkasan", "Per desa", "Ganda", "Di luar sheet", "Detail"])
    wb.close()
    data = mon.data_html(detail2, ganda2, luar2, {"judul": "Uji", "dibuat": "x", "sheet": ["batch21", "batch22"],
                                                 "snapshot": "-", "ada_snapshot": True})
    check("html: batch & kategori jadi indeks", (data["meta"]["batch"], data["detail"][0][0], data["detail"][0][5]),
          (["batch21", "batch22"], 0, 0))
    check("html: url dasar", data["meta"]["url_dasar"].endswith("/"), True)
    h = mon.tulis_html(Path(tmp) / "m.html", data)
    teks = h.read_text(encoding="utf-8")
    check("html: data tersuntik sekali", (teks.count('"url_dasar"'), "/*__DATA__*/null" in teks), (1, False))

print("\nSEMUA LULUS" if ok_all else "\nADA YANG GAGAL")
sys.exit(0 if ok_all else 1)
