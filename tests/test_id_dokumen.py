# -*- coding: utf-8 -*-
"""Uji inti/id_dokumen.py (ID dokumen di kolom sheet sumber) + pemakaiannya di
main_gabungan, sinkron_list & tulis_id_sumber — offline, tanpa browser.
Jalankan: python tests/test_id_dokumen.py
"""
import csv
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import os  # noqa: E402
os.environ.setdefault("FASIH_ABAIKAN_CONFIG_LOKAL", "1")  # hasil uji tidak bergantung inti/config_lokal.py

import openpyxl  # noqa: E402

import input_gabungan.main_gabungan as mg  # noqa: E402
from input_gabungan.sinkron_list import rencana_sinkron, url_entry  # noqa: E402
from input_gabungan.tulis_id_sumber import rencana_id  # noqa: E402
from inti import id_dokumen as idd  # noqa: E402
from inti.gabungan_loader import KOLOM, GabunganRow, Pemeriksaan, _cari_indeks  # noqa: E402
from inti.tahap2_loader import KOLOM_TAHAP2, KOLOM_TAHAP2_TAMBAHAN, _indeks_tahap2  # noqa: E402

ok_all = True


def check(label, got, want):
    global ok_all
    ok = got == want
    ok_all &= ok
    print(f"{'PASS' if ok else 'FAIL'} | {label}: got={got!r} want={want!r}")


I1 = "11111111-1111-4111-8111-111111111111"
I2 = "22222222-2222-4222-8222-222222222222"
I3 = "33333333-3333-4333-8333-333333333333"

# --- bentuk ID ---
check("ID polos", idd.id_dari_teks(I1), I1)
check("ID dari URL entry", idd.id_dari_teks(f"https://x/survey/s/p/{I2}/entry"), I2)
check("huruf besar & spasi dirapikan", idd.id_dari_teks(f"  {I3.upper()} "), I3)
check("bukan ID", idd.id_dari_teks("baris 12"), "")
check("URL bukan entry", idd.id_dari_teks("https://x/survey/s/p"), "")

# --- kolom tambahan tidak mengganggu loader ---
judul_std = [f"{a} judul" for a in KOLOM.values()]
check("format standar: kolom ID diabaikan loader",
      _cari_indeks(judul_std + [idd.JUDUL_ID]), _cari_indeks(judul_std))
judul_t2 = list(dict.fromkeys([*KOLOM_TAHAP2.values(), *KOLOM_TAHAP2_TAMBAHAN.values()]))
check("format tahap 2: kolom ID diabaikan loader",
      _indeks_tahap2(judul_t2 + [idd.JUDUL_ID]), _indeks_tahap2(judul_t2))


def row(baris, nama, pemilik="I MADE"):
    return GabunganRow(baris, {"akun_ppl": "ppl@gmail.com", "idsubsls": "5108010002000501", "nama": nama,
                               "pengusaha": pemilik, "kbli": "86201"})


# --- pasang_id: sel tidak valid & ID ganda di-skip offline ---
rs = [row(2, "A"), row(3, "B"), row(4, "C"), row(5, "D")]
hasil = {r.baris: Pemeriksaan() for r in rs}
idd.pasang_id(rs, hasil, idd.IsiKolomId(kolom=9, mentah={2: I1, 3: "salah ketik", 4: I2, 5: I2},
                                        sidik={2: "s2", 3: "s3", 4: "s4", 5: "s5"}))
check("ID valid dipasang", (rs[0].id_dokumen, rs[0].sidik_sumber), (I1, "s2"))
check("sel bukan ID -> skip", hasil[3].status, "SKIP_DATA_ID_DOKUMEN_TIDAK_VALID")
check("ID sama di 2 baris -> keduanya skip", (hasil[4].status, hasil[5].status),
      ("SKIP_DATA_ID_DOKUMEN_GANDA", "SKIP_DATA_ID_DOKUMEN_GANDA"))
check("baris bersih tetap SIAP", hasil[2].status, "SIAP")


# --- penulisan ke xlsx ---
def buat_xlsx(path, isi, judul=("nama", "pemilik", "kbli")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "input_usaha"
    ws.append(list(judul))
    for b in isi:
        ws.append(list(b))
    wb.save(path)


def baca_xlsx(path):
    wb = openpyxl.load_workbook(path)
    return [[c.value for c in r] for r in wb["input_usaha"].iter_rows()]


class R:  # baris tiruan: cukup baris + sidik
    def __init__(self, baris, sidik, id_dokumen=""):
        self.baris, self.sidik_sumber, self.id_dokumen = baris, sidik, id_dokumen


pesan = []
with tempfile.TemporaryDirectory() as d:
    p = Path(d) / "sumber.xlsx"
    buat_xlsx(p, [("APOTEK A", "I MADE", "47721"), ("APOTEK B", "I KETUT CONTOH", "47721")])
    isi = idd.baca_kolom_id(p)
    check("belum ada kolom ID", isi.kolom, None)
    a, b = R(2, isi.sidik[2]), R(3, isi.sidik[3])
    pc = idd.PencatatIdSumber(p, cetak=pesan.append)
    pc.catat(a, f"https://x/survey/s/p/{I1}/entry")
    data = baca_xlsx(p)
    check("judul kolom ID dibuat di kanan", data[0], ["nama", "pemilik", "kbli", idd.JUDUL_ID])
    check("ID ditulis di baris yang benar", (data[1][3], data[2][3]), (I1, None))
    check("sheet dibaca lagi: ID terbaca", idd.baca_kolom_id(p).mentah.get(2), I1)
    check("sidik tidak berubah krn kolom ID", idd.baca_kolom_id(p).sidik[2], isi.sidik[2])

    # Sheet diurutkan di Excel selagi batch jalan: baris B pindah ke baris 2.
    buat_xlsx(p, [("APOTEK B", "I KETUT CONTOH", "47721", None), ("APOTEK A", "I MADE", "47721", I1)],
              judul=("nama", "pemilik", "kbli", idd.JUDUL_ID))
    pc.catat(b, I2)
    data = baca_xlsx(p)
    check("baris bergeser -> ID ikut sidik, bukan nomor baris", (data[1][3], data[2][3]), (I2, I1))

    # Excel sedang membuka berkas -> ditunda, lalu tertulis setelah ditutup.
    buat_xlsx(p, [("APOTEK C", "I MADE", "47721")])
    c = R(2, idd.baca_kolom_id(p).sidik[2])
    kunci_excel = p.with_name("~$" + p.name)
    kunci_excel.write_text("x")
    pc2 = idd.PencatatIdSumber(p, cetak=pesan.append)
    pc2.catat(c, I3)
    check("Excel terbuka -> ditunda, sheet tidak disentuh", (len(pc2.tertunda), len(baca_xlsx(p)[0])), (1, 3))
    kunci_excel.unlink()
    check("Excel ditutup -> tertulis", (pc2.simpan(), baca_xlsx(p)[1][3]), (True, I3))

    # Sel berisi hal lain / ID lain tidak pernah ditimpa.
    buat_xlsx(p, [("APOTEK D", "I MADE", "47721", "catatan manual")],
              judul=("nama", "pemilik", "kbli", idd.JUDUL_ID))
    dd = R(2, idd.baca_kolom_id(p).sidik[2])
    idd.PencatatIdSumber(p, cetak=pesan.append).catat(dd, I1)
    check("sel berisi teks lain tidak ditimpa", baca_xlsx(p)[1][3], "catatan manual")

    # Isi baris berubah (bukan cuma bergeser) -> tidak ditulis ke mana pun.
    buat_xlsx(p, [("APOTEK E", "I MADE", "47721")])
    e = R(2, "sidik-lama-yang-tidak-ada")
    pc3 = idd.PencatatIdSumber(p, cetak=pesan.append)
    pc3.catat(e, I1)
    check("baris tidak ketemu -> tidak ditulis & tidak tertunda",
          (len(baca_xlsx(p)[0]), pc3.tertunda), (3, {}))

    # Sheet berisi rumus -> tidak ditulis sama sekali.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "input_usaha"
    ws.append(["nama", "angka"])
    ws.append(["APOTEK F", "=1+1"])
    wb.save(p)
    f = R(2, idd.baca_kolom_id(p).sidik[2])
    pc4 = idd.PencatatIdSumber(p, cetak=pesan.append)
    pc4.catat(f, I1)
    check("sheet ber-rumus -> penulisan dimatikan", (bool(pc4.dimatikan), baca_xlsx(p)[1][1]), (True, "=1+1"))

    # CSV
    pcsv = Path(d) / "sumber.csv"
    with pcsv.open("w", newline="", encoding="utf-8-sig") as fh:
        csv.writer(fh).writerows([["nama", "kbli"], ["APOTEK G", "47721"], ["APOTEK H", "47721"]])
    g = R(3, idd.baca_kolom_id(pcsv).sidik[3])
    idd.PencatatIdSumber(pcsv, cetak=pesan.append).catat(g, I2)
    with pcsv.open(newline="", encoding="utf-8-sig") as fh:
        isi_csv = list(csv.reader(fh))
    check("CSV: judul & ID baris yang benar", (isi_csv[0][-1], isi_csv[1][-1], isi_csv[2][-1]),
          (idd.JUDUL_ID, "", I2))
    check("tidak ada berkas kunci/tmp tertinggal",
          sorted(x.name for x in Path(d).iterdir()), ["sumber.csv", "sumber.xlsx"])

# --- sinkron_list: nama di sheet dikoreksi, dokumen tetap dikenali lewat ID ---
AKUN, SUBSLS, ASG = "m@mail.com", "5108060014000403", "asg"
r_baru = row(7, "APOTEK NAMA BARU")
r_baru.id_dokumen = I1
audit_s = [{"kunci": "kunci-lama", "status": mg.STATUS_DIBUAT, "akun_login": AKUN, "nama_usaha": "APOTEK NAMA LAMA",
            "dokumen_url": url_entry(I1, ASG)}]
items = [{"id": I1, "data1": "APOTEK NAMA LAMA (I MADE)", "assignmentStatusAlias": "SUBMITTED BY Pencacah",
          "dateCreated": "2026-09-14T13:51:06.000+00:00"}]
lap, tulis, tak = rencana_sinkron([("S", r_baru, "SIAP")], items, AKUN, SUBSLS, ASG, audit_s, lengkap=True)
check("ID sheet: dokumen bernama lama dikenali", lap[0]["kategori"], "TERKIRIM")
check("ID sheet: dicatat ulang utk kunci baru",
      [(t["status"], t["kunci"]) for t in tulis],
      [(mg.STATUS_DIBUAT, r_baru.kunci), ("TERKIRIM_TERVERIFIKASI", r_baru.kunci)])
check("ID sheet: dokumen tidak dilaporkan tak dikenal", tak, [])
r_tanpa = row(7, "APOTEK NAMA BARU")
lap2, _, _ = rencana_sinkron([("S", r_tanpa, "SIAP")], items, AKUN, SUBSLS, ASG, audit_s, lengkap=True)
check("tanpa ID sheet: perilaku lama (belum ada)", lap2[0]["kategori"], "BELUM_ADA")

# --- giliran: nama dipakai baris lain tidak menahan baris ber-ID ---
with tempfile.TemporaryDirectory() as d:
    mg.AUDIT_LOG_PATH = Path(d) / "audit.csv"
    mg.append_audit({"kunci": "k108", "status": mg.STATUS_DIBUAT, "akun_login": "a@mail.com",
                     "idsubsls_input": "1", "nama_usaha": "PANGKALAN GAS (NYOMAN)",
                     "dokumen_url": url_entry(I2, ASG)})
    check("tanpa ID sheet: nama dipakai baris lain -> lewati",
          mg.alasan_lewati_saat_giliran("k267", ("a@mail.com", "1"), set(), "PANGKALAN GAS (NYOMAN)")
          .startswith("SKIP_NAMA_DIPAKAI_BARIS_LAIN"), True)
    check("ber-ID sheet: dibuka lewat ID -> dikerjakan",
          mg.alasan_lewati_saat_giliran("k267", ("a@mail.com", "1"), set(), "PANGKALAN GAS (NYOMAN)",
                                        id_sheet=I3), "")

    # --- tulis_id_sumber.rencana_id ---
    ra, rb, rc, rd, re_ = row(2, "A"), row(3, "B"), row(4, "C"), row(5, "D"), row(6, "E")
    rb.id_dokumen, rc.id_dokumen, rd.id_dokumen = I2, I3, I1
    dok = {ra.kunci: ("a", "1", url_entry(I1, ASG)), rb.kunci: ("a", "1", url_entry(I2, ASG)),
           rc.kunci: ("a", "1", url_entry(I1, ASG))}
    check("rencana_id", [(r.baris, k) for r, k, _ in rencana_id([ra, rb, rc, rd, re_], dok)],
          [(2, "TULIS"), (3, "SUDAH"), (4, "BEDA"), (5, "HANYA_SHEET"), (6, "-")])
    rf, rg = row(7, "F"), row(8, "G")
    dok2 = {rf.kunci: ("a", "1", url_entry(I1, ASG)), rg.kunci: ("a", "1", url_entry(I1, ASG))}
    check("audit memberi ID sama ke 2 baris -> tidak ditulis",
          [k for _, k, _ in rencana_id([rf, rg], dok2)], ["ID_GANDA", "ID_GANDA"])

print("\nSEMUA LULUS" if ok_all else "\nADA YANG GAGAL")
sys.exit(0 if ok_all else 1)
