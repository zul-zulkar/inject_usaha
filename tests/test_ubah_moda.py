# -*- coding: utf-8 -*-
"""Uji logika murni ubah_moda.py — offline, tanpa browser/VPN.
Jalankan: python tests/test_ubah_moda.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import os  # noqa: E402
os.environ.setdefault("FASIH_ABAIKAN_CONFIG_LOKAL", "1")  # hasil uji tidak bergantung inti/config_lokal.py

from inti.gabungan_loader import GabunganRow, Pemeriksaan
from fasih_sm.ganti_moda.ubah_moda import (
    STATUS_BERHENTI_SEGERA, STATUS_TUNTAS_LIVE, BarisAssignment, Berhenti, Target, angka_item_menu, bangun_target,
    baca_daftar, baris_dari_tabel, dialog_sesuai, diulang_per_subsls, jeda_cek_verifikasi, jeda_rate_limit,
    masih_tuntas, normalisasi_kode, pilih_tombol_konfirmasi, pola_item_ganti_mode, putuskan_verifikasi, rencanakan,
    target_dari_daftar_kode, target_dari_daftar_subsls, tulis_console,
)

ok_all = True


def check(label, got, want):
    global ok_all
    ok = got == want
    ok_all &= ok
    print(f"{'PASS' if ok else 'FAIL'} | {label}: got={got!r} want={want!r}")


S = "5108070013000901"
PPL = "ppl@gmail.com"
T = Target(S, (PPL,), [2])


def b(no, mode="CAPI", petugas=PPL, sub=S, indeks=0):
    return BarisAssignment(kode=f"{sub} - UMK - {no}", mode=mode, petugas=petugas, indeks=indeks)


# --- tabel: judul & isi PERSIS seperti yang terbaca di fasih-sm 2026-09-13 ---
HEAD = ["", "Kode Identitas", "Nama Keluarga/Bangunan/Usaha", "Alamat Prelist", "Nomor Urut Bangunan / IDSBR",
        "NIB / No. KK", "Email", "Skala Usaha / Jenis Prelist", "Jumlah Usaha", "Kode Pos", "Perubahan SLS",
        "IDSBR UMKM SLS Sama", "Status", "Mode", "Petugas Saat Ini", "Keterangan", ""]
ROW = ["", "5108060003000402 - UMK - 4", "USAHA CONTOH", "DSN CONTOH", "58 / 00000000",
       "0000000000000", "usaha.contoh@gmail.com", "UMK / OSS PERORANGAN", "", "81119", "", "",
       "approved by pengawas", "CAPI", "pml.contoh@gmail.com", "Pengawas", ""]
hasil = baris_dari_tabel({"head": HEAD, "rows": [ROW, ["Tidak ada data"]]})
check("parse tabel: 1 baris, baris pesan dilewati", len(hasil), 1)
check("parse kolom lewat judul", (hasil[0].idsubsls, hasil[0].mode, hasil[0].petugas, hasil[0].keterangan),
      ("5108060003000402", "CAPI", "pml.contoh@gmail.com", "Pengawas"))
try:
    baris_dari_tabel({"head": [h for h in HEAD if h != "Mode"], "rows": []})
    check("kolom Mode disembunyikan -> berhenti", "tidak berhenti", "KOLOM_TIDAK_ADA")
except Berhenti as e:
    check("kolom Mode disembunyikan -> berhenti", e.kode, "KOLOM_TIDAK_ADA")

# --- rencanakan (cakupan satu) ---
LAIN = "5108010009000202"
check("tanpa assignment", rencanakan(T, []).status, "TIDAK_ADA_ASSIGNMENT")
check("mode aneh", rencanakan(T, [b(1, mode="CAWI")]).status, "MODE_TIDAK_DIKENAL")
check("sudah ada PAPI milik PPL", rencanakan(T, [b(1, "PAPI"), b(2)]).status, "SUDAH_ADA_PAPI")
r = rencanakan(T, [b(1, petugas="lain@gmail.com"), b(2), b(3)])
check("CAPI milik PPL didahulukan -> pilih tepat satu", (r.status, [x.kode for x in r.pilih]),
      ("PERLU_DIUBAH", [f"{S} - UMK - 2"]))
check("email petugas beda huruf besar tetap cocok",
      rencanakan(T, [b(1, petugas="x@gmail.com"), b(2, petugas="PPL@Gmail.com")]).pilih[0].kode, f"{S} - UMK - 2")
# Ketetapan user 2026-09-14: assignment MANA PUN di subsls itu boleh (dulu PETUGAS_BEDA / PAPI_ADA_PETUGAS_LAIN).
check("PAPI milik petugas lain sudah cukup",
      rencanakan(T, [b(1, "PAPI", "x@gmail.com"), b(2, petugas="y@gmail.com")]).status, "SUDAH_ADA_PAPI")
r = rencanakan(T, [b(1, petugas="x@gmail.com"), b(2, petugas="y@gmail.com")])
check("CAPI semua milik petugas lain -> tetap pilih satu", (r.status, [x.kode for x in r.pilih]),
      ("PERLU_DIUBAH", [f"{S} - UMK - 1"]))

# --- baris subsls lain di hasil pencarian (dryrun 2026-09-14) -> diabaikan, bukan berhenti ---
r = rencanakan(T, [b(1, sub=LAIN), b(2, petugas="x@gmail.com"), b(3, "PAPI", sub=LAIN)])
check("subsls lain diabaikan: pilih baris subsls target saja", (r.status, [x.kode for x in r.pilih]),
      ("PERLU_DIUBAH", [f"{S} - UMK - 2"]))
check("subsls lain dicatat di pesan", "2 baris subsls lain diabaikan" in r.pesan, True)
check("PAPI subsls lain tidak dihitung", rencanakan(T, [b(1, "PAPI", sub=LAIN), b(2)]).status, "PERLU_DIUBAH")
check("hanya subsls lain di halaman -> berhenti (pencarian tidak menyaring)",
      rencanakan(T, [b(1, sub=LAIN)]).status, "SUBSLS_TIDAK_TAMPIL")
check("halaman kosong tapi ada halaman lain -> berhenti",
      rencanakan(T, [], ada_halaman_lain=True).status, "SUBSLS_TIDAK_TAMPIL")

# --- paginasi TIDAK dipindah (temuan user 2026-09-14): hanya halaman tampil yang dibaca ---
r = rencanakan(T, [b(1)], ada_halaman_lain=True)
check("ada halaman lain tapi CAPI tampil -> tetap diubah, dicatat", (r.status, "hanya halaman tampil" in r.pesan),
      ("PERLU_DIUBAH", True))
check("ada halaman lain & PAPI tampil -> sudah cukup",
      rencanakan(T, [b(1), b(2, "PAPI")], ada_halaman_lain=True).status, "SUDAH_ADA_PAPI")

# --- rencanakan (cakupan semua) ---
r = rencanakan(T, [b(1), b(2, "PAPI"), b(3, petugas="x@gmail.com"), b(4, sub=LAIN)], cakupan="semua")
check("semua: seluruh CAPI subsls ini, PAPI & subsls lain dilewati", (r.status, len(r.pilih)), ("PERLU_DIUBAH", 2))
check("semua: sudah PAPI semua", rencanakan(T, [b(1, "PAPI")], cakupan="semua").status, "TIDAK_ADA_CAPI")
check("semua: tanpa CAPI tampil tapi ada halaman lain -> berhenti",
      rencanakan(T, [b(1, "PAPI")], cakupan="semua", ada_halaman_lain=True).status, "PERLU_HALAMAN_LAIN")

# --- list kode identitas milik user (2026-09-15): HANYA kode itu yang diubah ---
# Kasus kembar dgn tests/test_ubah_moda_console.js.
check("normalisasi kode: spasi & huruf kecil & nol depan", normalisasi_kode(f"{S}-umk-04"), f"{S} - UMK - 4")
check("normalisasi kode: bukan kode", normalisasi_kode(S), "")
check("normalisasi kode: 17 digit bukan idsubsls", normalisasi_kode(f"9{S} - UMK - 4"), "")
dk_t, dk_tidak, dk_ganda = target_dari_daftar_kode([
    "Kode Identitas\tNama",
    f"{S} - UMK - 4\tWARUNG A",
    f"{S} - UMK - 41",
    f"{LAIN} - UMK - 2 ; {S} - umk - 4",
    "5108070013000999\tNIK tanpa kode",
    f"{S} - WARUNG - BU SRI - 12",
    f"{S} - NON-UMK - 3",
])
check("daftar kode: SATU target per kode, urutan list dipertahankan",
      [(t.kode, t.idsubsls, t.baris_sheet) for t in dk_t],
      [(f"{S} - UMK - 4", S, [2]), (f"{S} - UMK - 41", S, [3]), (f"{LAIN} - UMK - 2", LAIN, [4]),
       (f"{S} - NON-UMK - 3", S, [7])])
check("daftar kode: ganda dilaporkan", dk_ganda, [(4, f"{S} - UMK - 4")])
check("daftar kode: 16 digit tanpa pola kode dilaporkan", [x[0] for x in dk_tidak], [5, 6])
# Bentuk nyata dari list user & tabel fasih-sm (2026-09-15)
check("kode nama keluarga dgn '/'", normalisasi_kode("5108060029000102 - I KETUT CONTOH / I KOMANG AGUS CONTOH - 46"),
      "5108060029000102 - I KETUT CONTOH / I KOMANG AGUS CONTOH - 46")
check("kode nama diakhiri '/'", normalisasi_kode("5108070005000601 - WAYAN CONTOH / - 21"),
      "5108070005000601 - WAYAN CONTOH / - 21")
check("kode nama berangka", normalisasi_kode("5108020014000104 - MUH UMAR FARIDL / 1 - 48"),
      "5108020014000104 - MUH UMAR FARIDL / 1 - 48")
check("kode nama ber-apostrof & titik", normalisasi_kode(f"{S} - WR. MAK'E (BU TUT) - 9"), f"{S} - WR. MAK'E (BU TUT) - 9")
check("sel tabel berakhiran '/ - 81119'", normalisasi_kode("5108060029000102 - BANGUNAN KOSONG - 6 / - 81119"),
      "5108060029000102 - BANGUNAN KOSONG - 6")
check("sel tabel berakhiran '/ - 0'", normalisasi_kode("5108060029000102 - I KADEK RIKI CONTOH / KETUT CONTOH - 46 / - 0"),
      "5108060029000102 - I KADEK RIKI CONTOH / KETUT CONTOH - 46")
check("kode dari xlsx (tab + nama)", normalisasi_kode(f"{S} - DTSEN - 44\tNAMA"), f"{S} - DTSEN - 44")

# Pencarian memakai KODE itu sendiri; hasil pencarian "…- UMK - 4" bisa ikut memuat "- 41", "- 40", dst.
TK = Target(S, (), [2], kode=f"{S} - UMK - 4")
check("kode: istilah cari = kode identitas", TK.istilah_cari, f"{S} - UMK - 4")
check("sheet: istilah cari = idsubsls", T.istilah_cari, S)
r = rencanakan(TK, [b(41), b(4, petugas="x@gmail.com"), b(40)])
check("kode: pilih PERSIS '- 4', bukan '- 41'/'- 40'", (r.status, [x.kode for x in r.pilih]),
      ("PERLU_DIUBAH", [f"{S} - UMK - 4"]))
check("kode: baris lain yg ikut tampil dicatat", "2 baris kode lain" in r.pesan, True)
check("kode: kode di tabel beda spasi/huruf tetap cocok",
      rencanakan(TK, [BarisAssignment(kode=f"{S}-umk-4", mode="CAPI")]).status, "PERLU_DIUBAH")
check("kode: PAPI lain di subsls TIDAK membuat kode ini dilewati", rencanakan(TK, [b(41, "PAPI"), b(4)]).status,
      "PERLU_DIUBAH")
check("kode: sudah PAPI", rencanakan(TK, [b(4, "PAPI")]).status, "KODE_SUDAH_PAPI")
check("kode: KODE_SUDAH_PAPI tuntas", "KODE_SUDAH_PAPI" in STATUS_TUNTAS_LIVE, True)
check("kode: hasil kosong -> KODE_TIDAK_ADA (lanjut)", rencanakan(TK, []).status, "KODE_TIDAK_ADA")
check("kode: hanya kode lain subsls sama -> KODE_TIDAK_ADA", rencanakan(TK, [b(41)]).status, "KODE_TIDAK_ADA")
check("kode: tidak ada di halaman tampil tapi >1 halaman -> KODE_TIDAK_TAMPIL",
      rencanakan(TK, [b(41)], ada_halaman_lain=True).status, "KODE_TIDAK_TAMPIL")
check("kode: KODE_TIDAK_ADA/TIDAK_TAMPIL tidak tuntas & tidak menghentikan",
      [s in STATUS_TUNTAS_LIVE or s in STATUS_BERHENTI_SEGERA for s in ("KODE_TIDAK_ADA", "KODE_TIDAK_TAMPIL")],
      [False, False])
r = rencanakan(TK, [b(1, sub=LAIN), b(2, sub=LAIN)])
check("kode: hasil berisi subsls lain -> PENCARIAN_TIDAK_MENYARING (berhenti)",
      (r.status, r.status in STATUS_BERHENTI_SEGERA), ("PENCARIAN_TIDAK_MENYARING", True))
check("kode: kode sendiri tampil + subsls lain -> tetap kode itu saja",
      [x.kode for x in rencanakan(TK, [b(4), b(1, sub=LAIN)]).pilih], [f"{S} - UMK - 4"])
r = rencanakan(TK, [b(4), b(4)])
check("kode: tampil 2x -> KODE_GANDA (berhenti)", (r.status, r.status in STATUS_BERHENTI_SEGERA), ("KODE_GANDA", True))
check("kode: cakupan 'semua' diabaikan", len(rencanakan(TK, [b(4), b(41)], cakupan="semua").pilih), 1)
check("kode: mode aneh pada kode ini", rencanakan(TK, [b(4, "CAWI")]).status, "MODE_TIDAK_DIKENAL")
check("kode: mode aneh pada kode LAIN tidak menghalangi", rencanakan(TK, [b(41, "CAWI"), b(4)]).status,
      "PERLU_DIUBAH")
check("kunci: target sheet = idsubsls", T.kunci, S)
check("kunci: target kode = kode identitas", TK.kunci, f"{S} - UMK - 4")
check("cocok kode: '- 41' bukan '- 4'", [TK.cocok(x) for x in (b(4), b(41))], [True, False])
check("cocok sheet: semua baris subsls", [T.cocok(x) for x in (b(4), b(41), b(1, sub=LAIN))], [True, True, False])

# --- verifikasi tertunda & rate limit (run user 2026-09-15: Mode baru terbaca belakangan,
#     cek beruntun memicu HTTP 429). Kasus kembar dgn tests/test_ubah_moda_console.js. ---
check("jadwal cek ulang", [jeda_cek_verifikasi(k) for k in (0, 1, 4, 5, 40)], [30_000, 45_000, 120_000, 180_000, 180_000])
check("verifikasi: semua PAPI (huruf kecil juga)", putuskan_verifikasi({"a": "PAPI", "b": "papi"}, 1000, 60_000),
      "TERVERIFIKASI")
check("verifikasi: PAPI setelah batas tetap terverifikasi", putuskan_verifikasi({"a": "PAPI"}, 999_999, 60_000),
      "TERVERIFIKASI")
check("verifikasi: masih CAPI sebelum batas -> menunggu", putuskan_verifikasi({"a": "CAPI"}, 59_999, 60_000), "MENUNGGU")
check("verifikasi: masih CAPI saat batas habis", putuskan_verifikasi({"a": "CAPI"}, 60_000, 60_000), "BELUM_TERVERIFIKASI")
check("verifikasi: sebagian PAPI -> menunggu", putuskan_verifikasi({"a": "PAPI", "b": "CAPI"}, 1000, 60_000), "MENUNGGU")
check("verifikasi: kode tidak tampil -> menunggu", putuskan_verifikasi({"a": "(hilang)"}, 1000, 60_000), "MENUNGGU")
check("verifikasi: tanpa kode tidak pernah terverifikasi",
      [putuskan_verifikasi({}, 1000, 60_000), putuskan_verifikasi({}, 60_000, 60_000)], ["MENUNGGU", "BELUM_TERVERIFIKASI"])
check("verifikasi: waktu klik tak diketahui (inf) & masih CAPI", putuskan_verifikasi({"a": "CAPI"}, float("inf"), 60_000),
      "BELUM_TERVERIFIKASI")
check("429: 15 dtk x 2^ke, maks 2 mnt", [jeda_rate_limit(k) for k in (0, 1, 2, 3, 6)], [15_000, 30_000, 60_000, 120_000, 120_000])
check("429: Retry-After dihormati (5 dtk - 5 mnt)", [jeda_rate_limit(0, ra) for ra in ("20", "1", "999", "", "abc")],
      [20_000, 5_000, 300_000, 15_000, 15_000])
check("RATE_LIMIT & DIUBAH_BELUM_TERVERIFIKASI menghentikan batch",
      [s in STATUS_BERHENTI_SEGERA for s in ("RATE_LIMIT", "DIUBAH_BELUM_TERVERIFIKASI")], [True, True])
check("DIUBAH_MENUNGGU: tidak tuntas & tidak menghentikan",
      ["DIUBAH_MENUNGGU" in STATUS_TUNTAS_LIVE, "DIUBAH_MENUNGGU" in STATUS_BERHENTI_SEGERA], [False, False])

# --- arah balik PAPI -> CAPI utk subsls tertentu (permintaan user 2026-09-22). Kembar dgn test_ubah_moda_console.js ---
TC = Target(S, (), [1], ke="CAPI")
TKC = Target(S, (), [2], kode=f"{S} - UMK - 4", ke="CAPI")
check("arah: target bawaan = PAPI", (T.ke, TC.ke), ("PAPI", "CAPI"))
check("kunci: arah CAPI berawalan, PAPI tetap", (TC.kunci, TKC.kunci, T.kunci), (f"CAPI:{S}", f"CAPI:{S} - UMK - 4", S))
check("pola item menu per arah", [bool(pola_item_ganti_mode("CAPI").search("Ganti Mode (Ke CAPI) (3)")),
                                  bool(pola_item_ganti_mode("CAPI").search("Ganti Mode (Ke PAPI) (3)")),
                                  bool(pola_item_ganti_mode("PAPI").search("Ganti Mode ( Ke PAPI )"))], [True, False, True])
ds_t, ds_tidak, ds_ganda, ds_kode = target_dari_daftar_subsls([
    "idsubsls\tnama SLS",
    f"{S}\tBANJAR A",
    f"{LAIN}, 5108070013000902;5108070013000903",
    f"{S}",
    f"{S} - UMK - 4",
    "5.10807E+15\t510807001300090",
    "3 orang",
], "CAPI")
check("daftar subsls: dimuat berurutan, ber-arah", [(t.idsubsls, t.ke, t.baris_sheet) for t in ds_t],
      [(S, "CAPI", [2]), (LAIN, "CAPI", [3]), ("5108070013000902", "CAPI", [3]), ("5108070013000903", "CAPI", [3])])
check("daftar subsls: ganda / kode identitas / tidak dikenali", (ds_ganda, [x[0] for x in ds_kode], ds_tidak),
      ([(4, S)], [5], [(6, "5.10807E+15"), (6, "510807001300090")]))
r = rencanakan(TC, [b(1, "PAPI"), b(2), b(3, "PAPI", "x@gmail.com"), b(4, "PAPI", sub=LAIN)])
check("CAPI subsls: SEMUA PAPI subsls ini (petugas siapa pun), subsls lain diabaikan",
      (r.status, [x.kode for x in r.pilih]), ("PERLU_DIUBAH", [f"{S} - UMK - 1", f"{S} - UMK - 3"]))
check("CAPI subsls: cakupan diabaikan", len(rencanakan(TC, [b(1, "PAPI"), b(2, "PAPI")], cakupan="satu").pilih), 2)
check("CAPI subsls: semua sudah CAPI -> TIDAK_ADA_PAPI (tuntas)",
      (rencanakan(TC, [b(1), b(2)]).status, "TIDAK_ADA_PAPI" in STATUS_TUNTAS_LIVE), ("TIDAK_ADA_PAPI", True))
r = rencanakan(TC, [b(1), b(2)], ada_halaman_lain=True)
check("CAPI subsls: >1 halaman tanpa PAPI tampil -> CEK_HALAMAN_LAIN (lanjut, tidak tuntas)",
      (r.status, r.status in STATUS_TUNTAS_LIVE, r.status in STATUS_BERHENTI_SEGERA), ("CEK_HALAMAN_LAIN", False, False))
check("CAPI subsls: >1 halaman & PAPI tampil -> tetap diubah",
      rencanakan(TC, [b(1, "PAPI")], ada_halaman_lain=True).status, "PERLU_DIUBAH")
check("CAPI subsls: hanya subsls lain -> SUBSLS_TIDAK_TAMPIL (berhenti)",
      rencanakan(TC, [b(1, "PAPI", sub=LAIN)]).status, "SUBSLS_TIDAK_TAMPIL")
check("CAPI subsls: kosong -> TIDAK_ADA_ASSIGNMENT", rencanakan(TC, []).status, "TIDAK_ADA_ASSIGNMENT")
check("CAPI subsls: mode aneh", rencanakan(TC, [b(1, "CAWI")]).status, "MODE_TIDAK_DIKENAL")
r = rencanakan(TKC, [b(41, "PAPI"), b(4, "PAPI")])
check("CAPI kode: PAPI persis -> diubah", (r.status, [x.kode for x in r.pilih]), ("PERLU_DIUBAH", [f"{S} - UMK - 4"]))
check("CAPI kode: sudah CAPI -> KODE_SUDAH_CAPI (tuntas)",
      (rencanakan(TKC, [b(4)]).status, "KODE_SUDAH_CAPI" in STATUS_TUNTAS_LIVE), ("KODE_SUDAH_CAPI", True))
check("diulang per subsls: CAPI subsls & cakupan semua saja",
      [diulang_per_subsls(TC, "satu"), diulang_per_subsls(T, "satu"), diulang_per_subsls(T, "semua"),
       diulang_per_subsls(TKC, "semua")], [True, False, True, False])
check("verifikasi arah CAPI", [putuskan_verifikasi({"a": "CAPI"}, 1000, 60_000, "CAPI"),
                               putuskan_verifikasi({"a": "PAPI"}, 1000, 60_000, "CAPI"),
                               putuskan_verifikasi({"a": "CAPI"}, 1000, 60_000)],
      ["TERVERIFIKASI", "MENUNGGU", "MENUNGGU"])
check("dialog sesuai arah", [
    dialog_sesuai("Apakah Anda yakin mengubah mode 3 assignment ke CAPI?", "CAPI"),
    dialog_sesuai("Ubah mode ke PAPI?", "CAPI"),
    dialog_sesuai("Ganti Mode (Ke PAPI)", "PAPI"),
    dialog_sesuai("Ubah mode dari PAPI menjadi CAPI?", "CAPI"),
    dialog_sesuai("Apakah Anda yakin mengganti mode assignment?", "CAPI"),
    dialog_sesuai("Assignment PAPI akan diubah", "CAPI"),
    dialog_sesuai("Hapus assignment?", "CAPI"),
], [True, False, True, True, True, False, False])
# masih_tuntas (padanan masihTuntas Console, di sini dgn urutan baris audit)
papi_ok = {"idsubsls": S, "status": "SUDAH_ADA_PAPI", "kode_target": "", "ke": ""}
check("masih tuntas: tanpa klik arah lawan", masih_tuntas(T, [papi_ok], STATUS_TUNTAS_LIVE), True)
check("masih tuntas: klik ke CAPI sesudahnya -> diperiksa lagi",
      masih_tuntas(T, [papi_ok, {"idsubsls": S, "status": "DIUBAH_TERVERIFIKASI", "ke": "CAPI"},
                       {"idsubsls": S, "status": "TIDAK_ADA_PAPI", "ke": "CAPI"}], STATUS_TUNTAS_LIVE), False)
check("masih tuntas: klik ke CAPI SEBELUMNYA tidak berpengaruh",
      masih_tuntas(T, [{"idsubsls": S, "status": "DIUBAH_TERVERIFIKASI", "ke": "CAPI"}, papi_ok], STATUS_TUNTAS_LIVE), True)
check("masih tuntas: klik ke CAPI di subsls LAIN tidak berpengaruh",
      masih_tuntas(T, [papi_ok, {"idsubsls": LAIN, "status": "DIUBAH_TERVERIFIKASI", "ke": "CAPI"}], STATUS_TUNTAS_LIVE), True)
check("masih tuntas: kunci arah CAPI terpisah dari PAPI",
      [masih_tuntas(TC, [papi_ok], STATUS_TUNTAS_LIVE),
       masih_tuntas(TC, [{"idsubsls": S, "status": "TIDAK_ADA_PAPI", "ke": "CAPI"}], STATUS_TUNTAS_LIVE)], [False, True])

# --- baca file daftar (.txt & .xlsx) ---
import tempfile
with tempfile.TemporaryDirectory() as tmp:
    txt = Path(tmp) / "list.txt"
    txt.write_text(f"kode\n{S} - UMK - 4\n", encoding="utf-8")
    check("baca daftar .txt", baca_daftar(txt), ["kode", f"{S} - UMK - 4"])
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["No", "Kode Identitas"])
    wb.active.append([1, f"{S} - UMK - 4"])
    wb.create_sheet("lain").append([f"{LAIN} - UMK - 2"])
    xl = Path(tmp) / "list.xlsx"
    wb.save(xl)
    check("baca daftar .xlsx: sheet pertama saja", [t.kode for t in target_dari_daftar_kode(baca_daftar(xl))[0]],
          [f"{S} - UMK - 4"])
    check("baca daftar .xlsx: --sheet", [t.idsubsls for t in target_dari_daftar_kode(baca_daftar(xl, "lain"))[0]],
          [LAIN])
    # idsubsls tersimpan sbg ANGKA di Excel -> TIDAK dimuat (Excel memotong jadi 15 digit: …0901 -> …0900,
    # openpyxl membacanya sbg int 16 digit yang tampak sah). Hanya sel TEKS yang dimuat.
    wb2 = openpyxl.Workbook()
    wb2.active.append(["idsubsls"])
    wb2.active.append([S])
    wb2.active.append([int(S)])
    wb2.active.append([5.1080700130009e15])
    xl2 = Path(tmp) / "subsls.xlsx"
    wb2.save(xl2)
    t2, tidak2, _, _ = target_dari_daftar_subsls(baca_daftar(xl2), "CAPI")
    check("baca subsls .xlsx: teks dimuat, sel angka dilaporkan", ([t.idsubsls for t in t2], [x[0] for x in tidak2]),
          ([S], [3, 4]))

    # file siap-tempel memuat arah per target (Console menolak jalan ke arah lain)
    import json
    import re as _re
    import fasih_sm.ganti_moda.ubah_moda as um
    um_siap_lama = um.KONSOL_SIAP
    um.KONSOL_SIAP = Path(tmp) / "siap.js"
    try:
        teks = tulis_console([TC]).read_text(encoding="utf-8")
        data = json.loads(_re.search(r"const TARGET = (\[.*?\]);", teks).group(1))
        check("siap.js: target subsls ber-arah CAPI", data, [{"idsubsls": S, "ppl": [], "baris": [1], "siap": False, "ke": "CAPI"}])
    finally:
        um.KONSOL_SIAP = um_siap_lama

# --- menu & dialog ---
check("angka item menu (spasi)", angka_item_menu("Ganti Mode (Ke PAPI) (3)"), 3)
check("angka item menu (rapat + baris baru)", angka_item_menu("Ganti Mode (Ke PAPI)\n(0)"), 0)
check("angka item menu tidak ada", angka_item_menu("Ganti Mode (Ke PAPI)"), None)
check("tombol konfirmasi tunggal", pilih_tombol_konfirmasi(["Batal", "Ya, Ganti Mode"]), 1)
check("tombol 'Tidak' bukan konfirmasi", pilih_tombol_konfirmasi(["Tidak", "Ya"]), 1)
check("tanpa tombol konfirmasi", pilih_tombol_konfirmasi(["Batal", "Tutup"]), None)
check("dua kandidat -> ambigu", pilih_tombol_konfirmasi(["Ya", "Simpan"]), None)

# --- target dari sheet ---
def g(baris, ids, pilih=None, ppl=PPL):
    v = {"idsubsls": ids, "akun_ppl": ppl, "nama": f"U{baris}", "pilih_prov": (pilih or ids)[:2],
         "pilih_kab": (pilih or ids)[2:4], "pilih_kec": (pilih or ids)[4:7], "pilih_desa": (pilih or ids)[7:10],
         "pilih_sls": (pilih or ids)[10:14], "pilih_subsls": (pilih or ids)[14:16]}
    return GabunganRow(baris, v)


rows = [g(2, S), g(3, S), g(4, "5108070013000902"), g(5, "5108010001200100", pilih="5108010001000302")]
cek = {2: Pemeriksaan(), 3: Pemeriksaan([("X", "")]), 4: Pemeriksaan([("X", "")]), 5: Pemeriksaan()}
targets, keluar = bangun_target(rows, cek)
check("satu target per idsubsls, baris digabung",
      [(t.idsubsls, t.baris_sheet, t.siap_input) for t in targets],
      [(S, [2, 3], True), ("5108070013000902", [4], False)])
check("idsubsls tidak konsisten dikeluarkan", keluar, [(5, "5108010001200100", "5108010001000302")])
check("--hanya-siap", [t.idsubsls for t in bangun_target(rows, cek, hanya_siap=True)[0]], [S])

print("\nSEMUA PASS" if ok_all else "\nADA YANG FAIL")
sys.exit(0 if ok_all else 1)
