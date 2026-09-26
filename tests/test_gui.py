#!/usr/bin/env python3
"""test_gui.py — uji OFFLINE GUI web lokal (gui/). Jalankan: python tests/test_gui.py

Yang dikunci:
  * setiap opsi CLI di gui/alat.py BENAR-BENAR ada di skrip alatnya (GUI tidak boleh
    menyusun perintah yang ditolak argparse, apalagi diam-diam salah opsi);
  * penyusun argumen: wajib / wajib_jika (termasuk @aksi), nilai Pengaturan, pola, centang_nilai;
  * setiap prompt "Ketik 'YA'" di skrip dikenali POLA_PROMPT_YA (dialog YA GUI);
  * gui/sisip/sitecustomize.py menimpakan pengaturan GUI ke inti.config, menggabungkan dict
    kodepos/wilayah, memakai password sesi, dan menghentikan alat kalau pengaturannya rusak;
  * proses + jawaban YA lewat stdin; clipboard KBLI/koordinat; validasi pengaturan;
  * server hanya melayani token + Host localhost.
Semua data contoh FIKTIF."""

import os
os.environ["FASIH_ABAIKAN_CONFIG_LOKAL"] = "1"   # uji pakai data contoh Buleleng di config.py

import sys as _sys, os as _os
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))

import json
import re
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path

AKAR = Path(__file__).resolve().parents[1]
SEMENTARA = Path(tempfile.mkdtemp(prefix="uji_gui_"))

from gui import alat as A  # noqa: E402
from gui import berkas_tabel as T  # noqa: E402
from gui import server as SV  # noqa: E402

# keluaran uji TIDAK boleh menyentuh gui/hasil/ & gui/pengaturan.json asli
SV.LOG = SEMENTARA / "log"
SV.SNAPSHOT = SEMENTARA / "proses"
SV.ISIAN_PATH = SEMENTARA / "isian_terakhir.json"
SV.PENG = SV.Pengaturan(SEMENTARA / "pengaturan.json")

gagal = 0


def cek(nama, dapat, harap):
    global gagal
    if dapat == harap:
        print(f"  OK  {nama}")
    else:
        gagal += 1
        print(f"FAIL  {nama}\n      dapat: {dapat!r}\n      harap: {harap!r}")


def galat(fungsi) -> str:
    try:
        fungsi()
    except (A.IsianSalah, T.TabelSalah, ValueError) as e:
        return str(e)
    return ""


# ------------------------------------------------------------------ opsi CLI ada di skrip
print("== opsi CLI gui/alat.py ada di skrip alatnya ==")
BERSAMA = (AKAR / "input_usaha" / "mesin.py").read_text(encoding="utf-8")   # opsi_audit / opsi_format
for a in A.ALAT:
    if a["skrip"].startswith("@kbli/"):
        skrip = AKAR.parent / "generate_kbli" / "generate_kbli.py"
        if not skrip.exists():
            print(f"  --  {a['id']}: proyek generate_kbli tidak ada di PC ini, dilewati")
            continue
    else:
        skrip = AKAR / a["skrip"]
    cek(f"{a['id']}: skrip ada", skrip.exists(), True)
    sumber = skrip.read_text(encoding="utf-8") + BERSAMA
    opsi = {i["arg"] for i in a["isian"] if i.get("arg") and i["arg"] != "@posisi"}
    opsi |= {x for ak in a["aksi"] for x in ak["tambah"]}
    opsi |= {x for i in a["isian"] for v in (i.get("arg_nilai") or {}).values() for x in v}
    hilang = sorted(o for o in opsi if f'"{o}"' not in sumber and f"'{o}'" not in sumber)
    cek(f"{a['id']}: {len(opsi)} opsi dikenal skrip", hilang, [])
    cek(f"{a['id']}: tidak pernah --headless / --ya", sorted(opsi & {"--headless", "--ya"}), [])

print("\n== konsistensi daftar alat ==")
for a in A.ALAT:
    nama = [i["nama"] for i in a["isian"]]
    cek(f"{a['id']}: nama isian unik", len(nama), len(set(nama)))
    for i in a["isian"]:
        for syarat in (i.get("tampil_jika") or {}, i.get("wajib_jika") or {}):
            for k in syarat:
                if k != "@aksi":
                    cek(f"{a['id']}.{i['nama']}: syarat '{k}' merujuk isian yang ada", k in nama, True)
        if i.get("aksi"):
            cek(f"{a['id']}.{i['nama']}: aksi terbatas dikenal", set(i["aksi"]) <= {x["id"] for x in a["aksi"]}, True)
    for ak in a["aksi"]:
        if ak["jenis"] == "bahaya":
            # tindakan irreversible: skripnya sendiri meminta YA (dialog GUI), atau GUI meminta YA dulu
            cek(f"{a['id']}.{ak['id']}: bahaya = butuh password (login)", bool(ak["password"]), True)

# ------------------------------------------------------------------ penyusun argumen
print("\n== susun_argumen ==")
inp = A.ALAT_PER_ID["input"]
cek("cek tanpa akun boleh", A.susun_argumen(inp, "cek", {"sumber": "bahan/a.xlsx", "lewati_selesai": False}),
    ["--sumber", "bahan/a.xlsx", "--cek"])
cek("dry-run tanpa akun & tanpa Pengaturan ditolak",
    galat(lambda: A.susun_argumen(inp, "dryrun", {"sumber": "a.xlsx"})),
    "Akun PPL (email): wajib diisi.\nSubsls wadah (16 digit): wajib diisi.")
cek("dry-run memakai akun & subsls dari Pengaturan",
    A.susun_argumen(inp, "dryrun", {"sumber": "a.xlsx", "lewati_selesai": False, "sinkron_dulu": False},
                    lambda n: {"GABUNGAN_AKUN_TUNGGAL": "ppl.contoh@mail.com",
                               "GABUNGAN_SUBSLS_TUNGGAL": "5108010010000105"}.get(n, "")),
    ["--sumber", "a.xlsx", "--akun-tunggal", "ppl.contoh@mail.com", "--subsls-tunggal", "5108010010000105"])
cek("alur per-baris tidak butuh akun tunggal",
    galat(lambda: A.susun_argumen(inp, "dryrun", {"sumber": "a.xlsx", "per_baris": True})), "")
cek("kirim menambah --submit",
    A.susun_argumen(inp, "kirim", {"sumber": "a.xlsx", "akun_tunggal": "ppl.contoh@mail.com",
                                   "subsls_tunggal": "5108010010000105"})[-1], "--submit")
cek("subsls bukan 16 digit ditolak",
    galat(lambda: A.susun_argumen(inp, "kirim", {"sumber": "a.xlsx", "akun_tunggal": "ppl.contoh@mail.com",
                                                 "subsls_tunggal": "5108"})),
    "Subsls wadah (16 digit): format tidak sesuai (dapat '5108').")
cek("baris 2,5,10-20 lolos", "--baris" in A.susun_argumen(inp, "cek", {"sumber": "a", "baris": "2, 5,10-20"}), True)
cek("baris ngawur ditolak", bool(galat(lambda: A.susun_argumen(inp, "cek", {"sumber": "a", "baris": "2;x"}))), True)
cek("angka bukan bilangan ditolak", bool(galat(lambda: A.susun_argumen(inp, "cek", {"sumber": "a", "limit": "1.5"}))), True)
cek("pilihan tak dikenal ditolak", bool(galat(lambda: A.susun_argumen(inp, "cek", {"sumber": "a", "koordinat": "x"}))), True)
kk = A.ALAT_PER_ID["kontrol"]
cek("centang_nilai tanpa nilai", A.susun_argumen(kk, "jalan", {"sumber": "a", "per_ppl": {"aktif": True, "nilai": ""}}),
    ["--sumber", "a", "--per-ppl"])
cek("centang_nilai dgn nilai", A.susun_argumen(kk, "jalan", {"sumber": "a", "per_ppl": {"aktif": True, "nilai": "D:/x"}}),
    ["--sumber", "a", "--per-ppl", "D:/x"])
bw = A.ALAT_PER_ID["buka_wilayah"]
cek("pilihan tanpa arg -> arg_nilai", A.susun_argumen(bw, "console", {"cakupan": "semua"}), ["--semua", "--console"])
cek("isian tersembunyi tidak diteruskan", A.susun_argumen(bw, "console", {"cakupan": "semua", "daftar": "x.txt"}),
    ["--semua", "--console"])
cek("daftar wajib kalau cakupan daftar", galat(lambda: A.susun_argumen(bw, "console", {"cakupan": "daftar"})),
    "Daftar idsubsls (satu per baris): wajib diisi.")
ga = A.ALAT_PER_ID["gabung_audit"]
cek("berkas banyak: koma di path TIDAK memecah",
    A.susun_argumen(ga, "laporan", {"sumber": ["audit/pc/pc 1, lama", "audit/pc/pc2"]}),
    ["--sumber", "audit/pc/pc 1, lama", "--sumber", "audit/pc/pc2"])
ap_ = A.ALAT_PER_ID["approve"]
cek("approve mode daftar", A.susun_argumen(ap_, "approve", {"target": "daftar", "daftar": "bahan/submit.xlsx", "limit": "1"}),
    ["--daftar", "bahan/submit.xlsx", "--limit", "1", "--eksekusi"])
cek("approve mode audit wajib akun PPL", "Akun PPL" in galat(lambda: A.susun_argumen(ap_, "cek", {"target": "audit", "akun_pml": "pml.satu@mail.com"})), True)
kb = A.ALAT_PER_ID["kbli"]
cek("KBLI: berkas = argumen posisi pertama", A.susun_argumen(kb, "jalan", {"berkas": "s.xlsx", "model": "kecil"}),
    ["s.xlsx", "--model", "kecil"])
cek("path keluaran bawaan KBLI", A.path_keluaran(kb["keluaran"][0], {"berkas": "bahan/sheet uji.xlsx"}, AKAR),
    AKAR / "gui/hasil/kbli/sheet uji_kbli.xlsx")
cek("path keluaran diganti isian", A.path_keluaran(kb["keluaran"][0], {"berkas": "a.xlsx", "keluaran": "D:/h.xlsx"}, AKAR),
    Path("D:/h.xlsx"))

# ------------------------------------------------------------------ prompt YA
print("\n== prompt YA di skrip dikenali GUI ==")
jumlah = 0
for f in list(AKAR.glob("*/*.py")) + list(AKAR.glob("*/*/*.py")):
    if "tests" in f.parts or "gui" in f.parts or "arsip" in f.parts or ".claude" in f.parts:
        continue
    for m in re.finditer(r"input\((f?)([\"'])(.*?)\2", f.read_text(encoding="utf-8")):
        if "YA" in m.group(3):
            jumlah += 1
            cek(f"{f.relative_to(AKAR).as_posix()}: {m.group(3)[:50]}", bool(A.POLA_PROMPT_YA.search(m.group(3))), True)
cek("ada prompt YA yang ditemukan", jumlah >= 3, True)
cek("baris biasa bukan prompt", bool(A.POLA_PROMPT_YA.search("DRY_RUN_SIAP_KIRIM baris 2")), False)

# ------------------------------------------------------------------ penimpa pengaturan (sitecustomize)
print("\n== gui/sisip/sitecustomize.py ==")


def jalan_anak(pengaturan, kode, password=""):
    path = SEMENTARA / f"p_{time.time_ns()}.json"
    path.write_text(pengaturan if isinstance(pengaturan, str) else json.dumps(pengaturan), encoding="utf-8")
    env = {**os.environ, "PYTHONPATH": str(AKAR / "gui" / "sisip"), "FASIH_GUI_PENGATURAN": str(path),
           "PYTHONIOENCODING": "utf-8", "FASIH_PASSWORD": password}
    return subprocess.run([sys.executable, "-c", f"import sys; sys.path.insert(0, r'{AKAR}');\n{kode}"],
                          capture_output=True, text=True, encoding="utf-8", env=env, cwd=AKAR)


p = jalan_anak({"timpa": {"KODE_KAB": "5201", "TAHAP2_KOTAK_KOORDINAT": {"__py__": "(-9.0, -8.5, 115.8, 116.3)"},
                          "GABUNGAN_BARIS_PER_SESI": 25},
                "tambah": {"KODEPOS_BY_DESA": {"5201010001": "83355"}}},
               "from inti.config import KODE_KAB, TAHAP2_KOTAK_KOORDINAT, KODEPOS_BY_DESA, FIXED_PASSWORD, GABUNGAN_BARIS_PER_SESI\n"
               "import inti.tahap2_loader as t\n"
               "print(KODE_KAB, TAHAP2_KOTAK_KOORDINAT, GABUNGAN_BARIS_PER_SESI, KODEPOS_BY_DESA.get('5201010001'),"
               " KODEPOS_BY_DESA.get('5108010010'), FIXED_PASSWORD, t.KODE_KAB)", password="sandi-uji")
cek("timpa + gabung dict + password + modul lain ikut", p.stdout.strip(),
    "5201 (-9.0, -8.5, 115.8, 116.3) 25 83355 81155 sandi-uji 5201")
p = jalan_anak({"timpa": {}}, "from inti.config import FIXED_PASSWORD; print(repr(FIXED_PASSWORD))")
cek("tanpa password sesi: password config tidak diubah", p.stdout.strip(), "''")
p = jalan_anak("{rusak", "print('TIDAK BOLEH JALAN')")
cek("pengaturan rusak -> alat berhenti (kode 3)", (p.returncode, "TIDAK BOLEH" in p.stdout), (3, False))
p = subprocess.run([sys.executable, "-c", f"import sys; sys.path.insert(0, r'{AKAR}'); from inti.config import KODE_KAB; print(KODE_KAB)"],
                   capture_output=True, text=True, env={**os.environ, "PYTHONPATH": str(AKAR / "gui" / "sisip")}, cwd=AKAR)
cek("tanpa FASIH_GUI_PENGATURAN tidak ada efek", p.stdout.strip(), "5108")

# ------------------------------------------------------------------ proses + jawaban YA
print("\n== proses & dialog YA ==")
SKRIP_UJI = SEMENTARA / "prompt.py"
SKRIP_UJI.write_text("print('ringkasan: 3 dokumen', flush=True)\n"
                     "j = input(\"Ketik 'YA' utk konfirmasi submit 3 dokumen SUNGGUHAN: \")\n"
                     "print('jawab=' + j)\n", encoding="utf-8")
pr = SV.Proses(1, {"id": "uji", "judul": "Uji", "keluaran": []}, {"id": "kirim", "label": "kirim", "jenis": "bahaya"},
               {}, [sys.executable, "-u", str(SKRIP_UJI)], SEMENTARA,
               {**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUNBUFFERED": "1"}, None)
pr.mulai_jalan()
batas = time.time() + 20
while not pr.menunggu_ya and time.time() < batas:
    time.sleep(0.1)
cek("prompt terdeteksi sbg menunggu YA", pr.menunggu_ya, True)
cek("ringkasan sebelum prompt tertangkap", pr.baris[:1], ["ringkasan: 3 dokumen"])
pr.masukan("YA")
while pr.status == "berjalan" and time.time() < batas:
    time.sleep(0.1)
cek("proses selesai kode 0", (pr.status, pr.kode), ("selesai", 0))
cek("jawaban sampai ke skrip", pr.baris[-1], "jawab=YA")
cek("log tertulis", "jawab=YA" in pr.log_path.read_text(encoding="utf-8"), True)
cek("proses selesai tidak menerima masukan", bool(galat(lambda: pr.masukan("YA"))), True)

# ------------------------------------------------------------------ clipboard KBLI & koordinat
print("\n== clipboard KBLI & koordinat ==")
from openpyxl import Workbook  # noqa: E402

wb = Workbook()
ws = wb.active
ws.title = "input_usaha"
ws.append(["8b.", "13a", "13f", "Latitude", "Longitude", "Kode KBLI", "Judul KBLI"])
ws.append(["WARUNG CONTOH", "menjual beras eceran", "beras", "-8.1", "115.1", None, None])
ws.append(["KANDANG CONTOH", "beternak ayam", "telur", "", "", "1464", "Peternakan ayam"])   # nol hilang
ws.append(["TOKO CONTOH", "jual pulsa", "pulsa", "-8,148,438", "115.2", "47411", "Judul petugas"])
ws.append([None] * 7)
wb.save(SEMENTARA / "sumber.xlsx")
wb = Workbook()
ws = wb.active
ws.append(["Kode KBLI", "Judul KBLI", "Sumber", "Skor", "Selisih", "Catatan", "Alternatif 2"])
ws.append(["47241", "Perdagangan Eceran Beras", "mesin", 0.9, 0.2, "", "47111 X"])
ws.append(["01464", "Peternakan Ayam Ras Petelur", "mesin", 0.7, 0.01, "PERIKSA: skor rendah", "01463 Y"])
ws.append(["61202", "Jasa Pulsa", "mesin", 0.8, 0.02, "PERIKSA: pilihan 1 & 2 rapat (0.020)", ""])
wb.save(SEMENTARA / "hasil_kbli.xlsx")
d = T.salin_kbli(SEMENTARA / "hasil_kbli.xlsx", SEMENTARA / "sumber.xlsx", "input_usaha", pertahankan=True)
cek("sel tujuan = kolom Kode KBLI baris 2", d["tujuan"]["kode"], "F2")
cek("kode & judul berdampingan", "gabung" in d["salin"], True)
cek("baris kode sheet dipertahankan", d["salin"]["gabung"]["teks"].split("\r\n")[:3],
    ["47241\tPerdagangan Eceran Beras", "1464\tPeternakan ayam", "47411\tJudul petugas"])
cek("ringkasan", d["ringkas"], {"baris": 3, "mesin": 1, "dipertahankan": 2, "kosong": 0, "periksa": 0})
cek("beda kode dilaporkan", [b["baris"] for b in d["beda"]], [3, 4])
cek("kode 4 digit diperingatkan", any("4 digit" in x for x in d["peringatan"]), True)
cek("HTML menandai kode sbg teks", "mso-number-format" in d["salin"]["gabung"]["html"], True)
d2 = T.salin_kbli(SEMENTARA / "hasil_kbli.xlsx", SEMENTARA / "sumber.xlsx", "input_usaha", pertahankan=False)
cek("tanpa pertahankan: kode mesin, nol di depan utuh", d2["salin"]["gabung"]["teks"].split("\r\n")[1].split("\t")[0], "01464")
cek("tanpa pertahankan: PERIKSA terdaftar", [x["baris"] for x in d2["periksa"]], [3, 4])
wb = Workbook()
ws = wb.active
ws.append(["Latitude", "Longitude", "Berubah"])
ws.append(["-8.1", "115.1", ""])
ws.append([None, None, None])                      # baris yang tidak diproses alat
ws.append(["-8.148438", "115.2", "YA"])
wb.save(SEMENTARA / "hasil_koordinat.xlsx")
k = T.salin_koordinat(SEMENTARA / "hasil_koordinat.xlsx", SEMENTARA / "sumber.xlsx", "input_usaha")
cek("koordinat: sel tujuan", k["tujuan"]["kode"], "D2")
cek("koordinat: ringkasan", k["ringkas"], {"baris": 3, "berubah": 1, "dipertahankan": 0})
k_isi = k["salin"]["gabung"]["teks"].split("\r\n")
cek("koordinat: baris kosong tetap kosong kalau sheet juga kosong", k_isi[1], "\t")
cek("berkas bukan hasil generate_kbli ditolak",
    "Judul kolom 'Kode KBLI' tidak ada" in galat(lambda: T.salin_kbli(SEMENTARA / "hasil_koordinat.xlsx")), True)
cek("berkas bukan hasil koordinat ditolak",
    "bukan hasil koordinat_pengganti" in galat(lambda: T.salin_koordinat(SEMENTARA / "hasil_kbli.xlsx")), True)
# baris yang dilewati alat koordinat tapi berisi di sheet -> nilai sheet dipertahankan (tidak terhapus saat ditempel)
wb = Workbook()
ws = wb.active
ws.append(["Latitude", "Longitude", "Berubah"])
ws.append([None, None, None])                      # baris 2 dilewati alat, di sheet berisi -8.1/115.1
ws.append(["-8.2", "115.3", "YA"])
wb.save(SEMENTARA / "hasil_koordinat2.xlsx")
k2 = T.salin_koordinat(SEMENTARA / "hasil_koordinat2.xlsx", SEMENTARA / "sumber.xlsx", "input_usaha")
cek("koordinat: baris tengah yang dilewati alat memakai nilai sheet",
    (k2["salin"]["gabung"]["teks"], k2["ringkas"]["dipertahankan"]), ("-8.1\t115.1\r\n-8.2\t115.3\r\n", 1))
p = T.impor_tabel
wb = Workbook()
ws = wb.active
ws.append(["iddesa", "nmdesa", "kodepos"])
ws.append(["5201010001", "DESA SATU", "83355"])
ws.append(["5201010002000101", "SUBSLS", "83356"])
ws.append(["52010100", "SALAH", "83357"])
wb.save(SEMENTARA / "kodepos.xlsx")
r = T.impor_tabel(SEMENTARA / "kodepos.xlsx", "kodepos", {"kode": 0, "kodepos": 2})
cek("impor kodepos desa & idsubsls", r["data"], {"desa": {"5201010001": "83355"}, "subsls": {"5201010002000101": "83356"}})
cek("impor kodepos: baris salah dilaporkan", len(r["dilewati"]), 1)

# ------------------------------------------------------------------ validasi pengaturan
print("\n== validasi pengaturan ==")
SV.CONFIG.segarkan()
cek("KODE_KAB 4 digit", SV.periksa_pengaturan({"KODE_KAB": "5201"}, {}), [])
cek("KODE_KAB salah", bool(SV.periksa_pengaturan({"KODE_KAB": "52"}, {})), True)
cek("int harus int", bool(SV.periksa_pengaturan({"GABUNGAN_BARIS_PER_SESI": "40"}, {})), True)
cek("kotak sah", SV.periksa_pengaturan({"TAHAP2_KOTAK_KOORDINAT": {"__py__": "(-9.0, -8.5, 115.8, 116.3)"}}, {}), [])
cek("kotak None boleh", SV.periksa_pengaturan({"TAHAP2_KOTAK_KOORDINAT": None}, {}), [])
cek("kotak terbalik ditolak", bool(SV.periksa_pengaturan({"TAHAP2_KOTAK_KOORDINAT": {"__py__": "(-8.0, -9.0, 115.8, 116.3)"}}, {})), True)
cek("tuple lewat literal Python", SV.periksa_pengaturan({"TAHAP2_BULAN_OPERASI": {"__py__": "('JULI',)"}}, {}), [])
cek("literal rusak ditolak", bool(SV.periksa_pengaturan({"TAHAP2_BULAN_OPERASI": {"__py__": "('JULI'"}}, {})), True)
cek("nama tak dikenal ditolak", bool(SV.periksa_pengaturan({"TIDAK_ADA": 1}, {})), True)
cek("password tidak bisa diatur lewat pengaturan", bool(SV.periksa_pengaturan({"FIXED_PASSWORD": "x"}, {})), True)
cek("tambah kodepos sah", SV.periksa_pengaturan({}, {"KODEPOS_BY_DESA": {"5201010001": "83355"}}), [])
cek("tambah kodepos salah", bool(SV.periksa_pengaturan({}, {"KODEPOS_BY_DESA": {"520101": "83355"}})), True)
cek("tambah selain kodepos/wilayah ditolak", bool(SV.periksa_pengaturan({}, {"TAHAP2_DEFAULT": {"a": "b"}})), True)
cek("config.py: password tidak pernah dikirim", "FIXED_PASSWORD" in SV.CONFIG.per_nama, False)

# ------------------------------------------------------------------ server HTTP
print("\n== server HTTP (token & Host) ==")
srv = SV.ThreadingHTTPServer(("127.0.0.1", 0), SV.Penangan)
SV.Penangan.token = "token-uji"
SV.Penangan.port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
url = f"http://127.0.0.1:{SV.Penangan.port}"


def minta(path, data=None, token="token-uji", host=None):
    kepala = {"X-Token": token}
    if host:
        kepala["Host"] = host
    req = urllib.request.Request(url + path, data=json.dumps(data).encode() if data is not None else None, headers=kepala)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read() or b"{}") if "json" in r.headers.get("Content-Type", "") else r.read()
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


cek("ping tanpa token", minta("/api/ping", token="")[0], 200)
cek("api tanpa token ditolak", minta("/api/proses", token="")[0], 403)
cek("api token salah ditolak", minta("/api/proses", token="salah")[0], 403)
cek("Host asing ditolak", minta("/api/proses", host="penyerang.contoh:80")[0], 403)
cek("POST tanpa token ditolak", minta("/api/jalankan", {"alat": "bungkus_pc", "aksi": "daftar"}, token="")[0], 403)
s, j = minta("/api/proses")
cek("api dgn token", (s, "proses" in j), (200, True))
s, j = minta("/")
cek("halaman memuat token", b"token-uji" in j, True)
s, j = minta("/api/baca", {"path": "gui/server.py"})
cek("berkas .py tidak boleh dibaca", s, 400)
s, j = minta("/api/jalankan", {"alat": "otomatis", "aksi": "jalan", "isian": {
    "sumber": "a.xlsx", "dari": "2", "sampai": "3", "akun": "ppl.contoh@mail.com", "subsls": "5108010010000105"}})
cek("otomatis tanpa YA di GUI ditolak", (s, j.get("galat")), (400, "Aksi ini wajib dikonfirmasi dgn mengetik YA."))
s, j = minta("/api/pratinjau", {"alat": "bungkus_pc", "aksi": "daftar", "isian": {}})
cek("pratinjau perintah", j.get("perintah"), "python antar_pc/bungkus_pc.py --daftar")
srv.shutdown()

print(f"\n{'SEMUA UJI LULUS' if not gagal else f'{gagal} UJI GAGAL'}")
_sys.exit(1 if gagal else 0)
