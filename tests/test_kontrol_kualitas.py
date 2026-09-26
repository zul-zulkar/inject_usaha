#!/usr/bin/env python3
"""
test_kontrol_kualitas.py — Uji OFFLINE (tanpa browser/VPN) program kontrol kualitas
sumber data (input_usaha/kontrol_kualitas.py). Jalankan:
    python tests/test_kontrol_kualitas.py

Yang dikunci:
- laporan TIDAK punya aturan sendiri: baris berkategori DITOLAK persis baris yang
  ditolak pemeriksaan main_gabungan --cek;
- pesan pemeriksaan -> kolom Excel (sel) & kategori yang benar, format standar & tahap 2;
- SETIAP pola di POLA_TANDA punya contoh pesan di sini, dan setiap pesan yang muncul dari
  data contoh dikenali. Kalau teks pesan di inti/gabungan_loader.py / inti/tahap2_loader.py
  berubah, uji ini merah -> sesuaikan POLA_TANDA supaya laporan tetap menunjuk sel yang benar;
- berkas Excel: lembar lengkap, sel bermasalah diwarnai + komentar, nomor baris & huruf
  kolom salinan SAMA dgn sheet asli; CLI (kode keluar, --per-ppl, --csv).

Semua data di sini FIKTIF.
"""

import os
os.environ["FASIH_ABAIKAN_CONFIG_LOKAL"] = "1"   # uji pakai data contoh Buleleng di config.py

import sys as _sys, os as _os
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))

import csv
import json
import tempfile
import warnings
from pathlib import Path

import openpyxl

import input_usaha.kontrol_kualitas as kk
from inti.gabungan_loader import KEY_26, KEY_27, KEY_29, KEY_PEKERJA, GabunganRow, _cari_indeks, _sel

warnings.simplefilter("ignore")   # openpyxl: "Data Validation extension is not supported"

gagal = 0


def cek(nama, dapat, harap):
    global gagal
    if dapat == harap:
        print(f"  OK  {nama}")
    else:
        gagal += 1
        print(f"  !!  {nama}\n        dapat : {dapat!r}\n        harap : {harap!r}")


def cek_benar(nama, syarat):
    cek(nama, bool(syarat), True)


AKAR = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp())
PILIH = ("pilih_prov", "pilih_kab", "pilih_kec", "pilih_desa", "pilih_sls", "pilih_subsls")


def satu(lap, baris, jenis):
    """Temuan `jenis` di `baris` (harus tepat satu)."""
    t = [x for x in lap.temuan[baris] if x.jenis == jenis]
    return t[0] if len(t) == 1 else None


def sel(t):
    return [f"{kk.huruf_kolom(i)}{t.baris}" for i in t.kolom] if t else None


def konsisten(lap):
    """Kategori DITOLAK persis baris yang ditolak pemeriksaan main_gabungan."""
    return all((kk.status_qc(lap.temuan[r.baris]) == "DITOLAK") == (not lap.hasil[r.baris].bisa_diproses)
               for r in lap.dipilih)


# --------------------------------------------------------------------------
print("\n== pola TANDA -> kategori & rincian ==")
CONTOH_TANDA = [
    ("KOORDINAT BELUM ADA -> disimpan sbg DRAFT tanpa geotag, TIDAK dikirim",
     "DRAFT", "KOORDINAT_BELUM_ADA", ("latitude", "longitude")),
    ("12c umur kosong -> 45 (nilai pengganti, tidak ada di sheet)", "DIGANTI", "UMUR_PENGGANTI", ("umur",)),
    ("25 tahun operasi kosong -> 2019 (nilai pengganti, tidak ada di sheet); usaha lain pemilik ini berbeda-beda: "
     "['2010', '2015']", "DIGANTI", "TAHUN_OPERASI_PENGGANTI", ("tahun_operasi",)),
    ("12c umur '4' -> '45' (koreksi per baris, ketetapan user)", "DIGANTI", "UMUR_DIKOREKSI_PER_BARIS", ("umur",)),
    ("Nama Jalan kosong -> nama wilayah 'DESA CONTOH' (ketetapan user)", "DIGANTI", "JALAN_KOSONG_JADI_WILAYAH",
     ("jalan_domisili",)),
    ("13a kosong -> judul KBLI 'PERDAGANGAN ECERAN ROKOK DAN TEMBAKAU' (ketetapan user)", "DIGANTI",
     "13A_KOSONG_JADI_JUDUL_KBLI", ("keg_utama",)),
    ("12c umur kosong -> 50 (disalin dari usaha lain pemilik yang sama)", "DIKOREKSI", "UMUR_DARI_USAHA_LAIN", ("umur",)),
    ("25 tahun operasi kosong -> 2015 (disalin dari usaha lain pemilik yang sama)", "DIKOREKSI",
     "TAHUN_DARI_USAHA_LAIN", ("tahun_operasi",)),
    ("no WA '8,13E+10' tidak valid -> '9999' (tidak ada/tidak bersedia)", "DIGANTI", "HP_TIDAK_VALID_JADI_9999", ("hp",)),
    ("no WA '' kosong -> '9999' (tidak ada/tidak bersedia)", "INFO", "HP_KOSONG_JADI_9999", ("hp",)),
    ("no WA '81234567890' -> '081234567890' (nol di depan dikembalikan)", "INFO", "HP_NOL_DEPAN_DIKEMBALIKAN", ("hp",)),
    ("12d NIK '123' tidak valid (bukan 16 digit) -> '9999'", "DIGANTI", "NIK_TIDAK_VALID_JADI_9999", ("nik_pengusaha",)),
    ("12a kosong/'-' -> 'NI KETUT CONTOH' (nama dalam kurung di nama usaha)", "DIKOREKSI",
     "PENGUSAHA_DARI_NAMA_USAHA", ("pengusaha",)),
    ("12a kosong/'-' -> 'PEMILIK KIOS CONTOH' (tidak ada nama dalam kurung di nama usaha)", "DIGANTI",
     "PENGUSAHA_PENGGANTI", ("pengusaha",)),
    ("24 kosong semua -> 1 pekerja perempuan (ikut pemilik), tidak dibayar (nilai minimal)", "DIGANTI",
     "PEKERJA_KOSONG_JADI_1", KEY_PEKERJA),
    ("26a diisi 200,000 (100,000 x 2 pekerja dibayar) (form menolak 26a = 0)", "DIGANTI", "GAJI_DIISI_OTOMATIS",
     ("gaji", "tk_dibayar")),
    ("26f 0 (semua pos kosong/nol) -> 26d diisi minimal 100,000 (DINAIKKAN)", "DIGANTI", "TOTAL_NOL_JADI_MINIMAL", KEY_26),
    ("27c 0 (semua pos kosong/nol) -> 27a diisi minimal 100,000 (DINAIKKAN)", "DIGANTI", "TOTAL_NOL_JADI_MINIMAL", KEY_27),
    ("26f 60,000 < minimal 100,000 -> biaya_pembelian ditambah 40,000 (DINAIKKAN)", "DIGANTI",
     "TOTAL_DINAIKKAN_KE_MINIMAL", KEY_26),
    ("16b1-b6 '1' (satu kolom) -> Ya hanya utk b1,b4,b5, sisanya Tidak", "DIGANTI", "16B_SATU_KODE_YA",
     ("internet_semua",)),
    ("16b1-b6 '2,1,2,2,2' berisi 5 nilai -> dianggap b1-b5, b6 Lainnya = '2. Tidak'", "DIGANTI", "16B_LIMA_NILAI",
     ("internet_semua",)),
    ("16a Ya tapi 16b '2' tanpa Ya -> 16b6 Lainnya = Ya", "DIGANTI", "16B_TANPA_YA_JADI_B6",
     ("internet", "internet_semua")),
    ("27d kosong -> 10% krn 16b1 (menerima pesanan) = Ya", "DIGANTI", "27D_DIISI_KRN_PESANAN",
     ("pendapatan_online", "internet_semua")),
    ("KBLI 98100 kategori U ditolak form -> 13g diisi rekomendasi GenAI pertama saat pengisian", "DIGANTI",
     "KBLI_DIGANTI_GENAI", ("kbli",)),
    ("koordinat sheet '-8.148.438' / '1.145.951' (format rusak Excel) -> -8.148438, 114.5951", "DIKOREKSI",
     "KOORDINAT_DIPULIHKAN", ("latitude", "longitude")),
    ("idsubsls '5100090007000901' -> '5108090007000901' (awalan kabupaten salah ketik; kecamatan 090 cocok dgn "
     "kolom Sumber/Kec.)", "DIKOREKSI", "IDSUBSLS_AWALAN_DIBETULKAN", ("idsubsls",)),
    ("tk_tdk_dibayar kosong -> 0 (total 2 - tk_dibayar 2)", "DIKOREKSI", "PEKERJA_DARI_TOTAL",
     ("tk_tdk_dibayar", "cek_tk_bayar")),
    ("24 laki/perempuan 1/0 -> 0/1 (ikut jenis kelamin pemilik; total = dibayar 0 + tidak dibayar 1)", "DIKOREKSI",
     "PEKERJA_IKUT_JK_PEMILIK", ("tk_laki", "tk_pr")),
    ("26a 150,000 terisi tapi 24a2 = 0 -> 1 pekerja tidak dibayar dipindah ke dibayar (24a2 1, 24b2 0)", "DIKOREKSI",
     "PEKERJA_JADI_DIBAYAR", ("tk_dibayar", "tk_tdk_dibayar", "gaji")),
    ("KBLI 01464 tanpa 26c di form: 26c 1,000 dijumlahkan ke 26b (0 -> 1,000)", "DIKOREKSI", "26C_DIPINDAH_KE_26B",
     ("biaya_pembelian", "biaya_produksi")),
    ("30c 0 (usaha dagang, varian bulanan) -> 30b 1,000 dipindah ke 30c", "DIKOREKSI", "POS_PENGELUARAN_DIPINDAH",
     ("biaya_pembelian", "biaya_produksi")),
    ("26b 0 (KBLI 10710 wajib biaya produksi) -> 26d 5,000 dipindah ke 26b", "DIKOREKSI", "POS_PENGELUARAN_DIPINDAH",
     ("biaya_produksi", "operasional")),
    ("27d '12,5' dibulatkan -> 13", "DIKOREKSI", "27D_DIBULATKAN", ("pendapatan_online",)),
    ("Nama Jalan dilengkapi nama wilayah: 'BR. X' -> 'BR. X, DESA CONTOH'", "DIKOREKSI", "JALAN_DILENGKAPI",
     ("jalan_domisili",)),
    ("Nama Jalan dilengkapi nama wilayah: '0' -> 'DESA CONTOH' (desa idsubsls dipakai; desa menurut idsubsls "
     "'CONTOH, GEROKGAK' vs 8c 'LAIN, GEROKGAK')", "TINJAU", "WILAYAH_BENTROK_ALAMAT",
     ("jalan_domisili", "idsubsls", "alamat_usaha_view")),
    ("13a dilengkapi judul KBLI: 'JUAL BERAS' -> 'JUAL BERAS (PERDAGANGAN)'", "DIKOREKSI", "13A_DILENGKAPI",
     ("keg_utama",)),
    ("13a 'JUAL BERAS' < 15 karakter -> dilengkapi judul KBLI 47241 dari Master KBLI saat pengisian", "DIKOREKSI",
     "13A_DILENGKAPI", ("keg_utama",)),
    ("24 (laki, perempuan, dibayar, tidak dibayar) 0/2/0/1 -> 0/2/0/2 (ketetapan user)", "DIKOREKSI",
     "PEKERJA_POLA_DIKOREKSI", KEY_PEKERJA),
    ("11a '1. Perseroan Terbatas (PT)/CV' -> '7. Persekutuan Komanditer (CV)' dari awalan nama (ketetapan user)",
     "DIKOREKSI", "BADAN_USAHA_DARI_AWALAN", ("badan_usaha",)),
    ("BUMDES: 11a/11d/29 ('13. Bukan Badan Usaha', '2. Tidak', '100/0/0/0/0/0') -> ('6. BUM Desa', '1. Ya', "
     "pemerintah 100) (validasi form, ketetapan user)", "DIKOREKSI", "BUMDES_DIKOREKSI",
     ("badan_usaha", "lap_keuangan", *KEY_29)),
    ("nama usaha diganti -> 'PUSTU CONTOH' (ketetapan user)", "DIKOREKSI", "NAMA_USAHA_DIGANTI", ("nama",)),
    ("nama kembar persis (baris [10, 11]) & wilayahnya sama -> nama dibedakan penomoran '1': WARUNG (I KETUT "
     "CONTOH) 1", "TINJAU", "NAMA_KEMBAR_DINOMORI", ("nama_komersial",)),
    ("nama kembar persis (baris [10, 11]) -> nama dibedakan desa 'CONTOH': WARUNG CONTOH (I KETUT CONTOH)",
     "INFO", "NAMA_DIBEDAKAN", ("nama_komersial",)),
    ("usaha pecahan bernama sama (baris [10, 11]) -> nama dibedakan 13f 'sembako': warung sembako (i ketut contoh)",
     "INFO", "NAMA_DIBEDAKAN", ("nama_komersial",)),
    ("kolom total 26a+26b+26c+26d+26e di sheet = 5,000, jumlah rincian 6,000 — dipakai rinciannya (total dihitung "
     "form)", "TINJAU", "TOTAL_BEDA_DGN_RINCIAN", ("cek_26f", *KEY_26)),
    ("kolom total 24a1+24b1 di sheet = 0 padahal rinciannya 3 — dianggap tidak diisi (total dihitung form)",
     "INFO", "TOTAL_TIDAK_DIISI", ("cek_tk_gender",)),
    ("kolom total '########' (27a+27b) tidak terbaca — dipakai rinciannya", "INFO", "TOTAL_TIDAK_TERBACA_DIABAIKAN",
     ("cek_27c",)),
    ("KBLI 43213 'Pemasangan Sistem Elektronika' tidak berbagi satu kata pun dgn 13a 'SEWA SOUND' — periksa "
     "apakah KBLI-nya keliru", "TINJAU", "KBLI_TIDAK_NYAMBUNG", ("kbli", "keg_utama", "produk")),
    ("12b '1. Laki-laki' tapi 24a1 (pekerja laki-laki)=0 dari 24c1=2 — form PERNAH menolak pola ini", "TINJAU",
     "JK_VS_PEKERJA", ("jk", "tk_laki", "tk_pr")),
    ("kodepos 81155 beda dgn mayoritas desa 5108010008 (81154, {'81154': 5, '81155': 1})", "TINJAU",
     "KODEPOS_MINORITAS", ("kodepos",)),
    ("wilayah tujuan ubah alokasi ambigu: idsubsls=5108060014000302 tapi kolom Pilih PROVINSI..SUBSLS = "
     "5108060014000301", "TINJAU", "WILAYAH_PILIH_BEDA", ("idsubsls", *PILIH)),
    ("varian bulanan (mulai beroperasi 2026): rincian 30-33 diisi dari kolom 26-29", "TINJAU",
     "VARIAN_BULANAN_DARI_KOLOM", ("tahun_operasi",)),
    ("nama 'WARUNG' terkandung di nama dokumen baris 5 — bisa SKIP_DOKUMEN_NAMA_LAMA kalau dokumen baris itu "
     "dibuat lebih dulu", "TINJAU", "NAMA_TERMUAT_NAMA_LAIN", ("nama",)),
    ("Nama Jalan KOSONG (belum pernah diuji dikosongkan)", "TINJAU", "JALAN_KOSONG", ("jalan_domisili",)),
    ("indikator ekonomi kosong dianggap 0: gaji, pendapatan_lain", "INFO", "UANG_KOSONG_JADI_0",
     ("gaji", "pendapatan_lain")),
    ("16b1-b6 diisi '2. Tidak' dari satu kolom '16b1-b6'", "INFO", "16B_SATU_KODE", ("internet_semua",)),
    ("16b1-b6 'B1,B3' -> Ya utk b1,b3", "INFO", "16B_DIURAI", ("internet_semua",)),
    ("16b1-b6 dari daftar '1,2,1,1,1,1'", "INFO", "16B_DIURAI", ("internet_semua",)),
    ("13b1/b2/b3 diturunkan dari golongan KBLI 47: 2/2/1", "INFO", "13B_DARI_KBLI",
     ("produk_sendiri", "layanan_mamin", "keg_penjualan")),
    ("13d/13e diisi dari judul KBLI: 'Tepung' / 'Industri roti'", "INFO", "13DE_DARI_JUDUL_KBLI",
     ("input_produksi", "proses_produksi")),
    ("13c default '5. Kedai, stan, tenda' (usaha makan-minum KBLI 56102: form menolak kode 1-4; tidak ada di "
     "kuesioner tahap 2)", "INFO", "13C_DEFAULT_MAKAN_MINUM", ("lokasi_usaha",)),
    ("jenis_kawasan default '10. Di luar kawasan' (tidak ada di kuesioner tahap 2)", "INFO", "DEFAULT_TAHAP2",
     ("jenis_kawasan",)),
    ("kodepos 81155 dari daftar wilayah (tidak ada kolomnya di sheet)", "INFO", "KODEPOS_DARI_DAFTAR", ("kodepos",)),
    ("13f disalin dari 13a (sheet tidak punya kolom 13f)", "INFO", "13F_DISALIN_DARI_13A", ("produk",)),
    ("nama dokumen tanpa (12a): format lengkap > 50 karakter", "INFO", "NAMA_TANPA_12A", ("nama",)),
    ("8b tanpa (12a): format lengkap > 50 karakter", "INFO", "NAMA_TANPA_12A", ("nama_komersial",)),
    # 2026-09-26 (kontrol kualitas input_tahap2_23)
    ("12d NIK kosong -> '9999' (tidak ada/tidak bersedia)", "DIGANTI", "NIK_KOSONG_JADI_9999", ("nik_pengusaha",)),
    ("26a 300,000 / 8 pekerja dibayar = 37,500 <= Rp 50,000 -> 26a 800,000 (100,000 x 8) (DINAIKKAN)", "DIGANTI",
     "GAJI_PER_PEKERJA_DINAIKKAN", ("gaji", "tk_dibayar")),
    ("25 tahun operasi 2099 di masa depan -> 2025 (nilai pengganti)", "DIGANTI", "TAHUN_MASA_DEPAN_PENGGANTI",
     ("tahun_operasi",)),
    ("25 tahun operasi 2099 di masa depan -> 2010 (disalin dari usaha lain pemilik yang sama)", "DIKOREKSI",
     "TAHUN_MASA_DEPAN_DARI_USAHA_LAIN", ("tahun_operasi",)),
    ("27d 150 > 100 persen -> dibatasi 100", "DIGANTI", "27D_DIBATASI_100", ("pendapatan_online",)),
    ("24 dibayar/tidak dibayar 0 padahal laki+perempuan 2 -> 2 pekerja tidak dibayar (26a 0)", "DIKOREKSI",
     "PEKERJA_STATUS_DARI_JK", KEY_PEKERJA),
    ("nilai uang < 1.000 dianggap ribuan (dikali 1.000): biaya_pembelian 950 -> 950000, operasional 50 -> 50000",
     "DIKOREKSI", "UANG_RIBUAN_DIKALI_1000", ("biaya_pembelian", "operasional")),
    ("kodepos desa 5108080099 tidak ada di daftar -> 81172 (semua desa lain di kecamatan 080 berkodepos sama)",
     "TINJAU", "KODEPOS_DARI_KECAMATAN", ("kodepos",)),
    ("nama dokumen 'WARUNG CONTOH' termuat di nama dokumen baris 12 'WARUNG CONTOH (BU MADE)' -> nama dibedakan "
     "dgn tetap memuat 12a: WARUNG CONTOH (I KETUT CONTOH)", "INFO", "NAMA_TERMUAT_DIBEDAKAN", ("nama_komersial",)),
]
diuji = set()
for teks, kategori, jenis, keys in CONTOH_TANDA:
    pola, dapat = kk.klasifikasi_tanda(teks)
    cek(f"{jenis}: {teks[:60]}", (pola and pola.kategori, pola and pola.jenis, dapat), (kategori, jenis, keys))
    diuji.add(jenis)

ada_13f, kosong_13f = GabunganRow(2, {"produk": "GAS"}), GabunganRow(2, {"keg_utama": "GAS"})
pola, keys = kk.klasifikasi_tanda("13f dilengkapi: 'GAS' -> 'GAS ECERAN'", ada_13f)
cek("13f dilengkapi -> kolom 13f kalau terisi", (pola.kategori, keys), ("DIKOREKSI", ("produk",)))
cek("13f dilengkapi -> kolom 13a kalau 13f kosong (disalin dari 13a)",
    kk.klasifikasi_tanda("13f dilengkapi: 'GAS' -> 'GAS ECERAN'", kosong_13f)[1], ("keg_utama",))
diuji.add(pola.jenis)
cek("SETIAP pola POLA_TANDA punya contoh pesan di uji ini", sorted({p.jenis for p in kk.POLA_TANDA} - diuji), [])
cek("pola: jenis unik", len({p.jenis for p in kk.POLA_TANDA}), len(kk.POLA_TANDA))
cek_benar("pola: kategori sah", all(p.kategori in kk.KATEGORI for p in kk.POLA_TANDA))
pola, keys = kk.klasifikasi_tanda("pesan baru yang belum dikenal: umur=5, 26a terlalu besar")
cek("pesan tak dikenal -> None + rincian yg jelas disebut", (pola, keys), (None, ("umur", "gaji")))

print("\n== kode MASALAH -> rincian & saran ==")
cek("WAJIB_KOSONG", kk.keys_masalah("WAJIB_KOSONG", "kolom kosong: hp, umur, produk (13f)"), ("hp", "umur", "produk"))
cek("ANGKA_TIDAK_VALID daftar", kk.keys_masalah("ANGKA_TIDAK_VALID", "kolom: umur, latitude/longitude (di luar "
                                                "Indonesia)"), ("umur", "latitude", "longitude"))
cek("ANGKA_TIDAK_VALID 27d", kk.keys_masalah("ANGKA_TIDAK_VALID", "27d=150 > 100 persen"), ("pendapatan_online",))
cek("ANGKA_TIDAK_VALID tahun", kk.keys_masalah("ANGKA_TIDAK_VALID", "tahun_operasi=2030 (di masa depan)"),
    ("tahun_operasi",))
cek("OPSI_TIDAK_ADA_DI_FORM", kk.keys_masalah("OPSI_TIDAK_ADA_DI_FORM", "jk='Laki' bukan salah satu opsi form"),
    ("jk",))
cek("NILAI_TIDAK_DIDUKUNG", kk.keys_masalah("NILAI_TIDAK_DIDUKUNG", "ubah_sls='1. Ya' (alur skrip hanya ...)"),
    ("ubah_sls",))
cek("DI_BAWAH_MINIMAL 27c", kk.keys_masalah("DI_BAWAH_MINIMAL", "27c=5000 < 100000"), KEY_27)
cek("DI_BAWAH_MINIMAL 30f (bulanan)", kk.keys_masalah("DI_BAWAH_MINIMAL", "30f=5000 < 10000"), KEY_26)
cek("TOTAL_TIDAK_COCOK", kk.keys_masalah("TOTAL_TIDAK_COCOK", "kolom total 27a+27b di sheet = 1,000 tapi jumlah "
                                         "rinciannya 2,000 — perbaiki di Excel"), ("cek_27c", *KEY_27))
cek("TOTAL_TIDAK_TERBACA", kk.keys_masalah("TOTAL_TIDAK_TERBACA", "kolom total 'abc' (28a+28b) tidak terbaca sbg "
                                           "angka"), ("cek_28c", "aset_usaha_thn", "aset_lain_thn"))
cek("UMUR_DI_LUAR_10_99", kk.keys_masalah("UMUR_DI_LUAR_10_99", "12c umur=5 (form: wajib 10-99)"), ("umur",))
cek("kode baru tak dikenal -> rincian yg disebut", kk.keys_masalah("KODE_BARU", "umur=5 aneh"), ("umur",))
cek_benar("saran opsi menyebut opsi form yang sah",
          "1. Laki-laki" in kk.saran_masalah("OPSI_TIDAK_ADA_DI_FORM", "jk='Laki' bukan salah satu opsi form"))
cek("nomor baris ringkas (format --baris)", kk.nomor_ringkas([9, 2, 3, 4, 7, 3]), "2-4,7,9")
cek("huruf kolom", [kk.huruf_kolom(i) for i in (0, 25, 26, 51, 52)], ["A", "Z", "AA", "AZ", "BA"])

# --------------------------------------------------------------------------
print("\n== format lama Agenda (dasar: tests/fixture_format_agenda.json, data fiktif) ==")
_fixture = json.loads((AKAR / "tests" / "fixture_format_agenda.json").read_text(encoding="utf-8"))
JUDUL_STD, DAGANG = _fixture["judul"], _fixture["contoh_dagang"]
IDX = _cari_indeks(JUDUL_STD)


def std(nama, **ubah):
    b = list(DAGANG)
    b[IDX["nama"]] = b[IDX["nama_komersial"]] = nama
    for key, nilai in ubah.items():
        b[IDX[key]] = nilai
    return b


BARIS_STD = [
    std("TOKO ALFA"),                                                      # 2 bersih
    std("TOKO BETA", umur="5"),                                            # 3 ditolak umur
    std("TOKO GAMMA", hp="12345"),                                         # 4 ditolak hp
    std("TOKO DELTA", jk="Laki"),                                          # 5 ditolak opsi
    std("TOKO EPSILON", kodepos=""),                                       # 6 ditolak wajib kosong
    std("CV. ZETA MAKMUR", badan_usaha="1. Perseroan Terbatas (PT)/CV"),   # 7 dikoreksi 11a
    std("TOKO ETA"), std("TOKO ETA"),                                      # 8, 9 baris ganda
    std("TOKO THETA", pilih_subsls="03"),                                  # 10 tinjau wilayah
    std("TOKO IOTA", tk_laki="0", tk_pr="2", tk_dibayar="0", tk_tdk_dibayar="1", gaji="0"),   # 11 koreksi 24
    std("TOKO KAPPA", keg_utama="JUAL SNACK"),                             # 12 13a dilengkapi
]
SUMBER_STD = TMP / "input_usaha.csv"
with SUMBER_STD.open("w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows([JUDUL_STD, *BARIS_STD])

lap = kk.periksa_sumber(str(SUMBER_STD), "standar", audit=[])
cek("semua baris diperiksa", [r.baris for r in lap.dipilih], list(range(2, 2 + len(BARIS_STD))))
cek_benar("DITOLAK == ditolak pemeriksaan main_gabungan (tidak ada aturan kedua)", konsisten(lap))
cek("semua pesan dikenali", [t.pesan for t in lap.semua_temuan(dgn_seragam=True) if t.jenis == kk.JENIS_TANDA_LAIN], [])
cek("baris 2 dasar bersih", kk.status_qc(lap.temuan[2]), "BERSIH")


def kolom_std(key, baris):
    return [f"{kk.huruf_kolom(IDX[key])}{baris}"]


cek("umur 5 -> DITOLAK di sel 12c", sel(satu(lap, 3, "UMUR_DI_LUAR_10_99")), kolom_std("umur", 3))
cek("hp -> DITOLAK di sel no HP", sel(satu(lap, 4, "HP_TIDAK_VALID")), kolom_std("hp", 4))
t = satu(lap, 5, "OPSI_TIDAK_ADA_DI_FORM")
cek("opsi jk -> DITOLAK di sel 12b", sel(t), kolom_std("jk", 5))
cek_benar("saran opsi jk menyebut '2. Perempuan'", t and "2. Perempuan" in t.saran)
cek("kodepos kosong -> WAJIB_KOSONG di sel kodepos", sel(satu(lap, 6, "WAJIB_KOSONG")), kolom_std("kodepos", 6))
t = satu(lap, 7, "BADAN_USAHA_DARI_AWALAN")
cek("11a PT/CV -> DIKOREKSI di sel 11a", (t and t.kategori, sel(t)), ("DIKOREKSI", kolom_std("badan_usaha", 7)))
cek_benar("nilai dipakai skrip = opsi CV", t and "7. Persekutuan Komanditer (CV)" in kk.nilai_skrip(lap.baris(7), t.keys))
cek("baris ganda -> kedua baris DITOLAK", [kk.status_qc(lap.temuan[b]) for b in (8, 9)], ["DITOLAK", "DITOLAK"])
cek_benar("baris ganda menunjuk sel nama", kolom_std("nama", 8)[0] in (sel(satu(lap, 8, "BARIS_GANDA")) or []))
t = satu(lap, 10, "WILAYAH_PILIH_BEDA")
cek("pilih subsls beda -> TINJAU (mode satu subsls)", t and t.kategori, "TINJAU")
cek_benar("... menunjuk sel idsubsls & Pilih SUBSLS",
          t and {kolom_std("idsubsls", 10)[0], kolom_std("pilih_subsls", 10)[0]} <= set(sel(t)))
cek("24 pola 0/2/0/1 -> DIKOREKSI 4 sel", sel(satu(lap, 11, "PEKERJA_POLA_DIKOREKSI")),
    [f"{kk.huruf_kolom(IDX[k])}11" for k in KEY_PEKERJA])
cek("13a pendek -> DIKOREKSI di sel 13a", sel(satu(lap, 12, "13A_DILENGKAPI")), kolom_std("keg_utama", 12))

lap_pb = kk.periksa_sumber(str(SUMBER_STD), "standar", mode_satu_subsls=False, audit=[])
cek("mode per baris: pilih subsls beda -> DITOLAK WILAYAH_TIDAK_KONSISTEN",
    satu(lap_pb, 10, "WILAYAH_TIDAK_KONSISTEN") and satu(lap_pb, 10, "WILAYAH_TIDAK_KONSISTEN").kategori, "DITOLAK")
cek_benar("mode per baris tetap konsisten dgn main_gabungan", konsisten(lap_pb))

lap_sebagian = kk.periksa_sumber(str(SUMBER_STD), "standar", dari=3, sampai=5, audit=[])
cek("--dari/--sampai hanya melaporkan rentang", [r.baris for r in lap_sebagian.dipilih], [3, 4, 5])
cek("... tapi baris ganda tetap diperiksa lintas SELURUH sheet",
    kk.periksa_sumber(str(SUMBER_STD), "standar", baris="8", audit=[]).hasil[8].status, "SKIP_DATA_BARIS_GANDA")

print("\n== berkas Excel ==")
KELUAR_STD = TMP / "qc_std.xlsx"
kk.tulis_excel(lap, KELUAR_STD)
wb = openpyxl.load_workbook(KELUAR_STD)
cek("lembar", wb.sheetnames, ["Ringkasan", "Temuan", "Per PPL", "Per kolom", "Data bertanda"])
ws = wb["Temuan"]
cek("judul lembar Temuan", [c.value for c in ws[1]], list(kk.KOLOM_TEMUAN))
cek("jumlah baris Temuan = temuan (tanpa INFO seragam)", ws.max_row - 1, len(lap.semua_temuan()))
cek("Temuan urut keparahan: baris pertama DITOLAK", ws.cell(2, 5).value, "DITOLAK")
tautan = next((c.hyperlink for c in ws["C"][1:] if c.hyperlink), None)
cek_benar("sel Temuan bertaut ke lembar 'Data bertanda'", tautan and tautan.location.startswith("'Data bertanda'!"))
d = wb["Data bertanda"]
cek("salinan: judul kolom sama dgn sheet asli", [c.value for c in d[1]][:len(JUDUL_STD)], JUDUL_STD)
umur = d.cell(3, IDX["umur"] + 1)
cek("salinan: nomor baris & huruf kolom sama (sel umur baris 3)", (umur.coordinate, umur.value),
    (kolom_std("umur", 3)[0], "5"))
cek("salinan: sel DITOLAK diwarnai", umur.fill.start_color.rgb[-6:], kk.WARNA["DITOLAK"])
cek_benar("salinan: komentar berisi keterangan & saran",
          umur.comment and "[DITOLAK]" in umur.comment.text and "->" in umur.comment.text)
lebar = len(JUDUL_STD)
cek("salinan: kolom QC_STATUS di paling kanan", (d.cell(1, lebar + 1).value, d.cell(3, lebar + 1).value),
    ("QC_STATUS", "DITOLAK"))
cek("salinan: baris bersih tidak diberi komentar", any(c.comment for c in d[2][:lebar]), False)
per_ppl = [r[0] for r in wb["Per PPL"].iter_rows(min_row=2, values_only=True)]
cek("Per PPL: satu PPL (akun contoh)", per_ppl, ["ppl.contoh@gmail.com"])
wb.close()

# --------------------------------------------------------------------------
print("\n== format tahap 2 ==")
JUDUL_T2 = ["Sumber/Kec.", "Periode", "Uraian:", "Nama PPL", "3", "4", "5", "8b.", "8c.", "no WA",
            "12a", "12b", "12c", "12d", "13a", "13f", "14a", "16a", "16b1-b6", "17b", "21", "22",
            "24.L", "24.P", "24.Total", "24.Dibayar", "24.Tidak dibayar", "24.Total", "25", "Rp26",
            "26a", "26b", "26c", "26d", "26e", "27a", "27b", "27c", "27d", "28a", "28b", "28c",
            "28c1", "28d", "Latitude", "Longitude", "Kode KBLI", "Judul KBLI"]
DASAR_T2 = ["010", "21Sep", "Responden 1", "I KETUT CONTOH", "GEROKGAK 510801", "PATAS 0010",
            "5108010008000101", "USAHA JUAL BERAS", "JALAN CONTOH NOMOR SATU", "81234567890",
            "NI LUH CONTOH", "1", "34", "9999", "MENJUAL BERAS ECERAN", "BERAS ECERAN (47241)",
            "1", "2", "2", "1", "2", "5",
            "1", "2", "3", "0", "3", "3", "2025", "Rp10.700.000",
            "Rp0", "Rp0", "Rp10.500.000", "Rp200.000", "Rp0", "Rp11.800.000", "Rp0", "Rp11.800.000",
            "0,00", "Rp1.000.000", "Rp2.000.000", "Rp3.000.000", "-", "4",
            "-8,2004731", "114,7987732", "47241", ""]


def t2(usaha, pemilik, **ubah):
    b = list(DASAR_T2)
    b[JUDUL_T2.index("8b.")], b[JUDUL_T2.index("12a")] = usaha, pemilik
    for judul, nilai in ubah.items():
        b[JUDUL_T2.index(judul)] = nilai
    return b


BARIS_T2 = [
    t2("USAHA JUAL BERAS", "NI LUH CONTOH"),                                      # 2 dasar
    t2("WARUNG SATU", "I KETUT SATU", **{"12c": ""}),                             # 3 umur kosong
    t2("WARUNG DUA", "I MADE DUA", **{"no WA": "8,13E+10"}),                      # 4 HP notasi ilmiah
    t2("WARUNG TIGA", "NI LUH TIGA", Latitude="", Longitude=""),                  # 5 tanpa koordinat
    t2("WARUNG EMPAT", "I WAYAN EMPAT", Latitude="-8.148.438", Longitude="1.145.951"),   # 6 koordinat rusak
    t2("WARUNG LIMA", "I NYOMAN LIMA", **{"16a": "1", "16b1-b6": "1"}),          # 7 16b satu kode
    t2("WARUNG ENAM", "KADEK ENAM", **{"12b": "7"}),                              # 8 kode 12b tak ada
    t2("WARUNG TUJUH", "KOMANG TUJUH", Rp26="Rp9.000.000"),                       # 9 total beda
    t2("WARUNG DELAPAN", "GEDE DELAPAN", **{"Kode KBLI": "98100"}),               # 10 KBLI kategori U
    t2("WARUNG SEMBILAN", "PUTU SEMBILAN", **{"12d": "123"}),                     # 11 NIK tidak valid
]
SUMBER_T2 = TMP / "input_tahap2.csv"
with SUMBER_T2.open("w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows([JUDUL_T2, *BARIS_T2])


def kolom_t2(judul, baris):
    return f"{kk.huruf_kolom(JUDUL_T2.index(judul))}{baris}"


kunci_2 = None
lap2 = kk.periksa_sumber(str(SUMBER_T2), "tahap2", audit=[])
kunci_2 = lap2.baris(2).kunci
cek_benar("tahap 2: DITOLAK == ditolak pemeriksaan main_gabungan", konsisten(lap2))
cek("tahap 2: semua pesan dikenali",
    [t.pesan for t in lap2.semua_temuan(dgn_seragam=True) if t.jenis == kk.JENIS_TANDA_LAIN], [])
t = satu(lap2, 3, "UMUR_PENGGANTI")
cek("umur kosong -> DIGANTI di sel 12c", (t and t.kategori, sel(t)), ("DIGANTI", [kolom_t2("12c", 3)]))
cek_benar("... nilai dipakai skrip = nilai pengganti config", t and "12c: 45" in kk.nilai_skrip(lap2.baris(3), t.keys))
cek("HP notasi ilmiah -> DIGANTI di sel no WA", sel(satu(lap2, 4, "HP_TIDAK_VALID_JADI_9999")), [kolom_t2("no WA", 4)])
t = satu(lap2, 5, "KOORDINAT_BELUM_ADA")
cek("tanpa koordinat -> DRAFT di sel Latitude & Longitude", (t and t.kategori, sel(t)),
    ("DRAFT", [kolom_t2("Latitude", 5), kolom_t2("Longitude", 5)]))
cek("... status input SIAP_TANPA_KOORDINAT", lap2.hasil[5].status, "SIAP_TANPA_KOORDINAT")
cek("koordinat rusak -> DIKOREKSI", (satu(lap2, 6, "KOORDINAT_DIPULIHKAN") or kk.Temuan(0, "", "", "")).kategori,
    "DIKOREKSI")
t = satu(lap2, 7, "16B_SATU_KODE_YA")
cek("16b satu kode Ya -> DIGANTI di sel 16b1-b6", (t and t.kategori, sel(t)), ("DIGANTI", [kolom_t2("16b1-b6", 7)]))
cek_benar("... nilai dipakai skrip merinci b1-b6", t and "16b1-b6: Ya,Tidak,Tidak,Ya,Ya,Tidak" in
          kk.nilai_skrip(lap2.baris(7), t.keys))
cek_benar("27d 0 + 16b1 Ya -> DIGANTI 27d", satu(lap2, 7, "27D_DIISI_KRN_PESANAN"))
cek("kode 12b tak dikenal -> DITOLAK di sel 12b", (kk.status_qc(lap2.temuan[8]), sel(satu(lap2, 8, "WAJIB_KOSONG"))),
    ("DITOLAK", [kolom_t2("12b", 8)]))
t = satu(lap2, 9, "TOTAL_BEDA_DGN_RINCIAN")
cek_benar("total Rp26 beda -> TINJAU menunjuk sel Rp26", t and t.kategori == "TINJAU" and kolom_t2("Rp26", 9) in sel(t))
cek("KBLI kategori U -> DIGANTI (GenAI) di sel Kode KBLI", sel(satu(lap2, 10, "KBLI_DIGANTI_GENAI")),
    [kolom_t2("Kode KBLI", 10)])
cek("NIK tidak valid -> DIGANTI di sel 12d", sel(satu(lap2, 11, "NIK_TIDAK_VALID_JADI_9999")), [kolom_t2("12d", 11)])
cek("default rincian tanpa kolom -> diringkas, tidak per baris",
    any(t.jenis == "DEFAULT_TAHAP2" for t in lap2.semua_temuan()), False)
cek_benar("... tapi tercatat di rekap asumsi seragam", any(g["jenis"] == "DEFAULT_TAHAP2" for g in kk.rekap_seragam(lap2)))
cek("PPL tahap 2 = kolom Nama PPL", kk.ppl_baris(lap2.baris(2)), "I KETUT CONTOH")

lap_audit = kk.periksa_sumber(str(SUMBER_T2), "tahap2", hanya_belum_terkirim=True,
                              audit=[{"kunci": kunci_2, "status": "TERKIRIM_TERVERIFIKASI"}])
cek("--hanya-belum-terkirim melewati baris terkirim di audit", 2 in [r.baris for r in lap_audit.dipilih], False)
isi = kk.baris_temuan(lap_audit, [t for t in lap_audit.semua_temuan() if t.baris == 3])
cek_benar("kolom Status audit terisi dari audit", all(len(b) == len(kk.KOLOM_TEMUAN) for b in isi))

# --------------------------------------------------------------------------
print("\n== CLI ==")
awal = os.getcwd()
os.chdir(TMP)   # audit_log_gabungan.csv di folder kerja proyek tidak ikut terbaca
try:
    keluar = TMP / "cli.xlsx"
    cek("ada baris ditolak -> kode keluar 1",
        kk.main(["--sumber", str(SUMBER_STD), "--format", "agenda", "--audit", str(TMP / "audit_cli.csv"),
                 "--keluaran", str(keluar), "--tanpa-salinan"]), 1)
    cek("--tanpa-salinan: tanpa lembar Data bertanda",
        "Data bertanda" in openpyxl.load_workbook(keluar, read_only=True).sheetnames, False)
    cek("hanya baris bersih -> kode keluar 0",
        kk.main(["--sumber", str(SUMBER_STD), "--format", "standar", "--baris", "2", "--keluaran", str(keluar)]), 0)
    folder = TMP / "per_ppl"
    csv_path = TMP / "temuan.csv"
    cek("tahap 2 + --per-ppl + --csv -> kode keluar 1 (baris 8 ditolak)",
        kk.main(["--sumber", str(SUMBER_T2), "--format", "tahap2", "--keluaran", str(keluar),
                 "--per-ppl", str(folder), "--csv", str(csv_path)]), 1)
    cek("berkas per PPL", sorted(p.name for p in folder.glob("*.xlsx")), ["I_KETUT_CONTOH.xlsx"])
    with csv_path.open(encoding="utf-8-sig") as f:
        cek("CSV temuan: judul kolom", next(csv.reader(f)), list(kk.KOLOM_TEMUAN))
    cek("berkas tidak ada -> kode keluar 2", kk.main(["--sumber", str(TMP / "tidak_ada.xlsx")]), 2)
    cek("sheet format lama dibaca sbg format bawaan (tahap2) -> kode keluar 2",
        kk.main(["--sumber", str(SUMBER_STD), "--keluaran", str(keluar)]), 2)
    cek("--dari > --sampai ditolak", kk.main(["--sumber", str(SUMBER_STD), "--format", "agenda", "--dari", "9", "--sampai", "3"]), 2)
finally:
    os.chdir(awal)

print(f"\n{'SEMUA LULUS' if not gagal else f'{gagal} GAGAL'}")
if gagal:
    print("FAIL")
sys_exit = 1 if gagal else 0
_sys.exit(sys_exit)
