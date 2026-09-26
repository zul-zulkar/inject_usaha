#!/usr/bin/env python3
"""
ubah_moda.py — Ganti mode assignment CAPI -> PAPI di fasih-sm untuk setiap
idsubsls di input_usaha.xlsx (tab input_usaha). Jalankan SEBELUM input_usaha/jalankan.py.

KENAPA
------
fasih-web hanya bisa "+ Dokumen Baru" di subsls yang sudah punya assignment
(gejala: SKIP_DOKUMEN_BELUM_ADA). Mengubah assignment CAPI di subsls itu
menjadi PAPI membuka jalannya.

YANG SUDAH DILIHAT LANGSUNG di fasih-sm (2026-09-13)
---------------------------------------------------
- List: /app/surveys/{SURVEY_ID}/{ASSIGNMENT_ID_GABUNGAN}/data?page=1&perPage=10
- Kolom tabel: "Kode Identitas" (format "5108060003000402 - UMK - 4" =
  idsubsls - jenis - nomor), "Status", "Mode" (CAPI/PAPI), "Petugas Saat Ini"
  (email), "Keterangan" (Pencacah/Pengawas). Tiap baris: checkbox "Pilih baris".
- Tombol "Aksi Lainnya" -> menu berisi item MASSAL "Ganti Mode (Ke PAPI) (N)",
  N = jumlah baris tercentang. Diklik saat N=0: tidak ada dialog/notifikasi.
- ⚠️ Menu yang masih terbuka MENELAN klik berikutnya: saat dipetakan, klik ke
  ikon filter justru mengenai item "Broadcast Status" di menu yang belum
  tertutup. Karena itu setiap menu WAJIB terbukti tertutup (tutup_menu())
  sebelum klik apa pun berikutnya — jangan hapus penjagaan ini.

BELUM DIPASTIKAN (makanya ada --petakan & dry-run)
-------------------------------------------------
- apakah kotak "Cari..." menyaring per idsubsls;
- apakah "Ganti Mode (Ke PAPI)" dgn N>=1 memunculkan dialog konfirmasi, dan
  apa teks tombolnya.
Semua pengecekan GAGAL-TERTUTUP: hasil pencarian yang memuat subsls lain,
angka N di menu yang tidak sama dgn yang dicentang, dialog yang tidak
menyebut PAPI/mode, tombol konfirmasi yang ambigu, atau hasil yang tidak
terverifikasi -> berhenti. Tidak ada yang ditebak.

LANGKAH
-------
1. Daftar target, tanpa browser:
       python ubah_moda.py --sumber input_usaha.xlsx --cek
2. Pemetaan 1 subsls (login MANUAL sekali; profil disimpan di .profil_fasih_sm/).
   Tidak mencentang & tidak mengubah apa pun; tabel/menu/API direkam ke log_fasih_sm/:
       python ubah_moda.py --sumber input_usaha.xlsx --petakan
3. Dry-run: cari, rencanakan, centang, cocokkan angka N di menu, lalu lepas
   centang. TIDAK mengklik "Ganti Mode":
       python ubah_moda.py --sumber input_usaha.xlsx --limit 5
4. LIVE untuk SATU subsls dulu (wajib ketik YA), cek hasilnya di fasih-sm & fasih-web:
       python ubah_moda.py --sumber input_usaha.xlsx --idsubsls 5108070013000901 --eksekusi
5. Sisanya:
       python ubah_moda.py --sumber input_usaha.xlsx --eksekusi --lewati-selesai

Skrip TIDAK PERNAH mengetik kata sandi fasih-sm — login dilakukan manusia di
jendela browser. Headless sengaja tidak disediakan (fasih-sm memakai
proteksi anti-bot F5/TSPD, sama seperti fasih-web).

⚠️ JALUR YANG DISARANKAN: Console Chrome biasa (ubah_moda_console.js), bukan
Playwright — fasih-sm mendeteksi browser otomatis. File siap-tempel dibuat dgn:
       python ubah_moda.py --sumber input_usaha.xlsx --console

LIST KODE IDENTITAS MILIK SENDIRI (2026-09-15)
---------------------------------------------
Ganti --sumber dgn --daftar: file .xlsx/.csv/.txt berisi kode identitas
("5108060003000402 - UMK - 4") di kolom/baris mana pun. HANYA kode itu yang
diubah — skrip tidak memilih sendiri & --cakupan diabaikan. Tiap kode diproses
sendiri: KODE itu diketik di kotak "Cari...", baris yang kodenya PERSIS sama
dicentang ("…- UMK - 4" bukan "…- UMK - 41"), "Ganti Mode (Ke PAPI) (1)", lalu
kode itu dicari ulang utk verifikasi. Sudah PAPI -> KODE_SUDAH_PAPI; tidak ada
-> KODE_TIDAK_ADA; hasil pencarian berisi subsls lain -> PENCARIAN_TIDAK_MENYARING
(berhenti).
       python ubah_moda.py --daftar list.xlsx --cek
       python ubah_moda.py --daftar list.xlsx --console
Jalur Playwright di file ini tetap ada sbg cadangan & sumber logika bersama.

ARAH BALIK: PAPI -> CAPI UNTUK SUBSLS TERTENTU (2026-09-22)
----------------------------------------------------------
--subsls <file/daftar koma> --ke CAPI: tiap subsls dicari, SEMUA assignment PAPI
subsls itu di halaman tampil dicentang -> "Ganti Mode (Ke CAPI) (N)", ditunggu
sampai terbaca CAPI, lalu dicari lagi sampai tidak ada PAPI tersisa (TIDAK_ADA_PAPI).
Hasil >1 halaman tanpa PAPI yang tampil -> CEK_HALAMAN_LAIN (paginasi tidak
dipindah; sisanya lewat --daftar kode identitas --ke CAPI). --ke WAJIB ditulis utk
--subsls; --sumber hanya ke PAPI. Arah ikut tertulis di file siap-tempel, jadi
Console tidak bisa jalan ke arah lain. Kunci hasil/audit arah CAPI berawalan "CAPI:".
       python fasih_sm/ganti_moda/ubah_moda.py --subsls daftar_subsls.txt --ke CAPI --cek
       python fasih_sm/ganti_moda/ubah_moda.py --subsls daftar_subsls.txt --ke CAPI --console
Ke CAPI dikerjakan PER BARIS lewat menu ⋮ (cara manual user): ⋮ -> "Ganti Mode" ->
dialog "Mode Pendataan" -> pilih CAPI -> "Ubah Mode Pendataan". Struktur itu dilihat
langsung 2026-09-22 dan HANYA diimplementasikan di ubah_moda_console.js; jalur
Playwright file ini menolak --ke CAPI selain --cek/--console. Mulai dari petakan ->
dryrun -> otomatis limit:1.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
import time
import traceback
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from playwright.sync_api import TimeoutError as PWTimeout

# -- jalankan dari root proyek; sisipkan root ke sys.path utk paket inti/ dll --
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))))
from inti import lokasi
from inti.config import ASSIGNMENT_ID_GABUNGAN, FASIH_SM_BASE, SURVEY_ID
from inti.gabungan_loader import load_gabungan, periksa_semua

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

HASIL = lokasi.hasil("fasih_sm", "ganti_moda")
PROFIL_DIR = HASIL / ".profil_fasih_sm"
LOG_DIR = HASIL / "log_fasih_sm"
AUDIT_PATH = HASIL / "audit_ubah_moda.csv"
TARGET_CSV = HASIL / "target_ubah_moda.csv"
KONSOL_TEMPLATE = Path(__file__).resolve().parent / "ubah_moda_console.js"
KONSOL_SIAP = HASIL / "ubah_moda_console.siap.js"
PENANDA_TARGET = "/*__TARGET__*/[]"
AUDIT_FIELDS = [
    "timestamp", "jalan", "idsubsls", "akun_ppl", "baris_sheet", "status",
    "jumlah_assignment", "capi", "papi", "dipilih", "petugas_dipilih", "pesan",
    "kode_target",  # di BELAKANG: baris audit lama tetap terbaca benar
    "ke",           # mode tujuan (kosong di baris lama = PAPI)
]

KOLOM_TABEL = {
    "kode": "Kode Identitas", "status": "Status", "mode": "Mode",
    "petugas": "Petugas Saat Ini", "keterangan": "Keterangan",
}
MODE = ("CAPI", "PAPI")
POLA_ANGKA_ITEM = re.compile(r"\(\s*(\d+)\s*\)\s*$")
POLA_TOMBOL_KONFIRMASI = re.compile(
    r"^\s*(ya|konfirmasi|ganti|ubah|lanjut|lanjutkan|simpan|ok|oke|proses)\b", re.I)
POLA_TOMBOL_BATAL = re.compile(r"batal|tutup|cancel|kembali|^\s*tidak\b", re.I)


def lawan_mode(m: str) -> str:
    return "PAPI" if m == "CAPI" else "CAPI"


def pola_item_ganti_mode(ke: str) -> re.Pattern:
    """Item massal menu "Aksi Lainnya": "Ganti Mode (Ke PAPI) (N)" / "Ganti Mode (Ke CAPI) (N)"."""
    return re.compile(rf"Ganti Mode\s*\(\s*Ke {ke}\s*\)", re.I)


STATUS_TUNTAS_LIVE = {"DIUBAH_TERVERIFIKASI", "SUDAH_ADA_PAPI", "TIDAK_ADA_CAPI", "KODE_SUDAH_PAPI",
                      "KODE_SUDAH_CAPI", "TIDAK_ADA_PAPI"}
STATUS_TUNTAS_DRY = STATUS_TUNTAS_LIVE | {"DRY_RUN_AKAN_DIUBAH"}
# Status yang membuktikan cara kerja skrip tidak cocok dgn halaman — batch
# berhenti SEKETIKA, bukan lanjut ke subsls berikutnya dgn asumsi yang sama.
STATUS_BERHENTI_SEGERA = {
    "SUBSLS_TIDAK_TAMPIL", "PERLU_HALAMAN_LAIN", "KOLOM_TIDAK_ADA",
    "JUMLAH_TERCENTANG_BEDA", "MENU_TIDAK_TERTUTUP", "TABEL_BERUBAH", "DIALOG_TIDAK_DIKENAL",
    "TOMBOL_KONFIRMASI_AMBIGU", "DIUBAH_BELUM_TERVERIFIKASI", "CENTANG_GAGAL",
    "ITEM_MENU_TIDAK_ADA", "BELUM_LOGIN", "TIDAK_ADA_AKSES",
    "PENCARIAN_TIDAK_MENYARING", "KODE_GANDA", "RATE_LIMIT",
}

# Perubahan mode TIDAK langsung terbaca di tabel (run user 2026-09-15: 3x cari ulang
# dlm ±10 dtk setelah konfirmasi tetap CAPI, dan cek beruntun itu memicu HTTP 429 di
# datatable). Jadwal cek ulang sejak diklik & batas tunggunya — HARUS sama dgn
# ubah_moda_console.js.
JADWAL_CEK_VERIFIKASI_MS = (30_000, 45_000, 60_000, 90_000, 120_000)
JEDA_CEK_VERIFIKASI_MAKS_MS = 180_000
BATAS_TUNGGU_VERIFIKASI_MS = 15 * 60_000
BATAS_ULANG_429 = 6
# Target subsls: kode yang diklik tapi tidak tampil di halaman hasil pencarian subsls dicari satu
# per satu, paling banyak sekian per cek; hasil >1 halaman -> klik per putaran juga dibatasi segini.
MAKS_CARI_PER_KODE = 10
# Status yang berarti item "Ganti Mode" sudah diklik (dipakai masih_tuntas).
STATUS_SETELAH_KLIK = {"DIUBAH_TERVERIFIKASI", "DIUBAH_BELUM_TERVERIFIKASI"}


class Berhenti(RuntimeError):
    def __init__(self, kode: str, pesan: str = ""):
        super().__init__(pesan or kode)
        self.kode = kode


# ---------------------------------------------------------------------------
# Logika murni (diuji offline: tests/test_ubah_moda.py)
# ---------------------------------------------------------------------------

@dataclass
class Target:
    idsubsls: str
    akun_ppl: tuple[str, ...]
    baris_sheet: list[int] = field(default_factory=list)
    siap_input: bool = False
    # Kode identitas (bentuk baku) yang HARUS diubah persis — target list milik
    # user, dicari dgn kode itu sendiri. Kosong = target sheet/subsls: dicari per
    # idsubsls, skrip memilih sendiri menurut cakupan (atau semua PAPI utk ke CAPI).
    kode: str = ""
    # Mode TUJUAN. "CAPI" = arah balik (permintaan user 2026-09-22).
    ke: str = "PAPI"

    @property
    def kunci(self) -> str:
        """Kunci hasil/audit: kode identitas, atau idsubsls utk target sheet/subsls.
        Arah CAPI diberi awalan "CAPI:" supaya hasil arah PAPI (mis. SUDAH_ADA_PAPI)
        tidak membuat subsls/kode yang sama dilewati. Kunci PAPI tetap."""
        return ("CAPI:" if self.ke == "CAPI" else "") + (self.kode or self.idsubsls)

    @property
    def istilah_cari(self) -> str:
        """Teks yang diketik di kotak "Cari..."."""
        return self.kode or self.idsubsls

    def cocok(self, b: "BarisAssignment") -> bool:
        """Baris tabel yang dicari: kode PERSIS ("…- UMK - 4" bukan "…- UMK - 41"),
        atau semua baris subsls-nya utk target sheet."""
        return normalisasi_kode(b.kode) == self.kode if self.kode else b.idsubsls == self.idsubsls


# "5108060003000402 - UMK - 4". Bagian tengah bisa berupa nama ("I KETUT CONTOH /
# I KOMANG AGUS CONTOH", "WAYAN CONTOH /") dan boleh memuat "-" yang TIDAK diikuti
# spasi ("NON-UMK"); " - " di tengah nama tidak dikenali -> dilaporkan, bukan ditebak.
# Nol di depan nomor dibuang. Sel tabel bisa berakhiran lain ("… - 6 / - 81119"):
# yang diambil hanya kode di depannya. HARUS sama dgn ubah_moda_console.js.
POLA_KODE_IDENTITAS = re.compile(
    r"(?<!\d)(\d{16})\s*-\s*([A-Za-z0-9](?:[A-Za-z0-9 ._/'&(),+]|-(?=\S))*?)\s*-\s*0*(\d+)(?!\d)")


def _kode_baku(m: re.Match) -> str:
    return f"{m.group(1)} - {' '.join(m.group(2).split()).upper()} - {m.group(3)}"


def normalisasi_kode(teks) -> str:
    """Bentuk baku kode identitas utk dicocokkan PERSIS; "" kalau bukan kode."""
    m = POLA_KODE_IDENTITAS.search(str(teks or ""))
    return _kode_baku(m) if m else ""


def target_dari_daftar_kode(baris_teks: list[str]):
    """List kode identitas milik user (satu teks = satu baris file) ->
    (targets, tidak_dikenali, ganda). SATU TARGET PER KODE: kodenya sendiri yang
    diketik di kotak "Cari...". Baris tanpa 16 digit (judul kolom dsb.)
    diabaikan; baris ber-16 digit tanpa pola kode dilaporkan (mis. NIK, atau
    jenis yang memuat '-')."""
    targets: list[Target] = []
    sudah: set[str] = set()
    tidak_dikenali, ganda = [], []
    for no, teks in enumerate(baris_teks, start=1):
        s = str(teks or "")
        kode = [_kode_baku(m) for m in POLA_KODE_IDENTITAS.finditer(s)]
        if not kode:
            if re.search(r"(?<!\d)\d{16}(?!\d)", s):
                tidak_dikenali.append((no, " ".join(s.split())[:80]))
            continue
        for k in kode:
            if k in sudah:
                ganda.append((no, k))
                continue
            sudah.add(k)
            targets.append(Target(k[:16], (), [no], kode=k))
    return targets, tidak_dikenali, ganda


_POLA_KODE_DI_BARIS = re.compile(r"\d{16}\s*-\s*\S")
_POLA_ILMIAH = re.compile(r"^\d(\.\d+)?e\+?\d+$", re.I)


def target_dari_daftar_subsls(baris_teks: list[str], ke: str):
    """List idsubsls (satu teks = satu baris file) -> (targets, tidak_dikenali, ganda,
    kode_identitas). Target subsls ber-arah `ke`; utk ke CAPI SEMUA assignment PAPI
    subsls itu yang tampil diubah. Hanya angka 16 digit UTUH yang dimuat. Baris berpola
    kode identitas ("<16 digit> - …") TIDAK dimuat: list kode yang tertempel ke sini
    akan melebar jadi satu subsls penuh. Angka panjang lain (15/17 digit, notasi ilmiah
    Excel "5.10806E+15") dilaporkan, tidak ditebak. HARUS sama dgn ubah_moda_console.js."""
    targets: list[Target] = []
    sudah: set[str] = set()
    tidak_dikenali, ganda, kode_identitas = [], [], []
    for no, teks in enumerate(baris_teks, start=1):
        s = str(teks or "")
        if _POLA_KODE_DI_BARIS.search(s):
            kode_identitas.append((no, " ".join(s.split())[:80]))
            continue
        for tok in (x for x in re.split(r"[\s,;]+", s) if x):
            if re.fullmatch(r"\d{16}", tok):
                if tok in sudah:
                    ganda.append((no, tok))
                else:
                    sudah.add(tok)
                    targets.append(Target(tok, (), [no], ke=ke))
            elif len(re.sub(r"\D", "", tok)) >= 10 or _POLA_ILMIAH.match(tok):
                tidak_dikenali.append((no, tok[:80]))
    return targets, tidak_dikenali, ganda, kode_identitas


def diulang_per_subsls(t: Target, cakupan: str) -> bool:
    """Target subsls yang diulang per putaran sampai habis (cakupan "semua", atau arah CAPI)."""
    return not t.kode and (t.ke == "CAPI" or cakupan == "semua")


def _teks_sel(v) -> str:
    """Sel xlsx -> teks. Angka panjang (idsubsls/NIK yang tersimpan sbg NUMBER) diberi awalan
    ANGKA_EXCEL: supaya TIDAK dimuat sbg idsubsls: Excel memotong angka jadi 15 digit
    ("5108070013000901" -> 5108070013000900), dan hasilnya bisa saja subsls LAIN yang ada."""
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) >= 1e9:
        return f"ANGKA_EXCEL:{v}"
    return str(v)


def baca_daftar(path: str | Path, sheet: Optional[str] = None) -> list[str]:
    """File list -> satu teks per baris. .xlsx: sheet pertama (atau `sheet`),
    sel sebaris digabung tab; lainnya dibaca sbg teks."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File daftar tidak ditemukan: {path}")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None:
                ws = wb.worksheets[0]
            elif sheet in wb.sheetnames:
                ws = wb[sheet]
            else:
                raise ValueError(f"sheet '{sheet}' tidak ada di {path.name} (ada: {wb.sheetnames})")
            print(f"Membaca sheet '{ws.title}' dari {path.name}")
            return ["\t".join(_teks_sel(v) for v in row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc).splitlines()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").splitlines()


def bangun_target(rows, hasil_cek=None, hanya_siap: bool = False):
    """-> (targets, dikeluarkan). Satu target per idsubsls.

    Baris yang idsubsls-nya TIDAK sama dgn kolom Pilih PROVINSI..SUBSLS
    dikeluarkan: mengubah mode di subsls yang salah adalah efek samping di
    sistem produksi, jadi lebih baik dilaporkan & diperbaiki di sheet dulu."""
    per: dict[str, Target] = {}
    dikeluarkan = []
    for r in rows:
        if len(r.idsubsls) != 16 or r.idsubsls != r.idsubsls_pilih:
            dikeluarkan.append((r.baris, r.idsubsls, r.idsubsls_pilih))
            continue
        t = per.setdefault(r.idsubsls, Target(r.idsubsls, ()))
        if r.akun_ppl and r.akun_ppl not in t.akun_ppl:
            t.akun_ppl = tuple(sorted({*t.akun_ppl, r.akun_ppl}))
        t.baris_sheet.append(r.baris)
        if hasil_cek is not None and hasil_cek[r.baris].status == "SIAP":
            t.siap_input = True
    targets = [t for t in per.values() if t.siap_input or not hanya_siap]
    return targets, dikeluarkan


@dataclass
class BarisAssignment:
    kode: str
    status: str = ""
    mode: str = ""
    petugas: str = ""
    keterangan: str = ""
    indeks: int = -1   # urutan baris data di tabel saat dibaca

    @property
    def idsubsls(self) -> str:
        m = re.match(r"\s*(\d{16})", self.kode)
        return m.group(1) if m else ""


def baris_dari_tabel(data: Optional[dict]) -> list[BarisAssignment]:
    """Ubah hasil _JS_TABEL ({head, rows}) jadi BarisAssignment, dipetakan
    lewat JUDUL kolom (bukan posisi) supaya tahan kolom ditambah/diurutkan."""
    if not data:
        raise Berhenti("KOLOM_TIDAK_ADA", "tabel berjudul 'Kode Identitas' tidak ditemukan")
    head = [" ".join((h or "").split()) for h in data.get("head", [])]
    idx = {}
    for key, judul in KOLOM_TABEL.items():
        cocok = [i for i, h in enumerate(head) if h.lower() == judul.lower()]
        if not cocok:
            if key == "keterangan":
                continue
            raise Berhenti("KOLOM_TIDAK_ADA",
                           f"kolom '{judul}' tidak tampil (judul terbaca: {head}). Tampilkan lewat tombol 'Kolom'.")
        idx[key] = cocok[0]
    out = []
    for i, sel in enumerate(data.get("rows", [])):
        if len(sel) <= max(idx.values()):
            continue  # baris pesan "tidak ada data" (colspan) dsb.
        b = BarisAssignment(indeks=i, **{k: " ".join((sel[j] or "").split()) for k, j in idx.items()})
        if b.idsubsls:
            out.append(b)
    return out


@dataclass
class Rencana:
    status: str
    pilih: list[BarisAssignment] = field(default_factory=list)
    pesan: str = ""


def rencanakan(target: Target, baris: list[BarisAssignment], cakupan: str = "satu",
               ada_halaman_lain: bool = False) -> Rencana:
    """Tentukan assignment mana yang diubah ke PAPI, dari baris HALAMAN YANG TAMPIL.

    cakupan "satu" (default): cukup SATU assignment PAPI mana pun per subsls —
    itu yang dibutuhkan fasih-web. Petugasnya TIDAK harus PPL sheet (ketetapan
    user 2026-09-14); milik PPL sheet hanya didahulukan kalau ada.
    "semua": seluruh assignment CAPI yang tampil.

    Baris subsls LAIN yang ikut terbawa pencarian diabaikan (tidak pernah
    dipilih). ada_halaman_lain = hasil pencarian > 1 halaman; halaman lain
    SENGAJA tidak dibaca (paginasi fasih-sm tidak memuat data dgn benar saat
    dipindah — temuan user 2026-09-14), jadi PAPI di sana tidak terlihat.

    Target kode identitas -> _rencanakan_kode (cakupan diabaikan).
    Target subsls arah CAPI -> _rencanakan_subsls_ke_capi (cakupan diabaikan)."""
    if target.kode:
        return _rencanakan_kode(target, baris, ada_halaman_lain)
    if target.ke == "CAPI":
        return _rencanakan_subsls_ke_capi(target, baris, ada_halaman_lain)
    milik = [b for b in baris if b.idsubsls == target.idsubsls]
    n_asing = len(baris) - len(milik)
    catatan = (f" | {n_asing} baris subsls lain diabaikan" if n_asing else "") + \
        (" | >1 halaman, hanya halaman tampil yang dibaca" if ada_halaman_lain else "")
    if not milik:
        if n_asing or ada_halaman_lain:
            return Rencana("SUBSLS_TIDAK_TAMPIL", pesan=(
                f"halaman hasil pencarian tidak memuat satu pun baris {target.idsubsls} — "
                f"pencarian tidak menyaring?{catatan}"))
        return Rencana("TIDAK_ADA_ASSIGNMENT", pesan=(
            f"tidak ada assignment subsls ini di hasil pencarian — tidak bisa dibantu dgn ganti mode{catatan}"))
    aneh = sorted({b.mode for b in milik if b.mode.upper() not in ("CAPI", "PAPI")})
    if aneh:
        return Rencana("MODE_TIDAK_DIKENAL", pesan=f"nilai kolom Mode: {aneh}{catatan}")

    capi = [b for b in milik if b.mode.upper() == "CAPI"]
    papi = [b for b in milik if b.mode.upper() == "PAPI"]

    if cakupan == "semua":
        if capi:
            return Rencana("PERLU_DIUBAH", capi, f"{len(capi)} assignment CAPI yang tampil{catatan}")
        if ada_halaman_lain:
            return Rencana("PERLU_HALAMAN_LAIN", pesan=(
                f"tidak ada CAPI di halaman tampil, sisa CAPI (kalau ada) di halaman lain — butuh filter Mode{catatan}"))
        return Rencana("TIDAK_ADA_CAPI", pesan=f"{len(papi)} assignment sudah PAPI{catatan}")

    if papi:
        return Rencana("SUDAH_ADA_PAPI", pesan=(
            f"{len(papi)} assignment sudah PAPI (mis. {papi[0].kode}, petugas {papi[0].petugas}){catatan}"))
    ppl = {p.lower() for p in target.akun_ppl}
    pilih = next((b for b in capi if b.petugas.strip().lower() in ppl), capi[0])
    punya_ppl = pilih.petugas.strip().lower() in ppl
    return Rencana("PERLU_DIUBAH", [pilih], (
        f"1 dari {len(capi)} assignment CAPI (petugas {pilih.petugas}{' = PPL sheet' if punya_ppl else ''}){catatan}"))


def _rencanakan_subsls_ke_capi(target: Target, baris: list[BarisAssignment],
                               ada_halaman_lain: bool = False) -> Rencana:
    """Target subsls arah CAPI (permintaan user 2026-09-22: "subsls tertentu saja, dari
    PAPI ke CAPI") -> SEMUA assignment PAPI subsls itu di halaman tampil. Baris subsls
    lain diabaikan. Paginasi tetap tidak dipindah: hasil >1 halaman tanpa PAPI yang
    tampil -> CEK_HALAMAN_LAIN (lanjut, tidak tuntas; ubah lewat list kode identitas)."""
    milik = [b for b in baris if b.idsubsls == target.idsubsls]
    n_asing = len(baris) - len(milik)
    catatan = (f" | {n_asing} baris subsls lain diabaikan" if n_asing else "") + \
        (" | >1 halaman, hanya halaman tampil yang dibaca" if ada_halaman_lain else "")
    if not milik:
        if n_asing or ada_halaman_lain:
            return Rencana("SUBSLS_TIDAK_TAMPIL", pesan=(
                f"halaman hasil pencarian tidak memuat satu pun baris {target.idsubsls} — "
                f"pencarian tidak menyaring?{catatan}"))
        return Rencana("TIDAK_ADA_ASSIGNMENT", pesan=f"tidak ada assignment subsls ini di hasil pencarian{catatan}")
    aneh = sorted({b.mode for b in milik if b.mode.upper() not in MODE})
    if aneh:
        return Rencana("MODE_TIDAK_DIKENAL", pesan=f"nilai kolom Mode: {aneh}{catatan}")
    papi = [b for b in milik if b.mode.upper() == "PAPI"]
    capi = len(milik) - len(papi)
    if papi:
        return Rencana("PERLU_DIUBAH", papi, f"{len(papi)} assignment PAPI yang tampil ({capi} sudah CAPI){catatan}")
    if ada_halaman_lain:
        return Rencana("CEK_HALAMAN_LAIN", pesan=(
            f"tidak ada PAPI di halaman tampil ({capi} CAPI), tapi hasil pencarian >1 halaman — PAPI di halaman "
            f"lain (kalau ada) tidak terlihat; ubah lewat list kode identitas{catatan}"))
    return Rencana("TIDAK_ADA_PAPI", pesan=f"{capi} assignment, semua sudah CAPI{catatan}")


def _rencanakan_kode(target: Target, baris: list[BarisAssignment], ada_halaman_lain: bool = False) -> Rencana:
    """Target kode identitas (hasil pencarian KODE itu) -> HANYA baris kode itu
    yang diubah ke mode tujuannya (cakupan diabaikan). Baris lain yang ikut tampil
    (mis. "…- UMK - 41" saat mencari "…- UMK - 4") diabaikan."""
    ke, asal = target.ke, lawan_mode(target.ke)
    cocok = [b for b in baris if target.cocok(b)]
    n_lain = len(baris) - len(cocok)
    catatan = f" | {n_lain} baris kode lain ikut tampil, diabaikan" if n_lain else ""
    if len(cocok) > 1:
        return Rencana("KODE_GANDA", pesan=(
            f"kode {target.kode} tampil {len(cocok)}x di hasil pencarian — tidak dipilih{catatan}"))
    if not cocok:
        subsls_lain = sum(b.idsubsls != target.idsubsls for b in baris)
        if subsls_lain:
            return Rencana("PENCARIAN_TIDAK_MENYARING", pesan=(
                f"hasil pencarian {target.kode} memuat {subsls_lain} baris subsls lain & kode itu tidak ada — "
                f"kotak Cari tidak menyaring per kode identitas?{catatan}"))
        if ada_halaman_lain:
            return Rencana("KODE_TIDAK_TAMPIL", pesan=f"kode tidak ada di halaman tampil, hasil pencarian >1 halaman{catatan}")
        return Rencana("KODE_TIDAK_ADA", pesan=f"kode tidak ditemukan di fasih-sm — cek penulisan kode / periode survei{catatan}")
    b = cocok[0]
    mode = b.mode.upper()
    if mode == ke:
        return Rencana(f"KODE_SUDAH_{ke}", pesan=f"sudah {ke} (petugas {b.petugas}){catatan}")
    if mode != asal:
        return Rencana("MODE_TIDAK_DIKENAL", pesan=f"nilai kolom Mode: {b.mode}{catatan}")
    return Rencana("PERLU_DIUBAH", [b], f"{asal} (petugas {b.petugas}){catatan}")


def dialog_sesuai(teks: str, ke: str) -> bool:
    """Teks dialog konfirmasi cocok dgn arah `ke`? Frasa "ke/to/menjadi <MODE>"
    menentukan (semuanya harus = ke); tanpa frasa itu dialog harus menyebut `ke`,
    atau menyebut "mode" tanpa menyebut mode lawan. HARUS sama dgn ubah_moda_console.js."""
    teks = teks or ""
    tujuan = [m.upper() for m in re.findall(r"\b(?:ke|to|menjadi)\s+(CAPI|PAPI)\b", teks, re.I)]
    if tujuan:
        return all(m == ke for m in tujuan)
    if re.search(rf"\b{ke}\b", teks, re.I):
        return True
    return bool(re.search(r"mode", teks, re.I)) and not re.search(rf"\b{lawan_mode(ke)}\b", teks, re.I)


def angka_item_menu(teks: str) -> Optional[int]:
    """"Ganti Mode (Ke PAPI) (3)" -> 3."""
    m = POLA_ANGKA_ITEM.search(" ".join((teks or "").split()))
    return int(m.group(1)) if m else None


def pilih_tombol_konfirmasi(teks_tombol: list[str]) -> Optional[int]:
    """Indeks SATU-SATUNYA tombol konfirmasi di dialog, atau None kalau tidak
    ada / lebih dari satu (ambigu -> jangan diklik)."""
    kandidat = [i for i, t in enumerate(teks_tombol)
                if POLA_TOMBOL_KONFIRMASI.search(t or "") and not POLA_TOMBOL_BATAL.search(t or "")]
    return kandidat[0] if len(kandidat) == 1 else None


def jeda_cek_verifikasi(ke: int) -> int:
    """ms sebelum cek ulang berikutnya (ke = jumlah cek yang sudah dilakukan)."""
    return JADWAL_CEK_VERIFIKASI_MS[ke] if ke < len(JADWAL_CEK_VERIFIKASI_MS) else JEDA_CEK_VERIFIKASI_MAKS_MS


def putuskan_verifikasi(mode: dict, sejak_klik_ms: float, batas_ms: int = BATAS_TUNGGU_VERIFIKASI_MS,
                        ke: str = "PAPI") -> str:
    """{kode: mode terbaca} -> 'TERVERIFIKASI' (semua = mode tujuan `ke`) | 'MENUNGGU' |
    'BELUM_TERVERIFIKASI' (batas tunggu habis). Kode yang tidak tampil
    ('(hilang)') dianggap belum berubah."""
    if mode and all(str(m).upper() == ke for m in mode.values()):
        return "TERVERIFIKASI"
    return "BELUM_TERVERIFIKASI" if sejak_klik_ms >= batas_ms else "MENUNGGU"


def kunci_audit(r: dict) -> str:
    """Kunci baris audit (sama dgn Target.kunci)."""
    return ("CAPI:" if (r.get("ke") or "").upper() == "CAPI" else "") + (r.get("kode_target") or r.get("idsubsls") or "")


def masih_tuntas(t: Target, riwayat: list[dict], tuntas: set) -> bool:
    """riwayat = baris audit berurutan (lama -> baru). Status terakhir kunci `t` harus
    tuntas DAN sesudahnya tidak ada klik ganti mode arah LAWAN di subsls yang sama
    (mis. subsls yang dulu SUDAH_ADA_PAPI lalu PAPI-nya dikembalikan ke CAPI -> periksa
    lagi). Padanan masihTuntas() di ubah_moda_console.js (di sana dgn waktu klik)."""
    idx = None
    for i, r in enumerate(riwayat):
        if kunci_audit(r) == t.kunci:
            idx = i
    if idx is None or riwayat[idx].get("status") not in tuntas:
        return False
    lawan = lawan_mode(t.ke)
    return not any((r.get("ke") or "PAPI").upper() == lawan and r.get("idsubsls") == t.idsubsls
                   and r.get("status") in STATUS_SETELAH_KLIK for r in riwayat[idx + 1:])


def jeda_rate_limit(ke: int, retry_after=None) -> int:
    """ms sebelum mengulang pencarian yang kena HTTP 429 (percobaan ke-0,1,..).
    Retry-After (detik) dihormati (5 dtk–5 mnt); tanpa itu 15 dtk x 2^ke, maks 2 menit."""
    try:
        detik = float(retry_after)
    except (TypeError, ValueError):
        detik = None
    if detik is not None and math.isfinite(detik) and detik >= 0:
        return int(min(max(detik * 1000, 5_000), 300_000))
    return min(15_000 * 2 ** ke, 120_000)


# ---------------------------------------------------------------------------
# Interaksi browser
# ---------------------------------------------------------------------------

_JS_TABEL = r"""
() => {
  const tampak = el => { const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const bersih = t => (t || '').replace(/\s+/g, ' ').trim();
  const tabel = [...document.querySelectorAll('table')].filter(tampak)
      .find(t => /Kode Identitas/i.test(t.innerText));
  if (!tabel) return null;
  const head = [...tabel.querySelectorAll('thead th')].map(h => bersih(h.innerText));
  const rows = [...tabel.querySelectorAll('tbody tr')]
      .map(tr => [...tr.querySelectorAll('td')].map(td => bersih(td.innerText)));
  const teks = document.body.innerText;
  const hal = teks.match(/Page\s+(\d+)\s+of\s+(\d+)/i);
  return {head, rows, halaman: hal ? [Number(hal[1]), Number(hal[2])] : null};
}
"""

_JS_DUMP = r"""
() => {
  const tampak = el => { const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const bagian = [];
  for (const sel of ['[role="menu"]', '[role="dialog"]', '[role="alertdialog"]', '[data-sonner-toast]', 'table']) {
    document.querySelectorAll(sel).forEach(el => {
      if (!tampak(el)) return;
      bagian.push('### ' + sel + '\n' + (el.innerText || '').slice(0, 6000)
                  + '\n--- html ---\n' + el.outerHTML.slice(0, 12000));
    });
  }
  return 'URL: ' + location.href + '\n\n' + bagian.join('\n\n');
}
"""


class FasihSm:
    def __init__(self, page, log_dir: Path = LOG_DIR):
        self.page = page
        self.log_dir = log_dir
        self.log_dir.mkdir(parents=True, exist_ok=True)
        self.rekaman = self.log_dir / f"api_{time.strftime('%Y%m%d_%H%M%S')}.jsonl"
        self._pasang_perekam()

    def _log(self, msg: str):
        print(f"[fasih-sm] {msg}")

    def _pasang_perekam(self):
        """Rekam SEMUA panggilan /app/api/ (termasuk body request & respons)
        ke log_fasih_sm/api_*.jsonl. Tujuannya: (1) saat --petakan, bukti
        bagaimana halaman menyaring & mengubah mode; (2) saat --eksekusi,
        jejak persis request ganti mode yang terkirim. Berisi data pribadi —
        folder ini masuk .gitignore."""
        def _on_response(resp):
            try:
                req = resp.request
                if "/app/api/" not in resp.url or req.resource_type not in ("xhr", "fetch"):
                    return
                try:
                    isi = resp.text()[:20_000]
                except Exception:
                    isi = "(tidak terbaca)"
                baris = {"t": time.strftime("%H:%M:%S"), "method": req.method, "url": resp.url,
                         "status": resp.status, "request": (req.post_data or "")[:20_000], "response": isi}
                with self.rekaman.open("a", encoding="utf-8") as f:
                    f.write(json.dumps(baris, ensure_ascii=False) + "\n")
            except Exception:
                pass
        self.page.on("response", _on_response)

    def dump(self, nama: str):
        stamp = f"{time.strftime('%Y%m%d_%H%M%S')}_{re.sub(r'[^A-Za-z0-9_-]', '_', nama)[:50]}"
        try:
            (self.log_dir / f"{stamp}.txt").write_text(self.page.evaluate(_JS_DUMP), encoding="utf-8")
            self.page.screenshot(path=str(self.log_dir / f"{stamp}.png"))
            self._log(f"📄 dump: log_fasih_sm/{stamp}.txt (+ .png)")
        except Exception as e:
            self._log(f"⚠️ dump gagal (tidak fatal): {e}")

    # --- halaman -----------------------------------------------------------
    def buka_list(self, url: str):
        self.page.goto(url, wait_until="domcontentloaded")
        judul = self.page.get_by_text("Kode Identitas", exact=True).locator("visible=true").first
        try:
            judul.wait_for(state="visible", timeout=30_000)
        except PWTimeout:
            print("\n>>> Tabel assignment belum tampil. LOGIN MANUAL di jendela browser "
                  "(akun yang berhak 'Ganti Mode').")
            print(">>> Setelah daftar assignment tampil, kembali ke terminal ini dan tekan Enter.")
            input()
            self.page.goto(url, wait_until="domcontentloaded")
            try:
                judul.wait_for(state="visible", timeout=60_000)
            except PWTimeout:
                self.dump("belum_login")
                raise Berhenti("BELUM_LOGIN", "tabel assignment tetap tidak tampil setelah login")
        try:
            self._tombol_aksi_lainnya().wait_for(state="visible", timeout=15_000)
        except PWTimeout:
            self.dump("tanpa_aksi_lainnya")
            raise Berhenti("TIDAK_ADA_AKSES", "tombol 'Aksi Lainnya' tidak ada — akun ini tidak berhak ganti mode")
        self._log("Daftar assignment tampil & tombol 'Aksi Lainnya' ada.")

    def _tombol_aksi_lainnya(self):
        return self.page.get_by_role("button", name=re.compile(r"Aksi Lainnya", re.I)).locator("visible=true").first

    def _tabel(self):
        return self.page.locator("table").filter(has_text="Kode Identitas").locator("visible=true").first

    def tutup_menu(self):
        """Pastikan TIDAK ada menu terbuka. wait_for(hidden) — bukan count():
        menu Radix menutup dgn animasi, count() bisa masih melihatnya."""
        menu = self.page.locator('[role="menu"]').locator("visible=true").first
        try:
            menu.wait_for(state="hidden", timeout=1_000)
            return
        except PWTimeout:
            pass
        self.page.keyboard.press("Escape")
        try:
            menu.wait_for(state="hidden", timeout=3_000)
            return
        except PWTimeout:
            pass
        self.page.get_by_role("heading", name="Data", exact=True).locator("visible=true").first.click()
        try:
            menu.wait_for(state="hidden", timeout=3_000)
        except PWTimeout:
            self.dump("menu_tidak_tertutup")
            raise Berhenti("MENU_TIDAK_TERTUTUP", "menu tetap terbuka — klik berikutnya bisa mengenai item menu")

    def baca_tabel(self) -> tuple[list[BarisAssignment], Optional[list]]:
        data = self.page.evaluate(_JS_TABEL)
        return baris_dari_tabel(data), (data or {}).get("halaman")

    def cari(self, istilah: str) -> tuple[list[BarisAssignment], bool]:
        """Saring daftar lewat kotak "Cari..." (kode identitas atau idsubsls —
        lihat Target.istilah_cari). Selalu dikosongkan dulu supaya pencarian
        ulang (verifikasi) benar-benar memuat data baru.
        -> (baris halaman yang tampil, ada_halaman_lain). Paginasi TIDAK PERNAH
        dipindah — halaman yang dipindah tidak memuat datanya dgn benar.
        Respons datatable HTTP 429 (rate limit, run user 2026-09-15) -> tabelnya
        tidak dipercaya: tunggu jeda_rate_limit lalu ulangi, maks BATAS_ULANG_429
        kali -> RATE_LIMIT."""
        for ulang in range(BATAS_ULANG_429 + 1):
            self.tutup_menu()
            kotak = self.page.locator('input[placeholder="Cari..."]').locator("visible=true").first
            kotak.wait_for(state="visible", timeout=15_000)
            respons_429 = None
            for nilai in ("", istilah):
                try:
                    with self.page.expect_response(lambda r: "datatable" in r.url, timeout=25_000) as info:
                        kotak.fill(nilai)
                        kotak.press("Enter")
                    if info.value.status == 429:
                        respons_429 = info.value
                        break
                except PWTimeout:
                    self._log(f"⚠️ tidak ada respons datatable setelah mengisi '{nilai}' — baca tabel apa adanya.")
                self.page.wait_for_timeout(1_500)
            if respons_429 is None:
                baris, halaman = self.baca_tabel()
                return baris, bool(halaman and halaman[1] > 1)
            if ulang >= BATAS_ULANG_429:
                break
            jeda = jeda_rate_limit(ulang, respons_429.headers.get("retry-after"))
            self._log(f"⏳ HTTP 429 (rate limit) saat mencari '{istilah}' — tunggu {jeda // 1000} dtk lalu cari ulang "
                      f"({ulang + 1}/{BATAS_ULANG_429}).")
            self.page.wait_for_timeout(jeda)
        raise Berhenti("RATE_LIMIT", f"pencarian '{istilah}' masih HTTP 429 setelah {BATAS_ULANG_429} kali menunggu")

    def _baris_tr(self, b: BarisAssignment):
        """<tr> utk baris `b`, DIVERIFIKASI kodenya persis sama. has_text tidak
        dipakai: "…- UMK - 4" juga cocok dgn "…- UMK - 41"."""
        tr = self._tabel().locator("tbody tr").nth(b.indeks)
        tr.wait_for(state="visible", timeout=10_000)
        data = self.page.evaluate(_JS_TABEL)
        sekarang = baris_dari_tabel(data)
        cocok = [x for x in sekarang if x.indeks == b.indeks]
        if not cocok or cocok[0].kode != b.kode:
            raise Berhenti("TABEL_BERUBAH", f"baris ke-{b.indeks} bukan lagi '{b.kode}'")
        return tr

    def _tercentang(self, cb) -> bool:
        return cb.get_attribute("aria-checked") == "true" or cb.get_attribute("data-state") == "checked"

    def atur_centang(self, daftar: list[BarisAssignment], centang: bool):
        for b in daftar:
            self.tutup_menu()
            cb = self._baris_tr(b).get_by_role("checkbox").first
            cb.wait_for(state="visible", timeout=10_000)
            if self._tercentang(cb) != centang:
                cb.click()
                self.page.wait_for_timeout(300)
            if self._tercentang(cb) != centang:
                raise Berhenti("CENTANG_GAGAL", f"checkbox '{b.kode}' tidak berubah jadi {centang}")
            self._log(f"  {'☑' if centang else '☐'} {b.kode} ({b.mode}, {b.petugas})")

    def lepas_centang_kode(self, kode: list[str]):
        sekarang = [b for b in self.baca_tabel()[0] if b.kode in kode]
        self.atur_centang(sekarang, False)

    def buka_menu_ganti_mode(self, ke: str = "PAPI", wajib: bool = True):
        """Buka 'Aksi Lainnya' -> (item 'Ganti Mode (Ke <ke>)', angka N). wajib=False
        (petakan): item yang tidak ada -> (None, None) & menu ditutup, bukan berhenti."""
        self.tutup_menu()
        self._tombol_aksi_lainnya().click()
        item = self.page.get_by_role("menuitem").filter(has_text=pola_item_ganti_mode(ke)).locator("visible=true").first
        try:
            item.wait_for(state="visible", timeout=8_000)
        except PWTimeout:
            try:
                terlihat = [" ".join(t.split()) for t in self.page.get_by_role("menuitem").all_inner_texts()][:30]
            except Exception:  # noqa: BLE001
                terlihat = []
            self.dump("menu_tanpa_ganti_mode")
            self.tutup_menu()
            if not wajib:
                self._log(f"Item 'Ganti Mode (Ke {ke})' tidak tampil (0 dicentang). Item terlihat: {terlihat}")
                return None, None
            raise Berhenti("ITEM_MENU_TIDAK_ADA",
                           f"item 'Ganti Mode (Ke {ke})' tidak ada di menu 'Aksi Lainnya'. Item yang terlihat: {terlihat}")
        return item, angka_item_menu(item.inner_text())

    def klik_ganti_mode(self, item, ke: str = "PAPI") -> str:
        """⚠️ IRREVERSIBLE. Klik item massal lalu tangani dialog konfirmasi
        kalau ada. -> 'DIKONFIRMASI' | 'TANPA_DIALOG'."""
        item.click()
        dialog = self.page.locator('[role="dialog"], [role="alertdialog"]').locator("visible=true").first
        try:
            dialog.wait_for(state="visible", timeout=8_000)
        except PWTimeout:
            self._log("Tidak ada dialog konfirmasi — aksi mungkin langsung diproses; diverifikasi lewat tabel.")
            self.dump("setelah_ganti_mode_tanpa_dialog")
            return "TANPA_DIALOG"
        teks = " ".join(dialog.inner_text().split())
        self.dump("dialog_ganti_mode")
        self._log(f"Dialog: {teks[:200]}")
        tombol = dialog.get_by_role("button")
        teks_tombol = [" ".join(t.split()) for t in tombol.all_inner_texts()]
        if not dialog_sesuai(teks, ke):
            self.page.keyboard.press("Escape")
            raise Berhenti("DIALOG_TIDAK_DIKENAL", f"dialog tidak cocok dgn arah ke {ke}: {teks[:200]}")
        i = pilih_tombol_konfirmasi(teks_tombol)
        if i is None:
            self.page.keyboard.press("Escape")
            raise Berhenti("TOMBOL_KONFIRMASI_AMBIGU", f"tombol dialog: {teks_tombol}")
        self._log(f"Klik tombol konfirmasi '{teks_tombol[i]}'.")
        tombol.nth(i).click()
        try:
            dialog.wait_for(state="hidden", timeout=20_000)
        except PWTimeout:
            self.dump("dialog_tidak_tertutup")
        self.page.wait_for_timeout(2_000)
        return "DIKONFIRMASI"

    def verifikasi_mode(self, istilah: str, kode: list[str], ke: str = "PAPI", per_kode: bool = False,
                        batas_ms: int = BATAS_TUNGGU_VERIFIKASI_MS) -> bool:
        """Mode baru TIDAK langsung terbaca (run user 2026-09-15). Cari ulang
        menurut jeda_cek_verifikasi sampai semua = `ke` atau batas_ms habis.
        per_kode (target subsls): kode yang tidak tampil di halaman hasil
        pencarian subsls dicari satu per satu (maks MAKS_CARI_PER_KODE).
        Berbeda dgn Console (antrean, batch lanjut), jalur cadangan ini
        MENUNGGU per target."""
        mulai = time.monotonic()
        n_cek = 0
        while True:
            self.page.wait_for_timeout(jeda_cek_verifikasi(n_cek))
            n_cek += 1
            baris = {b.kode: b for b in self.cari(istilah)[0]}
            mode = {k: (baris[k].mode if k in baris else "(hilang)") for k in kode}
            hilang = [k for k, m in mode.items() if m == "(hilang)"]
            if per_kode and hilang and len(hilang) <= MAKS_CARI_PER_KODE:
                for k in hilang:
                    ketemu = [b for b in self.cari(k)[0] if b.kode == k]
                    mode[k] = ketemu[0].mode if ketemu else "(hilang)"
            sejak = (time.monotonic() - mulai) * 1000
            keputusan = putuskan_verifikasi(mode, sejak, batas_ms, ke)
            self._log(f"Verifikasi ke-{n_cek} ({sejak / 60_000:.1f} mnt sejak diklik): {mode} -> {keputusan}")
            if keputusan != "MENUNGGU":
                return keputusan == "TERVERIFIKASI"


# ---------------------------------------------------------------------------
# Orkestrasi
# ---------------------------------------------------------------------------

def append_audit(row: dict):
    baru = not AUDIT_PATH.exists()
    with lokasi.siapkan(AUDIT_PATH).open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=AUDIT_FIELDS)
        if baru:
            w.writeheader()
        w.writerow({k: row.get(k, "") for k in AUDIT_FIELDS})


def riwayat_audit() -> list[dict]:
    """Seluruh baris audit, berurutan lama -> baru (bahan masih_tuntas)."""
    if not AUDIT_PATH.exists():
        return []
    with AUDIT_PATH.open(newline="", encoding="utf-8") as f:
        return [r for r in csv.DictReader(f) if r.get("idsubsls")]


def proses_target(sm: FasihSm, t: Target, args, jalan: str, maks_per_klik: int = 50) -> dict:
    hasil = {"timestamp": time.strftime("%Y-%m-%d %H:%M:%S"), "jalan": jalan, "idsubsls": t.idsubsls,
             "akun_ppl": ",".join(t.akun_ppl), "baris_sheet": ",".join(map(str, t.baris_sheet)),
             "kode_target": t.kode, "ke": t.ke}
    pilih: list[BarisAssignment] = []
    try:
        baris, ada_halaman_lain = sm.cari(t.istilah_cari)
        rencana = rencanakan(t, baris, args.cakupan, ada_halaman_lain)
        pilih, pesan = list(rencana.pilih), rencana.pesan
        # Hasil >1 halaman: baris yang diklik bisa pindah halaman setelah modenya berubah ->
        # diverifikasi per kode, jadi per klik paling banyak MAKS_CARI_PER_KODE.
        maks = min(maks_per_klik, MAKS_CARI_PER_KODE) if ada_halaman_lain else maks_per_klik
        if len(pilih) > maks:
            pesan += f" | {len(pilih) - maks} sisanya di putaran berikutnya (maks {maks} per klik)"
            pilih = pilih[:maks]
        # Target kode: baris kode persis itu saja; target sheet/subsls: semua baris subsls-nya.
        milik = [b for b in baris if t.cocok(b)]
        hasil.update(jumlah_assignment=len(milik), capi=sum(b.mode.upper() == "CAPI" for b in milik),
                     papi=sum(b.mode.upper() == "PAPI" for b in milik), pesan=pesan,
                     dipilih=" | ".join(b.kode for b in pilih),
                     petugas_dipilih=" | ".join(b.petugas for b in pilih))
        for b in baris:
            sm._log(f"  {b.kode:28s} {b.mode:5s} {b.status:24s} {b.petugas} ({b.keterangan})")
        if jalan == "petakan":
            sm.dump(f"petakan_tabel_{t.idsubsls}")
            pilih = []
            item, n = sm.buka_menu_ganti_mode(t.ke, wajib=False)
            sm.dump("petakan_menu_aksi_lainnya")
            sm.tutup_menu()
            hasil["status"] = f"PETAKAN_{rencana.status}"
            hasil["pesan"] = f"{pesan} | angka menu 'Ke {t.ke}' saat 0 dicentang = {n if item else '(item tidak tampil)'}"
            return hasil
        if rencana.status != "PERLU_DIUBAH":
            hasil["status"] = rencana.status
            return hasil

        sm.atur_centang(pilih, True)
        item, n = sm.buka_menu_ganti_mode(t.ke)
        if n != len(pilih):
            raise Berhenti("JUMLAH_TERCENTANG_BEDA", f"menu menunjukkan ({n}), yang dicentang {len(pilih)}")
        if jalan == "dry-run":
            sm.tutup_menu()
            sm.atur_centang(pilih, False)
            pilih = []
            hasil["status"] = "DRY_RUN_AKAN_DIUBAH"
            return hasil

        cara = sm.klik_ganti_mode(item, t.ke)
        kode = [b.kode for b in pilih]
        ok = sm.verifikasi_mode(t.istilah_cari, kode, t.ke, per_kode=not t.kode)
        hasil["status"] = "DIUBAH_TERVERIFIKASI" if ok else "DIUBAH_BELUM_TERVERIFIKASI"
        hasil["pesan"] = f"{pesan} | {cara}"
        # Centang yang tertinggal bisa ikut terkirim di aksi massal subsls
        # berikutnya. Indeks baris bisa bergeser setelah pencarian ulang,
        # jadi dicocokkan lewat kode.
        sm.tutup_menu()
        sm.lepas_centang_kode(kode)
        return hasil
    except Berhenti as e:
        hasil["status"], hasil["pesan"] = e.kode, str(e)
        sm.dump(f"BERHENTI_{e.kode}_{t.idsubsls}")
    except Exception as e:  # noqa: BLE001
        hasil["status"] = "ERROR_TAK_TERDUGA"
        hasil["pesan"] = f"{type(e).__name__}: {e} || {' ~ '.join(traceback.format_exc().splitlines())[-800:]}"
        sm.dump(f"ERROR_{t.idsubsls}")
    # Setelah kegagalan: jangan tinggalkan centang yang bisa ikut terkirim
    # di aksi massal subsls berikutnya.
    if pilih:
        try:
            sm.tutup_menu()
            sm.lepas_centang_kode([b.kode for b in pilih])
        except Exception as e:  # noqa: BLE001
            # Centang yang tertinggal bisa ikut terkirim di aksi massal
            # berikutnya -> paksa batch berhenti (CENTANG_GAGAL).
            hasil["pesan"] += f" | status asli {hasil['status']}; ⚠️ gagal melepas centang: {e}"
            hasil["status"] = "CENTANG_GAGAL"
    return hasil


def tulis_console(targets) -> Path:
    """Suntikkan daftar target ke template ubah_moda_console.js -> file siap
    tempel di DevTools Console Chrome (berisi email PPL -> .gitignore)."""
    teks = KONSOL_TEMPLATE.read_text(encoding="utf-8")
    if teks.count(PENANDA_TARGET) != 1:
        raise ValueError(f"Penanda {PENANDA_TARGET} harus muncul tepat 1x di {KONSOL_TEMPLATE.name}")
    data = [{"idsubsls": t.idsubsls, "ppl": list(t.akun_ppl), "baris": t.baris_sheet, "siap": t.siap_input,
             **({"kode": t.kode} if t.kode else {}), "ke": t.ke}
            for t in targets]
    lokasi.siapkan(KONSOL_SIAP).write_text(teks.replace(PENANDA_TARGET, json.dumps(data, ensure_ascii=False)), encoding="utf-8")
    return KONSOL_SIAP


def laporan_cek_subsls(targets, tidak_dikenali, ganda, kode_identitas, sumber: str, ke: str):
    print(f"=== TARGET GANTI MODE (idsubsls) -> {ke} dari {sumber} ===")
    if ke == "CAPI":
        print(f"  {len(targets)} subsls — SEMUA assignment PAPI yang tampil di tiap subsls diubah ke CAPI")
    else:
        print(f"  {len(targets)} subsls — diubah ke PAPI menurut --cakupan")
    if ganda:
        print(f"  {len(ganda)} subsls GANDA dilewati (hanya diproses sekali), mis.: "
              + ", ".join(f"baris {no}: {k}" for no, k in ganda[:5]))
    if kode_identitas:
        print(f"  !! {len(kode_identitas)} baris berisi KODE IDENTITAS, bukan idsubsls — TIDAK dimuat "
              "(kalau hanya kode itu yang ingin diubah, pakai --daftar):")
        for no, isi in kode_identitas[:10]:
            print(f"       baris {no}: {isi}")
    if tidak_dikenali:
        print(f"  !! {len(tidak_dikenali)} angka panjang BUKAN idsubsls 16 digit — TIDAK dimuat "
              "(kolom Excel berformat angka/notasi ilmiah? simpan sbg teks):")
        for no, isi in tidak_dikenali[:10]:
            print(f"       baris {no}: {isi}")
    with lokasi.siapkan(TARGET_CSV).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["idsubsls", "ke", "baris_daftar"])
        for t in targets:
            w.writerow([t.idsubsls, t.ke, ",".join(map(str, t.baris_sheet))])
    print(f"\nDaftar lengkap: {TARGET_CSV}")
    print("Berikutnya (file siap-tempel utk Console Chrome):")
    print(f"  python fasih_sm/ganti_moda/ubah_moda.py --subsls {sumber} --ke {ke} --console")


def laporan_cek_daftar(targets, tidak_dikenali, ganda, sumber: str, ke: str = "PAPI"):
    print(f"=== TARGET GANTI MODE (kode identitas) -> {ke} dari {sumber} ===")
    print(f"  {len(targets)} kode identitas (tiap kode dicari sendiri), "
          f"tersebar di {len({t.idsubsls for t in targets})} idsubsls")
    if ganda:
        print(f"  {len(ganda)} kode GANDA dilewati (hanya diproses sekali), mis.: "
              + ", ".join(f"baris {no}: {k}" for no, k in ganda[:5]))
    if tidak_dikenali:
        print(f"  !! {len(tidak_dikenali)} baris berisi 16 digit tapi BUKAN kode identitas — TIDAK dimuat, periksa:")
        for no, isi in tidak_dikenali[:10]:
            print(f"       baris {no}: {isi}")
    with lokasi.siapkan(TARGET_CSV).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["kode_identitas", "idsubsls", "baris_daftar"])
        for t in targets:
            w.writerow([t.kode, t.idsubsls, ",".join(map(str, t.baris_sheet))])
    print(f"\nDaftar lengkap: {TARGET_CSV}")
    print("Berikutnya (file siap-tempel utk Console Chrome):")
    print(f"  python fasih_sm/ganti_moda/ubah_moda.py --daftar {sumber}{' --ke CAPI' if ke == 'CAPI' else ''} --console")


def laporan_cek(targets, dikeluarkan, sumber: str):
    print(f"=== TARGET GANTI MODE dari {sumber} ===")
    print(f"  {len(targets)} idsubsls ({sum(t.siap_input for t in targets)} punya baris SIAP input), "
          f"{len({p for t in targets for p in t.akun_ppl})} PPL")
    if dikeluarkan:
        pasangan = Counter((a, b) for _, a, b in dikeluarkan)
        print(f"  !! {len(dikeluarkan)} baris DIKELUARKAN (idsubsls ≠ kolom Pilih; {len(pasangan)} pasangan) — "
              "perbaiki di sheet, lalu jalankan ulang:")
        for (a, b), n in pasangan.most_common(8):
            print(f"       idsubsls {a} vs Pilih {b}  ({n} baris)")
    with lokasi.siapkan(TARGET_CSV).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["idsubsls", "akun_ppl", "baris_sheet", "siap_input"])
        for t in targets:
            w.writerow([t.idsubsls, ",".join(t.akun_ppl), ",".join(map(str, t.baris_sheet)), t.siap_input])
    print(f"\nDaftar lengkap: {TARGET_CSV}")
    print("Berikutnya (pemetaan 1 subsls, tidak mengubah apa pun):")
    print(f"  python ubah_moda.py --sumber {sumber} --petakan")


def main():
    ap = argparse.ArgumentParser(description="Ganti mode assignment fasih-sm: CAPI -> PAPI (bawaan) atau PAPI -> CAPI (--ke CAPI)")
    sumber = ap.add_mutually_exclusive_group(required=True)
    sumber.add_argument("--sumber", help="input_usaha.xlsx (tab input_usaha) atau .csv-nya — skrip memilih 1 CAPI per subsls")
    sumber.add_argument("--daftar",
                        help="List KODE IDENTITAS milikmu (.xlsx/.csv/.txt, mis. '5108060003000402 - UMK - 4'); "
                             "HANYA kode itu yang diubah, --cakupan diabaikan")
    sumber.add_argument("--subsls",
                        help="List IDSUBSLS (.xlsx/.csv/.txt, atau langsung '5108...,5108...'); WAJIB dgn --ke. "
                             "--ke CAPI: SEMUA assignment PAPI tiap subsls diubah ke CAPI")
    ap.add_argument("--ke", choices=MODE, default=None,
                    help="Mode tujuan. Bawaan PAPI utk --sumber/--daftar; WAJIB ditulis utk --subsls")
    ap.add_argument("--sheet", default=None, help="Nama sheet utk --daftar/--subsls .xlsx (default: sheet pertama)")
    ap.add_argument("--cek", action="store_true", help="Hanya daftar target, tanpa browser")
    ap.add_argument("--console", action="store_true",
                    help=f"Tulis {KONSOL_SIAP} utk ditempel di DevTools Console Chrome biasa (DISARANKAN)")
    ap.add_argument("--petakan", action="store_true",
                    help="Buka browser, cari 1 subsls, rekam tabel/menu/API. Tidak mencentang/mengubah apa pun.")
    ap.add_argument("--eksekusi", action="store_true", help="LIVE — benar-benar klik Ganti Mode. Default: dry-run.")
    ap.add_argument("--cakupan", choices=("satu", "semua"), default="satu",
                    help="satu (default): pastikan ada 1 assignment PAPI (petugas siapa pun) per subsls; "
                         "semua: ubah seluruh CAPI")
    ap.add_argument("--hanya-siap", action="store_true",
                    help="Hanya subsls yang punya baris SIAP input (lihat input_usaha/jalankan.py --cek)")
    ap.add_argument("--idsubsls", default=None, help="Batasi ke idsubsls tertentu, pisah koma")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--lewati-selesai", action="store_true", help=f"Lewati idsubsls yang sudah tuntas di {AUDIT_PATH}")
    ap.add_argument("--assignment-id", default=ASSIGNMENT_ID_GABUNGAN, help="Segmen periode survei di URL list")
    ap.add_argument("--per-page", type=int, default=100)
    ap.add_argument("--jeda-detik", type=float, default=4.0, help="Jeda antar subsls (jangan terlalu cepat)")
    ap.add_argument("--maks-error-beruntun", type=int, default=3)
    ap.add_argument("--maks-per-klik", type=int, default=50,
                    help="Baris maksimal per klik Ganti Mode (target subsls). Sebelum ada satu perubahan terbukti "
                         "ke arah itu di audit: 1. Hasil pencarian >1 halaman: maks 10")
    args = ap.parse_args()

    # Arah melekat pada target. --subsls tanpa --ke ditolak: salah arah = seluruh subsls terbalik.
    if args.subsls and not args.ke:
        ap.error("--subsls wajib dgn --ke CAPI (PAPI -> CAPI) atau --ke PAPI")
    if args.sumber and args.ke == "CAPI":
        ap.error("--sumber hanya utk ke PAPI (membuka '+Dokumen Baru'); utk ke CAPI pakai --subsls atau --daftar")
    # Ke CAPI lewat menu ⋮ per baris (item massal "Ganti Mode (Ke CAPI)" tidak diketahui ada) — alur itu
    # hanya ada di Console. Jalur Playwright di file ini cuma mengenal item massal.
    if args.ke == "CAPI" and not (args.cek or args.console):
        ap.error("ke CAPI hanya lewat Console Chrome: tambahkan --cek atau --console")
    ke = args.ke or "PAPI"
    tidak_dikenali, ganda, kode_identitas, dikeluarkan = [], [], [], []
    if args.daftar:
        targets, tidak_dikenali, ganda = target_dari_daftar_kode(baca_daftar(args.daftar, args.sheet))
        for t in targets:
            t.ke = ke
    elif args.subsls:
        baris_teks = baca_daftar(args.subsls, args.sheet) if Path(args.subsls).exists() else [args.subsls]
        targets, tidak_dikenali, ganda, kode_identitas = target_dari_daftar_subsls(baris_teks, ke)
    else:
        rows = load_gabungan(args.sumber)
        targets, dikeluarkan = bangun_target(rows, periksa_semua(rows), args.hanya_siap)
    if args.idsubsls:
        ingin = {s.strip() for s in args.idsubsls.split(",") if s.strip()}
        tak_dikenal = sorted(ingin - {t.idsubsls for t in targets})
        if tak_dikenal:
            print(f"⚠️ idsubsls bukan target (tidak ada di sheet/daftar / dikeluarkan): {tak_dikenal}")
        targets = [t for t in targets if t.idsubsls in ingin]
    if args.cek:
        if args.daftar:
            laporan_cek_daftar(targets, tidak_dikenali, ganda, args.daftar, ke)
        elif args.subsls:
            laporan_cek_subsls(targets, tidak_dikenali, ganda, kode_identitas, args.subsls, ke)
        else:
            laporan_cek(targets, dikeluarkan, args.sumber)
        return 0
    if args.console:
        if not targets:
            print("Tidak ada target — file siap-tempel TIDAK ditulis. Periksa dgn --cek.")
            return 1
        path = tulis_console(targets)
        if args.daftar or args.subsls:
            satuan = "kode identitas" if args.daftar else "subsls"
            ditolak = len(tidak_dikenali) + len(kode_identitas)
            print(f"{path} ditulis: {len(targets)} {satuan} -> {ke}"
                  + (f" — ⚠️ {ditolak} baris tidak dimuat, lihat --cek" if ditolak else "") + ".")
        else:
            print(f"{path} ditulis: {len(targets)} subsls ({len(dikeluarkan)} baris sheet dikeluarkan — lihat --cek).")
        print("Chrome biasa -> login fasih-sm -> buka list dgn perPage=100 -> F12 Console -> tempel isi file itu ->")
        print('  await ubahModa.jalankan({mode: "petakan"})')
        return 0

    jalan = "petakan" if args.petakan else ("live" if args.eksekusi else "dry-run")
    riwayat = riwayat_audit()
    if args.lewati_selesai and jalan != "petakan":
        tuntas = STATUS_TUNTAS_LIVE if jalan == "live" else STATUS_TUNTAS_DRY
        sebelum = len(targets)
        targets = [t for t in targets if not masih_tuntas(t, riwayat, tuntas)]
        print(f"--lewati-selesai: {sebelum - len(targets)} target dilewati.")
    if args.petakan:
        targets = targets[:1]
    elif args.limit:
        targets = targets[: args.limit]
    if not targets:
        print("Tidak ada target yang perlu diproses.")
        return 0

    rincian = "SEMUA PAPI tiap subsls" if (ke == "CAPI" and not args.daftar) else f"cakupan '{args.cakupan}'"
    print(f"{jalan.upper()} — {len(targets)} target -> {ke}, {rincian}.")
    if jalan == "live":
        jawab = input(f"⚠️ Ketik 'YA' utk MENGUBAH MODE assignment ke {ke} di {len(targets)} target SUNGGUHAN: ")
        if jawab.strip().upper() != "YA":
            print("Dibatalkan.")
            return 1
    # Sebelum ada satu perubahan terbukti ke arah ini, tiap klik hanya 1 baris.
    bukti = any(r.get("jalan") == "live" and r.get("status") == "DIUBAH_TERVERIFIKASI"
                and (r.get("ke") or "PAPI").upper() == ke for r in riwayat)

    from playwright.sync_api import sync_playwright

    url = f"{FASIH_SM_BASE}/app/surveys/{SURVEY_ID}/{args.assignment_id}/data?page=1&perPage={args.per_page}"
    hitung: Counter = Counter()
    kode_keluar = 0
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(PROFIL_DIR), headless=False, viewport={"width": 1440, "height": 900})
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.set_default_timeout(15_000)
        sm = FasihSm(page)
        try:
            sm.buka_list(url)
            error_beruntun = 0
            for i, t in enumerate(targets, start=1):
                siapa = f"kode {t.kode}" if t.kode else f"{t.idsubsls}" + (f" — PPL {', '.join(t.akun_ppl)}" if t.akun_ppl else "")
                print(f"\n=== [{i}/{len(targets)}] {siapa} -> {t.ke} ===")
                # Target subsls diulang per putaran (tiap putaran sudah diverifikasi) sampai tidak ada lagi
                # yang perlu diubah di halaman tampil.
                for putaran in range(1, 51):
                    hasil = proses_target(sm, t, args, jalan, args.maks_per_klik if bukti else 1)
                    append_audit(hasil)
                    if hasil["status"] != "DIUBAH_TERVERIFIKASI":
                        break
                    bukti = True
                    if not (jalan == "live" and diulang_per_subsls(t, args.cakupan)):
                        break
                    print(f"  (putaran {putaran}) terverifikasi {t.ke} — cari lagi yang tersisa")
                    page.wait_for_timeout(int(args.jeda_detik * 1000))
                hitung[hasil["status"]] += 1
                print(f"  -> {hasil['status']}" + (f" ({hasil['pesan'][:250]})" if hasil.get("pesan") else ""))
                if hasil["status"] in STATUS_BERHENTI_SEGERA:
                    print(f"\n⛔ {hasil['status']} — batch DIHENTIKAN. Periksa {LOG_DIR}/ sebelum menjalankan ulang.")
                    kode_keluar = 1
                    break
                error_beruntun = error_beruntun + 1 if hasil["status"].startswith("ERROR_") else 0
                if args.maks_error_beruntun and error_beruntun >= args.maks_error_beruntun:
                    print(f"\n⛔ {error_beruntun} error berturut-turut — batch DIHENTIKAN (VPN/sesi?).")
                    kode_keluar = 1
                    break
                page.wait_for_timeout(int(args.jeda_detik * 1000))
        except Berhenti as e:
            print(f"\n⛔ {e.kode}: {e}")
            kode_keluar = 1
        finally:
            print(f"\nRingkasan: {dict(hitung)}")
            print(f"Audit: {AUDIT_PATH} | rekaman API & dump: {LOG_DIR}/")
            ctx.close()
    return kode_keluar


if __name__ == "__main__":
    sys.exit(main())
