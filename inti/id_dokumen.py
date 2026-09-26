"""ID dokumen fasih-web DI SHEET SUMBER (permintaan user 2026-09-25).

Kenapa: mencari dokumen lewat NAMA tidak bisa diandalkan — nama usaha di sheet
bisa dikoreksi sesudah dokumennya dibuat, lalu (a) `sinkron_list` tidak
menemukan dokumen bernama baru itu di list server dan (b) `kunci` baris ikut
berubah sehingga audit tidak mengenalinya lagi -> dokumen KEDUA dibuat. ID
dokumen (segmen URL entry, UUID) tidak pernah berubah, jadi ia ditulis ke kolom
`JUDUL_ID` di baris sheet itu sendiri begitu dokumen dibuat/dibuka, dan dibaca
lagi di run berikutnya:

- input_usaha: baris ber-ID yang tidak dikenali audit dibuka lewat URL ID itu,
  TIDAK PERNAH dibuat baru;
- sinkron_list: dokumen server dicocokkan lewat ID dulu, baru nama.

Kolom ini bukan kolom data kuesioner: loader format standar & tahap 2 hanya
membaca judul yang mereka kenal, jadi kolom tambahan di paling kanan tidak
mengganggu (dikunci tests/test_id_dokumen.py).

Penulisan ke sheet sengaja HATI-HATI karena sheet itu bahan kerja manusia:
- baris dicari lewat SIDIK isi barisnya (bukan cuma nomor baris) -> sheet yang
  diurutkan/disisipi selagi batch jalan tidak membuat ID masuk baris lain;
  sidik tidak ketemu/ganda -> TIDAK ditulis (audit tetap mencatat URL-nya);
- ID lain yang sudah ada di sel TIDAK ditimpa (dilaporkan);
- sheet berisi RUMUS tidak ditulis (openpyxl membuang nilai hasil rumus);
- berkas dibuka Excel (ada `~$<nama>`) / terkunci -> ditunda, dicoba lagi pada
  penulisan berikutnya; gagal menulis TIDAK PERNAH menghentikan batch;
- dua proses (dua akun) yang menulis sheet yang sama bergiliran lewat berkas
  kunci `<sumber>.idlock`, dan berkas selalu dibaca ulang tepat sebelum ditulis.
"""
from __future__ import annotations

import csv
import hashlib
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

from inti.config import FASIH_WEB_BASE, SURVEY_ID
from inti.gabungan_loader import NAMA_SHEET_DITERIMA, _norm_judul, _sel

JUDUL_ID = "ID Dokumen FASIH"
POLA_ID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
# Berkas kunci yang lebih tua dari ini dianggap sisa proses yang mati.
KUNCI_BASI_DTK = 120


def url_entry(doc_id: str, assignment_id: str) -> str:
    return f"{FASIH_WEB_BASE}/survey/{SURVEY_ID}/{assignment_id}/{doc_id}/entry"


def id_dari_teks(teks) -> str:
    """ID dari URL entry ATAU ID polos; "" kalau bukan ID dokumen."""
    t = _sel(teks).strip().strip("/")
    bagian = [p for p in t.split("/") if p]
    if len(bagian) >= 2 and bagian[-1] == "entry":
        t = bagian[-2]
    return t.lower() if POLA_ID.match(t) else ""


def _lembar_xlsx(wb, format_sumber: str):
    """Worksheet yang SAMA dgn yang dibaca loader format itu."""
    per_nama = {w.title.strip().lower(): w for w in wb.worksheets}
    if format_sumber == "tahap2":
        from inti.tahap2_loader import NAMA_SHEET_TAHAP2
        return next((per_nama[n] for n in NAMA_SHEET_TAHAP2 if n in per_nama), wb.worksheets[0])
    return next((per_nama[n] for n in NAMA_SHEET_DITERIMA if n in per_nama), None)


def _kolom_id(judul: list) -> int | None:
    target = _norm_judul(JUDUL_ID)
    return next((i for i, j in enumerate(judul) if _norm_judul(_sel(j)) == target), None)


def sidik(nilai: list, kolom_id: int | None) -> str:
    """Sidik isi satu baris (tanpa kolom ID, sel kosong di ujung dibuang)."""
    isi = [_sel(v) for i, v in enumerate(nilai) if i != kolom_id]
    while isi and not isi[-1]:
        isi.pop()
    return hashlib.sha1("\x1f".join(isi).encode("utf-8")).hexdigest()[:16]


def _baca_mentah(path: Path, format_sumber: str) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = _lembar_xlsx(wb, format_sumber)
            return [] if ws is None else [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [list(r) for r in csv.reader(f)]


@dataclass
class IsiKolomId:
    kolom: int | None = None                      # indeks kolom ID (None = belum ada)
    mentah: dict = field(default_factory=dict)     # {baris: teks sel ID apa adanya}
    sidik: dict = field(default_factory=dict)      # {baris: sidik isi baris}


def baca_kolom_id(path, format_sumber: str = "standar") -> IsiKolomId:
    mentah = _baca_mentah(Path(path), format_sumber)
    if not mentah:
        return IsiKolomId()
    kolom = _kolom_id(mentah[0])
    out = IsiKolomId(kolom=kolom)
    for nomor, baris in enumerate(mentah[1:], start=2):
        out.sidik[nomor] = sidik(baris, kolom)
        if kolom is not None and kolom < len(baris):
            out.mentah[nomor] = _sel(baris[kolom])
    return out


def pasang_id(rows, hasil: dict, isi: IsiKolomId) -> None:
    """Isi `row.id_dokumen`/`row.sidik_sumber` & tambahkan masalah offline:
    - sel terisi tapi bukan ID dokumen -> ID_DOKUMEN_TIDAK_VALID (di-skip; kalau
      diabaikan diam-diam, barisnya bisa dibuatkan dokumen baru = ganda);
    - ID sama di >1 baris (salah salin/fill-down) -> ID_DOKUMEN_GANDA (di-skip)."""
    per_id: dict[str, list[int]] = {}
    for r in rows:
        r.sidik_sumber = isi.sidik.get(r.baris, "")
        teks = isi.mentah.get(r.baris, "")
        if not teks:
            continue
        i = id_dari_teks(teks)
        if not i:
            if r.baris in hasil:
                hasil[r.baris].masalah.append(
                    ("ID_DOKUMEN_TIDAK_VALID", f"kolom '{JUDUL_ID}' berisi '{teks[:60]}' — bukan ID dokumen "
                                               "fasih-web (UUID); kosongkan atau perbaiki"))
            continue
        r.id_dokumen = i
        per_id.setdefault(i, []).append(r.baris)
    for i, daftar in per_id.items():
        if len(daftar) > 1:
            for r in rows:
                if r.baris in daftar and r.baris in hasil:
                    hasil[r.baris].masalah.append(
                        ("ID_DOKUMEN_GANDA", f"ID dokumen {i[:8]} tertulis di baris {daftar} — satu dokumen "
                                             "hanya milik satu baris; perbaiki kolom "
                                             f"'{JUDUL_ID}'"))


class PencatatIdSumber:
    """Menulis ID dokumen ke kolom `JUDUL_ID` sheet sumber. `catat()` tidak
    pernah melempar exception — gagal menulis cuma dicetak & ditunda."""

    def __init__(self, path, format_sumber: str = "standar", cetak=print):
        self.path = Path(path)
        self.format = format_sumber
        self.cetak = cetak
        self.tertunda: dict[int, tuple[str, str]] = {}   # baris -> (sidik, id)
        self.tertulis = 0
        self.dimatikan = ""       # alasan permanen (mis. sheet berisi rumus)
        self.beres: dict[int, str] = {}   # baris -> ID yang sudah ada di sel (tidak perlu ditulis lagi)
        self._pesan_tunda = 0.0

    # --- dipanggil input_usaha ---
    def catat(self, row, url_atau_id: str) -> None:
        try:
            i = id_dari_teks(url_atau_id)
            if not i or i in (getattr(row, "id_dokumen", ""), self.beres.get(row.baris)):
                return
            if getattr(row, "id_dokumen", ""):
                self.cetak(f"  ⚠ ID dokumen di sheet baris {row.baris} ({row.id_dokumen[:8]}) beda dgn dokumen "
                           f"yang dipakai ({i[:8]}) — sel TIDAK ditimpa, periksa manual.")
                return
            if self.dimatikan:
                return
            self.antre(row, i)
            self.simpan()
        except Exception as e:  # noqa: BLE001 — pencatatan ID tidak boleh mematikan batch
            self.cetak(f"  ⚠ Gagal mencatat ID dokumen ke {self.path.name}: {e}")

    def antre(self, row, doc_id: str) -> None:
        """Masukkan antrean TANPA langsung menulis (pengisian massal: satu simpan())."""
        self.tertunda[row.baris] = (getattr(row, "sidik_sumber", ""), doc_id)

    def simpan(self) -> bool:
        """Tulis semua yang tertunda. True kalau tidak ada lagi yang tertunda."""
        if not self.tertunda or self.dimatikan:
            return not self.tertunda
        excel = self.path.with_name("~$" + self.path.name)
        if excel.exists():
            self._tunda(f"{self.path.name} sedang dibuka di Excel (ada {excel.name})")
            return False
        kunci = self._kunci()
        if kunci is None:
            self._tunda(f"berkas kunci {self.path.name}.idlock dipegang proses lain")
            return False
        try:
            return self._tulis()
        except PermissionError as e:
            self._tunda(f"{self.path.name} tidak bisa ditulis ({e})")
            return False
        finally:
            kunci.unlink(missing_ok=True)

    def ringkasan(self) -> str:
        if self.dimatikan:
            return f"ID dokumen TIDAK ditulis ke {self.path.name}: {self.dimatikan}"
        teks = f"ID dokumen ditulis ke kolom '{JUDUL_ID}' {self.path.name}: {self.tertulis} baris"
        if self.tertunda:
            teks += (f"; {len(self.tertunda)} baris TERTUNDA (berkas terkunci?) — tutup Excel lalu jalankan "
                     f"input_usaha/tulis_id_sumber.py --sumber {self.path} --tulis")
        return teks

    # --- dalaman ---
    def _tunda(self, alasan: str) -> None:
        if time.time() - self._pesan_tunda > 300:   # jangan membanjiri log tiap baris
            self._pesan_tunda = time.time()
            self.cetak(f"  ⚠ ID dokumen ditunda ({len(self.tertunda)} baris): {alasan}. Dicoba lagi nanti; "
                       "audit tetap mencatat URL-nya.")

    def _kunci(self) -> Path | None:
        p = self.path.with_name(self.path.name + ".idlock")
        batas = time.time() + 20
        while True:
            try:
                fd = os.open(p, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(fd, str(os.getpid()).encode())
                os.close(fd)
                return p
            except FileExistsError:
                try:
                    if time.time() - p.stat().st_mtime > KUNCI_BASI_DTK:
                        p.unlink(missing_ok=True)
                        continue
                except FileNotFoundError:
                    continue
                if time.time() > batas:
                    return None
                time.sleep(0.5)

    def _tempatkan(self, baris_isi: dict[int, list], kolom: int | None) -> dict[int, int]:
        """{baris tertunda: nomor baris tujuan di berkas SEKARANG}. Baris yang
        isinya tidak ketemu (atau ketemu >1) tidak ditempatkan & dilaporkan."""
        per_sidik: dict[str, list[int]] = {}
        for n, isi in baris_isi.items():
            per_sidik.setdefault(sidik(isi, kolom), []).append(n)
        tujuan = {}
        for baris, (s, i) in list(self.tertunda.items()):
            if s and baris in baris_isi and sidik(baris_isi[baris], kolom) == s:
                tujuan[baris] = baris
            elif s and len(per_sidik.get(s, [])) == 1:
                tujuan[baris] = per_sidik[s][0]
                self.cetak(f"  ℹ Baris {baris} sheet sudah bergeser ke baris {tujuan[baris]} — ID ditulis di sana.")
            else:
                self.cetak(f"  ⚠ Baris {baris} tidak ketemu lagi di {self.path.name} (isinya berubah/diurutkan) — "
                           f"ID {i[:8]} TIDAK ditulis ke sheet (tetap ada di audit).")
                del self.tertunda[baris]
        return tujuan

    def _boleh_isi(self, n: int, lama: str, i: str) -> bool:
        """Sel kosong -> isi. Sudah berisi ID yang sama -> tidak perlu. Berisi apa pun
        yang lain -> TIDAK ditimpa (dilaporkan)."""
        if not lama:
            return True
        if id_dari_teks(lama) != i:
            self.cetak(f"  ⚠ Sel '{JUDUL_ID}' baris {n} sudah berisi '{lama[:40]}' (bukan {i[:8]}) — "
                       "TIDAK ditimpa, periksa manual.")
        return False

    def _tulis(self) -> bool:
        if self.path.suffix.lower() in (".xlsx", ".xlsm"):
            return self._tulis_xlsx()
        return self._tulis_csv()

    def _tulis_xlsx(self) -> bool:
        import openpyxl
        wb = openpyxl.load_workbook(self.path, keep_vba=self.path.suffix.lower() == ".xlsm")
        ws = _lembar_xlsx(wb, self.format)
        if ws is None:
            self.dimatikan = "tab sumber tidak ditemukan"
            return False
        if any(c.data_type == "f" for r in ws.iter_rows() for c in r):
            self.dimatikan = ("sheet berisi RUMUS — openpyxl akan membuang nilai hasil rumusnya; "
                              f"salin sebagai nilai (Paste Values) dulu, atau isi kolom '{JUDUL_ID}' manual")
            self.cetak(f"  ⚠ {self.ringkasan()}")
            return False
        judul = [c.value for c in ws[1]]
        kolom = _kolom_id(judul)
        baris_isi = {r[0].row: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
        tujuan = self._tempatkan(baris_isi, kolom)
        if not tujuan:
            return not self.tertunda
        if kolom is None:
            # Di KANAN seluruh sel terpakai (bukan cuma judul) — tidak menimpa data tanpa judul.
            kolom = ws.max_column
            ws.cell(row=1, column=kolom + 1, value=JUDUL_ID)
        ditulis = []
        for baris, n in tujuan.items():
            i = self.tertunda.pop(baris)[1]
            sel = ws.cell(row=n, column=kolom + 1)
            if self._boleh_isi(n, _sel(sel.value), i):
                sel.value = i
                ditulis.append(baris)
            self.beres[baris] = i
        if ditulis:
            tmp = self.path.with_name(f".{self.path.stem}.id-tmp{self.path.suffix}")
            wb.save(tmp)
            os.replace(tmp, self.path)
            self.tertulis += len(ditulis)
        return not self.tertunda

    def _tulis_csv(self) -> bool:
        with self.path.open(newline="", encoding="utf-8-sig") as f:
            data = [list(r) for r in csv.reader(f)]
        if not data:
            return False
        kolom = _kolom_id(data[0])
        baris_isi = {n: isi for n, isi in enumerate(data[1:], start=2)}
        tujuan = self._tempatkan(baris_isi, kolom)
        if not tujuan:
            return not self.tertunda
        if kolom is None:
            kolom = max(len(r) for r in data)
            for r in data:
                r.extend([""] * (kolom + 1 - len(r)))
            data[0][kolom] = JUDUL_ID
        ditulis = 0
        for baris, n in tujuan.items():
            i = self.tertunda.pop(baris)[1]
            isi = data[n - 1]
            isi.extend([""] * (kolom + 1 - len(isi)))
            if self._boleh_isi(n, _sel(isi[kolom]), i):
                isi[kolom] = i
                ditulis += 1
            self.beres[baris] = i
        if ditulis:
            tmp = self.path.with_name(f".{self.path.stem}.id-tmp{self.path.suffix}")
            with tmp.open("w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f).writerows(data)
            os.replace(tmp, self.path)
            self.tertulis += ditulis
        return not self.tertunda
