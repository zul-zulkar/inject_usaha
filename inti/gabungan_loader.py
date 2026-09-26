"""
gabungan_loader.py — DASAR pembaca & pemeriksa baris usaha: GabunganRow, periksa_semua,
OPSI_FORM (opsi form dikutip dari dump DOM), aturan nama dokumen / alamat / koordinat / HP / NIK.
Dipakai format input usaha yang sekarang jadi standar tunggal (inti/tahap2_loader.py membangun
Tahap2Row di atas GabunganRow), plus pembaca format LAMA Agenda (92 kolom, tab "gabungan" /
"input_usaha"; hanya lewat --format agenda) — nama modul ini berasal dari format lama itu.

Prinsip (kedua format):
- Setiap kolom sheet SUDAH jawaban final per rincian form — angka diketik APA ADANYA
  (penggantian nilai hanya lewat aturan TAHAP2_* yang tercatat di review_disarankan).
- Tidak ada kolom No/assignment_id. Identitas baris = `kunci` (hash akun
  PPL + idsubsls + nama dokumen), stabil walau urutan sheet berubah.

CARA EXPORT
-----------
Buka tab "gabungan", lalu File > Download > Microsoft Excel (.xlsx), atau
Comma Separated Values (.csv — hanya tab yang sedang aktif yang ikut).

Kolom dicocokkan lewat AWALAN judul (mis. "26. a."), bukan posisi huruf
kolom. Spasi, baris baru & spasi di sekitar titik dinormalkan dulu, jadi
"8.c. Alamat" dan "8. c. Alamat" sama-sama cocok.

PRINSIP PEMERIKSAAN: lebih baik skip sebelum dokumen dibuat drpd baru
ketahuan GALAT setelah dokumen terlanjur ada (hapus dokumen = soft-delete
lewat admin pusat). Pilihan opsi radio divalidasi thd daftar opsi yang
DIKUTIP dari dump DOM asli (log_screenshots/*.map.tsv), bukan tebakan.
"""

from __future__ import annotations

import csv
import datetime
import hashlib
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from inti.config import (
    GABUNGAN_13F_DARI_13A, GABUNGAN_IZINKAN_JALAN_KOSONG, GABUNGAN_MODE_MURNI, KBLI_DITOLAK_PAKAI_GENAI,
    LENGKAPI_13A_DGN_KBLI,
    LENGKAPI_13F_DGN, MIN_KARAKTER_13A, MIN_KARAKTER_13F, MINIMAL_TOTAL_RUPIAH,
    MINIMAL_TOTAL_RUPIAH_BULANAN, WILAYAH_BY_IDSUBSLS,
)

# Nama tab format standar. "gabungan" = nama lama (sheet "Agenda" BPS Buleleng) — tetap
# diterima supaya file lama tidak perlu diubah; kalau keduanya ada, "input_usaha" dipakai.
NAMA_SHEET = "input_usaha"
NAMA_SHEET_LAMA = ("gabungan",)
NAMA_SHEET_DITERIMA = (NAMA_SHEET, *NAMA_SHEET_LAMA)

# 8b "Nama komersial usaha/perusahaan": validasi form "Panjang maksimal 50"
# (GALAT ringkasan, run live 2026-09-14).
MAKS_8B = 50
# 12c Umur: GALAT "Wajib terisi 10-99" (run live 2026-09-14, Agenda1-1.xlsx
# berisi umur 0 di SEMUA baris -> tiap baris jadi DRAFT yang tak bisa dikirim).
UMUR_MIN, UMUR_MAKS = 10, 99
# 13c "Di mana usaha tersebut biasa dilakukan?": usaha makan-minum hanya boleh
# memakai kode 5-11 ("Kedai, stan, tenda" .. "Daring") — GALAT form "Usaha Makan
# Minum, maka lokasi hanya bisa diisi kode 5-11". Terjadi 8x di audit gabungan
# 22-23 Sep 2026, SEMUANYA KBLI golongan 56 (56102 warung nasi/soto, 56304 kedai
# minuman) dgn 13c sheet berkode 1-4. Dokumen sudah terbuat lalu nyangkut DRAFT
# ber-GALAT, jadi lebih murah dicegat di cek offline.
LOKASI_MAMIN_MIN = 5
# 26a "Total upah dan gaji, serta jaminan sosial pegawai" vs 24a2 (pekerja
# dibayar) — dua aturan file-validation `gaji` yang sama-sama muncul di audit
# 23 Sep 2026:
#   24a2 = 0  -> 26a WAJIB 0    (GALAT "Wajib terisi = 0, karena jumlah pekerja
#                                dibayar=0"; 3x — baris 1761/1762/1766)
#   24a2 > 0  -> 26a/24a2 WAJIB > Rp 50.000 (GALAT "Nilai R26a/R24a2 wajib >
#                                Rp 50.000 jika R24a2 (pekerja dibayar) > 0"; 1x
#                                — baris 1680)
# Aturan kedua sudah lama disebut di catatan KOREKSI_PEKERJA, tapi belum pernah
# diperiksa. Keduanya murni soal isian sheet: gaji/pekerja dibayar harus
# dibetulkan di Excel, jalan ulang tidak menolong.
GAJI_MIN_PER_PEKERJA_DIBAYAR = 50_000
# 13g: form menolak KBLI yang kategorinya P atau U — GALAT "Kategori tidak boleh
# berisi P atau U" (2x di audit 23 Sep 2026, baris 1490 KBLI 98100; log yang sama
# mencatat 13h terbaca "kategori U" utk kode itu). Kategori P & U bukan usaha yang
# dicakup SE2026, jadi KBLI barisnya yang salah — harus dibetulkan di Excel.
# Daftar di bawah HANYA golongan yang sudah terbukti/kategorinya tidak ambigu:
#   85 -> P (Pendidikan)
#   98 -> dibaca form sbg U (TERBUKTI live: KBLI 98100 ditolak)
#   99 -> U (Aktivitas badan internasional)
# Golongan 97 SENGAJA tidak dimasukkan: KBLI 2020 menaruhnya di kategori T
# bersama 98, tapi kita belum punya bukti bagaimana form membacanya — lebih baik
# lolos cek offline drpd men-skip baris yang sebenarnya sah.
KBLI_GOLONGAN_KATEGORI_DITOLAK = {"85": "P", "98": "U", "99": "U"}

# key internal -> AWALAN judul kolom (sesudah _norm_judul). Key sengaja
# disamakan dgn dataKey fasih-web kalau ada padanannya, supaya fill_gabungan
# bisa langsung memakainya. Pengecualian: "nib_nomor" (lihat config.DK).
KOLOM: dict[str, str] = {
    "akun_ppl": "akun ppl",
    "pilih_prov": "pilih provinsi",
    "pilih_kab": "pilih kabupaten",
    "pilih_kec": "pilih kecamatan",
    "pilih_desa": "pilih desa",
    "pilih_sls": "pilih sls",
    "pilih_subsls": "pilih subsls",
    "nama": "nama keluarga/bangunan/usaha",   # muncul 2x identik -> kolom pertama
    "ubah_sls": "8.apakah mengalami perubahan sls",
    "kodepos": "10.kodepos",
    "is_new": "tambah",
    "ada_bang_usaha": "keberadaan bangunan lainnya",
    "jalan_domisili": "nama jalan/gang",
    "nomor_domisili": "blok/nomor rumah",
    "no_bang": "nomor urut bangunan",          # TIDAK dipakai mengisi (aturan keselamatan #3)
    "kode_bang": "kode penggunaan bangunan",
    "latitude": "latitude",
    "longitude": "longitude",
    "pilih_umkm_sls": "pilih umkm dalam satu sls",
    "keberadaan_usaha": "keberadaan usaha",
    "nama_komersial": "8.b.",
    "alamat_usaha_view": "8.c.",               # auto dari SE2026-P, hanya informasi
    "hp": "nomor hp/whatsapp",
    "jenis_kawasan": "8.d.",
    "punya_nib": "10.a.",
    "nib_nomor": "10.b.",
    "tidak_nib": "10.c.",
    "badan_usaha": "11.a.",
    "lap_keuangan": "11.d.",
    "pengusaha": "12.a.",
    "jk": "12.b.",
    "umur": "12.c.",
    "nik_pengusaha": "12.d.",
    "keg_utama": "13.a.",
    "produk_sendiri": "13.b1.",
    "layanan_mamin": "13.b2.",
    "keg_penjualan": "13.b3.",
    "keg_jasa": "13.b4.",
    "lokasi_usaha": "13.c.",
    # 13d/13e dirender form kalau 13b1 = "1. Ya" (usaha memproduksi barang).
    "input_produksi": "13.d.",                  # OPSIONAL — wajib kalau 13b1 = Ya
    "proses_produksi": "13.e.",                 # OPSIONAL — wajib kalau 13b1 = Ya
    "produk": "13.f.",                          # OPSIONAL — kosong = salin 13a (GABUNGAN_13F_DARI_13A)
    "kbli": "pilih dari master kbli",
    "jaringan": "14.a.",
    "internet": "16.a.",
    "internet_pesanan": "16.b1.",
    "internet_produksi": "16.b2.",
    "internet_distribusi": "16.b3.",
    "internet_beli": "16.b4.",
    "internet_promosi": "16.b5.",
    "internet_lainnya": "16.b6.",
    "digital": "16.c.",
    "produksi_lingkungan": "17.a.",
    "perlindungan_lingkungan": "17.b.",
    "produk_seni": "18.",
    # 19 (halal BPJPH) & 20 (izin edar BPOM) hanya dirender utk kategori usaha
    # tertentu. Kolom kosong/tidak ada -> default config (DEFAULT_19A/19C/20B/20C,
    # 20a "3. Tidak"), dicatat ASUMSI di review_disarankan.
    "halal": "19.a.",                           # OPSIONAL
    "sudah_halal": "19.b.",                     # OPSIONAL — wajib kalau form merender 19b
    "belum_halal": "19.c.",                     # OPSIONAL
    "izin_edar": "20.a.",                       # OPSIONAL
    "sudah_bpom": "20.b.",                      # OPSIONAL
    "belum_bpom": "20.c.",                      # OPSIONAL
    "mitra_kdkmp": "21.",
    "peran_mbg": "22.",
    "barang_non_pddk": "23.a.",
    "jasa_non_pddk": "23.b.",
    "beli_jasa_non_pddk": "23.c.",
    "tk_laki": "24.a1.",
    "tk_pr": "24.b1.",
    "tk_dibayar": "24.a2.",
    "tk_tdk_dibayar": "24.b2.",
    "tahun_operasi": "25.",
    "idsubsls": "idsubsls",
    "gaji": "26.a.",
    "biaya_produksi": "26.b.",
    "biaya_pembelian": "26.c.",
    "operasional": "26.d.",
    "non_operasional": "26.e.",
    "nilai_pendapatan": "27.a.",
    "pendapatan_lain": "27.b.",
    "pendapatan_online": "27.d.",
    "aset_usaha_thn": "28.a.",   # label asli: aset TANAH & BANGUNAN
    "aset_lain_thn": "28.b.",    # label asli: aset SELAIN tanah & bangunan
    "luas_tanah_thn": "28.d.",
    "pribadi": "29.a.",
    "non_profit": "29.b.",
    "publik": "29.c.",
    "non_publik": "29.d.",
    "pemerintah": "29.e.",
    "asing": "29.f.",
    "nama_info_list": "nama pemberi informasi",
}
KOLOM_OPSIONAL = {"produk", "input_produksi", "proses_produksi", "halal", "sudah_halal", "belum_halal",
                  "izin_edar", "sudah_bpom", "belum_bpom"}
KEY_JUMLAH_19_20 = ("sudah_halal", "belum_halal", "sudah_bpom", "belum_bpom")

def kbli_26b_wajib_positif(kbli: str) -> bool:
    """Kategori B-F (golongan 05-43) & golongan 56: form mewajibkan 26b/30b > 0
    ("Biaya produksi harus>0 jika kategori usaha B-F dan I (gol 56)")."""
    return len(kbli) >= 2 and kbli[:2].isdigit() and (5 <= int(kbli[:2]) <= 43 or kbli[:2] == "56")


def kbli_tanpa_26c(kbli: str) -> bool:
    """26c (biaya pembelian barang yang dijual kembali) TIDAK dirender.

    26c mengikuti aturan yang sama dgn kembaran bulanannya 30c (enableCondition
    `biaya_pembelian_bln`): hanya perdagangan + valas + pulsa 61209. Dulu di sini
    dipakai aturan kategori B-F/56 saja, sehingga KBLI lain (mis. 01464 peternakan,
    61201 telekomunikasi) lolos pemeriksaan offline lalu baru gagal di tengah
    pengisian (`SKIP_26C_TIDAK_DIRENDER`) — dokumen terlanjur dibuat & nyangkut DRAFT.
    Terbukti dari run live 2026-09-23: 26c terisi & terkirim pada 433 dokumen KBLI 47xx,
    2 dokumen 46xx dan 1 dokumen 612xx; gagal pada 61201 (4 dokumen) & 01464.
    KBLI kosong -> False (tidak menebak)."""
    return bool(kbli) and not kbli_punya_30c(kbli)


def kbli_punya_30c(kbli: str) -> bool:
    """30c (varian bulanan) dirender HANYA utk perdagangan (kategori G, kecuali
    46100/47901/47909), valas (66125/64994) & pulsa (61209) — enableCondition
    template biaya_pembelian_bln — dan saat dirender WAJIB > 0."""
    k = kbli or ""
    if k in ("66125", "64994", "61209"):
        return True
    return k[:2] in ("45", "46", "47") and k not in ("46100", "47901", "47909")


def kbli_makan_minum(kbli: str) -> bool:
    """KBLI golongan 56 (penyediaan makanan & minuman) — dasar aturan 13c
    "Usaha Makan Minum" (lihat LOKASI_MAMIN_MIN). KBLI kosong -> False."""
    return bool(kbli) and kbli[:2] == "56"


def kbli_kategori_ditolak(kbli: str) -> str:
    """Kategori form ("P"/"U") kalau KBLI ini ditolak rincian 13g, "" kalau tidak.
    KBLI kosong -> "" (tidak menebak). Lihat KBLI_GOLONGAN_KATEGORI_DITOLAK."""
    return KBLI_GOLONGAN_KATEGORI_DITOLAK.get((kbli or "")[:2], "")


KATEGORI_KBLI_DITOLAK = tuple(sorted(set(KBLI_GOLONGAN_KATEGORI_DITOLAK.values())))   # ("P", "U")
# Opsi rekomendasi GenAI di radio 13g (#kbli_genai): "[G] 47112 Perdagangan Eceran ..."
# — label jawaban kbli_genai di export fasih-sm asli (value "1" + kode, urut skor
# result_gen_ai). Opsi lain di radio yang sama: "Pilih dari Master KBLI" (value 999999).
POLA_OPSI_GENAI = re.compile(r"^\s*\[([A-Z])\]\s*\[?(\d{5})\]?\s*(.*)$", re.S)


def opsi_kbli_genai(label: str) -> tuple[str, str, str] | None:
    """Label opsi rekomendasi GenAI -> (kategori, kode, judul); None kalau bukan
    rekomendasi (mis. "Pilih dari Master KBLI")."""
    m = POLA_OPSI_GENAI.match(" ".join(str(label or "").split()))
    return (m.group(1), m.group(2), m.group(3).strip()) if m else None


def pilih_opsi_genai(labels: list[str]) -> tuple[int, str]:
    """Indeks rekomendasi yang dipilih dari `labels` (urut tampil) -> (indeks, catatan).
    Yang PERTAMA (ketetapan user 2026-09-24), kecuali kategorinya P/U (form menolak
    13g) -> rekomendasi sah berikutnya, dicatat. Tidak ada yang sah -> (-1, sebab)."""
    lewat = []
    for i, label in enumerate(labels):
        o = opsi_kbli_genai(label)
        if not o:
            continue
        if o[0] in KATEGORI_KBLI_DITOLAK:
            lewat.append(f"[{o[0]}] {o[1]}")
            continue
        return i, (f"rekomendasi GenAI {', '.join(lewat)} dilewati (kategori ditolak form)" if lewat else "")
    return -1, ("semua rekomendasi GenAI berkategori P/U: " + ", ".join(lewat) if lewat
                else "tidak ada opsi rekomendasi GenAI")


def kode_opsi(nilai: str) -> int | None:
    """Angka di depan opsi form ("10. Keliling" -> 10), None kalau tidak ada.
    Perbandingan kode WAJIB numerik: startswith("5") salah utk "11. Daring"."""
    m = re.match(r"\s*(\d+)\s*\.", nilai or "")
    return int(m.group(1)) if m else None


YA_TIDAK = ("1. Ya", "2. Tidak")

# Opsi radio yang BENAR-BENAR ada di form — dikutip dari dump DOM asli
# (log_screenshots/*.map.tsv, label "-label" di bawah tiap komponen).
# Nilai sheet yang tidak persis salah satu dari ini PASTI gagal diklik, jadi
# baris itu di-skip sebelum dokumen dibuat. Contoh nyata 2026-09-13: 14 baris
# LPG ber-11a "1. Perseroan Terbatas (PT)/CV" — opsi itu tidak ada; form
# memisahkan "1.a. Perseroan (PT/NV, ...)" dan "7. Persekutuan Komanditer (CV)".
OPSI_FORM: dict[str, tuple] = {
    "jenis_kawasan": (
        "1. Kawasan Ekonomi Khusus (KEK)", "2. Kawasan Industri (KI)", "3. Stasiun", "4. Bandara",
        "5. Pelabuhan", "6. Terminal", "7. Rest area jalan tol",
        "8. Kawasan sentra ekonomi perdesaan/kelurahan", "9. Kawasan usaha lainnya",
        "10. Di luar kawasan",
    ),
    "punya_nib": YA_TIDAK,
    "tidak_nib": (
        "1. Dalam proses pembuatan NIB", "2. Pengurusan NIB rumit", "3. Tidak memerlukan NIB",
        "4. Tidak tahu tentang NIB", "5. Lainnya",
    ),
    "badan_usaha": (
        "1.a. Perseroan (PT/NV, PT Persero, PT Tbk, PT Persero Tbk, Perseroan Daerah",
        "1.b. Perseroan perorangan", "2. Yayasan", "3. Koperasi", "4. Dana Pensiun",
        "5. Perum/Perumda", "6. BUM Desa", "7. Persekutuan Komanditer (CV)",
        "8. Persekutuan Firma (Fa)", "9. Persekutuan Perdata (Maatschap)",
        "10. Kantor Perwakilan Luar Negeri", "11. Badan Usaha Luar Negeri",
        "12. Badan Usaha Lainnya (Contoh: BLU, PTN-BH)", "13. Bukan Badan Usaha",
    ),
    "lap_keuangan": YA_TIDAK,
    "jk": ("1. Laki-laki", "2. Perempuan"),
    "produk_sendiri": YA_TIDAK,
    "layanan_mamin": YA_TIDAK,
    "keg_penjualan": YA_TIDAK,
    "keg_jasa": ("1. Jasa", "2. Pertanian, Perikanan, dan Kehutanan"),
    "lokasi_usaha": (
        "1. Apotek", "2. Swalayan", "3. Los Pasar", "4. Toko, ruko, dan sejenisnya",
        "5. Kedai, stan, tenda", "6. Bar", "7. Kelab malam, diskotek", "8. Kafe",
        "9. Restoran, warung makan, dan sejenisnya", "10. Keliling", "11. Daring (online)",
    ),
    "jaringan": (
        "1. Tunggal", "2. Kantor pusat", "3. Cabang", "4. Perwakilan", "5. Pabrik",
        "6. Unit pembantu/penunjang",
    ),
    "internet": YA_TIDAK,
    "internet_pesanan": YA_TIDAK,
    "internet_produksi": YA_TIDAK,
    "internet_distribusi": YA_TIDAK,
    "internet_beli": YA_TIDAK,
    "internet_promosi": YA_TIDAK,
    "internet_lainnya": YA_TIDAK,
    "digital": YA_TIDAK,
    "produksi_lingkungan": ("1. Ya, seluruhnya", "2. Ya, sebagian", "3. Tidak sama sekali"),
    "perlindungan_lingkungan": YA_TIDAK,
    "produk_seni": YA_TIDAK,
    "mitra_kdkmp": YA_TIDAK,
    "peran_mbg": (
        "1. Ya, sebagai SATUAN PELAYANAN PEMENUHAN GIZI (SPPG)", "2. Ya, sebagai supplier",
        "3. Ya, sebagai penerima manfaat MBG (Sekolah, Puskesmas)", "4. Ya, peran lainnya",
        "5. Tidak terlibat MBG",
    ),
    "barang_non_pddk": YA_TIDAK,
    "jasa_non_pddk": YA_TIDAK,
    "beli_jasa_non_pddk": YA_TIDAK,
    # Opsional (kategori tertentu) — dikutip dari dump DOM 2026-09-07.
    "halal": ("1. Ya, oleh BPJPH", "2. Ya, bukan oleh BPJPH", "3. Tidak/Belum", "4. Dalam proses"),
    "izin_edar": ("1. Ya, oleh BPOM", "2. Ya, bukan oleh BPOM", "3. Tidak"),
}

KEY_16B = ("internet_pesanan", "internet_produksi", "internet_distribusi",
           "internet_beli", "internet_promosi", "internet_lainnya")
KEY_PEKERJA = ("tk_laki", "tk_pr", "tk_dibayar", "tk_tdk_dibayar")
# Pola 24 tidak konsisten yang PERSIS ini diganti (ketetapan user 2026-09-14):
# Agenda1-1.xlsx 66 baris & Agenda2.xlsx 57 baris berisi laki 0 + perempuan 2
# tapi dibayar 0 + tidak dibayar 1. Dibayar sengaja tetap 0: 26a di baris itu
# 0, dan form menolak 24a2 > 0 dgn 26a/24a2 <= Rp50.000 (file-validation `gaji`).
# Pola lain TIDAK ditebak -> tetap PEKERJA_24_TIDAK_KONSISTEN.
KOREKSI_PEKERJA = {("0", "2", "0", "1"): ("0", "2", "0", "2")}

# 11a di sheet "1. Perseroan Terbatas (PT)/CV" tidak ada di form. Ketetapan user
# 2026-09-15 (Agenda baris 33/37/40/41/53/64/74/99/103/177), diturunkan dari AWALAN
# nama usaha; awalan lain tidak ditebak (tetap OPSI_TIDAK_ADA_DI_FORM).
BADAN_USAHA_PT_CV = "1. Perseroan Terbatas (PT)/CV"
OPSI_PT = "1.a. Perseroan (PT/NV, PT Persero, PT Tbk, PT Persero Tbk, Perseroan Daerah"
OPSI_CV = "7. Persekutuan Komanditer (CV)"
KOREKSI_BADAN_DARI_AWALAN = (   # (regex awalan, opsi 11a, akhiran yang dipindah ke belakang nama)
    (r"PT\b\.?", OPSI_PT, "PT"),
    (r"CV\b\.?", OPSI_CV, "CV"),
    (r"UD\b\.?", "13. Bukan Badan Usaha", ""),
    (r"SPBU\b", OPSI_PT, ""),
)

# Nama memuat BUMDES -> validasi form `badan_usaha`: "Nama perusahaan mengandung Bum Desa/
# Bumdes/Badan Usaha Milik Desa maka status badan usaha harus berkode 6" (GALAT, Agenda2
# baris 239). Kode 6 sendiri menuntut 11d catatan keuangan Ya & 29e modal pemerintah > 0 dan
# > 29a pribadi. Ketetapan user 2026-09-15: 11a 6, 11d Ya, 29 = pemerintah 100%.
POLA_BUMDES = r"bum\s*des|badan usaha milik desa"
OPSI_BUMDES = "6. BUM Desa"


def koreksi_bumdes(v: dict) -> str:
    """Nama memuat BUMDES tapi 11a bukan kode 6 -> betulkan `v` di tempat, return
    catatan koreksinya ("" kalau tidak ada yang diubah).

    Ketiga isian dibetulkan sekaligus & memang harus bertiga: form menolak kode 6
    yang 11d-nya Tidak, dan menolak kode 6 yang modal pemerintah (29e) tidak
    dominan — membetulkan 11a saja hanya menukar satu GALAT dgn GALAT berikutnya.

    Dipakai format standar (mode non-murni) DAN format tahap 2, supaya aturannya
    tidak ditulis dua kali lalu menyimpang."""
    if (not re.search(POLA_BUMDES, f"{v.get('nama', '')} {v.get('nama_komersial', '')}", flags=re.I)
            or v.get("badan_usaha") == OPSI_BUMDES):
        return ""
    lama = (v.get("badan_usaha"), v.get("lap_keuangan"), "/".join(v.get(k, "") for k in KEY_29))
    v["badan_usaha"], v["lap_keuangan"] = OPSI_BUMDES, "1. Ya"
    v.update({k: "0" for k in KEY_29})
    v["pemerintah"] = "100"
    return (f"BUMDES: 11a/11d/29 {lama} -> ('{OPSI_BUMDES}', '1. Ya', pemerintah 100) "
            "(validasi form, ketetapan user)")

# Nama usaha yang diganti (ketetapan user 2026-09-15, Agenda1-1): baris 137 termuat
# di nama baris 60 ("PUSKESMAS PEMBANTU MUNDUK [BESTALA]", pencarian list bisa membuka
# dokumen yang salah), baris 90 55 karakter. Kunci UPPERCASE nama sheet; kunci baris
# (`GabunganRow.kunci`) tetap dari nama mentah.
KOREKSI_NAMA = {
    "PUSKESMAS PEMBANTU MUNDUK": "PUSKESMAS PEMBANTU DESA MUNDUK",
    "PUSKESMAS PEMBANTU DESA LOKAPAKSA DI BANJAR DINAS SORGA": "PUSTU DESA LOKAPAKSA BANJAR DINAS SORGA",
}


def nama_tampil(nama: str, akhiran_badan: str = "") -> str:
    """Nama sheet setelah KOREKSI_NAMA & pemindahan "PT."/"CV." ke belakang
    ("CV. WIRA ADITYA" -> "WIRA ADITYA, CV"): form menolak 8b yang diawali CV
    (GALAT file-validation `nama_komersial`) & panduannya meletakkan PT/CV di belakang."""
    nama = " ".join((nama or "").split())
    nama = KOREKSI_NAMA.get(nama.upper(), nama)
    if akhiran_badan:
        sisa = re.sub(rf"^\s*{akhiran_badan}\b\.?\s*", "", nama, flags=re.I).strip(" ,;:-/&")
        if sisa and sisa != nama:
            nama = f"{sisa}, {akhiran_badan}"
    return nama
KEY_26 = ("gaji", "biaya_produksi", "biaya_pembelian", "operasional", "non_operasional")
KEY_27 = ("nilai_pendapatan", "pendapatan_lain")
KEY_28 = ("aset_usaha_thn", "aset_lain_thn", "luas_tanah_thn")
KEY_29 = ("pribadi", "non_profit", "publik", "non_publik", "pemerintah", "asing")

# Jalur yang di-hardcode fasih_web.py / input_usaha/jalankan.py. Baris yang meminta
# nilai lain butuh alur berbeda (mis. rincian 9 kalau SLS berubah), jadi
# di-skip — bukan diam-diam diisi dgn nilai hardcode.
NILAI_TETAP: dict[str, tuple] = {
    "ubah_sls": ("2. Tidak",),
    "ada_bang_usaha": ("2. Baru",),
    "keberadaan_usaha": ("2. Baru",),
    "kode_bang": ("", "1. Bangunan Khusus Usaha"),
    "pilih_umkm_sls": ("", "Tidak Ada"),
    "nama_info_list": ("Lainnya",),
}


def format_nama_usaha(nama: str, pemilik: str) -> str:
    """Penamaan usaha format standar (mode normal): "<nama_usaha> (<nama_pemilik>)" —
    ketetapan user 2026-09-14, sama dgn pola 3 record manual backlog lama
    (mis. "WARUNG SEMBAKO (KETUT CONTOH)"). Pemilik = kolom 12a.

    Kalau nama pemilik SUDAH tertulis di nama usaha, yang di luar kurung
    dihapus & hanya yang di dalam kurung dicetak (ketetapan user 2026-09-14,
    supaya muat batas 50 karakter 8b): "PANGKALAN GAS BUDIMAN" + "BUDIMAN"
    -> "PANGKALAN GAS (BUDIMAN)". Cocok = nama pemilik UTUH sbg kata
    (bukan bagian kata lain), tanpa beda huruf besar.

    Hasilnya SELALU tepat satu pasang kurung, di belakang (ketetapan user
    2026-09-14): kurung di dalam nama/pemilik dibuang, isinya dipertahankan.
    Kasus nyata baris 393 — 12a "I Nyoman Contoh Putra Sp.P (K" (kurung tak
    tertutup) — dulu jadi "PRAKTIK DOKTER ) (I Nyoman Contoh Putra Sp.P (K)",
    kini "PRAKTIK DOKTER (I Nyoman Contoh Putra Sp.P K)".

    Dikembalikan apa adanya kalau 12a kosong (baris itu sudah di-skip)."""
    nama = " ".join(nama.split())
    pemilik = " ".join(re.sub(r"[()]", " ", pemilik).split())
    if not nama or not pemilik:
        return nama
    # Pemilik dicari per kata, toleran thd kurung/spasi di antaranya, supaya
    # "SP.P (K)" di nama tetap cocok dgn pemilik "Sp.P K".
    pola = r"[\s()]+".join(re.escape(k) for k in pemilik.split())
    sisa = re.sub(rf"\(?\s*(?<!\w){pola}(?!\w)\s*\)?", " ", nama, flags=re.I)
    # Kurung sisa & pemisah yang tertinggal di tepi ("UD. WIDE/ SUARDANTI").
    sisa = " ".join(re.sub(r"[()]", " ", sisa).split()).strip(" ,;:-/&")
    return f"{sisa} ({pemilik})" if sisa else f"({pemilik})"


def nama_muat(nama: str, pemilik: str) -> str:
    """format_nama_usaha, KECUALI hasilnya > MAKS_8B karakter: nama usaha saja
    TANPA kurung (ketetapan user 2026-09-14 — Agenda1-1.xlsx: 12a berisi jabatan
    "Bidan/Perawat Penanggung Jawab Pustu", bukan nama orang, sehingga 78 baris
    > 50). Kurung di nama dibuang, isinya dipertahankan; kolom 12a di form tetap
    diisi utuh. Masih > MAKS_8B -> baris di-skip 8B_TERLALU_PANJANG."""
    lengkap = format_nama_usaha(nama, pemilik)
    if len(lengkap) <= MAKS_8B:
        return lengkap
    return " ".join(re.sub(r"[()]", " ", nama).split())


def judul_dari_opsi_kbli(teks: str) -> str:
    """"[G][47241]Perdagangan Eceran Beras" (teks opsi/textarea Master KBLI) atau
    "[G] 47241 Perdagangan Eceran Beras" (opsi rekomendasi GenAI) -> "Perdagangan
    Eceran Beras". Teks tanpa awalan kode dikembalikan rapi."""
    t = " ".join(str(teks or "").split())
    m = re.match(r"^\[[A-Z]\]\s*\[?\d{5}\]?\s*(.*?)(?=\s*\[[A-Z]\]\s*\[?\d{5}\]?|$)", t)
    return m.group(1).strip() if m else t


def lengkapi_13f(produk: str, judul_kbli: str) -> str:
    """13f < MIN_KARAKTER_13F -> tambahkan LENGKAPI_13F_DGN ("GAS" -> "GAS ECERAN";
    huruf mengikuti isian asli). Masih kurang / kata itu dikosongkan -> pakai judul
    KBLI spt 13a. Judul juga tidak ada -> kembalikan apa adanya (pemanggil menolak)."""
    nilai = " ".join((produk or "").split())
    if not nilai or len(nilai) >= MIN_KARAKTER_13F:
        return nilai
    if LENGKAPI_13F_DGN:
        kata = LENGKAPI_13F_DGN if nilai == nilai.upper() else LENGKAPI_13F_DGN.title()
        if len(f"{nilai} {kata}") >= MIN_KARAKTER_13F:
            return f"{nilai} {kata}"
    return lengkapi_13a(nilai, judul_kbli, minimal=MIN_KARAKTER_13F)


def lengkapi_13a(keg: str, judul_kbli: str, minimal: int = MIN_KARAKTER_13A,
                 cara: str = LENGKAPI_13A_DGN_KBLI) -> str:
    """13a < `minimal` karakter -> "<13a> (<kata judul KBLI>)" (lihat
    LENGKAPI_13A_DGN_KBLI di config). 13a yang sudah cukup, judul kosong, atau
    `cara` kosong -> 13a apa adanya (pemanggil yang memutuskan berhenti).
    Judul ikut HURUF BESAR kalau 13a ditulis huruf besar."""
    keg = " ".join((keg or "").split())
    judul = judul_dari_opsi_kbli(judul_kbli)
    if len(keg) >= minimal or not judul or cara not in ("sedikit", "penuh"):
        return keg
    if keg and keg == keg.upper():
        judul = judul.upper()
    elif judul == judul.upper():
        judul = judul.title()   # "Menjual Rokok" + "PERDAGANGAN ..." -> "Menjual Rokok (Perdagangan)"
    kata = judul.split()
    ambil = kata if cara == "penuh" else []
    if cara == "sedikit":
        for k in kata:
            ambil.append(k)
            if len(f"{keg} ({' '.join(ambil)})") >= minimal:
                break
    return f"{keg} ({' '.join(ambil)})" if keg else " ".join(ambil)


# Nama Jalan (SE2026-P): isi selain kosong/"-" wajib memuat >= 10 huruf a-z
# (file-validation template, dataKey jalan_domisili).
MIN_HURUF_JALAN = 10


def jumlah_huruf(teks: str) -> int:
    return len(re.findall(r"[a-zA-Z]", teks or ""))


def _teks_kabkota(nama: str) -> str:
    """"BULELENG" -> "KABUPATEN BULELENG"; yang sudah berawalan KABUPATEN/KAB./KOTA dibiarkan."""
    nama = " ".join((nama or "").split()).upper()
    if not nama or re.match(r"^(KABUPATEN|KAB\.?|KOTA)\s", nama):
        return nama
    return f"KABUPATEN {nama}"


def lengkapi_alamat(jalan: str, wilayah: dict) -> str:
    """Nama Jalan yang kurang dari MIN_HURUF_JALAN huruf dilengkapi nama
    wilayah BARIS itu (ketetapan user 2026-09-14): banjar, lalu desa, lalu
    kecamatan, lalu kabupaten (2026-09-24: jalan 'GEROKGAK' di desa & kecamatan
    GEROKGAK -> 'GEROKGAK, KABUPATEN BULELENG'), lalu provinsi — berhenti begitu
    syarat huruf terpenuhi.
    Bagian yang namanya sudah tertulis tidak ditambah lagi ("DESA JULAH" tidak
    diberi "DESA JULAH" lagi). "0" bukan alamat -> dianggap kosong. Kosong
    atau "-" dibiarkan (lolos validasi form). Kalau nama wilayah tidak ada,
    hasilnya tetap pendek dan baris di-skip JALAN_KURANG_10_HURUF."""
    jalan = " ".join((jalan or "").split())
    if not jalan or jalan == "-" or jumlah_huruf(jalan) >= MIN_HURUF_JALAN:
        return jalan
    bagian = [] if jalan == "0" else [jalan]
    kandidat = [
        (wilayah.get("sls", ""), wilayah.get("sls", "")),
        (wilayah.get("desa", ""), f"DESA {wilayah.get('desa', '')}"),
        (wilayah.get("kecamatan", ""), f"KECAMATAN {wilayah.get('kecamatan', '')}"),
        (wilayah.get("kabkota", ""), _teks_kabkota(wilayah.get("kabkota", ""))),
        (wilayah.get("provinsi", ""), wilayah.get("provinsi", "")),
    ]
    for nama, teks in kandidat:
        if jumlah_huruf(", ".join(bagian)) >= MIN_HURUF_JALAN:
            break
        inti = re.sub(r"^(BANJAR|BR\.?|LINGKUNGAN|LINGK\.?|DUSUN|KABUPATEN|KAB\.?|KOTA)\s+", "",
                      nama.strip().upper())
        if not inti or inti in ", ".join(bagian).upper():
            continue
        bagian.append(teks.strip().upper())
    return ", ".join(bagian) if bagian else jalan


def _norm_judul(s) -> str:
    s = " ".join(str(s or "").split()).lower()
    return re.sub(r"\s*\.\s*", ".", s)


def _sel(v) -> str:
    """Nilai sel -> teks rapi. Angka bulat dari xlsx (15056000.0) -> "15056000"."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _cari_indeks(judul: list[str]) -> dict[str, int]:
    norm = [_norm_judul(j) for j in judul]
    idx: dict[str, int] = {}
    hilang, ambigu = [], []
    for key, awalan in KOLOM.items():
        cocok = [i for i, h in enumerate(norm) if h.startswith(awalan)]
        if not cocok:
            if key not in KOLOM_OPSIONAL:
                hilang.append(f"{key} (awalan '{awalan}')")
            continue
        if len({norm[i] for i in cocok}) > 1:
            ambigu.append(f"{key}: {[judul[i] for i in cocok]}")
            continue
        idx[key] = cocok[0]
    if hilang or ambigu:
        raise ValueError(
            f"Judul kolom sheet tidak sesuai. Hilang: {hilang}. Ambigu: {ambigu}. "
            "Sesuaikan KOLOM di gabungan_loader.py kalau judul sheet memang berubah."
        )
    return idx


def _baca_mentah(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl  # hanya perlu kalau sumbernya xlsx
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            per_nama = {w.title.strip().lower(): w for w in wb.worksheets}
            ws = next((per_nama[n] for n in NAMA_SHEET_DITERIMA if n in per_nama), None)
            if ws is None:
                raise ValueError(f"Tab '{NAMA_SHEET}' (atau nama lama {NAMA_SHEET_LAMA}) tidak ada di "
                                 f"{path.name}. Tersedia: {wb.sheetnames}")
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [list(r) for r in csv.reader(f)]


@dataclass
class GabunganRow:
    baris: int   # nomor baris di sheet (judul = baris 1) — yang dilihat user di Sheets
    v: dict
    # Nama wilayah asli baris {provinsi, kabkota, kecamatan, desa[, sls]} —
    # dari tab lain di xlsx (lihat _baca_nama_wilayah) & config. Bisa kosong.
    wilayah: dict = field(default_factory=dict)
    wilayah_bentrok: str = ""   # terisi kalau sumber nama wilayah saling bertentangan
    koreksi: list = field(default_factory=list)  # nilai sheet yang diubah ketetapan user (-> tanda review)
    akhiran_badan: str = ""  # "PT"/"CV" yang dipindah ke belakang nama (lihat nama_tampil)
    # MODE MURNI (config.GABUNGAN_MODE_MURNI): semua isian apa adanya dari sheet,
    # tanpa aturan penamaan/pelengkap/koreksi — lihat komentar di config.
    murni: bool = False
    # Kolom "ID Dokumen FASIH" sheet (inti/id_dokumen.py, diisi muat_sumber) — dokumen
    # baris ini dibuka lewat ID, bukan dicari lewat nama. `sidik_sumber` = sidik isi
    # baris saat dibaca, dipakai menemukan barisnya lagi waktu ID ditulis balik.
    id_dokumen: str = ""
    sidik_sumber: str = ""

    def __getitem__(self, key: str) -> str:
        return self.v.get(key, "")

    @property
    def nama(self) -> str:
        """Nama MENTAH kolom sheet — dipakai `kunci`, BUKAN nama di fasih-web."""
        return self["nama"]

    @property
    def nama_dokumen(self) -> str:
        """Nama yang diketik ke fasih-web ("+Dokumen Baru" & SE2026-P) dan
        dipakai mencari dokumen di list: "<nama> (<12a>)" (lihat nama_muat).
        Mode murni: nama sheet apa adanya."""
        if self.murni:
            return " ".join(self.nama.split())
        return nama_muat(nama_tampil(self.nama, self.akhiran_badan), self["pengusaha"])

    @property
    def nama_komersial(self) -> str:
        """8b dgn format yang sama dgn nama dokumen (ketetapan user: penamaan
        berlaku utk nama usaha DAN nama komersial). Masih lebih dari MAKS_8B
        karakter setelah nama_muat -> skip 8B_TERLALU_PANJANG, tidak dipotong.
        Mode murni: kolom 8b apa adanya."""
        if self.murni:
            return " ".join(self["nama_komersial"].split())
        return nama_muat(nama_tampil(self["nama_komersial"], self.akhiran_badan), self["pengusaha"])

    @property
    def nama_lama_dicari(self) -> str:
        """Nama yang dicari create_document sbg pengaman "dokumen bernama format
        LAMA sudah ada" (lihat DokumenNamaLamaAda). Format standar: nama mentah
        sheet, karena dokumen Buleleng sempat dibuat dgn nama itu sebelum aturan
        "<nama> (<12a>)" berlaku. Format yang belum pernah dipakai membuat
        dokumen mengembalikan "" -> pengaman dilewati (kalau tidak, dua usaha
        bernama sama di sheet saling menyandera)."""
        return self.nama

    @property
    def akun_ppl(self) -> str:
        return self["akun_ppl"].strip().lower()

    @property
    def idsubsls(self) -> str:
        return self["idsubsls"]

    @property
    def idsubsls_pilih(self) -> str:
        """idsubsls yang disusun dari kolom "Pilih PROVINSI..SUBSLS". Sheet
        membuang nol di depan pada sebagian baris ("60" utk "060"), jadi
        di-zero-pad ke lebar kodenya (2+2+3+3+4+2)."""
        lebar = (("pilih_prov", 2), ("pilih_kab", 2), ("pilih_kec", 3),
                 ("pilih_desa", 3), ("pilih_sls", 4), ("pilih_subsls", 2))
        return "".join(self[k].zfill(n) for k, n in lebar)

    @property
    def kunci(self) -> str:
        # Sengaja dari nama MENTAH, bukan nama_dokumen: aturan penamaan boleh
        # berubah tanpa memutus pencocokan --lewati-selesai di audit lama.
        teks = f"{self.akun_ppl}|{self.idsubsls}|{self.nama.upper()}"
        return hashlib.sha1(teks.encode("utf-8")).hexdigest()[:10]

    @property
    def jalan_lengkap(self) -> str:
        """Nama Jalan yang diketik ke SE2026-P (dilengkapi nama wilayah kalau
        kurang dari 10 huruf — lihat lengkapi_alamat). Mode murni: apa adanya."""
        if self.murni:
            return " ".join(self["jalan_domisili"].split())
        return lengkapi_alamat(self["jalan_domisili"], self.wilayah)

    @property
    def nomor_rumah(self) -> str:
        """Blok/Nomor Rumah SE2026-P. Kosong -> "-" (petunjuk form: "Jika tidak ada
        isikan -"); mode murni: apa adanya (kosong = tidak diisi)."""
        if self.murni:
            return self["nomor_domisili"]
        return self["nomor_domisili"] or "-"

    @property
    def produk_utama(self) -> str:
        """13f. Isian sependek "GAS" ditolak form (minimal 4 karakter), jadi
        dilengkapi judul KBLI seperti 13a. Mode murni: apa adanya."""
        nilai = self["produk"] if (self["produk"] or self.murni) else (
            self["keg_utama"] if GABUNGAN_13F_DARI_13A else "")
        return nilai if self.murni else lengkapi_13f(nilai, self.judul_kbli)

    @property
    def rincian_13b4_dirender(self) -> bool:
        return all(self[k].startswith("2") for k in ("produk_sendiri", "layanan_mamin", "keg_penjualan"))

    @property
    def judul_kbli(self) -> str:
        """Kolom "Judul KBLI" sheet (format tahap 2). Format standar tidak punya
        kolomnya -> "" (13a pendek dilengkapi judul opsi Master KBLI di form).
        KBLI sheet ditolak & diganti GenAI -> "" (judulnya milik KBLI yang salah)."""
        return "" if self.kbli_genai else " ".join(self["judul_kbli"].split())

    @property
    def kbli_genai(self) -> bool:
        """KBLI sheet berkategori P/U (ditolak 13g) -> 13g diisi rekomendasi GenAI
        pertama saat pengisian (KBLI_DITOLAK_PAKAI_GENAI). Mode murni: tidak."""
        return KBLI_DITOLAK_PAKAI_GENAI and not self.murni and bool(kbli_kategori_ditolak(self["kbli"]))

    @property
    def b13_dari_kbli(self) -> bool:
        """13b1-b3 diturunkan dari golongan KBLI sheet (aturan tahap 2) — kalau
        True dan 13g diisi GenAI, 13b disesuaikan dgn KBLI terpilih saat pengisian."""
        return False

    @property
    def pindah_26c_ke_26b(self) -> bool:
        """26c yang tidak dirender form dijumlahkan ke 26b (aturan tahap 2)."""
        return False

    @property
    def punya_koordinat(self) -> bool:
        """Latitude & longitude terisi DAN terbaca sbg titik di Indonesia.
        Koordinat yang rusak (mis. "-8.148.438" / "1.145.951" — titik ribuan
        dari Excel, data asli tahap 2) dianggap BELUM ADA: dengan --koordinat
        otomatis barisnya jadi DRAFT sampai koordinatnya diperbaiki di sheet."""
        return koordinat_valid(self["latitude"], self["longitude"])

    @property
    def bulanan_dari_kolom(self) -> bool:
        """True = varian bulanan (usaha mulai beroperasi tahun berjalan) boleh
        diisi dari kolom 26-29 apa adanya. Format standar: False (angka sheet
        Buleleng tahunan -> VARIAN_BULANAN, isi manual)."""
        return False

    def angka(self, key: str) -> int:
        return int(self[key])


KODE_NIK_KHUSUS = ("7777", "8888", "9999")


def nik_valid(nik: str) -> bool:
    """Aturan file-validation nik_pengusaha: 16 digit (tidak boleh digit sama
    semua) ATAU kode 7777 (NIK > 16 digit) / 8888 (belum punya) / 9999 (lainnya)."""
    n = str(nik or "").strip()
    if n in KODE_NIK_KHUSUS:
        return True
    return bool(re.fullmatch(r"\d{16}", n)) and len(set(n)) > 1


def hp_valid(hp: str) -> bool:
    """Aturan file-validation `hp`: "9999" (tidak ada/tidak bersedia) ATAU angka
    diawali 08, 10-13 digit, dan digit setelah 08 tidak sama semua."""
    h = str(hp or "")
    if h == "9999":
        return True
    return bool(re.fullmatch(r"08\d{8,11}", h)) and not re.fullmatch(r"(\d)\1+", h[2:])


def koordinat_valid(lat, lon) -> bool:
    if koordinat_kosong(lat) or koordinat_kosong(lon):
        return False
    try:
        a, b = float(str(lat).strip()), float(str(lon).strip())
    except ValueError:
        return False
    return -12 <= a <= 7 and 94 <= b <= 142


def koordinat_kosong(nilai) -> bool:
    """Kolom koordinat yang BELUM diisi: kosong, tanda strip, atau 0 (Excel
    kerap mengisi 0 utk sel kosong; titik 0 bukan lokasi di Indonesia)."""
    t = " ".join(str(nilai or "").split())
    if t in ("", "-", "–", "—"):
        return True
    try:
        return float(t.replace(",", ".")) == 0
    except ValueError:
        return False


def _baca_nama_wilayah(path: Path) -> dict[str, dict]:
    """{kode desa 10 digit: {provinsi, kabkota, kecamatan, desa}} dari tab LAIN
    di xlsx input usaha yang punya kolom kode "Pilih DESA" + nama "Desa/Kelurahan"
    (tab "Pangkalan Gas" & "Faskes", 2026-09-14). Tab gabungan sendiri hanya
    berisi kode. Tiap baris tab dicatat di bawah kode dari kolom `idsubsls`
    (10 digit pertama) DAN dari kolom "Pilih ...": sebagian baris sheet punya
    kode Pilih yang tidak cocok dgn idsubsls-nya, jadi per kode dipilih nama
    yang PALING SERING muncul (dulu kode bentrok dibuang -> 26 baris gagal)."""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return {}
    import openpyxl
    kolom = {"Pilih PROVINSI": 2, "Pilih KABUPATEN/KOTA": 2, "Pilih KECAMATAN": 3, "Pilih DESA": 3}
    nama = {"Provinsi": "provinsi", "Kabupaten/Kota": "kabkota", "Kecamatan": "kecamatan",
            "Desa/Kelurahan": "desa"}
    suara: dict[str, Counter] = defaultdict(Counter)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if ws.title.strip().lower() in NAMA_SHEET_DITERIMA:
                continue
            it = ws.iter_rows(values_only=True)
            judul = [" ".join(str(j or "").split()) for j in (next(it, None) or [])]
            if not all(k in judul for k in (*kolom, *nama)):
                continue
            ik = {k: judul.index(k) for k in (*kolom, *nama)}
            i_id = judul.index("idsubsls") if "idsubsls" in judul else None
            for sel in it:
                try:
                    kode_pilih = "".join(_sel(sel[ik[k]]).zfill(n) for k, n in kolom.items())
                    info = tuple(_sel(sel[ik[k]]).upper() for k in nama)
                    kode_id = _sel(sel[i_id])[:10] if i_id is not None and i_id < len(sel) else ""
                except IndexError:
                    continue
                if not info[3]:
                    continue
                for kode in {kode_pilih, kode_id}:
                    if len(kode) == 10 and kode.isdigit():
                        suara[kode][info] += 1
    finally:
        wb.close()
    return {kode: dict(zip(nama.values(), c.most_common(1)[0][0])) for kode, c in suara.items()}


def _nama_wilayah_dari_8c(alamat: str) -> dict:
    """Cadangan: "Desa/Kel. TINGA-TINGA, Kec. GEROKGAK, BULELENG, BALI, 81155"
    (kolom 8c baris itu sendiri) -> {desa, kecamatan, kabkota, provinsi}."""
    m = re.search(r"Desa/Kel\.\s*([^,]+),\s*Kec\.\s*([^,]+),\s*([^,]+),\s*([^,]+)", alamat or "", re.I)
    if not m:
        return {}
    desa, kec, kab, prov = (" ".join(x.split()).upper() for x in m.groups())
    return {"desa": desa, "kecamatan": kec, "kabkota": kab, "provinsi": prov}


def load_gabungan(path: str | Path, murni: bool | None = None) -> list[GabunganRow]:
    """`murni` None = config.GABUNGAN_MODE_MURNI. Mode murni: koreksi data
    (KOREKSI_PEKERJA, badan usaha dari awalan, BUMDES) TIDAK diterapkan."""
    murni = GABUNGAN_MODE_MURNI if murni is None else murni
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File sumber tidak ditemukan: {path}")
    mentah = _baca_mentah(path)
    if not mentah:
        raise ValueError(f"{path.name} kosong.")
    idx = _cari_indeks([_sel(j) for j in mentah[0]])
    nama_desa = _baca_nama_wilayah(path)
    out = []
    for nomor, sel in enumerate(mentah[1:], start=2):
        if not any(_sel(x) for x in sel):
            continue
        v = {k: (_sel(sel[i]) if i < len(sel) else "") for k, i in idx.items()}
        row = GabunganRow(nomor, v, murni=murni)
        peta = nama_desa.get(row.idsubsls[:10]) or nama_desa.get(row.idsubsls_pilih[:10]) or {}
        dari_8c = _nama_wilayah_dari_8c(row["alamat_usaha_view"])
        if peta and dari_8c and (peta["desa"], peta["kecamatan"]) != (dari_8c["desa"], dari_8c["kecamatan"]):
            # Kode wilayah & alamat 8c baris ini menunjuk desa berbeda (Agenda baris
            # 278/366, Agenda1-1 baris 41 dst.). Ketetapan user 2026-09-14: pakai desa
            # menurut idsubsls; bentroknya tetap dicatat sbg tanda review.
            row.wilayah_bentrok = (f"desa menurut idsubsls '{peta['desa']}, {peta['kecamatan']}' vs "
                                   f"8c '{dari_8c['desa']}, {dari_8c['kecamatan']}'")
        wil = dict(peta or dari_8c)
        pekerja = tuple(v.get(k, "") for k in KEY_PEKERJA)
        if murni:
            pass  # apa adanya: pelanggaran aturan form dilaporkan periksa_baris
        elif pekerja in KOREKSI_PEKERJA:
            baru = KOREKSI_PEKERJA[pekerja]
            v.update(zip(KEY_PEKERJA, baru))
            row.koreksi.append(f"24 (laki, perempuan, dibayar, tidak dibayar) {'/'.join(pekerja)} -> "
                               f"{'/'.join(baru)} (ketetapan user)")
        if not murni and v.get("badan_usaha") == BADAN_USAHA_PT_CV:
            for pola, opsi, akhiran in KOREKSI_BADAN_DARI_AWALAN:
                if re.match(rf"\s*{pola}", v.get("nama", ""), flags=re.I):
                    v["badan_usaha"] = opsi
                    row.akhiran_badan = akhiran
                    row.koreksi.append(f"11a '{BADAN_USAHA_PT_CV}' -> '{opsi}' dari awalan nama (ketetapan user)")
                    break
        if not murni and (ket_bumdes := koreksi_bumdes(v)):
            row.koreksi.append(ket_bumdes)
        if not murni and " ".join(v.get("nama", "").split()).upper() in KOREKSI_NAMA:
            row.koreksi.append(f"nama usaha diganti -> '{nama_tampil(v['nama'])}' (ketetapan user)")
        ref = WILAYAH_BY_IDSUBSLS.get(row.idsubsls) or {}
        if ref.get("sls"):
            wil["sls"] = ref["sls"]
        row.wilayah = wil
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# Pemeriksaan offline
# ---------------------------------------------------------------------------

# Baris lolos pemeriksaan tapi koordinatnya belum ada (izinkan_tanpa_koordinat):
# dokumen dibuat & diisi lengkap KECUALI geotag, lalu DITAHAN sbg DRAFT — tidak
# pernah dikirim sampai koordinatnya dilengkapi di sheet & skrip dijalankan ulang.
STATUS_SIAP_TANPA_KOORDINAT = "SIAP_TANPA_KOORDINAT"
STATUS_BISA_DIPROSES = ("SIAP", STATUS_SIAP_TANPA_KOORDINAT)


@dataclass
class Pemeriksaan:
    masalah: list[tuple[str, str]] = field(default_factory=list)  # (kode, pesan) -> skip
    tanda: list[str] = field(default_factory=list)                # review, tidak skip
    tanpa_koordinat: bool = False                                 # -> DRAFT, tidak dikirim

    @property
    def status(self) -> str:
        if self.masalah:
            return f"SKIP_DATA_{self.masalah[0][0]}"
        return STATUS_SIAP_TANPA_KOORDINAT if self.tanpa_koordinat else "SIAP"

    @property
    def bisa_diproses(self) -> bool:
        return self.status in STATUS_BISA_DIPROSES

    @property
    def pesan(self) -> str:
        return " | ".join(f"{k}: {p}" for k, p in self.masalah)


def _bulat(s: str) -> bool:
    return s.isdigit()


def periksa_baris(row: GabunganRow, tahun_berjalan: int | None = None,
                  mode_satu_subsls: bool = False, izinkan_tanpa_koordinat: bool = False) -> Pemeriksaan:
    """`mode_satu_subsls` = semua dokumen dibuat di satu subsls: dokumen TIDAK dibuat di
    idsubsls baris, jadi ketidakcocokan wilayah baris baru berarti saat
    ubah alokasi wilayah nanti -> tanda, bukan skip.

    `izinkan_tanpa_koordinat` (--koordinat otomatis): latitude/longitude yang
    belum diisi TIDAK men-skip baris -> status SIAP_TANPA_KOORDINAT (dokumen
    jadi DRAFT tanpa geotag). Tanpa flag ini koordinat tetap WAJIB."""
    tahun_berjalan = tahun_berjalan or datetime.date.today().year
    hasil = Pemeriksaan()
    salah = hasil.masalah.append
    hasil.tanpa_koordinat = izinkan_tanpa_koordinat and not row.punya_koordinat
    if hasil.tanpa_koordinat:
        terisi = [k for k in ("latitude", "longitude") if not koordinat_kosong(row[k])]
        if len(terisi) == 2:
            ket = (f" (koordinat TIDAK TERBACA '{row['latitude']}' / '{row['longitude']}' — "
                   "perbaiki di Excel, jalankan ulang)")
        else:
            ket = f" (hanya {terisi[0]} yang terisi)" if terisi else ""
        hasil.tanda.append("KOORDINAT BELUM ADA -> disimpan sbg DRAFT tanpa geotag, TIDAK dikirim" + ket)

    wajib = [
        "akun_ppl", "idsubsls", "nama", "kodepos", *(() if hasil.tanpa_koordinat else ("latitude", "longitude")),
        "nama_komersial",
        "hp", "jenis_kawasan", "punya_nib", "badan_usaha", "lap_keuangan", "pengusaha", "jk",
        "umur", "nik_pengusaha", "keg_utama", "produk_sendiri", "layanan_mamin", "keg_penjualan",
        "lokasi_usaha", "kbli", "jaringan", "internet", "produksi_lingkungan",
        "perlindungan_lingkungan", "produk_seni", "mitra_kdkmp", "peran_mbg", "barang_non_pddk",
        "jasa_non_pddk", "beli_jasa_non_pddk", "tahun_operasi",
        *KEY_PEKERJA, *KEY_26, *KEY_27, *KEY_28, *KEY_29,
    ]
    # 27d/31d (persen pendapatan online) HANYA dirender kalau 16a = Ya (enableCondition
    # template pendapatan_online_bln; fill_gabungan: "27d tidak dirender" utk 16a Tidak).
    # Dulu selalu diwajibkan -> 95 baris data asli tahap 2 (16a Tidak, 27d kosong) di-skip.
    if row["internet"].startswith("1"):
        wajib.append("pendapatan_online")
    wajib.append("nib_nomor" if row["punya_nib"].startswith("1") else "tidak_nib")
    if row["internet"].startswith("1"):
        wajib += [*KEY_16B, "digital"]
    if not GABUNGAN_IZINKAN_JALAN_KOSONG:
        wajib.append("jalan_domisili")
    if row["produk_sendiri"].startswith("1"):
        wajib += ["input_produksi", "proses_produksi"]   # form merender 13d & 13e
    kosong = [k for k in wajib if not row[k]]
    if not row.produk_utama:
        kosong.append("produk (13f)")
    if kosong:
        salah(("WAJIB_KOSONG", f"kolom kosong: {', '.join(kosong)}"))

    # Validasi form (file-validation template, dataKey jalan_domisili): isi
    # selain kosong/"-" wajib memuat >= 10 huruf a-z — GALAT "Minimal terisi
    # 10 huruf" (baris 17, run 2026-09-14). Sheet berisi "0", "BR. KAJANAN", dst.
    jalan = row.jalan_lengkap
    if jalan and jalan != "-" and jumlah_huruf(jalan) < MIN_HURUF_JALAN:
        sebab = ("mode murni: tidak dilengkapi nama wilayah" if row.murni
                 else f"nama wilayah bertentangan: {row.wilayah_bentrok}" if row.wilayah_bentrok
                 else "nama wilayah baris tidak ditemukan")
        salah(("JALAN_KURANG_10_HURUF", f"Nama Jalan '{jalan}' kurang dari {MIN_HURUF_JALAN} huruf & tidak bisa "
                                        f"dilengkapi — {sebab} (form menolak)"))
    elif jalan != " ".join(row["jalan_domisili"].split()):
        hasil.tanda.append(f"Nama Jalan dilengkapi nama wilayah: '{row['jalan_domisili']}' -> '{jalan}'"
                           + (f" (desa idsubsls dipakai; {row.wilayah_bentrok})" if row.wilayah_bentrok else ""))
    hasil.tanda.extend(row.koreksi)
    keg = " ".join(row["keg_utama"].split())
    if keg and len(keg) < MIN_KARAKTER_13A:
        if row.murni or LENGKAPI_13A_DGN_KBLI not in ("sedikit", "penuh"):
            salah(("13A_KURANG_15_KARAKTER", f"13a '{keg}' kurang dari {MIN_KARAKTER_13A} karakter (form menolak)"))
        elif row.judul_kbli:
            baru = lengkapi_13a(keg, row.judul_kbli)
            if len(baru) < MIN_KARAKTER_13A:
                salah(("13A_KURANG_15_KARAKTER", f"13a '{keg}' + judul KBLI '{row.judul_kbli}' masih kurang dari "
                                                 f"{MIN_KARAKTER_13A} karakter"))
            else:
                hasil.tanda.append(f"13a dilengkapi judul KBLI: '{keg}' -> '{baru}'")
        else:
            hasil.tanda.append(f"13a '{keg}' < {MIN_KARAKTER_13A} karakter -> dilengkapi judul KBLI "
                               f"{row['kbli']} dari Master KBLI saat pengisian")
    for label, nama in (("nama dokumen", row.nama_dokumen), ("8b", row.nama_komersial)):
        if not row.murni and nama and "(" not in nama and row["pengusaha"]:
            hasil.tanda.append(f"{label} tanpa (12a): format lengkap > {MAKS_8B} karakter")
    if row.murni:
        # Aturan form yang di mode non-murni dikoreksi skrip -> di sini dilaporkan.
        if (re.search(POLA_BUMDES, f"{row.nama} {row['nama_komersial']}", flags=re.I)
                and row["badan_usaha"] != OPSI_BUMDES):
            salah(("BUMDES_BUKAN_KODE_6", f"nama memuat BUMDES tapi 11a='{row['badan_usaha']}' — form "
                                          f"mewajibkan '{OPSI_BUMDES}', 11d Ya & modal pemerintah (29e) dominan"))
        if re.match(r"\s*CV\b", row.nama_komersial, flags=re.I):
            salah(("8B_DIAWALI_CV", f"8b '{row.nama_komersial}' diawali CV — form menolak; tulis 'NAMA, CV'"))

    for key, boleh in NILAI_TETAP.items():
        if row[key].lower() not in {b.lower() for b in boleh}:
            salah(("NILAI_TIDAK_DIDUKUNG", f"{key}='{row[key]}' (alur skrip hanya mendukung {boleh})"))
    if not row["is_new"].lower().startswith("bangunan lainnya"):
        salah(("NILAI_TIDAK_DIDUKUNG", f"is_new='{row['is_new']}' (hanya 'Bangunan Lainnya ...')"))

    for key, opsi in OPSI_FORM.items():
        nilai = row[key]
        if not nilai or (key == "keg_jasa" and not row.rincian_13b4_dirender):
            continue  # kosong sudah dilaporkan di atas; 13b4 tidak dirender -> tidak dipakai
        if key == "tidak_nib" and row["punya_nib"].startswith("1"):
            continue  # 10c hanya dirender kalau 10a "2. Tidak" (fill_gabungan tidak mengisinya)
        if key in KEY_16B + ("digital",) and not row["internet"].startswith("1"):
            continue
        if nilai not in opsi:
            salah(("OPSI_TIDAK_ADA_DI_FORM", f"{key}='{nilai}' bukan salah satu opsi form (lihat OPSI_FORM)"))

    # 13c usaha makan-minum: hanya kode 5-11. Dicek dari 13b2 = Ya ATAU KBLI
    # golongan 56 — audit 22-23 Sep 2026 selalu punya keduanya, tapi form bisa
    # menyimpulkan "makan minum" dari salah satunya saja, jadi keduanya dipakai.
    kode_13c = kode_opsi(row["lokasi_usaha"])
    mamin = row["layanan_mamin"].startswith("1") or kbli_makan_minum(row["kbli"])
    if mamin and kode_13c is not None and kode_13c < LOKASI_MAMIN_MIN:
        sebab = "13b2 = Ya" if row["layanan_mamin"].startswith("1") else f"KBLI {row['kbli']} golongan 56"
        salah(("13C_MAMIN_BUKAN_5_11",
               f"13c '{row['lokasi_usaha']}' berkode {kode_13c} tapi usaha makan-minum ({sebab}): form "
               f"hanya menerima kode {LOKASI_MAMIN_MIN}-11 — perbaiki 13c di sheet"))
    if row.kbli_genai:
        hasil.tanda.append(f"KBLI {row['kbli']} kategori {kbli_kategori_ditolak(row['kbli'])} ditolak form -> "
                           "13g diisi rekomendasi GenAI pertama saat pengisian")
    elif kat_ditolak := kbli_kategori_ditolak(row["kbli"]):
        salah(("KBLI_KATEGORI_DITOLAK",
               f"KBLI {row['kbli']} masuk kategori {kat_ditolak}: form menolak 13g dgn 'Kategori tidak boleh "
               f"berisi P atau U' — pilih KBLI usaha yang sesuai di sheet, bukan kategori {kat_ditolak}"))

    tidak_valid = [k for k in (*KEY_PEKERJA, *KEY_26, *KEY_27, "pendapatan_online", *KEY_28, *KEY_29, "umur",
                               *KEY_JUMLAH_19_20)
                   if row[k] and not _bulat(row[k])]
    for key, pola in (("idsubsls", r"\d{16}"), ("kodepos", r"\d{5}"), ("kbli", r"\d{5}"),
                      ("tahun_operasi", r"\d{4}")):
        if row[key] and not re.fullmatch(pola, row[key]):
            tidak_valid.append(key)
    try:
        if not hasil.tanpa_koordinat:
            lat, lon = float(row["latitude"]), float(row["longitude"])
            if not (-12 <= lat <= 7 and 94 <= lon <= 142):
                tidak_valid.append("latitude/longitude (di luar Indonesia)")
    except ValueError:
        if row["latitude"] and row["longitude"]:
            tidak_valid.append("latitude/longitude")
    if tidak_valid:
        salah(("ANGKA_TIDAK_VALID", f"kolom: {', '.join(tidak_valid)}"))
        return hasil  # pemeriksaan angka di bawah butuh angka yang valid

    if row["idsubsls"] and row.idsubsls_pilih != row.idsubsls:
        pesan = f"idsubsls={row.idsubsls} tapi kolom Pilih PROVINSI..SUBSLS = {row.idsubsls_pilih}"
        if mode_satu_subsls:
            hasil.tanda.append(f"wilayah tujuan ubah alokasi ambigu: {pesan}")
        else:
            salah(("WILAYAH_TIDAK_KONSISTEN", pesan))

    # Varian bulanan (30-33) = usaha mulai beroperasi TAHUN BERJALAN (ec_usaha_bulan
    # template: tahun_operasi == 2026). Minimal totalnya 10.000, bukan 100.000.
    bulanan = (row.bulanan_dari_kolom and row["tahun_operasi"].isdigit()
               and row.angka("tahun_operasi") == tahun_berjalan)

    if all(row[k] for k in KEY_PEKERJA):
        l, p, d, td = (row.angka(k) for k in KEY_PEKERJA)
        if l + p != d + td:
            salah(("PEKERJA_24_TIDAK_KONSISTEN",
                   f"24a1+24b1={l + p} (laki {l}, perempuan {p}) != 24a2+24b2={d + td} "
                   f"(dibayar {d}, tidak dibayar {td})"))
        # 24c1 "Cek konsistensi jenis kelamin pengusaha" (GALAT 22 Sep 2026 baris 77):
        # dugaan terkuatnya pengusaha ikut dihitung di 24, jadi kolom 24 sesuai jenis
        # kelaminnya (12b) tidak boleh 0. TANDA, BUKAN skip — sengaja: pola 24 laki 0 +
        # perempuan 2 ada di 123 baris Agenda1-1/Agenda2 (lihat KOREKSI_PEKERJA) dan
        # ratusan baris sejenis TERKIRIM tanpa GALAT ini, sementara GALAT-nya cuma
        # muncul 1x dari ~6.600 baris audit. Men-skip semuanya berdasarkan satu bukti
        # jelas lebih mahal drpd menampilkannya utk ditinjau.
        kode_jk = kode_opsi(row["jk"])
        if l + p >= 1 and kode_jk in (1, 2):
            punya, label = ((l, "24a1 (pekerja laki-laki)") if kode_jk == 1
                            else (p, "24b1 (pekerja perempuan)"))
            if punya == 0:
                hasil.tanda.append(
                    f"12b '{row['jk']}' tapi {label}=0 dari 24c1={l + p} — form PERNAH menolak pola ini "
                    f"('Cek konsistensi jenis kelamin pengusaha', 1x 22 Sep 2026); cek 24/12b di sheet "
                    f"kalau dokumen ini nanti ber-GALAT")
        # 26a vs 24a2 (dua aturan file-validation `gaji`) — lihat GAJI_MIN_PER_PEKERJA_DIBAYAR.
        if row["gaji"]:
            gaji = row.angka("gaji")
            r26a = "30a" if bulanan else "26a"
            if d == 0 and gaji != 0:
                salah(("26A_HARUS_0_TANPA_PEKERJA_DIBAYAR",
                       f"{r26a}={gaji:,} padahal 24a2 (pekerja dibayar)=0: form mewajibkan {r26a}=0 — "
                       f"betulkan {r26a} atau 24a2 di sheet"))
            elif d > 0 and gaji // d <= GAJI_MIN_PER_PEKERJA_DIBAYAR:
                salah(("26A_PER_PEKERJA_DI_BAWAH_MINIMAL",
                       f"{r26a}/24a2 = {gaji:,}/{d} = {gaji // d:,} <= Rp {GAJI_MIN_PER_PEKERJA_DIBAYAR:,}: "
                       f"form mewajibkan > Rp {GAJI_MIN_PER_PEKERJA_DIBAYAR:,} kalau 24a2 > 0 — "
                       f"betulkan {r26a} atau 24a2 di sheet"))
    if all(row[k] for k in KEY_29) and sum(row.angka(k) for k in KEY_29) != 100:
        salah(("MODAL_29_BUKAN_100", f"jumlah 29a-29f = {sum(row.angka(k) for k in KEY_29)}"))
    minimal, r_peng, r_pend = ((MINIMAL_TOTAL_RUPIAH_BULANAN, "30f", "31c") if bulanan
                               else (MINIMAL_TOTAL_RUPIAH, "26f", "27c"))
    if all(row[k] for k in KEY_26) and sum(row.angka(k) for k in KEY_26) < minimal:
        salah(("DI_BAWAH_MINIMAL", f"{r_peng}={sum(row.angka(k) for k in KEY_26)} < {minimal}"))
    if all(row[k] for k in KEY_27) and sum(row.angka(k) for k in KEY_27) < minimal:
        salah(("DI_BAWAH_MINIMAL", f"{r_pend}={sum(row.angka(k) for k in KEY_27)} < {minimal}"))
    if row["hp"] and not hp_valid(row["hp"]):
        salah(("HP_TIDAK_VALID", f"No HP/WA '{row['hp']}' bukan 08 + 10-13 digit / 9999 (form menolak)"))
    if row["nik_pengusaha"] and not nik_valid(row["nik_pengusaha"]):
        salah(("NIK_TIDAK_VALID", f"12d NIK '{row['nik_pengusaha']}' bukan 16 digit / 7777 / 8888 / 9999 "
                                  "(form mengosongkannya -> GALAT)"))
    if row["umur"] and not UMUR_MIN <= row.angka("umur") <= UMUR_MAKS:
        salah(("UMUR_DI_LUAR_10_99", f"12c umur={row['umur']} (form: wajib {UMUR_MIN}-{UMUR_MAKS}) "
                                     "— perbaiki umur di sheet, jangan ditebak"))
    if row["pendapatan_online"] and row.angka("pendapatan_online") > 100:
        salah(("ANGKA_TIDAK_VALID", f"27d={row['pendapatan_online']} > 100 persen"))
    if row["tahun_operasi"]:
        th = row.angka("tahun_operasi")
        if bulanan:
            hasil.tanda.append(f"varian bulanan (mulai beroperasi {th}): rincian 30-33 diisi dari kolom 26-29")
        elif th > tahun_berjalan:
            salah(("ANGKA_TIDAK_VALID", f"tahun_operasi={th} (di masa depan)"))
        elif th == tahun_berjalan:
            # Form mengganti 26-29 dgn 30-33 (angka SATU BULAN) utk usaha yang
            # mulai beroperasi tahun berjalan; angka sheet ini tahunan.
            salah(("VARIAN_BULANAN", f"tahun_operasi={th} -> form pakai rincian 30-33 bulanan"))
        elif th < 1900:
            salah(("ANGKA_TIDAK_VALID", f"tahun_operasi={th}"))
    if row["internet"].startswith("1") and not any(row[k].startswith("1") for k in KEY_16B):
        salah(("16B_TANPA_YA", "16a = Ya tapi 16b1-16b6 tidak ada yang Ya (form menolak)"))
    if row.kbli_genai:
        pass   # KBLI sebenarnya baru diketahui saat GenAI memilih — 26b/26c/30c dicek form.
    elif kbli_tanpa_26c(row["kbli"]) and row["biaya_pembelian"] and row.angka("biaya_pembelian") > 0:
        salah(("26C_KATEGORI_TANPA_26C",
               f"KBLI {row['kbli']} bukan perdagangan: form tidak punya 26c — pindahkan "
               f"26c={row['biaya_pembelian']} ke 26b di sheet"))
    elif kbli_26b_wajib_positif(row["kbli"]) and row["biaya_produksi"] and row.angka("biaya_produksi") == 0:
        salah(("26B_HARUS_LEBIH_0", f"KBLI {row['kbli']} (kategori B-F / golongan 56): form mewajibkan "
                                    f"{'30b' if bulanan else '26b'} biaya produksi > 0"))
    if (bulanan and not row.kbli_genai and kbli_punya_30c(row["kbli"]) and row["biaya_pembelian"]
            and row.angka("biaya_pembelian") == 0):
        salah(("30C_HARUS_LEBIH_0", f"varian bulanan KBLI {row['kbli']}: form mewajibkan 30c (biaya pembelian "
                                    "barang yang terjual) > 0"))

    if not row["produk"] and row.produk_utama:
        hasil.tanda.append("13f disalin dari 13a (sheet tidak punya kolom 13f)")
    asli_13f = " ".join(str(row["produk"] or row["keg_utama"] or "").split())
    if asli_13f and len(asli_13f) < MIN_KARAKTER_13F:
        if len(row.produk_utama) >= MIN_KARAKTER_13F:
            hasil.tanda.append(f"13f dilengkapi: '{asli_13f}' -> '{row.produk_utama}'")
        else:
            salah(("13F_KURANG_4_KARAKTER", f"13f '{asli_13f}' kurang dari {MIN_KARAKTER_13F} karakter "
                                            "(form menolak) & tidak bisa dilengkapi (KBLI kosong)"))
    if len(row.nama_komersial) > MAKS_8B:
        salah(("8B_TERLALU_PANJANG", f"8b '{row.nama_komersial}' {len(row.nama_komersial)} karakter > "
                                     f"{MAKS_8B} (form menolak) — singkatkan nama usaha/12a di sheet"))
    if not row["jalan_domisili"] and GABUNGAN_IZINKAN_JALAN_KOSONG:
        hasil.tanda.append("Nama Jalan KOSONG (belum pernah diuji dikosongkan)")
    return hasil


def _pasangan_termuat(jarum: list[str], jerami: list[str], n: int = 4) -> list[tuple[int, int]]:
    """[(i, j)] utk setiap `jarum[i]` yang jadi SUBSTRING `jerami[j]` (i == j ikut —
    penyaringnya urusan pemanggil).

    Membandingkan semua pasangan langsung itu O(n²): 4.099 baris tahap 2 = 16,8 juta
    pembandingan = ±9,5 menit sebelum prompt "YA" (diukur 2026-09-23). Di sini tiap
    jerami diindeks per potongan `n` huruf, lalu satu jarum cuma diadu dgn jerami yang
    memuat `n` huruf PERTAMA-nya — kalau jarum termuat, potongan itu pasti ada di sana,
    jadi hasilnya SAMA PERSIS, cuma kandidatnya jauh lebih sedikit (±2 detik).
    """
    indeks: dict[str, set[int]] = defaultdict(set)
    for j, h in enumerate(jerami):
        for p in range(len(h) - n + 1):
            indeks[h[p:p + n]].add(j)
    semua = list(range(len(jerami)))
    singgah: dict[str, list[int]] = {}
    hasil: list[tuple[int, int]] = []
    for i, na in enumerate(jarum):
        if na not in singgah:
            # Jarum < n huruf tidak punya potongan utuh utk dicari di indeks -> adu ke
            # semua jerami (jumlahnya sedikit; nama sependek itu jarang).
            kandidat = semua if len(na) < n else sorted(indeks.get(na[:n], ()))
            singgah[na] = [j for j in kandidat if na in jerami[j]]
        hasil.extend((i, j) for j in singgah[na])
    return hasil


def periksa_semua(rows: list[GabunganRow], tahun_berjalan: int | None = None,
                  mode_satu_subsls: bool = False,
                  izinkan_tanpa_koordinat: bool = False) -> dict[int, Pemeriksaan]:
    """{baris: Pemeriksaan} — per baris + pemeriksaan LINTAS baris.
    Mode satu subsls: SEMUA dokumen masuk list PENDATAAN satu akun, jadi
    bentrok nama diperiksa lintas SELURUH sheet, bukan per akun PPL."""
    hasil = {r.baris: periksa_baris(r, tahun_berjalan, mode_satu_subsls, izinkan_tanpa_koordinat)
             for r in rows}

    per_kunci = defaultdict(list)
    for r in rows:
        per_kunci[r.kunci].append(r.baris)
    for kunci, daftar in per_kunci.items():
        if len(daftar) > 1:
            for b in daftar:
                hasil[b].masalah.append(("BARIS_GANDA", f"akun+idsubsls+nama sama dgn baris {daftar}"))

    # Cek "dokumen sudah ada" di fasih_web mencocokkan nama sbg SUBSTRING.
    # Kalau satu PPL punya "APOTEK SEHAT" & "APOTEK SEHAT 2", mencari yang
    # pertama bisa menemukan yang kedua -> skrip mengisi dokumen yang SALAH.
    # Diperiksa pada nama_dokumen (yang dicari di list) DAN nama mentah (yang
    # dicari pengaman dokumen-bernama-lama di create_document).
    # Nama SAMA PERSIS dgn kunci beda (mis. subsls lain) juga bentrok —
    # pencarian membuka dokumen pertama yang ketemu. Kunci sama = BARIS_GANDA.
    per_list = defaultdict(list)
    for r in rows:
        per_list["" if mode_satu_subsls else r.akun_ppl].append(r)
    lingkup = "list satu akun" if mode_satu_subsls else "PPL sama"
    for anggota in per_list.values():
        nama_dok = [a.nama_dokumen.upper() for a in anggota]
        nama_mentah = [a.nama.upper() for a in anggota]
        # Mode satu subsls: nama mentah tidak pernah dipakai mencari (dokumen di list
        # satu akun ini selalu bernama nama_dokumen) -> cukup tanda di bawah.
        pakai_nama_lama = (not mode_satu_subsls) and any(a.nama_lama_dicari for a in anggota)

        def layak(i: int, j: int) -> bool:
            a, b = anggota[i], anggota[j]
            return a.baris != b.baris and a.kunci != b.kunci

        cocok: dict[tuple[int, int], tuple[str, str]] = {}
        for i, j in _pasangan_termuat(nama_dok, nama_dok):
            if layak(i, j) and (nama_dok[i] != nama_dok[j] or anggota[i].baris < anggota[j].baris):
                cocok[(i, j)] = (anggota[i].nama_dokumen, anggota[j].nama_dokumen)
        if pakai_nama_lama:
            for i, j in _pasangan_termuat(nama_mentah, nama_mentah):
                if ((i, j) not in cocok and anggota[i].nama_lama_dicari and layak(i, j)
                        and (nama_mentah[i] != nama_mentah[j] or anggota[i].baris < anggota[j].baris)):
                    cocok[(i, j)] = (anggota[i].nama, anggota[j].nama)
        for (i, j), (na, nb) in sorted(cocok.items()):
            for x, y in ((anggota[i], anggota[j]), (anggota[j], anggota[i])):
                hasil[x.baris].masalah.append((
                    "NAMA_TUMPANG_TINDIH",
                    f"'{na}' terkandung di '{nb}' (baris {y.baris}, {lingkup}) — "
                    "pencarian dokumen bisa membuka dokumen yang salah"))
        if mode_satu_subsls and any(a.nama_lama_dicari for a in anggota):
            # Pengaman dokumen-bernama-lama di create_document mencari nama MENTAH
            # sbg substring; kalau dokumen baris b sudah dibuat duluan, baris a akan
            # berhenti SKIP_DOKUMEN_NAMA_LAMA.
            for i, j in sorted(_pasangan_termuat(nama_mentah, nama_dok)):
                if (i, j) not in cocok and anggota[i].nama_lama_dicari and layak(i, j):
                    hasil[anggota[i].baris].tanda.append(
                        f"nama '{anggota[i].nama}' terkandung di nama dokumen baris "
                        f"{anggota[j].baris} — bisa SKIP_DOKUMEN_NAMA_LAMA kalau dokumen "
                        "baris itu dibuat lebih dulu")

    # Kodepos dialokasikan per DESA; nilai minoritas dalam satu desa patut dicek.
    per_desa = defaultdict(Counter)
    for r in rows:
        if len(r.idsubsls) == 16 and r["kodepos"]:
            per_desa[r.idsubsls[:10]][r["kodepos"]] += 1
    for r in rows:
        c = per_desa.get(r.idsubsls[:10])
        if c and len(c) > 1:
            mayoritas, n = c.most_common(1)[0]
            if r["kodepos"] != mayoritas and c[r["kodepos"]] < n:
                hasil[r.baris].tanda.append(
                    f"kodepos {r['kodepos']} beda dgn mayoritas desa {r.idsubsls[:10]} ({mayoritas}, {dict(c)})")
    return hasil


def idsubsls_dari_wilayah(nilai: dict) -> str:
    """Rincian 1-6 BLOK I -> kode wilayah dokumen ("[51] BALI"/"[060] BULELENG"/...
    -> 51 08 060 006 000116). "" kalau ada bagian yang tidak terbaca.

    Field kode SLS kadang berisi 6 digit (SLS + subsls) & kadang 4 digit (SLS
    saja, subsls tidak ditampilkan) — keduanya diterima, jadi hasilnya bisa 16
    ATAU 14 digit. Pemanggil yang memutuskan: 14 digit cukup utk memastikan
    dokumen masih di kabupaten yang sama, tapi TIDAK cukup utk dicatat sbg
    subsls di audit (lihat --izinkan-wilayah-beda)."""
    panjang = {"prov": 2, "kab": 2, "kec": 3, "desa": 3, "kode_sls": (4, 6)}
    keluar = ""
    for key, n in panjang.items():
        teks = " ".join(str(nilai.get(key) or "").split())
        cocok = re.match(r"^\[?(\d+)\]?", teks)
        boleh = n if isinstance(n, tuple) else (n,)
        if not cocok or len(cocok.group(1)) not in boleh:
            return ""
        keluar += cocok.group(1)
    return keluar


def kode_wilayah_api(obj) -> str:
    """Subsls 16 digit dari respons API dokumen, "" kalau tidak ada / tidak tunggal.

    Item list PENDATAAN (terverifikasi): region.level1..level6 {fullCode} — level6 =
    subsls. Respons detail (get-by-id-with-data) belum pernah terekam bagian
    region-nya, jadi dicari generik: semua `fullCode`/`full_code` 16 digit, di bawah
    kunci `region` kalau ada. Lebih dari satu kode berbeda -> "" (tidak ditebak)."""
    def kumpul(o, keluar: set):
        if isinstance(o, dict):
            for k, v in o.items():
                if k in ("fullCode", "full_code") and re.fullmatch(r"\d{16}", str(v or "")):
                    keluar.add(str(v))
                else:
                    kumpul(v, keluar)
        elif isinstance(o, list):
            for v in o:
                kumpul(v, keluar)
        return keluar

    if isinstance(obj, dict) and isinstance(obj.get("data"), dict):
        obj = obj["data"]
    if isinstance(obj, dict) and isinstance(obj.get("region"), dict):
        obj = obj["region"]
    kode = kumpul(obj, set())
    return kode.pop() if len(kode) == 1 else ""


def cocokkan_wilayah_dokumen(nilai: dict, idsubsls: str, nama_ref: dict | None = None) -> tuple[str, str]:
    """Bandingkan rincian 1-6 BLOK I dokumen yang terbuka (auto dari wilayah
    dokumen) dgn subsls tujuan input. -> ('COCOK'|'BEDA'|'TIDAK_TERBACA', ringkasan).

    Format nilai field belum pernah terekam (dump DOM tidak memuat value
    input), jadi yang diterima: kode saja ("014"), kode panjang
    ("5108060014"), atau "[014] NAMA". Field TANPA angka dibandingkan lewat
    nama (`nama_ref` = WILAYAH_BY_IDSUBSLS[idsubsls]) kalau ada.
    Kode SLS WAJIB terbukti cocok; tanpa itu hasilnya TIDAK_TERBACA."""
    i = idsubsls
    kandidat = {
        "prov": {i[:2]},
        "kab": {i[2:4], i[:4]},
        "kec": {i[4:7], i[:7]},
        "desa": {i[7:10], i[:10]},
        "kode_sls": {i[10:16], i[10:14], i[:14], i[7:16], i},
    }
    kolom_ref = {"prov": "provinsi", "kab": "kabkota", "kec": "kecamatan", "desa": "desa"}
    ringkas = "; ".join(f"{k}='{nilai.get(k, '')}'" for k in kandidat)
    beda, sls_cocok = [], False
    for key, boleh in kandidat.items():
        teks = " ".join(str(nilai.get(key) or "").split())
        # Hanya kode di AWAL teks — angka di dalam nama ("BANJAR 2") bukan kode.
        m = re.match(r"\[?([\d.\-\s]*\d)", teks)
        angka = re.sub(r"\D", "", m.group(1)) if m else ""
        if angka:
            if angka in boleh:
                sls_cocok |= key == "kode_sls"
            else:
                beda.append(f"{key}='{teks}'")
        elif teks and nama_ref and key in kolom_ref:
            if teks.upper() != str(nama_ref.get(kolom_ref[key], "")).upper():
                beda.append(f"{key}='{teks}' (referensi '{nama_ref.get(kolom_ref[key])}')")
    if beda:
        return "BEDA", f"wilayah dokumen BUKAN {idsubsls}: {', '.join(beda)} | {ringkas}"
    if not sls_cocok:
        return "TIDAK_TERBACA", f"kode SLS dokumen tidak terbaca sbg angka | {ringkas}"
    return "COCOK", ringkas


def kelompok_per_akun(rows: list[GabunganRow]) -> dict[str, list[GabunganRow]]:
    """Satu sesi login per akun PPL, urutan kemunculan di sheet dipertahankan."""
    grup: dict[str, list[GabunganRow]] = {}
    for r in rows:
        grup.setdefault(r.akun_ppl, []).append(r)
    return grup


def parse_pilihan_baris(teks: str) -> set[int]:
    """"2,5,10-20" -> {2, 5, 10, ..., 20} (nomor baris seperti di Google Sheets)."""
    hasil: set[int] = set()
    for bagian in (teks or "").split(","):
        bagian = bagian.strip()
        if not bagian:
            continue
        if "-" in bagian:
            a, b = (int(x) for x in bagian.split("-", 1))
            hasil.update(range(min(a, b), max(a, b) + 1))
        else:
            hasil.add(int(bagian))
    return hasil
