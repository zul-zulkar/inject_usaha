#!/usr/bin/env python3
"""
koordinat_pengganti.py — rapikan & ganti koordinat sheet tahap 2 yang jatuh jauh
dari subsls-nya. Hasilnya berkas xlsx BARU; berkas sumber tidak disentuh.

Ketetapan user 2026-09-24:
  1. Koordinat dianggap SALAH kalau tidak terbaca, di luar kabupaten, atau
     > 500 m (--batas-m) di luar poligon subsls-nya. Yang benar TETAP dipakai.
  2. Penggantinya titik ACAK di dalam subsls: diambil dari titik geotag LISTING
     (= bangunan/keluarga yang benar-benar didatangi -> pasti pemukiman), yang
     dekat JALAN kalau data jalan tersedia (--jalan, <= --jarak-jalan-m).
  3. Usaha dgn PEMILIK (12a) & ALAMAT (8c) sama = satu kelompok = SATU koordinat,
     di subsls dgn baris terbanyak di kelompok itu (seri -> subsls tempat
     koordinat aslinya jatuh terbanyak -> kode terkecil).
  4. Format derajat ("8°7'35,424\"S", "-8.127S", "8°12887404S") dikonversi ke
     desimal DULU, baru dinilai dgn aturan 1-3. Begitu juga derajat-menit-detik
     TANPA simbol ("-8,747" = 8°7'47") & salah ketik satu digit ("144,59" ->
     "114,59") — HANYA dipakai kalau hasilnya jatuh di subsls itu (<= batas).

Acak tapi TETAP: benih = kelompok + subsls, jadi menjalankan ulang memberi titik
yang sama (draft yang sudah di-geotag tidak berpindah). Titik listing digeser
5–15 m (tetap di poligon) supaya tidak persis di titik rumah keluarga.

    python koordinat/koordinat_pengganti.py --sumber bahan/input_tahap2_22.xlsx
    -> koordinat/hasil/input_tahap2_22_koordinat.xlsx

Keluaran utk SALIN-TEMPEL: baris ke-N = baris ke-N sheet sumber; kolom A:B =
Latitude/Longitude (salin A2:B<akhir>, tempel ke sel Latitude baris 2 sheet). Baris
yang tidak berubah berisi teks ASLINYA persis; yang berubah diwarnai kuning & diberi
keterangan (sumber, nilai asli, jarak asli ke subsls, kelompok) di kolom C dst.
"""

from __future__ import annotations

import os as _os, sys as _sys  # noqa: E401
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))

import argparse  # noqa: E402
import csv  # noqa: E402
import hashlib  # noqa: E402
import json  # noqa: E402
import math  # noqa: E402
import random  # noqa: E402
import re  # noqa: E402
from collections import Counter, defaultdict  # noqa: E402
from pathlib import Path  # noqa: E402

from inti import lokasi  # noqa: E402
from inti.config import JALAN_PATH, PETA_SLS_PATH, TAHAP2_KOTAK_KOORDINAT, TITIK_LISTING_PATH  # noqa: E402

for _stream in (_sys.stdout, _sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

DERAJAT = "°"
HASIL_DIR = lokasi.hasil("koordinat")
GESER_M = (5.0, 15.0)


# --------------------------------------------------------------------------
# Membaca teks koordinat (logika murni, diuji tests/test_koordinat_pengganti.py)
# --------------------------------------------------------------------------
def baca_koordinat(nilai, sumbu: str) -> float | None:
    """Teks sel -> derajat desimal (lintang selalu negatif / selatan, bujur positif).
    Derajat-menit-detik bersimbol dihitung; selain itu SEMUA tanda dibuang lalu titik
    desimal dipasang sesudah digit ke-1 (lintang) / ke-3 (bujur) — sama dgn formula
    spreadsheet yang diberikan ke user. Menangani koma/titik desimal, akhiran S/E,
    ° sbg titik desimal, titik desimal hilang, koma dobel, spasi, lintang tanpa minus."""
    if nilai is None:
        return None
    s = str(nilai).strip().upper()
    if not s:
        return None
    if "'" in s and DERAJAT in s:
        m = re.match(r"^-?\s*(\d+)\s*" + DERAJAT + r"\s*(\d+)\s*'\s*(\d+(?:[.,]\d+)?)?", s)
        if not m:
            return None
        x = int(m.group(1)) + int(m.group(2)) / 60 + float((m.group(3) or "0").replace(",", ".")) / 3600
    else:
        digit = re.sub(r"\D", "", s)
        lebar = 1 if sumbu == "lat" else 3
        if len(digit) < lebar:
            return None
        x = int(digit) / 10 ** (len(digit) - lebar)
    return round(-x if sumbu == "lat" else x, 8)


def dms_tanpa_simbol(nilai, sumbu: str) -> float | None:
    """"-8,747" -> 8°7'47" ; "155,2058" -> 115°20'58" (155 = salah ketik 115).
    Lintang: 1 digit derajat + 1 digit menit; bujur: 3 digit derajat + 2 digit menit;
    sisanya detik (2 digit, lebihnya desimal). Hanya DIAJUKAN — pemanggil menerimanya
    kalau hasilnya jatuh di subsls (Tejakula 2026-09-24: 24 baris, 0–1,4 km)."""
    if nilai is None or "'" in str(nilai):
        return None
    digit = re.sub(r"\D", "", str(nilai))
    d, m = (1, 1) if sumbu == "lat" else (3, 2)
    if len(digit) < d + m:
        return None
    derajat, menit, sisa = int(digit[:d]), int(digit[d:d + m]), digit[d + m:]
    if sumbu == "lon" and derajat == 155:
        derajat = 115
    if menit >= 60:
        return None
    detik = float(f"{sisa[:2]}.{sisa[2:]}") if len(sisa) > 2 else float(sisa or 0)
    if detik >= 60:
        return None
    x = derajat + menit / 60 + detik / 3600
    return round(-x if sumbu == "lat" else x, 8)


def varian_salah_ketik(nilai) -> list[str]:
    """Satu digit diganti, atau dua digit bersebelahan bertukar."""
    t = str(nilai or "")
    out = []
    for i, c in enumerate(t):
        if c.isdigit():
            out += [t[:i] + d + t[i + 1:] for d in "0123456789" if d != c]
            if i + 1 < len(t) and t[i + 1].isdigit() and t[i + 1] != c:
                out.append(t[:i] + t[i + 1] + c + t[i + 2:])
    return out


def norm_teks(teks) -> str:
    """Pemilik/alamat utk pengelompokan: huruf besar, tanda baca & spasi berlebih dibuang."""
    t = re.sub(r"[^0-9A-Z]+", " ", str(teks or "").upper()).strip()
    return "" if t in ("", "0") else t


def kunci_kelompok(baris: int, pemilik, alamat) -> str:
    """Pemilik & alamat sama -> kunci sama. Salah satu kosong -> baris sendiri
    (nama pemilik saja terlalu umum utk disamakan: "KETUT")."""
    p, a = norm_teks(pemilik), norm_teks(alamat)
    return f"{p}|{a}" if p and a else f"baris:{baris}"


def pilih_subsls(anggota: list[dict], dalam: dict) -> str:
    """Subsls dgn baris terbanyak; seri -> subsls tempat koordinat asli anggota
    jatuh terbanyak (`dalam`: {baris: idsubsls yang memuat titiknya}); lalu kode terkecil."""
    jumlah = Counter(a["idsubsls"] for a in anggota if a["idsubsls"])
    if not jumlah:
        return ""
    titik = Counter(dalam.get(a["baris"], "") for a in anggota)
    return min(jumlah, key=lambda s: (-jumlah[s], -titik.get(s, 0), s))


def benih(*bagian) -> random.Random:
    return random.Random(int(hashlib.sha1("|".join(map(str, bagian)).encode("utf-8")).hexdigest()[:16], 16))


def meter(lat1, lon1, lat2, lon2) -> float:
    kx = 111320 * math.cos(math.radians((lat1 + lat2) / 2))
    return math.hypot((lon2 - lon1) * kx, (lat2 - lat1) * 110574)


def geser(lat: float, lon: float, jarak_m: float, arah: float) -> tuple[float, float]:
    kx = 111320 * math.cos(math.radians(lat))
    return (round(lat + jarak_m * math.sin(arah) / 110574, 7), round(lon + jarak_m * math.cos(arah) / kx, 7))


# --------------------------------------------------------------------------
# Indeks spasial sederhana (grid) utk titik listing & ruas jalan
# --------------------------------------------------------------------------
SEL = 0.005   # derajat (±550 m)


def _sel(lat, lon):
    return (math.floor(lat / SEL), math.floor(lon / SEL))


class TitikListing:
    def __init__(self, titik: list[tuple[float, float]]):
        self.grid: dict = defaultdict(list)
        for la, lo in titik:
            self.grid[_sel(la, lo)].append((la, lo))

    def dalam_kotak(self, miny, minx, maxy, maxx):
        a, b = _sel(miny, minx), _sel(maxy, maxx)
        for i in range(a[0], b[0] + 1):
            for j in range(a[1], b[1] + 1):
                yield from self.grid.get((i, j), ())


class Jalan:
    """Ruas jalan (pasangan titik). Kosong = tidak ada data jalan."""

    def __init__(self, ruas: list[tuple[float, float, float, float]]):
        self.ruas = ruas
        self.grid: dict = defaultdict(list)
        for k, (la1, lo1, la2, lo2) in enumerate(ruas):
            a, b = _sel(min(la1, la2), min(lo1, lo2)), _sel(max(la1, la2), max(lo1, lo2))
            for i in range(a[0], b[0] + 1):
                for j in range(a[1], b[1] + 1):
                    self.grid[(i, j)].append(k)

    def __bool__(self):
        return bool(self.ruas)

    def jarak_m(self, lat: float, lon: float) -> float:
        i0, j0 = _sel(lat, lon)
        kx, ky = 111320 * math.cos(math.radians(lat)), 110574
        terdekat = math.inf
        for i in (i0 - 1, i0, i0 + 1):
            for j in (j0 - 1, j0, j0 + 1):
                for k in self.grid.get((i, j), ()):
                    la1, lo1, la2, lo2 = self.ruas[k]
                    ax, ay = (lo1 - lon) * kx, (la1 - lat) * ky
                    dx, dy = (lo2 - lo1) * kx, (la2 - la1) * ky
                    p = dx * dx + dy * dy
                    t = 0.0 if p == 0 else max(0.0, min(1.0, -(ax * dx + ay * dy) / p))
                    terdekat = min(terdekat, math.hypot(ax + t * dx, ay + t * dy))
        return terdekat

    def dalam_kotak(self, miny, minx, maxy, maxx):
        a, b = _sel(miny, minx), _sel(maxy, maxx)
        lihat = set()
        for i in range(a[0], b[0] + 1):
            for j in range(a[1], b[1] + 1):
                lihat.update(self.grid.get((i, j), ()))
        return [self.ruas[k] for k in sorted(lihat)]


def muat_listing(path: str | Path, kotak=None) -> TitikListing:
    """CSV berkolom latitude/longitude (idsubsls tidak dipakai: yang menentukan poligon)."""
    titik = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            try:
                la, lo = float(r.get("latitude") or ""), float(r.get("longitude") or "")
            except ValueError:
                continue
            if kotak is None or (kotak[0] <= la <= kotak[1] and kotak[2] <= lo <= kotak[3]):
                titik.append((la, lo))
    return TitikListing(titik)


def muat_jalan(path: str | Path) -> Jalan:
    """GeoJSON (LineString/MultiLineString) atau JSON Overpass (`out geom`)."""
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    garis = []
    if "elements" in d:
        garis = [[(p["lat"], p["lon"]) for p in el.get("geometry") or []]
                 for el in d["elements"] if el.get("type") == "way"]
    else:
        for f in d.get("features", []):
            g = f.get("geometry") or {}
            if g.get("type") == "LineString":
                garis.append([(y, x) for x, y, *_ in g["coordinates"]])
            elif g.get("type") == "MultiLineString":
                garis += [[(y, x) for x, y, *_ in bagian] for bagian in g["coordinates"]]
    ruas = [(a[0], a[1], b[0], b[1]) for g in garis for a, b in zip(g, g[1:])]
    return Jalan(ruas)


# --------------------------------------------------------------------------
# Memilih titik pengganti
# --------------------------------------------------------------------------
def titik_acak(idsubsls: str, peta, listing: TitikListing | None, jalan: Jalan | None, rng: random.Random,
               jarak_jalan_m: float = 50.0) -> tuple[float, float, str] | None:
    """(lat, lon, sumber) acak di dalam poligon subsls. Urutan: titik listing dekat
    jalan -> titik listing -> titik di tepi ruas jalan di poligon -> titik bebas di
    poligon. None kalau subsls tidak ada di peta."""
    from koordinat.peta import _dalam_poligon
    s = peta.sls.get(idsubsls)
    if s is None:
        return None
    minx, miny, maxx, maxy = s["bbox"]
    dalam = lambda la, lo: any(_dalam_poligon(lo, la, p) for p in s["poligon"])  # noqa: E731
    dekat_jalan = (lambda la, lo: jalan.jarak_m(la, lo) <= jarak_jalan_m) if jalan else None

    def selesaikan(la, lo, sumber, syarat):
        # Geser 5–15 m supaya tidak persis di titik rumah; syarat tetap dipenuhi.
        for _ in range(12):
            g = geser(la, lo, rng.uniform(*GESER_M), rng.uniform(0, 2 * math.pi))
            if syarat(*g):
                return (*g, sumber)
        return (round(la, 7), round(lo, 7), sumber)

    kandidat = sorted(p for p in (listing.dalam_kotak(miny, minx, maxy, maxx) if listing else ()) if dalam(*p))
    if kandidat:
        if dekat_jalan:
            dekat = [p for p in kandidat if dekat_jalan(*p)]
            if dekat:
                la, lo = rng.choice(dekat)
                return selesaikan(la, lo, "ACAK_LISTING_DEKAT_JALAN", lambda a, b: dalam(a, b) and dekat_jalan(a, b))
        la, lo = rng.choice(kandidat)
        return selesaikan(la, lo, "ACAK_LISTING", dalam)
    if jalan:
        ruas = [r for r in jalan.dalam_kotak(miny, minx, maxy, maxx)
                if dalam((r[0] + r[2]) / 2, (r[1] + r[3]) / 2)]
        for _ in range(200 if ruas else 0):
            la1, lo1, la2, lo2 = rng.choice(ruas)
            t = rng.random()
            g = geser(la1 + t * (la2 - la1), lo1 + t * (lo2 - lo1), rng.uniform(*GESER_M), rng.uniform(0, 2 * math.pi))
            if dalam(*g):
                return (*g, "ACAK_POLIGON_DEKAT_JALAN")
    for _ in range(5000):
        la, lo = rng.uniform(miny, maxy), rng.uniform(minx, maxx)
        if dalam(la, lo):
            return (round(la, 7), round(lo, 7), "ACAK_POLIGON")
    return None


def rencana(baris: list[dict], peta, listing=None, jalan=None, batas_m: float = 500.0,
            jarak_jalan_m: float = 50.0, kotak=TAHAP2_KOTAK_KOORDINAT) -> list[dict]:
    """baris: [{baris, idsubsls, pemilik, alamat, lat_mentah, lon_mentah}] ->
    [{...masukan, lat, lon, sumber, jarak_asli_m, kelompok, subsls_koordinat}] sejajar."""
    from koordinat.peta import _dalam_poligon
    dikabupaten = (lambda la, lo: kotak is None or (kotak[0] <= la <= kotak[1] and kotak[2] <= lo <= kotak[3]))

    def jarak(idsubsls, la, lo):
        if la is None or lo is None:
            return None
        return peta.jarak_m(idsubsls, lo, la) if idsubsls in peta.sls else None

    # 1. Baca + perbaiki format / derajat tanpa simbol / salah ketik (relatif ke subsls baris).
    hasil = []
    for b in baris:
        la, lo = baca_koordinat(b["lat_mentah"], "lat"), baca_koordinat(b["lon_mentah"], "lon")
        cara = "" if la is None or lo is None else "ASLI"
        j = jarak(b["idsubsls"], la, lo)
        # Perbaikan derajat-tanpa-simbol & salah ketik HANYA utk bacaan yang tidak terbaca /
        # di LUAR kabupaten: titik yang cuma meleset 1-5 km gampang "masuk" poligon dgn
        # mengganti satu digit desimal — itu tebakan, bukan perbaikan (uji 2026-09-24:
        # 62 baris "diperbaiki" padahal pola salah ketik yang terbukti cuma ±35).
        rusak = la is None or lo is None or not dikabupaten(la, lo)
        if b["idsubsls"] in peta.sls and rusak and (j is None or j > batas_m):
            d_la, d_lon = dms_tanpa_simbol(b["lat_mentah"], "lat"), dms_tanpa_simbol(b["lon_mentah"], "lon")
            for cla, clo, nama in ((d_la, d_lon, "DERAJAT_TANPA_SIMBOL"), (d_la, lo, "DERAJAT_TANPA_SIMBOL"),
                                   (la, d_lon, "DERAJAT_TANPA_SIMBOL")):
                jj = jarak(b["idsubsls"], cla, clo)
                if jj is not None and jj <= batas_m:
                    la, lo, cara, j = cla, clo, nama, jj
                    break
            else:
                cocok = set()
                for sumbu, teks in (("lat", b["lat_mentah"]), ("lon", b["lon_mentah"])):
                    for v in varian_salah_ketik(teks):
                        cla = baca_koordinat(v, "lat") if sumbu == "lat" else la
                        clo = baca_koordinat(v, "lon") if sumbu == "lon" else lo
                        if cla is not None and clo is not None and jarak(b["idsubsls"], cla, clo) == 0:
                            cocok.add((cla, clo))
                if len(cocok) == 1:
                    la, lo = cocok.pop()
                    cara, j = "SALAH_KETIK_DIPERBAIKI", 0.0
        hasil.append({**b, "lat_baca": la, "lon_baca": lo, "cara_baca": cara,
                      "jarak_asli_m": None if j is None else round(j)})

    # 2. Kelompok pemilik+alamat -> subsls terbanyak -> satu koordinat.
    kelompok = defaultdict(list)
    for h in hasil:
        h["kelompok"] = kunci_kelompok(h["baris"], h["pemilik"], h["alamat"])
        kelompok[h["kelompok"]].append(h)
    for kunci, anggota in kelompok.items():
        dalam = {}
        for h in anggota:
            if h["lat_baca"] is not None and h["lon_baca"] is not None:
                hit = [s for s in (h["idsubsls"],) if s in peta.sls
                       and any(_dalam_poligon(h["lon_baca"], h["lat_baca"], p) for p in peta.sls[s]["poligon"])]
                dalam[h["baris"]] = hit[0] if hit else ""
        s = pilih_subsls(anggota, dalam)
        layak = []   # koordinat anggota yang sah utk subsls terpilih
        for h in anggota:
            j = jarak(s, h["lat_baca"], h["lon_baca"])
            if j is not None and j <= batas_m:
                layak.append((round(h["lat_baca"], 7), round(h["lon_baca"], 7), j, h["baris"], h["cara_baca"]))
        if s not in peta.sls:
            for h in anggota:     # subsls tak ada di peta: tidak bisa dinilai, pakai bacaan kalau di kabupaten
                ok = h["lat_baca"] is not None and h["lon_baca"] is not None and dikabupaten(h["lat_baca"], h["lon_baca"])
                h.update(lat=h["lat_baca"] if ok else None, lon=h["lon_baca"] if ok else None, subsls_koordinat=s,
                         sumber="SUBSLS_TIDAK_ADA_DI_PETA" + ("" if ok else " (koordinat dikosongkan)"))
            continue
        if layak:
            frek = Counter((a, b) for a, b, *_ in layak)
            pilih = min(layak, key=lambda x: (-frek[(x[0], x[1])], x[2], int(x[3])))
            lat, lon, sumber_baris = pilih[0], pilih[1], pilih[3]
            for h in anggota:
                sama = h["lat_baca"] is not None and (round(h["lat_baca"], 7), round(h["lon_baca"], 7)) == (lat, lon)
                sumber = (h["cara_baca"] if sama and h["cara_baca"] != "ASLI" else "ASLI" if sama
                          else f"KELOMPOK (dari baris {sumber_baris})")
                h.update(lat=lat, lon=lon, sumber=sumber, subsls_koordinat=s)
        else:
            t = titik_acak(s, peta, listing, jalan, benih(kunci, s), jarak_jalan_m)
            for h in anggota:
                h.update(lat=t[0] if t else None, lon=t[1] if t else None, sumber=t[2] if t else "TIDAK_ADA_TITIK",
                         subsls_koordinat=s)
    return hasil


# --------------------------------------------------------------------------
# I/O & CLI
# --------------------------------------------------------------------------
def baca_sheet(path: Path) -> tuple[list[dict], int]:
    """([{baris, idsubsls, pemilik, alamat, lat_mentah, lon_mentah, dinilai}], nomor baris
    terakhir). idsubsls/pemilik/alamat lewat load_tahap2 (idsubsls sudah diperbaiki spt
    batch); koordinat MENTAH dari sel supaya format aslinya bisa dinilai sendiri. Baris
    yang dilewati loader (kosong) tidak dinilai tapi posisinya tetap dihitung."""
    from inti.gabungan_loader import _sel as sel_teks
    from inti.tahap2_loader import _baca_mentah_tahap2, _indeks_tahap2, load_tahap2
    mentah = _baca_mentah_tahap2(path)
    idx = _indeks_tahap2([sel_teks(j) for j in mentah[0]])
    rows = {r.baris: r for r in load_tahap2(path)}
    out = []
    for nomor, isi in enumerate(mentah[1:], start=2):
        ambil = lambda k: isi[idx[k]] if k in idx and idx[k] < len(isi) else None  # noqa: E731
        r = rows.get(nomor)
        out.append({"baris": nomor, "idsubsls": r.idsubsls if r else "", "pemilik": r["pengusaha"] if r else "",
                    "alamat": r["jalan_domisili"] if r else "", "lat_mentah": ambil("latitude"),
                    "lon_mentah": ambil("longitude"), "dinilai": r is not None})
    return out, len(mentah)


def teks_sel(v) -> str:
    return "" if v is None else str(v).strip()


def nilai_tempel(h: dict) -> tuple[str, str, bool]:
    """(lat, lon, berubah) utk ditempel ke kolom Latitude/Longitude sheet. Kalau koordinat
    hasil = bacaan skrip input atas teks aslinya, teks ASLI dikembalikan persis (menempel
    ulang tidak mengubah apa pun di baris itu)."""
    from inti.tahap2_loader import desimal_ke_titik
    a, b = teks_sel(h["lat_mentah"]), teks_sel(h["lon_mentah"])
    if not h.get("dinilai", True):
        return a, b, False
    if h.get("lat") is None or h.get("lon") is None:
        return "", "", bool(a or b)
    try:
        sama = (abs(float(desimal_ke_titik(a)) - h["lat"]) < 1e-9
                and abs(float(desimal_ke_titik(b)) - h["lon"]) < 1e-9)
    except ValueError:
        sama = False
    if sama:
        return a, b, False
    f = lambda x: f"{x:.7f}".rstrip("0").rstrip(".")  # noqa: E731
    return f(h["lat"]), f(h["lon"]), True


JUDUL_TEMPEL = ["Latitude", "Longitude", "Berubah", "Sumber koordinat", "Latitude asli", "Longitude asli",
                "Jarak asli ke subsls (m)", "idsubsls", "Kelompok"]


def tulis_tempel(keluaran: Path, hasil: list[dict], terakhir: int) -> int:
    """Baris ke-N berkas ini = baris ke-N sheet sumber. Kolom A:B = Latitude/Longitude
    (disalin ke kolom Latitude/Longitude sheet mulai baris 2); C dst = keterangan."""
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "koordinat"
    for k, judul in enumerate(JUDUL_TEMPEL, start=1):
        ws.cell(1, k, judul)
    kuning = PatternFill("solid", fgColor="FFF2CC")
    per_baris = {h["baris"]: h for h in hasil}
    berubah = 0
    for r in range(2, terakhir + 1):
        h = per_baris.get(r)
        if h is None:
            continue
        lat, lon, ubah = nilai_tempel(h)
        berubah += ubah
        isi = [lat, lon, "YA" if ubah else "", h.get("sumber", ""), teks_sel(h["lat_mentah"]),
               teks_sel(h["lon_mentah"]), "" if h.get("jarak_asli_m") is None else h["jarak_asli_m"],
               h["idsubsls"], "" if h.get("kelompok", "baris:").startswith("baris:") else h["kelompok"]]
        for k, v in enumerate(isi, start=1):
            c = ws.cell(r, k, v)
            if k <= 2:
                c.number_format = "@"          # teks: Excel tidak mengubah "-8,14435" jadi angka/tanggal
                if ubah:
                    c.fill = kuning
    ws.freeze_panes = "A2"
    keluaran.parent.mkdir(parents=True, exist_ok=True)
    wb.save(keluaran)
    return berubah


# Server utama + cadangan (server Overpass sering sibuk: 429/504).
OVERPASS_URL = ("https://overpass-api.de/api/interpreter", "https://overpass.kumi.systems/api/interpreter")
# Tanpa User-Agent, overpass-api.de menolak urllib dgn HTTP 406 Not Acceptable (2026-09-24).
OVERPASS_HEADER = {"User-Agent": "split_usaha-koordinat/1.0 (python urllib)", "Accept": "application/json",
                   "Content-Type": "application/x-www-form-urlencoded"}
# Jalan yang bisa dilewati kendaraan + jalan tani (track); gang setapak (footway/path) tidak.
JENIS_JALAN = ("motorway|trunk|primary|secondary|tertiary|unclassified|residential|living_street|service|track")


def unduh_jalan(tujuan: Path, kotak=TAHAP2_KOTAK_KOORDINAT) -> None:
    """Unduh ruas jalan OpenStreetMap di kotak kabupaten (Overpass, `out geom`) ke `tujuan`."""
    import urllib.parse
    import urllib.request
    lat_min, lat_maks, lon_min, lon_maks = kotak
    kueri = (f'[out:json][timeout:300];way[highway~"^({JENIS_JALAN})(_link)?$"]'
             f"({lat_min},{lon_min},{lat_maks},{lon_maks});out geom;")
    data = urllib.parse.urlencode({"data": kueri}).encode()
    tujuan.parent.mkdir(parents=True, exist_ok=True)
    sementara = tujuan.with_name(tujuan.name + ".unduh")
    import time
    import urllib.error
    galat, berhasil = [], False
    for url in OVERPASS_URL:
        for ke in range(2):
            print(f"Mengunduh jalan OSM dari {url} (bisa 1-3 menit) ...")
            try:
                with urllib.request.urlopen(urllib.request.Request(url, data=data, headers=OVERPASS_HEADER),
                                            timeout=400) as r, open(sementara, "wb") as f:
                    while True:
                        blok = r.read(1 << 20)
                        if not blok:
                            break
                        f.write(blok)
                berhasil = True
                break
            except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
                galat.append(f"{url}: {e}")
                print(f"  ⚠️ {e}")
                if not (isinstance(e, urllib.error.HTTPError) and e.code in (429, 502, 503, 504)):
                    break                      # bukan "server sibuk": langsung coba server berikutnya
                if ke == 0:
                    time.sleep(30)
        if berhasil:
            break
    if not berhasil:
        raise SystemExit("❌ Unduhan jalan gagal di semua server:\n  " + "\n  ".join(galat))
    jalan = muat_jalan(sementara)          # pastikan isinya benar sebelum menggantikan berkas lama
    if not jalan.ruas:
        raise SystemExit(f"❌ Unduhan tidak berisi ruas jalan — lihat {sementara}")
    sementara.replace(tujuan)
    print(f"✅ {len(jalan.ruas)} ruas jalan -> {tujuan} ({tujuan.stat().st_size / 1_048_576:.1f} MB)")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sumber", required=True, help="sheet tahap 2 (.xlsx)")
    ap.add_argument("--keluaran", default="",
                    help="xlsx hasil (default koordinat/hasil/<nama sumber>_koordinat.xlsx)")
    ap.add_argument("--peta", default=PETA_SLS_PATH, help="GeoJSON poligon subsls (default PETA_SLS_PATH)")
    ap.add_argument("--listing", default=TITIK_LISTING_PATH, help="CSV titik listing (default TITIK_LISTING_PATH)")
    ap.add_argument("--jalan", default=JALAN_PATH, help="jaringan jalan GeoJSON/Overpass (default JALAN_PATH)")
    ap.add_argument("--batas-m", type=float, default=500.0, help="koordinat > sekian meter di luar subsls diganti")
    ap.add_argument("--jarak-jalan-m", type=float, default=50.0, help="'dekat jalan' = sejauh ini dari ruas jalan")
    ap.add_argument("--unduh-jalan", action="store_true",
                    help="unduh dulu jaringan jalan OpenStreetMap ke --jalan (sekali saja, butuh internet)")
    args = ap.parse_args(argv)

    if args.unduh_jalan:
        if not args.jalan:
            print("❌ Isi JALAN_PATH di inti/config_lokal.py atau --jalan <berkas.json>.")
            return 2
        unduh_jalan(Path(args.jalan))
    from koordinat.peta import muat_peta
    sumber = Path(args.sumber)
    keluaran = Path(args.keluaran) if args.keluaran else HASIL_DIR / f"{sumber.stem}_koordinat.xlsx"
    if keluaran.resolve() == sumber.resolve():
        print("❌ --keluaran tidak boleh sama dgn --sumber (berkas sumber tidak ditimpa).")
        return 2
    if not args.peta or not Path(args.peta).exists():
        print("❌ Peta poligon subsls tidak ada — isi PETA_SLS_PATH di inti/config_lokal.py atau --peta.")
        return 2
    peta = muat_peta(args.peta)
    listing = None
    if args.listing and Path(args.listing).exists():
        listing = muat_listing(args.listing, TAHAP2_KOTAK_KOORDINAT)
        print(f"Titik listing (pemukiman): {sum(len(v) for v in listing.grid.values())} dari {args.listing}")
    else:
        print("⚠️  Tanpa titik listing — pengganti diambil acak di poligon (tidak dijamin di pemukiman).")
    jalan = None
    if args.jalan and Path(args.jalan).exists():
        jalan = muat_jalan(args.jalan)
        print(f"Jaringan jalan: {len(jalan.ruas)} ruas dari {args.jalan}")
    else:
        print("⚠️  Tanpa data jalan — syarat 'dekat jalan' tidak dipakai.")

    semua, terakhir = baca_sheet(sumber)
    hasil = rencana([b for b in semua if b["dinilai"]], peta, listing, jalan, args.batas_m, args.jarak_jalan_m)
    hasil += [b for b in semua if not b["dinilai"]]          # baris kosong: posisi tetap, isi apa adanya

    ringkas = Counter(re.sub(r" \(dari baris \d+\)", "", h["sumber"]) for h in hasil if h.get("sumber"))
    print(f"\n=== {len(hasil)} baris — asal koordinat hasil ===")
    for k, n in ringkas.most_common():
        print(f"  {n:5d}  {k}")
    kel = Counter(h["kelompok"] for h in hasil if not h.get("kelompok", "baris:").startswith("baris:"))
    print(f"Kelompok pemilik+alamat (>1 baris): {sum(1 for n in kel.values() if n > 1)} kelompok, "
          f"{sum(n for n in kel.values() if n > 1)} baris")
    beda = sum(1 for h in hasil if h.get("subsls_koordinat") and h["subsls_koordinat"] != h["idsubsls"])
    if beda:
        print(f"Baris yang koordinatnya ikut subsls mayoritas kelompoknya (bukan subsls barisnya): {beda}")

    berubah = tulis_tempel(keluaran, hasil, terakhir)
    print(f"\n✅ {keluaran}  — {berubah} baris berubah (kuning); baris lain = teks aslinya.")
    print(f"   Salin A2:B{terakhir} -> tempel (Paste Values) ke sel Latitude baris 2 di {sumber.name}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
