#!/usr/bin/env python3
"""
monitoring.py — monitoring hasil input usaha: SEMUA status (approved, terkirim, draft,
belum diinput, ditolak pemeriksaan data, ...) termasuk dokumen GANDA, direkap per
daerah (kecamatan, desa, SLS), per PPL & per batch.

OFFLINE & READ-ONLY: tidak membuka browser, tidak mengubah audit/sheet. Aman dijalankan
kapan saja, termasuk saat bot input/approve berjalan.

Sumber (ketetapan user 2026-09-29):
  1. Sheet input usaha (bawaan: semua bahan/input_tahap2*.xlsx) — satu baris = satu USAHA.
     Pemeriksaan data sama persis dgn `jalankan.py --cek` (baris ditolak = tidak akan diinput).
  2. Semua audit input di audit/** (tanpa cadangan .bak-*) + audit/audit_approve_pml.csv +
     audit/ganda_dihapus*.csv — dokumen tiap baris & status menurut catatan skrip.
  3. SNAPSHOT fasih-sm (opsional tapi disarankan): CSV unduhan monitoring_console.js dari akun
     admin fasih-sm = status SERVER (DRAFT/SUBMITTED/APPROVED/REJECTED), wilayah dokumen sekarang,
     petugas. Kalau ada, status server MENANG atas audit; dokumen server bernama sama yang tidak
     tercatat di audit ikut dikenali (ganda di luar audit).

Daerah = lokasi usaha di sheet (kolom 5 idsubsls), bukan subsls wadah tempat dokumen dibuat.
Nama kecamatan/desa/SLS dari peta SLS (config PETA_SLS_PATH), cadangan kolom 3/4 sheet.

Keluaran (monitoring/hasil/):
  monitoring.xlsx   Ringkasan, Per kecamatan, Per desa, Per SLS, Per PPL, Ganda, Di luar sheet, Detail
  monitoring.html   dasbor satu berkas (grafik, filter batch/kecamatan/desa, cari) — buka di browser

Contoh:
    python monitoring/monitoring.py
    python monitoring/monitoring.py --snapshot bahan/snapshot_fasih_sm_20260929-1000.csv
    python monitoring/monitoring.py --sumber bahan/input_tahap2_24.xlsx --sumber bahan/input_tahap2_25.xlsx
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from inti import lokasi  # noqa: E402
from inti.config import (  # noqa: E402
    ASSIGNMENT_ID_GABUNGAN, FASIH_WEB_BASE, KODE_KAB, PETA_SLS_PATH, SURVEY_ID, WILAYAH_BY_IDSUBSLS,
)
import input_usaha.mesin as mg  # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

HASIL = lokasi.hasil("monitoring")
POLA_SUMBER = str(lokasi.BAHAN / "input_tahap2*.xlsx")
NAMA_SNAPSHOT = "snapshot_fasih_sm*.csv"
POLA_AUDIT = lokasi.AUDIT / "**" / "audit_log_gabungan*.csv"
POLA_KODE = re.compile(rf"{re.escape(KODE_KAB)}\d{{12}}")
# Kode identitas prelist ("<subsls> - UMK - 19", "- DTSEN - 3"): bukan dokumen input usaha.
POLA_PRELIST = re.compile(r"^\d{16}\s*-\s*[A-Z]+\s*-\s*\d+$")

# Kategori per USAHA (urut tampil). Kode -> label.
KATEGORI = {
    "APPROVED": "Approved",
    "SUBMITTED": "Terkirim (belum approve)",
    "REJECTED": "Ditolak PML (rejected)",
    "DRAFT": "Draft",
    "LAIN": "Status server lain",
    "PERLU_CEK": "Dokumen tanpa URL (perlu cek)",
    "TIDAK_DI_SERVER": "Dokumen tidak ada di server",
    "BELUM_INPUT": "Belum diinput",
    "DITOLAK_DATA": "Ditolak pemeriksaan data",
}
# Peringkat status DOKUMEN utk memilih dokumen utama satu usaha (besar = lebih maju).
PERINGKAT = {"APPROVED": 6, "SUBMITTED": 5, "REJECTED": 4, "LAIN": 3, "DRAFT": 2, "TIDAK_DI_SERVER": 0}
KATEGORI_TERKIRIM = ("APPROVED", "SUBMITTED")
STATUS_APPROVE_OK = {"APPROVED_TERVERIFIKASI", "SUDAH_APPROVED"}

KOLOM_DETAIL = [
    "batch", "baris", "kunci", "nama_usaha", "pengusaha", "ppl", "idsubsls",
    "kec_kode", "kecamatan", "desa_kode", "desa", "sls_kode", "sls",
    "kategori", "keterangan", "status_data", "pesan_data",
    "jumlah_dokumen", "ganda", "id_dokumen", "status_server", "status_audit", "sumber_status",
    "subsls_dokumen", "di_wilayah", "petugas", "galat_server", "dokumen_url",
]
KOLOM_GANDA = ["grup", "jenis", "keterangan", "batch", "baris", "nama_usaha", "kecamatan", "desa", "idsubsls",
               "id_dokumen", "status", "sumber_status", "subsls_dokumen", "petugas", "dicatat_di", "dokumen_url"]
KOLOM_LUAR = ["id_dokumen", "kode_identitas", "nama", "status", "subsls_dokumen", "petugas", "dibuat", "dokumen_url"]
JENIS_GANDA = {
    "USAHA_DOKUMEN_GANDA": "Satu baris usaha punya >= 2 dokumen hidup — pasti ganda (lihat fasih_sm/hapus_ganda)",
    "DOKUMEN_DIPAKAI_BANYAK_BARIS": "Satu dokumen tercatat utk >= 2 baris sheet (baris identik / kunci sama)",
    "NAMA_SAMA_SATU_BATCH": "Dokumen bernama sama dari baris BERBEDA di batch yang sama — periksa",
    "NAMA_SAMA_BEDA_BATCH": "Dokumen bernama sama dari batch berbeda — bawaan dipertahankan (bukan ganda)",
    "NAMA_SAMA_LUAR_SHEET": "Dokumen server bernama sama, tidak dikenal sheet mana pun",
}


def norm(s) -> str:
    return " ".join(str(s or "").split()).upper()


def id_dari_url(url: str) -> str:
    bagian = [p for p in str(url or "").split("/") if p]
    return bagian[-2] if len(bagian) >= 2 and bagian[-1] == "entry" else ""


def url_entry(did: str) -> str:
    return f"{FASIH_WEB_BASE}/survey/{SURVEY_ID}/{ASSIGNMENT_ID_GABUNGAN}/{did}/entry" if did else ""


def kategori_server(alias: str) -> str:
    """assignmentStatusAlias -> kategori dokumen."""
    a = (alias or "").strip().upper()
    if a.startswith("APPROVED"):
        return "APPROVED"
    if a.startswith("SUBMITTED"):
        return "SUBMITTED"
    if a.startswith("REJECTED"):
        return "REJECTED"
    if not a or a.startswith("DRAFT") or a.startswith("OPEN") or a.startswith("ASSIGNED"):
        return "DRAFT"
    return "LAIN"


def kategori_audit(status: str) -> str:
    """Status audit input dokumen yang MASIH ADA -> kategori (terkirim menurut skrip / draft)."""
    return "SUBMITTED" if (status or "").strip() in mg.STATUS_TERKIRIM else "DRAFT"


# ------------------------------------------------------------------------------------------
# Audit
# ------------------------------------------------------------------------------------------
@dataclass
class InfoAudit:
    """Isi satu audit input yang relevan utk monitoring."""
    path: Path
    ids_kunci: dict = field(default_factory=dict)     # kunci -> {id} dokumen hidup menurut audit ini
    dokumen: dict = field(default_factory=dict)       # id -> {status, ts, akun, subsls, kunci}
    dihapus: set = field(default_factory=set)         # id yang digugurkan DOKUMEN_DIHAPUS
    tanpa_url: set = field(default_factory=set)       # kunci bertanda DOKUMEN_TANPA_URL (belum ber-URL)
    status_kunci: dict = field(default_factory=dict)  # kunci -> status terakhir (aturan mesin)


def ringkas_audit(path: Path, baris: list[dict]) -> InfoAudit:
    """Fungsi murni: baris audit (urutan berkas) -> InfoAudit. DOKUMEN_DIHAPUS menggugurkan SEMUA
    dokumen kunci itu (sama dgn mesin.dokumen_dari & gabung_audit.daftar_ganda)."""
    info = InfoAudit(Path(path))
    for b in baris:
        kunci = (b.get("kunci") or "").strip()
        if not kunci:
            continue
        status = (b.get("status") or "").strip()
        if status == mg.STATUS_DIHAPUS:
            info.dihapus |= info.ids_kunci.pop(kunci, set())
            info.tanpa_url.discard(kunci)
            continue
        did = id_dari_url(b.get("dokumen_url") or "")
        if not did:
            if status in mg.STATUS_TANPA_URL_SEMUA:
                info.tanpa_url.add(kunci)
            continue
        info.tanpa_url.discard(kunci)
        info.ids_kunci.setdefault(kunci, set()).add(did)
        lama = info.dokumen.get(did, {})
        # DOKUMEN_DIBUAT sesudah status lain utk dokumen yang sama = catatan sinkron, bukan kemunduran.
        if status == mg.STATUS_DIBUAT and lama.get("status") not in (None, "", mg.STATUS_DIBUAT):
            status = lama["status"]
        info.dokumen[did] = {
            "status": status, "ts": (b.get("timestamp") or "").strip() or lama.get("ts", ""),
            "akun": (b.get("akun_login") or "").strip().lower() or lama.get("akun", ""),
            "subsls": (b.get("idsubsls_input") or "").strip() or lama.get("subsls", ""),
            "kunci": kunci, "dibuat": lama.get("dibuat") or (b.get("timestamp") or "").strip(),
        }
    info.status_kunci = mg.status_terakhir_dari(baris)
    return info


def audit_kandidat() -> list[Path]:
    """Semua audit input di audit/** (termasuk audit/pc/** kiriman PC lain), tanpa cadangan .bak-*."""
    return [f for f in lokasi.cari(POLA_AUDIT) if ".bak-" not in f.name]


def baca_audit_semua(paths: list[Path]) -> list[InfoAudit]:
    hasil = []
    for p in paths:
        try:
            with p.open(newline="", encoding="utf-8-sig") as f:
                pembaca = csv.DictReader(f)
                if not {"kunci", "status"} <= set(pembaca.fieldnames or []):
                    continue
                baris = list(pembaca)
        except (OSError, UnicodeDecodeError) as e:
            print(f"⚠️ {p}: tidak terbaca ({e}) — dilewati.")
            continue
        rusak = mg.kerusakan_excel(baris)
        if rusak:
            print(f"⚠️ {p}: pernah disimpan Excel (kunci rusak) — DILEWATI. Pulihkan dgn antar_pc/pulihkan_excel.py.")
            continue
        hasil.append(ringkas_audit(p, baris))
    return hasil


def audit_milik(kunci_sheet: set, audits: list[InfoAudit]) -> list[InfoAudit]:
    """Audit yang 'memiliki' satu sheet: mengenal dokumen >= min(20, terbaik) kunci sheet DAN >= 10%
    dari audit yang paling mengenalnya. Menyaring audit batch lain yang kebetulan berbagi beberapa
    baris identik (batch 21 & 22 berbagi 1 kunci) supaya dokumen batch lain tidak dikira ganda."""
    kenal = [(a, len(kunci_sheet & set(a.ids_kunci))) for a in audits]
    terbaik = max((n for _, n in kenal), default=0)
    if not terbaik:
        return []
    batas = max(min(20, terbaik), terbaik * 0.1)
    return [a for a, n in kenal if n and n >= batas]


def dokumen_global(audits: list[InfoAudit]) -> tuple[dict, set]:
    """({id: info terbaru di audit mana pun}, {id dihapus}). Antar-audit: baris ber-timestamp terbaru."""
    per_id: dict = {}
    dihapus: set = set()
    for a in audits:
        dihapus |= a.dihapus
        for did, info in a.dokumen.items():
            lama = per_id.get(did)
            if lama is None or info["ts"] > lama["ts"]:
                per_id[did] = {**info, "audit": a.path}
    return per_id, dihapus


def baca_approve(path: Path = lokasi.AUDIT_APPROVE) -> dict:
    """{id: status approve terakhir} dari audit_approve_pml.csv."""
    out: dict = {}
    if not Path(path).exists():
        return out
    with Path(path).open(newline="", encoding="utf-8-sig") as f:
        for b in csv.DictReader(f):
            did = (b.get("id") or "").strip() or id_dari_url(b.get("dokumen_url") or "")
            if did:
                out[did] = (b.get("status") or "").strip()
    return out


def baca_ganda_dihapus() -> set:
    out: set = set()
    for f in lokasi.cari(lokasi.POLA_GANDA_DIHAPUS):
        with f.open(newline="", encoding="utf-8-sig") as fh:
            for b in csv.DictReader(fh):
                if (b.get("id") or "").strip() and "TERVERIFIKASI" in (b.get("status") or ""):
                    out.add(b["id"].strip())
    return out


# ------------------------------------------------------------------------------------------
# Snapshot fasih-sm
# ------------------------------------------------------------------------------------------
def _kode_region(region) -> str:
    """Kode wilayah terkecil dari region datatable (level1.fullCode ...) / detail (level_1.full_code)."""
    kode, node = "", region if isinstance(region, dict) else {}
    for n in range(1, 11):
        node = node.get(f"level{n}") or node.get(f"level_{n}") if isinstance(node, dict) else None
        if not isinstance(node, dict):
            break
        kode = str(node.get("fullCode") or node.get("full_code") or kode)
    return kode


def _teks(v) -> str:
    return "" if v is None else str(v).strip()


def ringkas_snapshot(it: dict) -> dict | None:
    """Satu baris CSV monitoring_console.js ATAU satu item datatable (list_api_*.json) -> bentuk baku."""
    did = (it.get("id") or "").strip()
    if not did:
        return None
    mode = it.get("mode")
    if isinstance(mode, list):
        mode = ",".join(mode)
    return {
        "id": did,
        "kode_identitas": (it.get("kode_identitas") or it.get("codeIdentity") or "").strip(),
        "nama": norm(it.get("nama") if "nama" in it else it.get("data1")),
        "status": (it.get("status") or it.get("assignmentStatusAlias") or "").strip(),
        "mode": (mode or "").strip(),
        "subsls": (it.get("subsls") or _kode_region(it.get("region")) or "").strip(),
        "petugas": (it.get("petugas") or it.get("currentUserUsername") or "").strip().lower(),
        "galat": _teks(it.get("galat") if "galat" in it else it.get("sumError")),
        "dibuat": (it.get("dibuat") or it.get("dateCreated") or "").strip(),
        "diambil": (it.get("diambil") or "").strip(),
        "total_server": str(it.get("total_server") or "").strip(),
    }


def nama_dari_kode(kode: str) -> str:
    return norm(re.sub(r"^\s*\d{16}\s*-\s*", "", kode or ""))


def baca_snapshot(paths: list[Path]) -> tuple[dict, dict]:
    """({id: dokumen}, info). Beberapa berkas digabung (yang terakhir menang)."""
    dok: dict = {}
    info = {"berkas": [], "diambil": "", "total_server": 0}
    for p in paths:
        if p.suffix.lower() == ".json":
            items = json.loads(p.read_text(encoding="utf-8"))
        else:
            with p.open(newline="", encoding="utf-8-sig") as f:
                items = list(csv.DictReader(f))
        n = 0
        for it in items:
            d = ringkas_snapshot(it)
            if d:
                dok[d["id"]] = d
                n += 1
                info["diambil"] = max(info["diambil"], d["diambil"])
                if d["total_server"].isdigit():
                    info["total_server"] = max(info["total_server"], int(d["total_server"]))
        info["berkas"].append((p, n))
    return dok, info


def cari_snapshot() -> list[Path]:
    """Snapshot terbaru di bahan/ atau folder Unduhan (Downloads) pengguna."""
    calon = lokasi.cari(lokasi.BAHAN / NAMA_SNAPSHOT) + lokasi.cari(HASIL / NAMA_SNAPSHOT)
    unduhan = Path.home() / "Downloads"
    if unduhan.is_dir():
        calon += lokasi.cari(unduhan / NAMA_SNAPSHOT)
    return [max(calon, key=lambda p: p.stat().st_mtime)] if calon else []


def waktu_lokal(iso: str) -> str:
    """ISO UTC (…Z / +00:00) -> 'YYYY-MM-DD HH:MM:SS' zona PC ini (sama dgn timestamp audit)."""
    try:
        t = dt.datetime.fromisoformat(iso.replace("Z", "+00:00"))
    except ValueError:
        return ""
    if t.tzinfo is None:
        return t.strftime("%Y-%m-%d %H:%M:%S")
    return t.astimezone().strftime("%Y-%m-%d %H:%M:%S")


# ------------------------------------------------------------------------------------------
# Wilayah
# ------------------------------------------------------------------------------------------
def baca_nama_wilayah(path: str | Path = "") -> dict:
    """{kode kec (7) / desa (10) / sls (14): nama} dari peta SLS GeoJSON; {} kalau tidak ada."""
    p = Path(path or PETA_SLS_PATH or "")
    if not str(path or PETA_SLS_PATH or "") or not p.exists():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        print(f"⚠️ Peta SLS {p} tidak terbaca ({e}) — nama wilayah dari sheet.")
        return {}
    nama: dict = {}
    for f in data.get("features", []):
        pr = f.get("properties") or {}
        ids = str(pr.get("idsubsls") or "")
        if len(ids) != 16:
            continue
        nama.setdefault(ids[:7], str(pr.get("nmkec") or "").strip())
        nama.setdefault(ids[:10], str(pr.get("nmdesa") or "").strip())
        nama.setdefault(ids[:14], str(pr.get("nmsls") or "").strip())
    return nama


def wilayah_baris(idsubsls: str, nama: dict, info: dict | None = None) -> dict:
    """Kode & nama kecamatan/desa/SLS lokasi usaha. Cadangan nama: WILAYAH_BY_IDSUBSLS, kolom 3/4 sheet."""
    kode = idsubsls if POLA_KODE.fullmatch(idsubsls or "") else ""
    w = WILAYAH_BY_IDSUBSLS.get(kode, {}) if kode else {}
    info = info or {}

    def pilih(k, *cadangan):
        return (nama.get(k) if k else "") or next((c for c in cadangan if c and not str(c).isdigit()), "")
    return {
        "kec_kode": kode[:7], "kecamatan": pilih(kode[:7], w.get("kecamatan"), info.get("kec")).upper(),
        "desa_kode": kode[:10], "desa": pilih(kode[:10], w.get("desa"), info.get("desa")).upper(),
        "sls_kode": kode[:14], "sls": pilih(kode[:14], w.get("sls")).upper(),
    }


# ------------------------------------------------------------------------------------------
# Inti: satu baris per usaha + daftar ganda
# ------------------------------------------------------------------------------------------
@dataclass
class Sheet:
    label: str
    rows: list
    periksa: dict
    audits: list  # InfoAudit milik sheet ini


def status_dokumen(did: str, audit_global: dict, approve: dict, snapshot: dict | None,
                   snapshot_lengkap: bool, snapshot_waktu: str) -> dict:
    """Status satu dokumen hidup: server (snapshot) > approve > audit input."""
    a = audit_global.get(did, {})
    ap = approve.get(did, "")
    if snapshot is not None and did in snapshot:
        s = snapshot[did]
        return {"kategori": kategori_server(s["status"]), "status_server": s["status"], "status_audit": a.get("status", ""),
                "sumber": "server", "subsls": s["subsls"], "petugas": s["petugas"], "galat": s["galat"]}
    hasil = {"status_server": "", "status_audit": a.get("status", ""), "subsls": a.get("subsls", ""),
             "petugas": a.get("akun", ""), "galat": ""}
    # Dokumen tidak ada di snapshot yang LENGKAP & dibuat sebelum snapshot diambil = hilang dari server.
    if snapshot is not None and snapshot_lengkap and (not a.get("dibuat") or a["dibuat"] < snapshot_waktu):
        return {**hasil, "kategori": "TIDAK_DI_SERVER", "sumber": "server"}
    if ap in STATUS_APPROVE_OK:
        return {**hasil, "kategori": "APPROVED", "sumber": "audit approve"}
    if not a:
        return {**hasil, "kategori": "DRAFT", "sumber": "ID sheet"}
    return {**hasil, "kategori": kategori_audit(a["status"]), "sumber": "audit"}


def _nama_asal(asal) -> str:
    """Path audit -> relatif thd akar proyek (lebih pendek di laporan); teks lain apa adanya."""
    if isinstance(asal, Path):
        try:
            return str(asal.resolve().relative_to(lokasi.AKAR))
        except ValueError:
            return str(asal)
    return str(asal)


def bangun(sheets: list[Sheet], audit_global: dict, dihapus: set, approve: dict, snapshot: dict | None,
           snapshot_info: dict | None, nama_wilayah: dict) -> tuple[list[dict], list[dict], list[dict]]:
    """Fungsi murni -> (detail per usaha, ganda per dokumen, dokumen server di luar sheet)."""
    snap_info = snapshot_info or {}
    # Lengkap HANYA kalau jumlah server tercatat (CSV monitoring_console.js) & semuanya terbaca. Tanpa itu
    # (list_api_*.json satu akun, unduhan terpotong) dokumen yang tak terbaca TIDAK dianggap hilang.
    snapshot_lengkap = bool(snapshot) and bool(snap_info.get("total_server"))         and len(snapshot) >= snap_info["total_server"]
    snapshot_waktu = waktu_lokal(snap_info.get("diambil", "")) if snapshot else ""

    # 1) dokumen per baris: ID sheet + audit milik sheet
    baris_info: list[dict] = []
    pemilik_id: dict = defaultdict(list)          # id -> [indeks baris_info]
    for sh in sheets:
        for row in sh.rows:
            ids = {row.id_dokumen} if getattr(row, "id_dokumen", "") else set()
            tercatat: dict = {}                    # id -> path audit
            for a in sh.audits:
                for did in a.ids_kunci.get(row.kunci, ()):
                    ids.add(did)
                    tercatat.setdefault(did, a.path)
            hidup = {i for i in ids if i not in dihapus}
            b = {"sheet": sh, "row": row, "ids": hidup, "ids_semua": ids, "tercatat": tercatat,
                 "tanpa_url": any(row.kunci in a.tanpa_url for a in sh.audits),
                 "status_audit": next((a.status_kunci[row.kunci] for a in reversed(sh.audits)
                                       if row.kunci in a.status_kunci), "")}
            for did in hidup:
                pemilik_id[did].append(len(baris_info))
            baris_info.append(b)

    # 2) dokumen server bernama sama yang belum dikenal -> baris pemilik nama (kalau tunggal)
    luar: list[dict] = []
    nama_luar: dict = defaultdict(list)
    if snapshot:
        per_nama: dict = defaultdict(list)
        for i, b in enumerate(baris_info):
            per_nama[norm(b["row"].nama_dokumen)].append(i)
        for did, s in snapshot.items():
            if did in pemilik_id or did in dihapus:
                continue
            nama = s["nama"] or nama_dari_kode(s["kode_identitas"])
            if POLA_PRELIST.match(s["kode_identitas"]) and not s["nama"]:
                continue
            calon = per_nama.get(nama, [])
            if len(calon) == 1:
                i = calon[0]
                baris_info[i]["ids"].add(did)
                baris_info[i]["tercatat"].setdefault(did, "server (nama sama, di luar audit)")
                pemilik_id[did].append(i)
            elif not calon:
                if POLA_PRELIST.match(s["kode_identitas"]):
                    continue
                luar.append({"id_dokumen": did, "kode_identitas": s["kode_identitas"], "nama": nama,
                             "status": s["status"], "subsls_dokumen": s["subsls"], "petugas": s["petugas"],
                             "dibuat": waktu_lokal(s["dibuat"]) or s["dibuat"], "dokumen_url": url_entry(did)})
                nama_luar[nama].append(did)
            # calon > 1: nama dipakai beberapa baris -> tidak ditebak (dilaporkan lewat NAMA_SAMA_*)

    # 3) status tiap dokumen & kategori per usaha
    st_cache: dict = {}

    def st(did):
        if did not in st_cache:
            st_cache[did] = status_dokumen(did, audit_global, approve, snapshot, snapshot_lengkap, snapshot_waktu)
        return st_cache[did]

    detail: list[dict] = []
    for b in baris_info:
        row, sh = b["row"], b["sheet"]
        pr = sh.periksa.get(row.baris)
        hidup = [d for d in b["ids"] if st(d)["kategori"] != "TIDAK_DI_SERVER"]
        id_sheet = getattr(row, "id_dokumen", "") or ""
        urut = sorted(b["ids"], key=lambda d: (-PERINGKAT.get(st(d)["kategori"], 1), d != id_sheet, d))
        utama = urut[0] if urut else ""
        ket = []
        if utama:
            s = st(utama)
            kat = s["kategori"]
            if kat == "DRAFT":
                if s["status_audit"] == mg.STATUS_DRAFT_TANPA_KOORDINAT:
                    ket.append("tanpa koordinat")
                if str(s["galat"]).isdigit() and int(s["galat"]) > 0:
                    ket.append(f"galat server {s['galat']}")
                elif s["status_audit"] == mg.STATUS_DRAFT_GALAT:
                    ket.append("draft ber-galat")
            if kat == "TIDAK_DI_SERVER":
                ket.append("dokumen tercatat tapi tidak ada di snapshot server")
        elif b["tanpa_url"]:
            kat, s = "PERLU_CEK", {}
            ket.append("audit: dokumen mungkin terbuat tanpa URL — jalankan sinkron_list")
        else:
            s = {}
            bisa = pr is None or pr.bisa_diproses
            kat = "BELUM_INPUT" if bisa else "DITOLAK_DATA"
            if b["ids_semua"]:
                ket.append("dokumen lama sudah dihapus")
            if pr is not None and pr.tanpa_koordinat and bisa:
                ket.append("tanpa koordinat")
        if len(hidup) > 1:
            ket.append(f"{len(hidup)} dokumen hidup (GANDA)")
        w = wilayah_baris(row.idsubsls, nama_wilayah, getattr(row, "info", {}))
        subsls_dok = s.get("subsls", "")
        detail.append({
            "batch": sh.label, "baris": row.baris, "kunci": row.kunci, "nama_usaha": row.nama_dokumen,
            "pengusaha": (row.v.get("pengusaha") or "").strip(),
            "ppl": (getattr(row, "info", {}).get("nama_ppl") or row.v.get("akun_ppl") or "").strip().upper(),
            "idsubsls": row.idsubsls, **w,
            "kategori": kat, "keterangan": "; ".join(ket),
            "status_data": pr.status if pr is not None else "", "pesan_data": pr.pesan if pr is not None else "",
            "jumlah_dokumen": len(hidup), "ganda": "ya" if len(hidup) > 1 else "",
            "id_dokumen": utama, "status_server": s.get("status_server", ""), "status_audit": b["status_audit"],
            "sumber_status": s.get("sumber", ""), "subsls_dokumen": subsls_dok,
            "di_wilayah": ("" if not subsls_dok or not utama else "ya" if subsls_dok == row.idsubsls else "tidak"),
            "petugas": s.get("petugas", ""), "galat_server": s.get("galat", ""),
            "dokumen_url": url_entry(utama),
        })

    # 4) ganda
    ganda: list[dict] = []
    no = 0

    def tambah(jenis, anggota):
        nonlocal no
        no += 1
        for i, did in anggota:
            d = detail[i] if i is not None else {}
            s = st(did) if i is not None else {}
            sv = (snapshot or {}).get(did, {})
            ganda.append({
                "grup": no, "jenis": jenis, "keterangan": JENIS_GANDA[jenis],
                "batch": d.get("batch", ""), "baris": d.get("baris", ""),
                "nama_usaha": d.get("nama_usaha", "") or sv.get("nama", ""),
                "kecamatan": d.get("kecamatan", ""), "desa": d.get("desa", ""), "idsubsls": d.get("idsubsls", ""),
                "id_dokumen": did, "status": s.get("status_server") or s.get("status_audit") or sv.get("status", ""),
                "sumber_status": s.get("sumber", "server" if sv else ""),
                "subsls_dokumen": s.get("subsls", "") or sv.get("subsls", ""),
                "petugas": s.get("petugas", "") or sv.get("petugas", ""),
                "dicatat_di": _nama_asal(baris_info[i]["tercatat"].get(did, "ID sheet")) if i is not None else "",
                "dokumen_url": url_entry(did),
            })

    for i, b in enumerate(baris_info):
        hidup = sorted((d for d in b["ids"] if st(d)["kategori"] != "TIDAK_DI_SERVER"),
                       key=lambda d: (-PERINGKAT.get(st(d)["kategori"], 1), d))
        if len(hidup) > 1:
            tambah("USAHA_DOKUMEN_GANDA", [(i, d) for d in hidup])
    for did, pemilik in sorted(pemilik_id.items()):
        if len(set(pemilik)) > 1 and st(did)["kategori"] != "TIDAK_DI_SERVER":
            tambah("DOKUMEN_DIPAKAI_BANYAK_BARIS", [(i, did) for i in sorted(set(pemilik))])
    # nama sama antar baris berbeda (nama server kalau ada, kalau tidak nama dokumen sheet)
    per_nama_dok: dict = defaultdict(list)
    for i, d in enumerate(detail):
        for did in baris_info[i]["ids"]:
            if st(did)["kategori"] == "TIDAK_DI_SERVER":
                continue
            nama = (snapshot or {}).get(did, {}).get("nama") or norm(d["nama_usaha"])
            per_nama_dok[nama].append((i, did))
    for nama, anggota in sorted(per_nama_dok.items()):
        baris_beda = {detail[i]["batch"] + "|" + str(detail[i]["baris"]) for i, _ in anggota}
        ids_beda = {did for _, did in anggota}
        if len(baris_beda) < 2 or len(ids_beda) < 2:
            continue
        batch_beda = {detail[i]["batch"] for i, _ in anggota}
        tambah("NAMA_SAMA_SATU_BATCH" if len(batch_beda) == 1 else "NAMA_SAMA_BEDA_BATCH", anggota)
    for nama, ids in sorted(nama_luar.items()):
        if len(ids) > 1:
            tambah("NAMA_SAMA_LUAR_SHEET", [(None, d) for d in ids])
    return detail, ganda, luar


# ------------------------------------------------------------------------------------------
# Rekap
# ------------------------------------------------------------------------------------------
def rekap(detail: list[dict], kunci: tuple[str, ...], nama: tuple[str, ...] = ()) -> list[dict]:
    """Jumlah usaha per kategori utk pengelompokan `kunci` (mis. ("kec_kode",)). Fungsi murni."""
    per: dict = {}
    for d in detail:
        k = tuple(d.get(x, "") for x in kunci)
        r = per.get(k)
        if r is None:
            r = per[k] = {**{x: d.get(x, "") for x in kunci}, **{x: d.get(x, "") for x in nama},
                          "total": 0, **{kat: 0 for kat in KATEGORI}, "usaha_ganda": 0, "dokumen_berlebih": 0,
                          "di_wilayah": 0, "belum_di_wilayah": 0}
        r["total"] += 1
        r[d["kategori"]] += 1
        if d["jumlah_dokumen"] > 1:
            r["usaha_ganda"] += 1
            r["dokumen_berlebih"] += d["jumlah_dokumen"] - 1
        if d["di_wilayah"] == "ya":
            r["di_wilayah"] += 1
        elif d["di_wilayah"] == "tidak":
            r["belum_di_wilayah"] += 1
    keluar = []
    for r in per.values():
        terkirim = sum(r[k] for k in KATEGORI_TERKIRIM)
        target = r["total"] - r["DITOLAK_DATA"]
        r["pct_terkirim"] = round(terkirim / target, 4) if target else 0.0
        r["pct_approved"] = round(r["APPROVED"] / target, 4) if target else 0.0
        keluar.append(r)
    keluar.sort(key=lambda r: tuple(str(r[x]) for x in kunci))
    return keluar


# ------------------------------------------------------------------------------------------
# Keluaran
# ------------------------------------------------------------------------------------------
KOLOM_REKAP_ANGKA = ["total", *KATEGORI, "pct_terkirim", "pct_approved", "usaha_ganda", "dokumen_berlebih",
                     "di_wilayah", "belum_di_wilayah"]
JUDUL_KOLOM = {
    "total": "Total usaha", **KATEGORI, "pct_terkirim": "% terkirim", "pct_approved": "% approved",
    "usaha_ganda": "Usaha ganda", "dokumen_berlebih": "Dokumen berlebih", "di_wilayah": "Dok. di wilayah usaha",
    "belum_di_wilayah": "Dok. masih di wadah", "kec_kode": "Kode kec", "kecamatan": "Kecamatan",
    "desa_kode": "Kode desa", "desa": "Desa", "sls_kode": "Kode SLS", "sls": "SLS", "ppl": "PPL", "batch": "Batch",
}
WARNA_KATEGORI = {"APPROVED": "C8DCF5", "SUBMITTED": "C5EBDD", "REJECTED": "F6C6D8", "DRAFT": "FBE3AE",
                  "LAIN": "E0DDF5", "PERLU_CEK": "F6C6D8", "TIDAK_DI_SERVER": "F6C6D8", "BELUM_INPUT": "E7E6E2",
                  "DITOLAK_DATA": "F6C6D8"}


def tulis_excel(path: Path, meta: list[tuple[str, str]], lembar_rekap: list[tuple[str, list[dict], list[str]]],
                ganda: list[dict], luar: list[dict], detail: list[dict]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    tebal = Font(bold=True)
    judul_isi = PatternFill("solid", fgColor="DDE7F3")

    def tabel(ws, baris0: int, kolom: list[str], data: list[dict], judul=None, persen=()):
        for j, k in enumerate(kolom, start=1):
            c = ws.cell(baris0, j, (judul or {}).get(k, k))
            c.font, c.fill = tebal, judul_isi
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, d in enumerate(data, start=baris0 + 1):
            for j, k in enumerate(kolom, start=1):
                v = d.get(k, "")
                c = ws.cell(i, j, v)
                if k in persen:
                    c.number_format = "0.0%"
        return baris0 + len(data)

    def lebar(ws, kolom: list[str], lebar_maks=45):
        for j, k in enumerate(kolom, start=1):
            ws.column_dimensions[get_column_letter(j)].width = min(lebar_maks, max(9, len(JUDUL_KOLOM.get(k, k)) * 0.9 + 2))

    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Monitoring input usaha"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for k, v in meta:
        ws.cell(r, 1, k).font = tebal
        ws.cell(r, 2, v)
        r += 1
    r += 1
    for judul, data, kolom in lembar_rekap[:1]:
        ws.cell(r, 1, judul).font = Font(bold=True, size=12)
        r = tabel(ws, r + 1, kolom, data, JUDUL_KOLOM, persen=("pct_terkirim", "pct_approved")) + 2
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16

    for judul, data, kolom in lembar_rekap[1:]:
        ws = wb.create_sheet(judul[:31])
        tabel(ws, 1, kolom, data, JUDUL_KOLOM, persen=("pct_terkirim", "pct_approved"))
        ws.freeze_panes = ws.cell(2, 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{max(1, len(data)) + 1}"
        lebar(ws, kolom)
        for j, k in enumerate(kolom, start=1):
            if k in ("kecamatan", "desa", "sls", "ppl"):
                ws.column_dimensions[get_column_letter(j)].width = 26

    for judul, data, kolom in (("Ganda", ganda, KOLOM_GANDA), ("Di luar sheet", luar, KOLOM_LUAR),
                               ("Detail", detail, KOLOM_DETAIL)):
        ws = wb.create_sheet(judul)
        tabel(ws, 1, kolom, data)
        ws.freeze_panes = ws.cell(2, 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{max(1, len(data)) + 1}"
        for j, k in enumerate(kolom, start=1):
            ws.column_dimensions[get_column_letter(j)].width = (
                40 if k in ("nama_usaha", "nama", "keterangan", "pesan_data", "dicatat_di") else
                18 if k in ("idsubsls", "subsls_dokumen", "kunci", "status_server", "status_audit") else 14)
        if "kategori" in kolom:
            jk = kolom.index("kategori") + 1
            for i in range(2, len(data) + 2):
                c = ws.cell(i, jk)
                warna = WARNA_KATEGORI.get(c.value)
                if warna:
                    c.fill = PatternFill("solid", fgColor=warna)
    lokasi.siapkan(path)
    wb.save(path)
    return path


def tulis_html(path: Path, data: dict) -> Path:
    templat = (Path(__file__).resolve().parent / "dasbor.html").read_text(encoding="utf-8")
    muatan = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    isi = templat.replace("/*__DATA__*/null", muatan, 1).replace("__JUDUL__", html.escape(data["meta"]["judul"]))
    lokasi.siapkan(path)
    path.write_text(isi, encoding="utf-8")
    return path


def data_html(detail: list[dict], ganda: list[dict], luar: list[dict], meta: dict) -> dict:
    """Muatan dasbor, dipadatkan (±13 MB -> ±3 MB utk 14 rb usaha): kolom sekali saja, batch &
    kategori jadi indeks, nama wilayah satu tabel per kode (JS menurunkan kode kec/desa/SLS dari
    idsubsls), URL dokumen dibangun di JS dari ID."""
    batch = list(dict.fromkeys(d["batch"] for d in detail))
    kategori = list(KATEGORI)
    wil: dict = {}
    for d in detail:
        for k, n in (("kec_kode", "kecamatan"), ("desa_kode", "desa"), ("sls_kode", "sls")):
            if d[k] and d[n]:
                wil.setdefault(d[k], d[n])
    kolom = ["batch", "baris", "nama_usaha", "ppl", "idsubsls", "kategori", "keterangan", "jumlah_dokumen",
             "id_dokumen", "status_server", "sumber_status", "subsls_dokumen", "di_wilayah", "pesan_data"]
    padat = []
    for d in detail:
        r = [d.get(k, "") for k in kolom]
        r[0], r[5] = batch.index(d["batch"]), kategori.index(d["kategori"])
        padat.append(r)
    pendek = [{k: v for k, v in g.items() if k not in ("keterangan", "dokumen_url")} for g in ganda]
    return {
        "meta": {**meta, "url_dasar": url_entry("@").replace("@/entry", ""), "batch": batch},
        "kategori": [{"kode": k, "label": v} for k, v in KATEGORI.items()],
        "kolom": kolom, "detail": padat, "wil": wil,
        "ganda": pendek, "jenis_ganda": JENIS_GANDA, "luar": luar,
    }


# ------------------------------------------------------------------------------------------
# CLI
# ------------------------------------------------------------------------------------------
def label_sheet(path: Path) -> str:
    return path.stem


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Monitoring hasil input usaha: semua status + ganda, rekap per daerah")
    ap.add_argument("--sumber", action="append", default=[],
                    help=f"sheet input usaha (boleh berulang; bawaan semua {POLA_SUMBER}, tanpa .bak-*)")
    ap.add_argument("--audit", action="append", default=[],
                    help="audit input (berkas/folder, boleh berulang; bawaan semua audit/**/audit_log_gabungan*.csv)")
    ap.add_argument("--snapshot", action="append", default=[],
                    help="CSV unduhan monitoring_console.js (atau list_api_*.json); bawaan snapshot_fasih_sm*.csv "
                         "terbaru di bahan/ atau folder Downloads")
    ap.add_argument("--tanpa-snapshot", action="store_true", help="abaikan snapshot server (murni dari audit)")
    ap.add_argument("--peta", default="", help="GeoJSON peta SLS utk nama wilayah (bawaan config PETA_SLS_PATH)")
    ap.add_argument("--keluaran", default="", help=f"folder keluaran (bawaan {HASIL})")
    ap.add_argument("--tanpa-html", action="store_true", help="hanya Excel")
    mg.opsi_format(ap)
    args = ap.parse_args(argv)
    lokasi.cek_struktur_lama()
    format_sumber = args.format

    sumber = [Path(s) for s in args.sumber] or [p for p in lokasi.cari(POLA_SUMBER)
                                                if ".bak-" not in p.name and not p.name.startswith("~$")]
    if not sumber:
        ap.error(f"tidak ada sheet: isi --sumber (bawaan {POLA_SUMBER})")
    hilang = [str(s) for s in sumber if not s.exists()]
    if hilang:
        ap.error(f"sheet tidak ada: {hilang}")

    paths_audit = ([mg.pakai_audit(a) if not Path(a).is_file() else Path(a) for a in args.audit]
                   if args.audit else audit_kandidat())
    audits = baca_audit_semua([p for p in paths_audit if p.exists()])
    print(f"Audit terbaca: {len(audits)} berkas")
    audit_global, dihapus = dokumen_global(audits)
    dihapus |= baca_ganda_dihapus()
    approve = baca_approve()
    print(f"Audit approve: {sum(1 for s in approve.values() if s in STATUS_APPROVE_OK)} dokumen APPROVED tercatat")

    snapshot, snap_info = None, None
    if not args.tanpa_snapshot:
        paths_snap = [Path(s) for s in args.snapshot] or cari_snapshot()
        if paths_snap:
            snapshot, snap_info = baca_snapshot(paths_snap)
            for p, n in snap_info["berkas"]:
                umur = (dt.datetime.now() - dt.datetime.fromtimestamp(p.stat().st_mtime))
                print(f"Snapshot server: {p} ({n} dokumen, umur berkas {umur.days} hari "
                      f"{umur.seconds // 3600} jam)")
            if not snap_info["total_server"]:
                print("ℹ️ Snapshot tanpa jumlah dokumen server (bukan CSV monitoring_console.js) — dianggap SEBAGIAN: "
                      "dokumen yang tidak ada di dalamnya memakai status audit.")
            elif len(snapshot) < snap_info["total_server"]:
                print(f"⚠️ Snapshot TIDAK LENGKAP: {len(snapshot)} dari {snap_info['total_server']} dokumen — dokumen "
                      "yang tidak terbaca memakai status audit (tidak dianggap hilang).")
        else:
            print("ℹ️ Tidak ada snapshot fasih-sm (snapshot_fasih_sm*.csv) — status dari audit saja. "
                  "Lihat monitoring/README.md cara mengambilnya.")

    sheets: list[Sheet] = []
    for p in sumber:
        rows, periksa = mg.muat_sumber(str(p), format_sumber, mode_satu_subsls=True, izinkan_tanpa_koordinat=True)
        milik = audit_milik({r.kunci for r in rows}, audits)
        print(f"{p.name}: {len(rows)} baris | audit: " + (", ".join(str(a.path.relative_to(lokasi.AKAR))
                                                           if a.path.is_relative_to(lokasi.AKAR) else str(a.path)
                                                           for a in milik) or "tidak ada yang mengenal sheet ini"))
        sheets.append(Sheet(label_sheet(p), rows, periksa, milik))

    nama_wilayah = baca_nama_wilayah(args.peta)
    if not nama_wilayah:
        print("ℹ️ Peta SLS tidak ada — nama kecamatan/desa dari kolom sheet, nama SLS kosong.")
    detail, ganda, luar = bangun(sheets, audit_global, dihapus, approve, snapshot, snap_info, nama_wilayah)

    kolom_angka = KOLOM_REKAP_ANGKA if snapshot else [k for k in KOLOM_REKAP_ANGKA if k != "TIDAK_DI_SERVER"]
    per_batch = rekap(detail, ("batch",))
    total = rekap([{**d, "batch": "TOTAL"} for d in detail], ("batch",))
    lembar = [
        ("Per batch", per_batch + total, ["batch", *kolom_angka]),
        ("Per kecamatan", rekap(detail, ("kec_kode",), ("kecamatan",)), ["kec_kode", "kecamatan", *kolom_angka]),
        ("Per desa", rekap(detail, ("desa_kode",), ("kecamatan", "desa")),
         ["desa_kode", "kecamatan", "desa", *kolom_angka]),
        ("Per SLS", rekap(detail, ("sls_kode",), ("kecamatan", "desa", "sls")),
         ["sls_kode", "kecamatan", "desa", "sls", *kolom_angka]),
        ("Per PPL", rekap(detail, ("ppl",)), ["ppl", *kolom_angka]),
    ]
    sekarang = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    snap_teks = (f"{', '.join(p.name for p, _ in snap_info['berkas'])} (diambil {waktu_lokal(snap_info['diambil']) or '?'})"
                 if snap_info else "tidak dipakai (status dari audit)")
    meta = [("Dibuat", sekarang), ("Sheet", ", ".join(s.label for s in sheets)),
            ("Snapshot server", snap_teks), ("Total usaha", str(len(detail))),
            ("Usaha ganda", str(sum(1 for d in detail if d["jumlah_dokumen"] > 1))),
            ("Grup ganda/nama sama", str(len({g['grup'] for g in ganda}))),
            ("Dokumen server di luar sheet", str(len(luar)) if snapshot else "-")]

    folder = Path(args.keluaran) if args.keluaran else HASIL
    xlsx = tulis_excel(folder / "monitoring.xlsx", meta, lembar, ganda, luar, detail)
    print(f"\nExcel: {xlsx}")
    if not args.tanpa_html:
        pagina = tulis_html(folder / "monitoring.html", data_html(detail, ganda, luar, {
            "judul": "Monitoring input usaha", "dibuat": sekarang, "sheet": [s.label for s in sheets],
            "snapshot": snap_teks, "ada_snapshot": bool(snapshot)}))
        print(f"HTML : {pagina}")

    print("\nRINGKASAN")
    lebar = max(len(v) for v in KATEGORI.values())
    t = total[0]
    for k, v in KATEGORI.items():
        if t[k]:
            print(f"  {v:<{lebar}} {t[k]:>7}")
    print(f"  {'TOTAL usaha':<{lebar}} {t['total']:>7}   | terkirim {t['pct_terkirim']:.1%}, approved {t['pct_approved']:.1%}"
          " (dari usaha yang lolos pemeriksaan data)")
    jenis = Counter(g["jenis"] for g in {g["grup"]: g for g in ganda}.values())
    if jenis:
        print("\nGANDA / NAMA SAMA (grup)")
        for k, n in jenis.most_common():
            print(f"  {k:<30} {n:>5}  {JENIS_GANDA[k]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
