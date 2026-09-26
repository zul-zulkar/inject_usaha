#!/usr/bin/env python3
"""
pindah_struktur.py — pindahkan berkas PC ini dari struktur folder LAMA (sebelum 2026-09-25,
semua berkas di akar proyek) ke struktur baru. Sekali per PC, sesudah kode baru dipasang.

TIDAK PERNAH menghapus atau menimpa apa pun: berkas dipindah (rename); kalau tujuannya
sudah ada, berkas itu DIBIARKAN di tempatnya & dilaporkan. Tanpa --jalankan hanya
menampilkan rencana.

    python antar_pc/pindah_struktur.py              # lihat rencana
    python antar_pc/pindah_struktur.py --jalankan   # pindahkan

Ke mana:
  audit_log_gabungan.csv (+ .bak-*, salinan _<label>)  -> audit/batch21/   (--audit-lama-ke)
      audit bawaan LAMA = batch 21 September di Buleleng. Batch 22 sudah di audit/ (--audit audit/).
  audit_approve_pml.csv, ganda_dihapus*.csv             -> audit/
  audit_pc/, bahan_pc/                                  -> audit/pc/
  laporan, list_api_*.json, log & screenshot input      -> input_usaha/hasil/
  *_console.siap.js, target_*.csv, sesi & log alat      -> <alat>/hasil/
  daftar*.txt, list*.txt, *tahap2*.xlsx                  -> bahan/
  kode lama (input_gabungan/, input_tahap2/, gabung_audit/, bungkus_pc/, input_fasihweb/,
  ganti_moda/ dst. di akar, config.py/main.py lama)     -> arsip/kode_lama/
  data lama lain (Agenda*.xlsx, LKpenyalinan, export/, audit_log.csv, *.xlsx/*.csv lain) -> arsip/
Yang tidak dikenali dibiarkan & dicetak.
"""

from __future__ import annotations

import argparse
import fnmatch
import os
import re
import shutil
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

AKAR = Path(__file__).resolve().parents[1]

# Dibiarkan di akar (bagian struktur baru / milik pengguna).
TETAP = {
    ".git", ".gitignore", ".claude", ".claude.zip", ".venv", "venv", "README.md", "LICENSE", "CLAUDE.md",
    "requirements.txt", "bahan", "audit", "arsip", "inti", "input_usaha", "approve_pml", "fasih_sm",
    "reset_mitra", "koordinat", "antar_pc", "templates", "docs", "tests", "nul",
}
KODE_LAMA = {
    "input_gabungan", "input_tahap2", "gabung_audit", "bungkus_pc", "input_fasihweb", "ganti_moda",
    "buka_wilayah", "tandai_selesai", "pindah_wilayah", "hapus_ganda", "otomatisasi_se2026", "__pycache__",
    "config.py", "data_loader.py", "fasih_web.py", "fill_blok2.py", "main.py", "scrape_source.py",
}


def aturan(audit_lama_ke: str) -> list[tuple[str, str]]:
    """(pola nama di akar, folder tujuan relatif akar). Urutan = prioritas; pola pertama menang."""
    return [
        ("audit_log_gabungan*.csv*", audit_lama_ke),
        ("audit_approve_pml.csv", "audit"),
        ("ganda_dihapus*.csv", "audit"),
        ("audit_pc", "audit/pc"),
        ("bahan_pc", "audit/pc"),
        # input_usaha
        ("cek_gabungan.csv", "input_usaha/hasil"), ("dokumen_tanpa_url.csv", "input_usaha/hasil"),
        ("sinkron_list.csv", "input_usaha/hasil"), ("rangkum_audit.csv", "input_usaha/hasil"),
        ("bersihkan_error.csv", "input_usaha/hasil"), ("kontrol_kualitas*", "input_usaha/hasil"),
        ("list_api_*.json", "input_usaha/hasil"), ("log_otomatis*.txt", "input_usaha/hasil"),
        (".status_otomatis*.json", "input_usaha/hasil"), ("log_screenshots", "input_usaha/hasil"),
        (".proses_*.lock", "input_usaha/hasil"),
        # antar_pc
        ("laporan_gabung*.csv", "antar_pc/hasil"), ("daftar_ganda.csv", "antar_pc/hasil"),
        ("split_usaha_pc_*.zip", "antar_pc/hasil"),
        # approve_pml
        (".sesi_fasih_web_*.json", "approve_pml/hasil"), (".profil_fasih_web_*", "approve_pml/hasil"),
        ("log_approve", "approve_pml/hasil"),
        # fasih_sm & reset_mitra
        ("ubah_moda_console.siap.js", "fasih_sm/ganti_moda/hasil"), ("target_ubah_moda.csv", "fasih_sm/ganti_moda/hasil"),
        ("audit_ubah_moda.csv", "fasih_sm/ganti_moda/hasil"), (".profil_fasih_sm", "fasih_sm/ganti_moda/hasil"),
        ("log_fasih_sm", "fasih_sm/ganti_moda/hasil"),
        ("buka_wilayah_console.siap.js", "fasih_sm/buka_wilayah/hasil"),
        ("tandai_selesai_console.siap.js", "fasih_sm/tandai_selesai/hasil"),
        ("pindah_wilayah_console.siap.js", "fasih_sm/pindah_wilayah/hasil"),
        ("hapus_ganda_console.siap.js", "fasih_sm/hapus_ganda/hasil"),
        ("reset_mitra_console.siap.js", "reset_mitra/hasil"), ("target_reset_mitra.csv", "reset_mitra/hasil"),
        # bahan milik pengguna
        ("daftar*.txt", "bahan"), ("list*.txt", "bahan"), ("*tahap2*.xlsx", "bahan"),
        # data lama
        ("export", "arsip"), ("*.xlsx", "arsip"), ("*.csv", "arsip"), ("*.siap.js", "arsip"),
    ]


def nama_tujuan(nama: str) -> str:
    """Nama baru di tujuan (hanya beberapa berkas yang namanya ikut dirapikan)."""
    if nama == "cek_gabungan.csv":
        return "cek_input.csv"
    m = re.fullmatch(r"(\.status_otomatis|log_otomatis)_tahap2(.*)", nama)
    return m.group(1) + m.group(2) if m else nama


def rencana(akar: Path, audit_lama_ke: str) -> tuple[list[tuple[Path, Path]], list[Path]]:
    """-> ([(asal, tujuan)], [berkas tak dikenal yang dibiarkan])."""
    pindah, tak_dikenal = [], []
    for p in sorted(akar.iterdir()):
        nama = p.name
        if nama in TETAP or nama.startswith("run_"):
            continue
        if nama in KODE_LAMA:
            pindah.append((p, akar / "arsip" / "kode_lama" / nama))
            continue
        tujuan = next((d for pola, d in aturan(audit_lama_ke) if fnmatch.fnmatch(nama, pola)), None)
        if tujuan is None:
            tak_dikenal.append(p)
            continue
        if nama in ("audit_pc", "bahan_pc"):
            # isi folder dipindah satu per satu ke audit/pc/ (folder tujuan boleh sudah ada)
            for anak in sorted(p.iterdir()):
                sub = "bahan_pc" if nama == "bahan_pc" else ""
                pindah.append((anak, akar / tujuan / sub / anak.name))
            continue
        pindah.append((p, akar / tujuan / nama_tujuan(nama)))
    # bahan/input_usaha.xlsx format LAMA (92 kolom) bentrok dgn nama baku format input usaha
    # sekarang (format tahap 2, ±48 kolom) -> arsip, supaya contoh perintah di README tidak
    # diam-diam membaca berkas yang salah.
    lama = akar / "bahan" / "input_usaha.xlsx"
    if lama.exists() and jumlah_kolom(lama) > 70:
        pindah.append((lama, akar / "arsip" / "input_usaha_format_lama.xlsx"))
    return pindah, tak_dikenal


def jumlah_kolom(xlsx: Path) -> int:
    """Jumlah judul kolom terisi di baris 1 tab pertama (0 kalau tidak terbaca)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True)
        ws = wb[wb.sheetnames[0]]
        judul = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        wb.close()
        return sum(1 for v in judul if v not in (None, ""))
    except Exception:  # noqa: BLE001
        return 0


def kunci_hidup(akar: Path) -> list[Path]:
    """.proses_*.lock yang prosesnya MASIH jalan (batch input sedang berjalan)."""
    hidup = []
    for f in list(akar.glob(".proses_*.lock")) + list((akar / "input_usaha" / "hasil").glob(".proses_*.lock")):
        try:
            pid = int(f.read_text(encoding="utf-8").split()[0])
        except (OSError, ValueError, IndexError):
            continue
        try:
            sys.path.insert(0, str(akar))
            from input_usaha.mesin import _pid_hidup
            if _pid_hidup(pid) and pid != os.getpid():
                hidup.append(f)
        except Exception:  # noqa: BLE001 — tanpa playwright dsb: anggap bisa hidup, lebih aman
            hidup.append(f)
    return hidup


def pindahkan(akar: Path, audit_lama_ke: str, jalankan: bool) -> tuple[list, list, list]:
    """Laksanakan (atau tampilkan kalau jalankan=False) rencana. -> (dipindah, konflik, tak dikenal).
    Tujuan yang sudah ada TIDAK ditimpa: asalnya dibiarkan & masuk daftar konflik."""
    pindah, tak_dikenal = rencana(akar, audit_lama_ke)
    dipindah, konflik = [], []
    for asal, tujuan in pindah:
        rel_a, rel_t = asal.relative_to(akar).as_posix(), tujuan.relative_to(akar).as_posix()
        if tujuan.exists():
            konflik.append((rel_a, rel_t))
            continue
        print(f"  {rel_a:55s} -> {rel_t}")
        dipindah.append((rel_a, rel_t))
        if jalankan:
            tujuan.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(asal), str(tujuan))
    if jalankan:
        for folder in ("audit_pc", "bahan_pc"):
            f = akar / folder
            if f.is_dir() and not any(f.iterdir()):
                f.rmdir()   # folder yang isinya sudah dipindah semua (kosong)
    return dipindah, konflik, tak_dikenal


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--jalankan", action="store_true", help="benar-benar memindahkan (tanpa ini: rencana saja)")
    ap.add_argument("--audit-lama-ke", default="audit/batch21",
                    help="folder tujuan audit LAMA dari akar proyek (bawaan audit/batch21)")
    args = ap.parse_args(argv)

    hidup = kunci_hidup(AKAR)
    if hidup:
        print("⛔ Ada batch input yang MASIH berjalan (" + ", ".join(f.name for f in hidup) + ").")
        print("   Tunggu/tutup prosesnya dulu, lalu jalankan skrip ini lagi.")
        return 2

    pindah, konflik, tak_dikenal = pindahkan(AKAR, args.audit_lama_ke, args.jalankan)
    if not pindah and not konflik:
        print("✅ Tidak ada yang perlu dipindah — PC ini sudah memakai struktur baru.")
    if konflik:
        print("\n⚠️ TIDAK dipindah — tujuannya sudah ada (periksa & gabungkan manual, tidak ada yang ditimpa):")
        for a, tj in konflik:
            print(f"  {a}  (tujuan {tj})")
    if tak_dikenal:
        print("\nDibiarkan (tidak dikenali): " + ", ".join(p.name for p in tak_dikenal))
    if not args.jalankan and pindah:
        print("\n(rencana saja — ulangi dgn --jalankan untuk memindahkan)")
    elif args.jalankan and pindah:
        print("\n✅ Selesai. Audit batch lama sekarang di "
              f"{args.audit_lama_ke}/ — pakai --audit {args.audit_lama_ke} utk melanjutkan batch itu.")
    return 1 if konflik else 0


if __name__ == "__main__":
    sys.exit(main())
