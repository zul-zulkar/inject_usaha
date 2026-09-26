#!/usr/bin/env python3
"""
approve_pml.py — Approve (oleh PML/Pengawas) dokumen SE2026 yang sudah dikirim PPL
lewat input_usaha/jalankan.py. Dipetakan langsung 2026-09-15 (akun PML pml.dua, PPL ppl.contoh).

YANG SUDAH DILIHAT DI fasih-web (akun Pengawas)
----------------------------------------------
- Membuka dokumen PML = URL entry yang sama dgn PPL (…/{id}/entry). Bar bawah kuesioner
  berisi tombol "Reject" (ikon x) & "Approve" (ikon checklist).
- Klik "Approve" -> dialog "Konfirmasi Approve" ("Apakah Anda yakin ingin melakukan approve
  dokumen ini?") dgn tombol "Batal" / "Approve" / "Close". Klik "Approve" DI DIALOG itulah
  yang mengeksekusi (irreversible).
- Status per dokumen: GET /api/assignment-general/api/assignment/web-entry/get-by-id-with-data?id=
  -> data.assignment_status_alias ("SUBMITTED BY Pencacah" / "APPROVED BY Pengawas") dan
  data.data (JSON string) berisi createdBy/updatedBy = email PPL pembuat.
- List PENDATAAN akun PML berisi ribuan dokumen (pml.dua: 4.870) -> API datatable dgn
  length 100 membalas 504. Karena itu target TIDAK diambil dari list penuh, melainkan dari
  audit_log_gabungan.csv (+ opsional pencarian list per subsls, length kecil).

PENGAMAN (gagal-tertutup)
------------------------
Satu dokumen hanya di-approve kalau, tepat sebelum diklik, API menyatakan
SUBMITTED BY Pencacah DAN createdBy/updatedBy = --akun-ppl; tombol Approve di bar & di
dialog masing-masing tepat satu. Sukses = status API berubah jadi APPROVED (bukan toast).
Dialog konfirmasi yang tidak muncul / tombol ambigu -> batch berhenti.

LANGKAH — satu PML, target dari audit_log_gabungan.csv
-----------------------------------------------------
1. Petakan / cek (READ-ONLY, tanpa klik): status semua target lewat API
       python approve_pml/approve_pml.py --akun-pml pml.dua@gmail.com --akun-ppl ppl.contoh@mail.com
2. Eksekusi 1 dokumen dulu, cek hasilnya:
       python approve_pml/approve_pml.py --akun-pml ... --akun-ppl ... --eksekusi --limit 1
3. Sisanya (dokumen yang sudah APPROVED otomatis dilewati):
       python approve_pml/approve_pml.py --akun-pml ... --akun-ppl ... --eksekusi

LANGKAH — MULTI PML, target dari file SQL Lab (--rencana, 2026-09-15)
--------------------------------------------------------------------
File .xlsx/.csv berkolom "Email PML", "Email PPL", "assignment_id" (= id dokumen), opsional
"PML", "PPL", "data1", "code_identity", "assignment_status_alias". Dokumen dikelompokkan per
Email PML (urut kemunculan di file); tiap kelompok: login PML itu -> proses -> simpan sesi ->
LOGOUT (UI "Keluar" + hapus cookie semua domain termasuk sso.bps.go.id) -> PML berikutnya.
Tiap dokumen tetap wajib createdBy/updatedBy = Email PPL BARIS ITU.
   0. Rencana saja, tanpa browser:     ... --rencana approve_pml/sqllab_....xlsx --cek
   1. Dry-run semua PML:               ... --rencana approve_pml/sqllab_....xlsx
   2. 1 dokumen PER PML dulu:          ... --rencana ... --eksekusi --limit 1
   3. Sisanya:                         ... --rencana ... --eksekusi
   Batasi PML: --akun-pml a@x.com --akun-pml b@y.com (boleh diulang / dipisah koma).
   --limit berlaku PER PML. Login PML gagal -> PML itu dilewati (ERROR_LOGIN_PML), lanjut PML
   berikutnya; STOP_* / 3 ERROR beruntun menghentikan SELURUH run (anomali UI/VPN).
   ⚠️ Dry-run 2026-09-15: 765/768 dokumen file SQL Lab = CAPI -> SKIP_TIDAK_ADA_AKSES (web-entry
   hanya membuka dokumen PAPI). Dokumen hasil input_usaha (PAPI) aman.

LANGKAH — MULTI PML, target dari ekspor tabel Data fasih-sm (--daftar bahan/submit.xlsx, 2026-09-15)
--------------------------------------------------------------------------------------------
File = salinan tabel list fasih-sm (Kode Identitas, Nama, ..., Status, Mode, Petugas Saat Ini).
Header-nya BERGESER, jadi kolom dibaca lewat JANGKAR sel "PAPI"/"CAPI": sel sebelumnya = Status,
sesudahnya = Petugas Saat Ini (= akun PML yang login). Hanya PAPI + "submitted by pencacah".
File tidak memuat id dokumen & email PPL: per PML, list PENDATAAN dicari per subsls (16 digit
awal kode) lalu codeIdentity dicocokkan PERSIS -> id; item harus bermode PAPI & petugas saat
ini = PML itu. createdBy/updatedBy TIDAK dicek (tidak ada PPL di file) tapi dicatat di kolom
akun_ppl audit. Kode yang terakhir APPROVED_TERVERIFIKASI/SUDAH_APPROVED di audit dilewati.
   0. Rencana saja, tanpa browser:     python approve_pml/approve_pml.py --daftar bahan/submit.xlsx --cek
   1. Dry-run semua PML:               python approve_pml/approve_pml.py --daftar bahan/submit.xlsx
   2. 1 dokumen PER PML dulu:          ... --daftar bahan/submit.xlsx --eksekusi --limit 1
   3. Sisanya:                         ... --daftar bahan/submit.xlsx --eksekusi
   --akun-pml / --limit / --abaikan-audit-approve sama dgn --rencana.

Login (sama dgn input_usaha): tiap PML = browser context BARU (cookie SSO kosong), login
otomatis dgn FIXED_PASSWORD (inti/config_lokal.py) — diulang maks 3x utk gangguan
transien, akun salah tidak diulang — lalu akun aktif WAJIB terbaca = PML itu. Selesai satu PML:
logout + tutup context. Satu PML saja / --login-manual: sesi disimpan di
approve_pml/hasil/.sesi_fasih_web_<akun>.json & dipakai ulang tanpa logout, supaya run berikutnya
tidak perlu login lagi. --login-manual = tunggu manusia login di jendela browser.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from inti import lokasi
from inti.config import (ASSIGNMENT_ID_GABUNGAN, FASIH_WEB_BASE, FASIH_WEB_LOGIN_URL, FIXED_PASSWORD, L,
                         SEL, SURVEY_ID)

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

AUDIT_GABUNGAN = lokasi.audit_dari_lingkungan()   # sama dgn input_usaha/mesin.py; --audit menimpa
AUDIT_APPROVE = lokasi.AUDIT_APPROVE              # audit/audit_approve_pml.csv
AUDIT_FIELDS = ["timestamp", "akun_pml", "akun_ppl", "id", "baris", "kunci", "nama", "sumber",
                "status_sebelum", "status", "pesan", "dokumen_url"]
HASIL = lokasi.HASIL_APPROVE                      # approve_pml/hasil/: screenshot & sesi login
API_DETAIL = "/api/assignment-general/api/assignment/web-entry/get-by-id-with-data?id="

ST_SIAP = "SIAP_APPROVE"
ST_SUDAH = "SUDAH_APPROVED"
ST_OK = "APPROVED_TERVERIFIKASI"
ST_TANPA_AKSES = "SKIP_TIDAK_ADA_AKSES"
LOGIN_PERCOBAAN = 3
LOGIN_JEDA_DTK = 30
# Kejanggalan yang pasti berulang di dokumen berikutnya -> hentikan batch.
STATUS_BERHENTI_SEGERA = {"STOP_DIALOG_TIDAK_MUNCUL", "STOP_TOMBOL_AMBIGU", "STOP_AKUN_BERUBAH",
                          "APPROVE_TIDAK_TERVERIFIKASI"}
# Nama kolom file SQL Lab (--rencana), dibandingkan setelah lower + spasi/_ diseragamkan.
KOLOM_RENCANA = {"email pml": "akun_pml", "email ppl": "akun_ppl", "assignment id": "id",
                 "pml": "nama_pml", "ppl": "nama_ppl", "data1": "nama", "code identity": "code_identity",
                 "assignment status alias": "status_file"}


def slug_akun(akun: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", akun.lower()).strip("_")


def id_dari_url(url: str) -> str:
    bagian = (url or "").rstrip("/").split("/")
    return bagian[-2] if len(bagian) >= 2 and bagian[-1] == "entry" else ""


def url_entry(doc_id: str, assignment_id: str) -> str:
    return f"{FASIH_WEB_BASE}/survey/{SURVEY_ID}/{assignment_id}/{doc_id}/entry"


# ----------------------------------------------------------------------
# Logika murni (tanpa browser) — diuji tests/test_approve_pml.py
# ----------------------------------------------------------------------
def dokumen_audit_ppl(audit: list[dict], akun_ppl: str) -> dict[str, dict]:
    """id dokumen -> baris audit TERAKHIR yang menyebut dokumen itu, utk akun PPL ini.
    Dokumen yang dicatat DOKUMEN_DIHAPUS digugurkan."""
    akun_ppl = akun_ppl.lower()
    hasil: dict[str, dict] = {}
    for b in audit:
        if (b.get("akun_login") or "").lower() != akun_ppl:
            continue
        i = id_dari_url(b.get("dokumen_url"))
        if not i:
            continue
        if b.get("status") == "DOKUMEN_DIHAPUS":
            hasil.pop(i, None)
            continue
        hasil[i] = b
    return hasil


def subsls_audit_ppl(audit: list[dict], akun_ppl: str) -> list[str]:
    """Subsls tempat dokumen PPL ini dibuat (kolom idsubsls_input)."""
    akun_ppl = akun_ppl.lower()
    return sorted({(b.get("idsubsls_input") or "").strip() for b in audit
                   if (b.get("akun_login") or "").lower() == akun_ppl} - {""})


def petugas_dokumen(detail: dict) -> set[str]:
    """{createdBy, updatedBy} dari data.data (JSON string) respons get-by-id-with-data."""
    isi = detail.get("data")
    if isinstance(isi, str):
        try:
            isi = json.loads(isi)
        except ValueError:
            isi = {}
    if not isinstance(isi, dict):
        return set()
    return {str(isi.get(k) or "").strip().lower() for k in ("createdBy", "updatedBy")} - {""}


def nilai_dokumen(detail: dict | None, akun_ppl: str | None) -> tuple[str, str]:
    """-> (kategori, pesan) satu dokumen menurut API detail. Hanya ST_SIAP yang boleh di-approve.
    akun_ppl None = createdBy/updatedBy TIDAK dicek (khusus --daftar); "" tetap gagal-tertutup."""
    if not detail:
        return "SKIP_DETAIL_TIDAK_TERBACA", "respons get-by-id-with-data kosong"
    alias = str(detail.get("assignment_status_alias") or "")
    up = alias.upper()
    if "APPROVED" in up:
        return ST_SUDAH, alias
    if not up.startswith("SUBMITTED BY PENCACAH"):
        return "SKIP_STATUS_" + (re.sub(r"[^A-Z]+", "_", up).strip("_") or "KOSONG"), alias
    if akun_ppl is None:
        return ST_SIAP, alias
    petugas = petugas_dokumen(detail)
    if akun_ppl.lower() not in petugas:
        return "SKIP_BUKAN_PPL", f"createdBy/updatedBy = {sorted(petugas) or '-'}"
    return ST_SIAP, alias


def gabung_target(audit_docs: dict[str, dict], items_list: list[dict]) -> list[dict]:
    """Target = dokumen audit PPL (urut audit) + dokumen list tersaring yang tidak ada di audit."""
    target = [{"id": i, "baris": b.get("baris", ""), "kunci": b.get("kunci", ""),
               "nama": b.get("nama_usaha", ""), "sumber": "audit"} for i, b in audit_docs.items()]
    sudah = set(audit_docs)
    for it in items_list:
        i = it.get("id")
        if i and i not in sudah:
            sudah.add(i)
            target.append({"id": i, "baris": "", "kunci": "", "nama": it.get("data1") or "", "sumber": "list"})
    return target


def _norm_kolom(k) -> str:
    return re.sub(r"[\s_]+", " ", str(k or "").strip().lower())


_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def target_dari_rencana(rows: list[dict]) -> tuple[list[dict], list[str]]:
    """Baris file SQL Lab -> (target, masalah). `baris` = nomor baris di file (header = 1).
    Baris tanpa id / email PML / email PPL yang valid dilewati & dilaporkan. Id yang muncul
    lebih dari sekali dgn pasangan PML/PPL BERBEDA digugurkan seluruhnya (tidak ditebak mana
    yang benar); duplikat identik cukup diproses sekali."""
    per_id: dict[str, dict | None] = {}
    urut: list[dict] = []
    masalah: list[str] = []
    for n, r in enumerate(rows, start=2):
        if all(v in (None, "") for v in r.values()):
            continue
        b = {KOLOM_RENCANA[_norm_kolom(k)]: ("" if v is None else str(v).strip())
             for k, v in r.items() if _norm_kolom(k) in KOLOM_RENCANA}
        pml, ppl, i = b.get("akun_pml", "").lower(), b.get("akun_ppl", "").lower(), b.get("id", "")
        if not i or not _EMAIL.match(pml) or not _EMAIL.match(ppl):
            masalah.append(f"baris {n}: id/Email PML/Email PPL kosong atau tidak valid "
                           f"(id={i or '-'}, pml={pml or '-'}, ppl={ppl or '-'}) — dilewati")
            continue
        if i in per_id:
            lama = per_id[i]
            if lama is not None and (lama["akun_pml"], lama["akun_ppl"]) != (pml, ppl):
                masalah.append(f"id {i}: pasangan PML/PPL berbeda di baris {lama['baris']} & {n} — digugurkan")
                per_id[i] = None
            continue
        t = {"id": i, "baris": str(n), "kunci": "", "nama": b.get("nama") or b.get("code_identity", ""),
             "sumber": "rencana", "akun_pml": pml, "akun_ppl": ppl, "nama_pml": b.get("nama_pml", "")}
        per_id[i] = t
        urut.append(t)
    return [t for t in urut if per_id.get(t["id"]) is t], masalah


def kelompokkan_per_pml(target: list[dict], pilih: list[str] | None = None) -> list[tuple[str, list[dict]]]:
    """[(akun_pml, target...)] urut kemunculan pertama; `pilih` (kosong = semua) menyaring PML."""
    pilih_set = {p.lower() for p in (pilih or [])}
    hasil: dict[str, list[dict]] = {}
    for t in target:
        if pilih_set and t["akun_pml"] not in pilih_set:
            continue
        hasil.setdefault(t["akun_pml"], []).append(t)
    return list(hasil.items())


def id_sudah_approved(audit_approve: list[dict]) -> set[str]:
    """Id dokumen yang status TERAKHIR-nya di audit approve = APPROVED_TERVERIFIKASI
    (status API sudah terbukti APPROVED) -> tidak perlu dicek ulang saat run diulang."""
    akhir: dict[str, str] = {}
    for b in audit_approve:
        if b.get("id"):
            akhir[b["id"]] = b.get("status", "")
    return {i for i, s in akhir.items() if s == ST_OK}


_KODE_IDENTITAS = re.compile(r"^\d{16}\s*-")


def baku_kode(teks) -> str:
    """Kode identitas utk dicocokkan PERSIS (spasi diseragamkan, besar/kecil huruf diabaikan)."""
    return " ".join(str(teks or "").split()).casefold()


def baris_daftar_sm(sel: list) -> dict | None:
    """Satu baris salinan tabel Data fasih-sm -> {kode, nama, status, mode, pml}, None kalau bukan
    baris data. Header submit.xlsx 2026-09-15 BERGESER (judul "Kode Identitas" hilang), jadi kolom
    tidak dibaca lewat judul: kode = sel pertama berawalan 16 digit; JANGKAR = sel "PAPI"/"CAPI"
    (kolom Mode) -> sel sebelumnya Status, sesudahnya Petugas Saat Ini. Kolom Email USAHA juga bisa
    berisi email (2 baris di submit.xlsx), jadi "email pertama" BUKAN petugas."""
    teks = ["" if v is None else str(v).strip() for v in sel]
    i_kode = next((i for i, v in enumerate(teks) if _KODE_IDENTITAS.match(v)), None)
    if i_kode is None:
        return None
    b = {"kode": " ".join(teks[i_kode].split()), "nama": teks[i_kode + 1] if i_kode + 1 < len(teks) else "",
         "status": "", "mode": "", "pml": ""}
    jangkar = [i for i, v in enumerate(teks) if i > i_kode and v.upper() in ("PAPI", "CAPI")]
    if len(jangkar) == 1:
        j = jangkar[0]
        b.update(status=teks[j - 1], mode=teks[j].upper(), pml=(teks[j + 1] if j + 1 < len(teks) else "").lower())
    return b


def target_dari_daftar(rows: list[list]) -> tuple[list[dict], list[str]]:
    """Baris salinan tabel fasih-sm (termasuk header) -> (target, masalah). `baris` = nomor baris
    file. Target hanya Mode PAPI + Status "submitted by pencacah" + Petugas Saat Ini berupa email;
    sisanya dilaporkan. Kode ganda dgn Petugas beda digugurkan seluruhnya (tidak ditebak).
    `id` diisi belakangan dari list PENDATAAN (cocokkan_list); `kunci` = kode (jejak audit)."""
    per_kode: dict[str, dict | None] = {}
    urut: list[dict] = []
    masalah: list[str] = []
    for n, sel in enumerate(rows, start=1):
        b = baris_daftar_sm(sel)
        if b is None:
            continue
        sebab = ""
        if not b["mode"]:
            sebab = "kolom Mode PAPI/CAPI tidak dikenali"
        elif b["mode"] != "PAPI":
            sebab = f"mode {b['mode']} (web-entry hanya membuka dokumen PAPI)"
        elif not b["status"].lower().startswith("submitted by pencacah"):
            sebab = f"status '{b['status']}'"
        elif not _EMAIL.match(b["pml"]):
            sebab = f"Petugas Saat Ini '{b['pml'] or '-'}' bukan email"
        if sebab:
            masalah.append(f"baris {n} ({b['kode']}): {sebab} — dilewati")
            continue
        k = baku_kode(b["kode"])
        if k in per_kode:
            lama = per_kode[k]
            if lama is not None and lama["akun_pml"] != b["pml"]:
                masalah.append(f"kode {b['kode']}: Petugas Saat Ini beda di baris {lama['baris']} & {n} — digugurkan")
                per_kode[k] = None
            continue
        t = {"id": "", "kode": b["kode"], "baris": str(n), "kunci": b["kode"], "nama": b["nama"],
             "sumber": "daftar", "akun_pml": b["pml"], "akun_ppl": None}
        per_kode[k] = t
        urut.append(t)
    return [t for t in urut if per_kode.get(baku_kode(t["kode"])) is t], masalah


def cocokkan_list(target: list[dict], items: list[dict], akun_pml: str) -> list[dict]:
    """Salinan target --daftar dgn `id` dari item list PENDATAAN (API datatable) lewat codeIdentity
    PERSIS. Yang tidak bisa dipastikan diberi `status`/`pesan` & tidak dibuka dokumennya:
    SKIP_KODE_TIDAK_DI_LIST / SKIP_KODE_GANDA / SKIP_BUKAN_PML_SAAT_INI / SKIP_MODE_BUKAN_PAPI.
    Item yang sudah APPROVED tetap diteruskan: statusnya diputuskan API detail (SUDAH_APPROVED)."""
    per_kode: dict[str, dict[str, dict]] = {}
    for it in items:
        per_kode.setdefault(baku_kode(it.get("codeIdentity")), {})[str(it.get("id") or "")] = it
    hasil = []
    for t in target:
        t = dict(t)
        cocok = list(per_kode.get(baku_kode(t["kode"]), {}).values())
        if not cocok:
            t.update(status="SKIP_KODE_TIDAK_DI_LIST", pesan="kode tidak ada di list PENDATAAN PML (dicari per subsls)")
        elif len(cocok) > 1:
            t.update(status="SKIP_KODE_GANDA",
                     pesan=f"{len(cocok)} dokumen berkode sama: {[str(it.get('id'))[:8] for it in cocok]}")
        else:
            it = cocok[0]
            t.update(id=str(it.get("id") or ""), nama=t["nama"] or str(it.get("data1") or ""))
            pemegang = str(it.get("currentUserUsername") or "").strip().lower()
            mode = [str(m).upper() for m in (it.get("mode") or [])]
            if "APPROVED" in str(it.get("assignmentStatusAlias") or "").upper():
                pass
            elif pemegang and pemegang != akun_pml:
                t.update(status="SKIP_BUKAN_PML_SAAT_INI", pesan=f"petugas saat ini di list = {pemegang}")
            elif "PAPI" not in mode:
                t.update(status="SKIP_MODE_BUKAN_PAPI", pesan=f"mode di list = {mode or '-'}")
        hasil.append(t)
    return hasil


def kode_sudah_approved(audit_approve: list[dict]) -> set[str]:
    """Kode identitas (baku) sumber --daftar yang status TERAKHIR-nya di audit approve sudah
    APPROVED_TERVERIFIKASI / SUDAH_APPROVED -> dilewati tanpa login saat run diulang."""
    akhir: dict[str, str] = {}
    for b in audit_approve:
        if b.get("sumber") == "daftar" and b.get("kunci"):
            akhir[baku_kode(b["kunci"])] = b.get("status", "")
    return {k for k, s in akhir.items() if s in (ST_OK, ST_SUDAH)}


def kode_halaman_error(teks: str) -> int:
    """Kode HTTP halaman galat fasih-web ("Terjadi Kesalahan (504)" / "Status Code: 504"), 0 kalau bukan."""
    m = re.search(r"Terjadi Kesalahan\s*\((\d{3})\)|Status Code:\s*(\d{3})", teks or "", re.I)
    return int(m.group(1) or m.group(2)) if m else 0


def baca_daftar(path: Path) -> list[list]:
    """Salinan tabel Data fasih-sm (.xlsx sheet pertama / .csv) -> baris mentah (header ikut, supaya
    nomor baris akurat). ValueError kalau tidak ada satu pun baris berkode identitas."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    if not any(baris_daftar_sm(r) for r in rows):
        raise ValueError(f"{path}: tidak ada baris berkode identitas (16 digit + ' - ')")
    return rows


def baca_rencana(path: Path) -> list[dict]:
    """File SQL Lab (.xlsx sheet pertama / .csv) -> list dict per baris. Kolom wajib dicek."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows(values_only=True)
            header = [str(h or "") for h in next(it, [])]
            rows = [dict(zip(header, r)) for r in it]  # baris kosong tetap ada -> nomor baris akurat
        finally:
            wb.close()
    else:
        with path.open(newline="", encoding="utf-8-sig") as f:
            rd = csv.DictReader(f)
            header = list(rd.fieldnames or [])
            rows = list(rd)
    ada = {KOLOM_RENCANA.get(_norm_kolom(h)) for h in header}
    kurang = [k for k, v in (("Email PML", "akun_pml"), ("Email PPL", "akun_ppl"), ("assignment_id", "id"))
              if v not in ada]
    if kurang:
        raise ValueError(f"{path}: kolom wajib tidak ada: {kurang} (header: {header})")
    return rows


def baca_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def append_audit(row: dict):
    baru = not AUDIT_APPROVE.exists()
    with lokasi.siapkan(AUDIT_APPROVE).open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=AUDIT_FIELDS)
        if baru:
            w.writeheader()
        w.writerow({k: row.get(k, "") for k in AUDIT_FIELDS})


# ----------------------------------------------------------------------
# Browser
# ----------------------------------------------------------------------
def penolakan_akses(status: int, text: str) -> str:
    """Pesan kalau respons get-by-id-with-data = PENOLAKAN pasti (bukan gangguan transien), '' kalau bukan.
    Diagnosis 2026-09-15: dokumen rencana SQL Lab (DTSEN/UMK, akun pml.satu) -> 200
    {"success":false,"message":"Anda tidak memiliki akses ke dalam survey","errorCode":23}; akun pml.delapan
    -> 403 kosong utk dokumen yang sama & utk id palsu, padahal dokumen PAPI-nya sendiri 200
    ("mode":["PAPI"]). Web-entry hanya melayani dokumen yang boleh dibuka akun itu (dugaan: mode PAPI)."""
    if status in (401, 403, 404):
        return f"HTTP {status} {text[:120]}".strip()
    if status == 200:
        try:
            j = json.loads(text)
        except ValueError:
            return ""
        if isinstance(j, dict) and j.get("success") is False and (
                j.get("errorCode") == 23 or "akses" in str(j.get("message") or "").lower()):
            return f"{j.get('message')} (errorCode {j.get('errorCode')})"
    return ""


def api_get(sess, path: str, info: dict | None = None) -> dict | None:
    """GET JSON dari halaman (cookie sesi). Endpoint get-by-id-with-data cukup dgn cookie
    (terlihat dari request aplikasi sendiri: tanpa header Authorization).
    `info` (opsional) diisi status HTTP & potongan teks respons terakhir."""
    try:
        hasil = sess.page.evaluate("""async (url) => {
          try {
            const r = await fetch(url, {credentials: 'include', headers: {'content-type': 'application/json'}});
            return {status: r.status, text: await r.text()};
          } catch (e) { return {status: 0, text: String(e)}; }
        }""", path)
    except Exception as e:
        hasil = {"status": 0, "text": f"evaluate gagal: {e}"}
    if info is not None:
        info.update(status=(hasil or {}).get("status", 0), text=str((hasil or {}).get("text") or "")[:300])
    if not hasil or hasil.get("status") != 200:
        return None
    try:
        return json.loads(hasil["text"])
    except ValueError:
        return None


def detail_dokumen(sess, doc_id: str, percobaan: int = 4, info: dict | None = None) -> dict | None:
    """Detail dokumen, diulang kalau gagal: run 2026-09-15 empat dokumen berturut-turut
    "kosong" tepat setelah approve dokumen sebelumnya (halaman sedang redirect -> evaluate
    gagal), padahal API-nya normal saat dicek ulang. PENOLAKAN akses tidak diulang;
    `info["ditolak"]` berisi pesannya."""
    info = info if info is not None else {}
    info["ditolak"] = ""
    for ke in range(1, percobaan + 1):
        j = api_get(sess, API_DETAIL + doc_id, info)
        data = (j or {}).get("data")
        if isinstance(data, dict):
            return data
        info["ditolak"] = penolakan_akses(info.get("status", 0), info.get("text", ""))
        if info["ditolak"]:
            return None
        if ke < percobaan:
            try:
                sess.page.wait_for_load_state("domcontentloaded", timeout=15_000)
            except Exception:
                pass
            sess.page.wait_for_timeout(2_000 * ke)
    return None


def email_aktif(sess) -> str:
    """Email akun yang sedang login (fetch check-user), '' kalau tidak terbaca.
    Catatan: aplikasi memanggil check-user dgn POST (201); GET bisa gagal -> andalkan sadapan."""
    for metode in ("POST", "GET"):
        try:
            hasil = sess.page.evaluate("""async (m) => {
              try {
                const r = await fetch('/api/survey/api/v1/users/check-user', {method: m, credentials: 'include',
                                      headers: {'content-type': 'application/json'}});
                if (!r.ok) return null;
                const j = await r.json();
                return (j && j.data) ? j.data : null;
              } catch (e) { return null; }
            }""", metode)
        except Exception:
            continue
        if isinstance(hasil, dict):
            e = str(hasil.get("email") or hasil.get("username") or "").strip().lower()
            if e:
                return e
    return ""


def akun_terbaca(sess, tunggu_ms: int) -> str:
    """Email akun aktif: sadapan respons check-user (FasihWebSession) — fetch langsung
    TIDAK cukup (run 2026-09-15: login manual sukses tapi fetch GET tetap kosong) —
    dgn fetch sbg cadangan."""
    batas = time.time() + tunggu_ms / 1000.0
    while True:
        if sess.akun_api.get("email"):
            return sess.akun_api["email"].lower()
        if time.time() >= batas:
            return email_aktif(sess)
        sess.page.wait_for_timeout(500)


def pastikan_login(sess, akun: str, manual: bool, file_sesi: Path | None):
    """Pakai sesi tersimpan kalau masih hidup & akunnya benar; kalau tidak, login.
    Cookie sesi (SSO & fasih-web tanpa tanggal kedaluwarsa -> TIDAK disimpan profil
    Chromium, terbukti 2026-09-15) disimpan di `file_sesi` & dipasang ulang saat mulai.
    file_sesi None = langsung login di context baru (multi PML, seperti input_usaha)."""
    akun = akun.lower()
    ctx = sess.page.context
    if file_sesi and file_sesi.exists():
        try:
            ctx.add_cookies(json.loads(file_sesi.read_text(encoding="utf-8")).get("cookies", []))
        except Exception as e:
            sess._log(f"⚠️ cookie sesi tersimpan tidak bisa dipasang: {e}")
        sess.akun_api = {}
        sess.page.goto(FASIH_WEB_BASE, wait_until="domcontentloaded", timeout=45_000)
        aktif = akun_terbaca(sess, 15_000)
        if aktif == akun:
            sess._log(f"✅ Sesi tersimpan masih hidup: {aktif}")
            return
        if aktif:
            sess._log(f"Sesi tersimpan milik akun LAIN ({aktif}) — diputus.")
            sess.logout()
    if not manual:
        sess.login(akun, FIXED_PASSWORD)
    else:
        sess.page.goto(FASIH_WEB_LOGIN_URL, wait_until="domcontentloaded", timeout=45_000)
        try:
            sess.page.get_by_text(L["sso_eksternal_btn"], exact=False).first.click(timeout=15_000)
        except Exception:
            pass
        print(f"\n>>> LOGIN MANUAL di jendela browser sbg {akun} (menunggu maks 15 menit) ...", flush=True)
        batas = time.time() + 15 * 60
        terakhir_lapor = 0.0
        while time.time() < batas:
            sess.page.wait_for_timeout(3_000)
            u = sess.page.url.lower()
            if time.time() - terakhir_lapor > 30:
                print(f"    (menunggu login — url sekarang {sess.page.url[:100]})", flush=True)
                terakhir_lapor = time.time()
            if "fasih-web.bps.go.id" in u and "/login" not in u and "/auth" not in u:
                if akun_terbaca(sess, 3_000):
                    break
        else:
            raise RuntimeError("Login manual tidak selesai dalam 15 menit")
    aktif = akun_terbaca(sess, 10_000)
    if aktif != akun:
        raise RuntimeError(f"AKUN SALAH/TIDAK TERVERIFIKASI: diminta '{akun}', aktif '{aktif or '-'}'")
    sess._log(f"✅ Login terverifikasi: {aktif}")
    simpan_sesi(sess, file_sesi)


def simpan_sesi(sess, file_sesi: Path | None):
    if not file_sesi:
        return
    try:
        file_sesi.write_text(json.dumps(sess.page.context.storage_state()), encoding="utf-8")
    except Exception as e:
        sess._log(f"⚠️ sesi tidak tersimpan: {e}")


def mulai_sesi_pml(browser, akun: str, manual: bool, file_sesi: Path | None):
    """Seperti input_usaha: context BARU (cookie SSO kosong -> tidak mungkin tembus sbg PML
    sebelumnya), login dgn FIXED_PASSWORD & verifikasi akun. Gangguan transien (goto timeout,
    "Execution context was destroyed") diulang maks LOGIN_PERCOBAAN x dgn jeda LOGIN_JEDA_DTK;
    akun salah & login manual TIDAK diulang. -> (ctx, sess); context sudah ditutup kalau melempar."""
    from inti.fasih_web import FasihWebSession
    for percobaan in range(1, LOGIN_PERCOBAAN + 1):
        ctx = browser.new_context(viewport={"width": 1440, "height": 900})
        try:
            sess = FasihWebSession(ctx.new_page())
            pastikan_login(sess, akun, manual, file_sesi)
            return ctx, sess
        except Exception as e:
            try:
                ctx.close()
            except Exception:
                pass
            if manual or "AKUN" in str(e) or percobaan == LOGIN_PERCOBAAN:
                raise
            print(f"⚠️ Login {akun} gagal (percobaan {percobaan}/{LOGIN_PERCOBAAN}): {str(e)[:150]} "
                  f"— ulangi {LOGIN_JEDA_DTK} dtk lagi.", flush=True)
            time.sleep(LOGIN_JEDA_DTK)
    raise RuntimeError(f"Login {akun} gagal")


def siapkan_target_daftar(sess, target: list[dict], akun_pml: str, assignment_id: str) -> list[dict]:
    """--daftar: isi `id` target dari list PENDATAAN PML, dicari per subsls (16 digit awal kode).
    Subsls yang list-nya gagal dibaca -> targetnya SKIP_LIST_TIDAK_TERBACA (diulang di run berikut)."""
    items: list[dict] = []
    gagal: dict[str, str] = {}
    for subsls in sorted({t["kode"][:16] for t in target}):
        try:
            items += cari_list(sess, assignment_id, subsls)
        except Exception as e:
            gagal[subsls] = str(e)[:150]
            print(f"⚠️ List subsls {subsls} gagal dibaca: {gagal[subsls]}", flush=True)
    hasil = cocokkan_list([t for t in target if t["kode"][:16] not in gagal], items, akun_pml)
    hasil += [{**t, "status": "SKIP_LIST_TIDAK_TERBACA", "pesan": gagal[t["kode"][:16]]}
              for t in target if t["kode"][:16] in gagal]
    siap = sum(1 for t in hasil if not t.get("status"))
    print(f"Kode dicocokkan ke list: {siap} siap, {dict(Counter(t['status'] for t in hasil if t.get('status')))}",
          flush=True)
    return hasil


def cari_list(sess, assignment_id: str, kata: str, per_halaman: int = 50) -> list[dict]:
    """Item list PENDATAAN yang cocok kata cari (mis. subsls), lewat API datatable yang
    dipanggil list sendiri (body + header x-* disadap), length KECIL: list PML penuh
    (length 100 tanpa saring) membalas 504. READ-ONLY."""
    tangkap: dict = {}

    def _on_request(req):
        if "datatable-all-user-survey-periode" in req.url and req.method == "POST" and "body" not in tangkap:
            tangkap.update(url=req.url, body=req.post_data, headers=req.headers)
    sess.page.on("request", _on_request)
    try:
        sess.goto_pendataan(assignment_id)
        batas = time.time() + 60
        while "body" not in tangkap and time.time() < batas:
            sess.page.wait_for_timeout(300)
    finally:
        sess.page.remove_listener("request", _on_request)
    if "body" not in tangkap:
        raise RuntimeError("Request datatable list PENDATAAN tidak tertangkap dalam 60 dtk")
    body = json.loads(tangkap["body"])
    hdr = {k: v for k, v in tangkap["headers"].items()
           if k.lower() in ("content-type", "accept") or k.lower().startswith("x-")}
    body["search"] = {"value": kata, "regex": False}
    per_id: dict = {}
    start = 0
    while True:
        body["start"], body["length"] = start, per_halaman
        hasil = None
        for percobaan in range(1, 4):
            hasil = sess.page.evaluate("""async ([url, body, hdr]) => {
                const r = await fetch(url, {method: 'POST', credentials: 'include', headers: hdr,
                                            body: JSON.stringify(body)});
                return {status: r.status, text: await r.text()};
            }""", [tangkap["url"], body, hdr])
            if hasil["status"] == 200:
                break
            sess._log(f"⚠️ API list '{kata}' start={start} status {hasil['status']} (percobaan {percobaan}/3)")
            sess.page.wait_for_timeout(5_000)
        if hasil["status"] != 200:
            raise RuntimeError(f"API list status {hasil['status']}: {hasil['text'][:200]}")
        data = json.loads(hasil["text"])
        halaman = data.get("searchData") or []
        total = int(data.get("totalHit") or 0)
        for it in halaman:
            per_id[it.get("id")] = it
        start += per_halaman
        if not halaman or start >= total:
            break
    sess._log(f"List tersaring '{kata}': {len(per_id)} dokumen (totalHit {total}).")
    return list(per_id.values())


def _dialog_konfirmasi(sess):
    return sess._visible(sess.page.locator('[role="dialog"], [role="alertdialog"]').filter(
        has_text=re.compile(r"yakin ingin melakukan approve", re.I)))


def teks_halaman(sess) -> str:
    try:
        return sess.page.evaluate("() => ((document.body && document.body.innerText) || '').slice(0, 3000)") or ""
    except Exception:
        return ""


def buka_dokumen(sess, url: str, doc_id: str, percobaan: int = 3) -> str:
    """Buka URL entry sampai form-engine mount -> '', atau pesan galat setelah `percobaan` kali.
    Run pml.delapan 2026-09-15: screenshot ERROR_FORM_TIDAK_MOUNT berisi halaman galat fasih-web
    "Terjadi Kesalahan (504)" = gangguan server sesaat. Dibuka ulang — aman, PML hanya MEMBACA
    dokumen (aturan keselamatan #2 soal reload menyangkut PPL yang sedang mengisi)."""
    form = sess.page.locator(SEL["form_root"])
    halaman_galat = sess.page.get_by_text(re.compile(r"Terjadi Kesalahan", re.I))
    pesan = ""
    for ke in range(1, percobaan + 1):
        try:
            sess.page.goto(url, wait_until="domcontentloaded", timeout=45_000)
            form.or_(halaman_galat).first.wait_for(state="attached", timeout=45_000)
            try:
                form.first.wait_for(state="attached", timeout=5_000)
                return ""
            except Exception:
                pesan = f"halaman galat fasih-web HTTP {kode_halaman_error(teks_halaman(sess)) or '?'}"
        except Exception as e:
            pesan = f"form-engine tidak mount dlm 45 dtk ({(str(e).splitlines() or [''])[0][:80]})"
        sess._log(f"⚠️ {doc_id[:8]}: {pesan} (buka {ke}/{percobaan})")
        if ke < percobaan:
            sess.page.wait_for_timeout(10_000 * ke)
    sess._shot(f"approve_form_tidak_mount_{doc_id[:8]}")
    return f"{pesan} — sudah {percobaan}x dibuka"


def approve_satu(sess, t: dict, assignment_id: str, akun_ppl: str | None, eksekusi: bool) -> dict:
    """Proses satu dokumen. Mengembalikan baris audit (tanpa timestamp/akun_pml).
    akun_ppl None (--daftar) = createdBy/updatedBy tidak dicek; petugas dokumen dicatat di `akun_ppl`."""
    res = {**t, "dokumen_url": url_entry(t["id"], assignment_id)}
    info: dict = {}
    detail = detail_dokumen(sess, t["id"], info=info)
    if detail is None and not info.get("ditolak"):
        # Halaman bisa tersangkut di keadaan yang membuat fetch gagal (mis. redirect pasca-approve
        # dokumen sebelumnya). Buka dokumennya dulu — hanya membuka, belum mengklik — lalu baca ulang,
        # supaya dokumen ini tetap dieksekusi di run yang sama, bukan dilewati.
        sess._log(f"⚠️ Detail {t['id'][:8]} tidak terbaca — buka dokumennya lalu baca ulang.")
        try:
            sess.page.goto(res["dokumen_url"], wait_until="domcontentloaded", timeout=45_000)
            sess.page.locator(SEL["form_root"]).first.wait_for(state="attached", timeout=45_000)
        except Exception:
            pass
        detail = detail_dokumen(sess, t["id"], info=info)
    kategori, pesan = nilai_dokumen(detail, akun_ppl)
    if detail is None:
        if info.get("ditolak"):
            kategori, pesan = ST_TANPA_AKSES, info["ditolak"]
        else:
            pesan += f" | HTTP {info.get('status')} {info.get('text', '')[:120]}"
    res["status_sebelum"] = (detail or {}).get("assignment_status_alias", "")
    if akun_ppl is None and detail:
        res["akun_ppl"] = ",".join(sorted(petugas_dokumen(detail)))
    if not res["nama"] and detail:
        res["nama"] = detail.get("data1") or ""
    if kategori != ST_SIAP:
        return {**res, "status": kategori, "pesan": pesan}

    galat = buka_dokumen(sess, res["dokumen_url"], t["id"])
    if galat:
        return {**res, "status": "ERROR_FORM_TIDAK_MOUNT", "pesan": galat}
    tombol = sess._visible(sess.page.get_by_role("button", name=re.compile(r"^\s*Approve\s*$")))
    try:
        tombol.first.wait_for(state="visible", timeout=30_000)
    except Exception:
        # Screenshot run 2026-09-15: bar bawah berisi "Revoke" = dokumen SUDAH di-approve (klik
        # sebelumnya terbaca terlambat). Status API yang memutuskan, bukan ada/tidaknya tombol.
        d2 = detail_dokumen(sess, t["id"], percobaan=2)
        alias = str((d2 or {}).get("assignment_status_alias") or "")
        if "APPROVED" in alias.upper():
            return {**res, "status": ST_SUDAH, "pesan": f"tombol Approve tidak ada; status API {alias}"}
        sess._shot(f"approve_tombol_tidak_ada_{t['id'][:8]}")
        return {**res, "status": "ERROR_TOMBOL_APPROVE_TIDAK_ADA",
                "pesan": f"tombol Approve di bar bawah tidak tampil dlm 30 dtk (status API '{alias or '-'}')"}
    if tombol.count() != 1:
        return {**res, "status": "STOP_TOMBOL_AMBIGU", "pesan": f"{tombol.count()} tombol Approve terlihat"}
    if not eksekusi:
        return {**res, "status": "DRY_RUN_SIAP_APPROVE", "pesan": "tombol Approve ada (tidak diklik)"}

    sess.klik_tahan(tombol.first, "Approve (bar bawah)")
    dialog = _dialog_konfirmasi(sess)
    try:
        dialog.first.wait_for(state="visible", timeout=15_000)
    except Exception:
        sess.debug_dialog("approve_tanpa_dialog")
        sess._shot(f"approve_dialog_tidak_muncul_{t['id'][:8]}")
        # Bisa jadi aplikasi langsung meng-approve tanpa dialog: status tetap diperiksa.
        d2 = detail_dokumen(sess, t["id"])
        if d2 and "APPROVED" in str(d2.get("assignment_status_alias") or "").upper():
            return {**res, "status": ST_OK, "pesan": "dialog tidak muncul, tapi status API sudah APPROVED"}
        return {**res, "status": "STOP_DIALOG_TIDAK_MUNCUL", "pesan": "dialog 'Konfirmasi Approve' tidak muncul"}
    konfirmasi = dialog.first.get_by_role("button", name=re.compile(r"^\s*Approve\s*$"))
    if konfirmasi.count() != 1:
        sess.debug_dialog("approve_konfirmasi_ambigu")
        return {**res, "status": "STOP_TOMBOL_AMBIGU", "pesan": f"{konfirmasi.count()} tombol Approve di dialog"}
    sess.page.wait_for_timeout(500)  # animasi buka dialog

    jejak: list[str] = []

    def _rekam(resp):
        try:
            if resp.request.method != "GET" and "/api/" in resp.url and "check-user" not in resp.url:
                isi = ""
                try:
                    isi = (resp.text() or "")[:200]
                except Exception:
                    pass
                jejak.append(f"{resp.request.method} {resp.status} {resp.url.split('/api/', 1)[-1][:90]} {isi}")
        except Exception:
            pass
    sess.page.on("response", _rekam)
    try:
        sess.klik_tahan(konfirmasi, "Approve (dialog Konfirmasi Approve)")
        # Sukses = status API berubah (toast/redirect bukan bukti: toast pernah muncul padahal server tetap DRAFT).
        status_akhir = ""
        for _ in range(30):
            sess.page.wait_for_timeout(2_000)
            d2 = detail_dokumen(sess, t["id"])
            status_akhir = str((d2 or {}).get("assignment_status_alias") or "")
            if "APPROVED" in status_akhir.upper():
                sess._log(f"✅ APPROVED: {res['nama']} ({t['id'][:8]}) -> {status_akhir}")
                return {**res, "status": ST_OK, "pesan": f"{status_akhir} | api={jejak[-3:]}"}
        sess._shot(f"approve_tidak_terverifikasi_{t['id'][:8]}")
        sess.debug_dialog("pasca_approve")
        return {**res, "status": "APPROVE_TIDAK_TERVERIFIKASI",
                "pesan": f"status API tetap '{status_akhir}' 60 dtk setelah konfirmasi | api={jejak[-5:]}"}
    finally:
        try:
            sess.page.remove_listener("response", _rekam)
        except Exception:
            pass


def proses_kelompok(sess, akun_pml: str, target: list[dict], args, hitung: Counter, file_sesi: Path) -> str:
    """Approve/dry-run semua target satu PML (sudah login). -> '' atau alasan SELURUH run dihentikan."""
    lokal: Counter = Counter()
    try:
        return _proses_kelompok(sess, akun_pml, target, args, lokal, file_sesi)
    finally:
        hitung.update(lokal)
        print(f"Ringkasan PML {akun_pml}: {dict(lokal)}", flush=True)


def _proses_kelompok(sess, akun_pml: str, target: list[dict], args, hitung: Counter, file_sesi: Path | None) -> str:
    diproses = 0
    error_beruntun = 0
    for n, t in enumerate(target, start=1):
        if args.limit and diproses >= args.limit:
            print(f"(--limit {args.limit} tercapai utk {akun_pml})")
            break
        if t.get("status"):
            # Sudah diputuskan sebelum dokumen dibuka (--daftar: kode tidak di list, ganda, dst.).
            res = {**t, "dokumen_url": url_entry(t["id"], args.assignment_id) if t["id"] else ""}
        else:
            res = approve_satu(sess, t, args.assignment_id, t["akun_ppl"], args.eksekusi)
        res.update(timestamp=time.strftime("%Y-%m-%d %H:%M:%S"), akun_pml=akun_pml)
        if res["status"] != ST_SUDAH and not res["status"].startswith("SKIP_"):
            diproses += 1
        # SUDAH_APPROVED dicatat khusus --daftar: kodenya jadi dilewati offline di run berikut.
        if res["status"] != ST_SUDAH or t.get("sumber") == "daftar":
            append_audit(res)
        hitung[res["status"]] += 1
        print(f"[{n}/{len(target)}] {res['status']:28s} {(t['id'] or '-')[:8]:8s} baris {t['baris'] or '-':>4} "
              f"{t.get('kode') or res['nama']}"
              + (f" | {str(res.get('pesan', ''))[:160]}" if res["status"] not in (ST_OK, ST_SUDAH) else ""),
              flush=True)
        if res["status"] in STATUS_BERHENTI_SEGERA:
            return res["status"]
        error_beruntun = error_beruntun + 1 if res["status"].startswith("ERROR_") else 0
        if args.maks_error_beruntun and error_beruntun >= args.maks_error_beruntun:
            return f"{error_beruntun} ERROR berturut-turut (VPN/sesi?)"
        if n % 10 == 0:
            if akun_terbaca(sess, 0) not in ("", akun_pml):
                return "STOP_AKUN_BERUBAH — akun aktif bukan PML yang diminta"
            simpan_sesi(sess, file_sesi)
    return ""


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--akun-pml", action="append", default=[],
                    help="mode audit: akun PML (wajib, satu). Mode --rencana: saring PML yang diproses "
                         "(boleh diulang / dipisah koma; kosong = semua PML di file)")
    ap.add_argument("--akun-ppl", help="mode audit: akun_login PPL di audit_log_gabungan.csv")
    ap.add_argument("--rencana", help="MULTI PML: file SQL Lab .xlsx/.csv (kolom Email PML, Email PPL, assignment_id)")
    ap.add_argument("--daftar", help="MULTI PML: salinan tabel Data fasih-sm .xlsx/.csv (mis. submit.xlsx: Kode "
                                     "Identitas, Status, Mode, Petugas Saat Ini); id dicari di list PENDATAAN PML")
    ap.add_argument("--cek", action="store_true", help="tampilkan rencana per PML saja, tanpa browser")
    ap.add_argument("--abaikan-audit-approve", action="store_true",
                    help="--rencana/--daftar: cek ulang juga dokumen yang di audit sudah APPROVED")
    ap.add_argument("--assignment-id", default=ASSIGNMENT_ID_GABUNGAN)
    ap.add_argument("--eksekusi", action="store_true", help="SUNGGUHAN klik Approve (irreversible)")
    ap.add_argument("--ya", action="store_true", help="lewati prompt ketik YA (izin sudah diberikan)")
    ap.add_argument("--limit", type=int, default=0, help="maks dokumen yang di-approve/di-dry-run PER PML")
    ap.add_argument("--termasuk-di-luar-audit", action="store_true",
                    help="mode audit: ikut proses dokumen list (dicari per subsls audit) yang tidak tercatat di "
                         "audit; tetap wajib createdBy/updatedBy = akun PPL")
    ap.add_argument("--login-manual", action="store_true", help="tunggu manusia login di jendela browser (tiap PML)")
    ap.add_argument("--maks-error-beruntun", type=int, default=3)
    ap.add_argument("--audit", default="", metavar="BERKAS",
                    help="audit input (default audit/audit_log_gabungan.csv; folder -> <folder>/audit_log_gabungan.csv)")
    args = ap.parse_args()
    lokasi.cek_struktur_lama()
    global AUDIT_GABUNGAN
    if args.audit:
        AUDIT_GABUNGAN = lokasi.jalur_audit(args.audit)
    from inti import fasih_web as _fw
    _fw.SCREENSHOT_DIR = HASIL / "log_screenshots"
    pml_dipilih = [a.strip().lower() for s in args.akun_pml for a in s.split(",") if a.strip()]
    akun_ppl = (args.akun_ppl or "").strip().lower()

    if args.rencana and args.daftar:
        ap.error("pilih salah satu: --rencana atau --daftar")
    audit: list[dict] = []
    audit_docs: dict[str, dict] = {}
    if args.rencana or args.daftar:
        try:
            if args.rencana:
                target, masalah = target_dari_rencana(baca_rencana(Path(args.rencana)))
            else:
                target, masalah = target_dari_daftar(baca_daftar(Path(args.daftar)))
        except (OSError, ValueError) as e:
            print(f"❌ {e}", file=sys.stderr)
            return 2
        for m in masalah[:30]:
            print(f"⚠️ {m}")
        if len(masalah) > 30:
            print(f"⚠️ ... {len(masalah) - 30} masalah lain")
        audit_approve = [] if args.abaikan_audit_approve else baca_csv(AUDIT_APPROVE)
        if args.rencana:
            sudah, kunci_target = id_sudah_approved(audit_approve), (lambda t: t["id"])
        else:
            sudah, kunci_target = kode_sudah_approved(audit_approve), (lambda t: baku_kode(t["kode"]))
        dilewati_audit = sum(1 for t in target if kunci_target(t) in sudah)
        kelompok = kelompokkan_per_pml([t for t in target if kunci_target(t) not in sudah], pml_dipilih)
        tak_ada = [p for p in pml_dipilih if p not in {k for k, _ in kelompok}]
        if tak_ada:
            print(f"⚠️ PML diminta tapi tidak punya dokumen tersisa di file: {tak_ada}")
        print(f"{'Rencana' if args.rencana else 'Daftar'} {args.rencana or args.daftar}: {len(target)} dokumen "
              f"valid, {len(masalah)} baris bermasalah, {dilewati_audit} sudah APPROVED di {AUDIT_APPROVE} (dilewati).")
    else:
        if len(pml_dipilih) != 1 or not akun_ppl:
            ap.error("mode audit butuh tepat satu --akun-pml dan --akun-ppl (multi PML: pakai --daftar / --rencana)")
        audit = baca_csv(AUDIT_GABUNGAN)
        audit_docs = dokumen_audit_ppl(audit, akun_ppl)
        if not audit_docs:
            print(f"❌ Tidak ada dokumen akun {akun_ppl} di {AUDIT_GABUNGAN}.", file=sys.stderr)
            return 2
        kelompok = [(pml_dipilih[0], [{**t, "akun_pml": pml_dipilih[0], "akun_ppl": akun_ppl}
                                      for t in gabung_target(audit_docs, [])])]
    if not kelompok:
        print("Tidak ada dokumen utk diproses.")
        return 0

    print(f"\n{'⚠️ EKSEKUSI (klik Approve sungguhan)' if args.eksekusi else 'DRY-RUN (tanpa klik Approve)'} — "
          f"{len(kelompok)} PML, {sum(len(tg) for _, tg in kelompok)} dokumen"
          + (f", maks {args.limit} per PML" if args.limit else "") + ":")
    for k, (pml, tg) in enumerate(kelompok, start=1):
        nama = tg[0].get("nama_pml") or ""
        rinci = (f"subsls {dict(Counter(t['kode'][:16] for t in tg))}" if args.daftar
                 else f"PPL {dict(Counter(t['akun_ppl'] for t in tg))}")
        print(f"  {k}. {pml} {('(' + nama + ')') if nama else ''} — {len(tg)} dokumen, {rinci}")
    if args.cek:
        return 0
    if args.eksekusi and not args.ya:
        if input("Ketik 'YA' utk APPROVE sungguhan (irreversible) utk SEMUA PML di atas: ").strip().upper() != "YA":
            print("Dibatalkan.")
            return 1

    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        try:
            hitung, kode = jalankan_semua_pml(browser, kelompok, args, audit, audit_docs, akun_ppl)
        finally:
            try:
                browser.close()
            except Exception:
                pass
    print(f"\nRingkasan: {dict(hitung)}\nAudit: {AUDIT_APPROVE}")
    return kode


def tutup_sesi_pml(ctx, sess, logout: bool):
    """Akhiri sesi satu PML. Multi PML (seperti input_usaha): logout UI + hapus cookie SEMUA domain
    (termasuk sso.bps.go.id), lalu tutup context. Satu PML / login manual: tanpa logout supaya sesi
    tersimpan tetap bisa dipakai run berikutnya."""
    if logout and sess is not None:
        try:
            sess.logout()
        except Exception as e:
            print(f"⚠️ Logout gagal (tidak fatal, context tetap ditutup): {str(e)[:120]}", flush=True)
    try:
        ctx.close()
    except Exception:
        pass


def jalankan_semua_pml(browser, kelompok: list[tuple[str, list[dict]]], args, audit: list[dict],
                       audit_docs: dict[str, dict], akun_ppl: str) -> tuple[Counter, int]:
    """Loop PML: login (context baru) -> siapkan target -> proses -> simpan sesi -> logout & tutup
    context -> PML berikutnya. Login gagal = PML itu dilewati (ERROR_LOGIN_PML); STOP_*/error beruntun/
    exception tak terduga = SELURUH run berhenti. -> (hitung status, kode keluar)."""
    hitung: Counter = Counter()
    kode = 0
    # Satu PML / login manual: sesi disimpan & dipakai ulang (tanpa logout). Multi PML: seperti
    # input_usaha — login FIXED_PASSWORD di context baru, logout + tutup context di akhir tiap PML.
    pakai_sesi = len(kelompok) == 1 or args.login_manual
    for k, (akun_pml, target) in enumerate(kelompok, start=1):
        print(f"\n===== PML {k}/{len(kelompok)}: {akun_pml} — {len(target)} dokumen =====", flush=True)
        file_sesi = lokasi.siapkan(HASIL / f".sesi_fasih_web_{slug_akun(akun_pml)}.json") if pakai_sesi else None
        try:
            ctx, sess = mulai_sesi_pml(browser, akun_pml, args.login_manual, file_sesi)
        except Exception as e:
            pesan = str(e)[:300]
            print(f"❌ Login {akun_pml} gagal — PML ini DILEWATI: {pesan}", flush=True)
            append_audit({"timestamp": time.strftime("%Y-%m-%d %H:%M:%S"), "akun_pml": akun_pml,
                          "status": "ERROR_LOGIN_PML", "pesan": pesan})
            hitung["ERROR_LOGIN_PML"] += 1
            kode = 1
            continue
        alasan = ""
        try:
            if args.daftar:
                # File fasih-sm tidak memuat id dokumen: cari di list PENDATAAN PML yang SEDANG login.
                target = siapkan_target_daftar(sess, target, akun_pml, args.assignment_id)
            elif not args.rencana:
                items_list: list[dict] = []
                for subsls in subsls_audit_ppl(audit, akun_ppl):
                    try:
                        items_list += cari_list(sess, args.assignment_id, subsls)
                    except Exception as e:
                        print(f"⚠️ Pencarian list subsls {subsls} gagal (tidak fatal): {str(e)[:150]}")
                luar = [it for it in items_list if it.get("id") not in audit_docs]
                print(f"Dokumen list subsls PPL yg TIDAK di audit: {len(luar)} "
                      f"{dict(Counter(i.get('assignmentStatusAlias') for i in luar))}"
                      + ("" if args.termasuk_di_luar_audit else " — tidak diproses (pakai --termasuk-di-luar-audit)"))
                if args.termasuk_di_luar_audit:
                    target = [{**t, "akun_pml": akun_pml, "akun_ppl": akun_ppl}
                              for t in gabung_target(audit_docs, items_list)]
                print(f"Target: {len(target)} dokumen.\n")
            alasan = proses_kelompok(sess, akun_pml, target, args, hitung, file_sesi)
            simpan_sesi(sess, file_sesi)
        except Exception as e:
            alasan = f"ERROR_TAK_TERDUGA di PML {akun_pml}: {type(e).__name__}: {str(e)[:200]}"
            append_audit({"timestamp": time.strftime("%Y-%m-%d %H:%M:%S"), "akun_pml": akun_pml,
                          "status": "ERROR_TAK_TERDUGA", "pesan": alasan})
            hitung["ERROR_TAK_TERDUGA"] += 1
        finally:
            tutup_sesi_pml(ctx, sess, logout=not pakai_sesi)
        if alasan:
            print(f"\n⛔ {alasan} — SELURUH run DIHENTIKAN (di PML {akun_pml}). "
                  f"Lihat {AUDIT_APPROVE} & log_screenshots/.", flush=True)
            kode = 1
            break
    return hitung, kode


if __name__ == "__main__":
    sys.exit(main())
