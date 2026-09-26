"""
berkas_tabel.py — baca Excel/CSV utk GUI: nama sheet, pratinjau tabel, impor kodepos/wilayah,
dan isi clipboard hasil generate KBLI & koordinat pengganti.

Clipboard dibuat dalam DUA bentuk: teks (tab-separated) dan HTML ber-`mso-number-format:'\\@'`
supaya Excel menempelkan kode KBLI sbg TEKS — "01464" tidak berubah jadi 1464.
Tidak ada yang ditebak: kolom dicari lewat JUDUL kolom; tidak ketemu = pesan jelas.
"""

from __future__ import annotations

import csv
import html
import re
from pathlib import Path


class TabelSalah(ValueError):
    """Berkas/tabel tidak sesuai; pesannya ditampilkan apa adanya di GUI."""


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _teks(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _huruf(kolom: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(kolom)


def _buka_xlsx(path: str | Path):
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise TabelSalah("Paket openpyxl belum terpasang (Persiapan > Pasang paket).") from e
    p = Path(path)
    if not p.exists():
        raise TabelSalah(f"Berkas tidak ada: {p}")
    try:
        return load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise TabelSalah(f"Berkas tidak bisa dibaca sbg Excel ({p.name}): {e}") from e


def daftar_lembar(path: str | Path) -> list[str]:
    if Path(path).suffix.lower() == ".csv":
        return []
    wb = _buka_xlsx(path)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _lembar(wb, lembar: str | None):
    if lembar:
        if lembar not in wb.sheetnames:
            raise TabelSalah(f"Sheet '{lembar}' tidak ada. Yang ada: {', '.join(wb.sheetnames)}")
        return wb[lembar]
    return wb.active


def baca_baris(path: str | Path, lembar: str | None = None) -> list[list[str]]:
    """Seluruh isi sheet/CSV sbg teks (baris ke-i daftar = baris ke-(i+1) berkas)."""
    p = Path(path)
    if p.suffix.lower() == ".csv":
        if not p.exists():
            raise TabelSalah(f"Berkas tidak ada: {p}")
        teks = p.read_bytes()
        for kode in ("utf-8-sig", "cp1252"):
            try:
                isi = teks.decode(kode)
                break
            except UnicodeDecodeError:
                continue
        contoh = isi[:4096]
        try:
            dialek = csv.Sniffer().sniff(contoh, delimiters=",;\t")
        except csv.Error:
            dialek = csv.excel
        return [[c.strip() for c in r] for r in csv.reader(isi.splitlines(), dialek)]
    wb = _buka_xlsx(p)
    try:
        ws = _lembar(wb, lembar)
        return [[_teks(v) for v in r] for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def pratinjau(path: str | Path, lembar: str | None = None, n: int = 8) -> dict:
    baris = baca_baris(path, lembar)
    if not baris:
        raise TabelSalah("Tabel kosong.")
    return {"judul": baris[0], "contoh": baris[1:1 + n], "jumlah": max(0, len(baris) - 1),
            "lembar": daftar_lembar(path)}


# ------------------------------------------------------------------ impor kodepos / wilayah
KOLOM_WILAYAH = ("provinsi", "kabkota", "kecamatan", "desa", "sls", "subsls")


def impor_tabel(path: str | Path, jenis: str, kolom: dict, lembar: str | None = None) -> dict:
    """Tabel -> {"data": {...}, "dilewati": [..]}. kolom = {peran: indeks kolom (0-based)}.
    kodepos : peran "kode" (10 digit desa ATAU 16 digit idsubsls) & "kodepos" (5 digit)
              -> data = {"desa": {10 digit: kp}, "subsls": {16 digit: kp}}
    wilayah : peran "idsubsls" + (opsional) provinsi..subsls -> data = {16 digit: {..}}"""
    baris = baca_baris(path, lembar)[1:]
    dilewati: list[str] = []

    def sel(r, peran):
        i = kolom.get(peran)
        if i is None or i == "":
            return ""
        i = int(i)
        return r[i].strip() if i < len(r) else ""

    if jenis == "kodepos":
        if kolom.get("kode") in (None, "") or kolom.get("kodepos") in (None, ""):
            raise TabelSalah("Pilih kolom kode desa/idsubsls DAN kolom kodepos.")
        desa, subsls = {}, {}
        for n, r in enumerate(baris, start=2):
            kode, kp = re.sub(r"\D", "", sel(r, "kode")), re.sub(r"\D", "", sel(r, "kodepos"))
            if not kode and not kp:
                continue
            if len(kp) != 5 or len(kode) not in (10, 16):
                dilewati.append(f"baris {n}: kode '{sel(r, 'kode')}' / kodepos '{sel(r, 'kodepos')}'")
                continue
            tujuan = desa if len(kode) == 10 else subsls
            if kode in tujuan and tujuan[kode] != kp:
                dilewati.append(f"baris {n}: {kode} punya dua kodepos ({tujuan[kode]} & {kp}) — dua-duanya dilewati")
                tujuan[kode] = None
                continue
            tujuan[kode] = kp
        return {"data": {"desa": {k: v for k, v in desa.items() if v},
                         "subsls": {k: v for k, v in subsls.items() if v}}, "dilewati": dilewati}

    if jenis == "wilayah":
        if kolom.get("idsubsls") in (None, ""):
            raise TabelSalah("Pilih kolom idsubsls.")
        data = {}
        for n, r in enumerate(baris, start=2):
            ids = re.sub(r"\D", "", sel(r, "idsubsls"))
            if not ids:
                continue
            if len(ids) != 16:
                dilewati.append(f"baris {n}: idsubsls '{sel(r, 'idsubsls')}' bukan 16 digit")
                continue
            data[ids] = {k: sel(r, k).upper() for k in KOLOM_WILAYAH}
        return {"data": data, "dilewati": dilewati}

    raise TabelSalah(f"Jenis impor '{jenis}' tidak dikenal.")


# ------------------------------------------------------------------ clipboard
def _html(baris: list[list[str]], kolom_teks: set[int]) -> str:
    gaya = "mso-number-format:'\\@'"
    isi = []
    for r in baris:
        sel = "".join(f'<td style="{gaya}">{html.escape(v)}</td>' if i in kolom_teks else f"<td>{html.escape(v)}</td>"
                      for i, v in enumerate(r))
        isi.append(f"<tr>{sel}</tr>")
    return ('<html><head><meta charset="utf-8"></head><body><table>' + "".join(isi) + "</table></body></html>")


def _tsv(baris: list[list[str]]) -> str:
    return "\r\n".join("\t".join(v.replace("\t", " ").replace("\n", " ") for v in r) for r in baris) + "\r\n"


def _cari_judul(baris: list[list[str]], nama: list[str], batas: int = 20) -> tuple[int, dict[str, int]] | None:
    """(indeks baris judul 0-based, {nama: indeks kolom}) — baris pertama (<= batas) yang memuat
    nama[0]. Nama lain dicari di baris yang sama (tidak ada -> tidak masuk dict)."""
    for i, r in enumerate(baris[:batas]):
        norm = [_norm(v) for v in r]
        if nama[0] in norm:
            return i, {n: norm.index(n) for n in nama if n in norm}
    return None


def salin_kbli(hasil: str | Path, sumber: str | Path | None = None, lembar: str | None = None,
               pertahankan: bool = True) -> dict:
    """Isi clipboard Kode + Judul KBLI dari berkas hasil generate_kbli.py.

    Baris hasil = baris sheet asal (generate_kbli menjaga nomor baris). Kalau `sumber` diberikan:
    sel tujuan tempel dicari dari judul kolom 'Kode KBLI' sheet asal, dan (pertahankan=True)
    baris yang kode KBLI-nya SUDAH terisi di sheet memakai isi sheet itu apa adanya — jadi
    menempel tidak mengubah baris tsb (bedanya dilaporkan, tidak ditimpa)."""
    b_hasil = baca_baris(hasil)
    ketemu = _cari_judul(b_hasil, ["kode kbli", "judul kbli", "sumber", "catatan"])
    if not ketemu:
        raise TabelSalah(f"Judul kolom 'Kode KBLI' tidak ada di {Path(hasil).name} — bukan hasil generate_kbli?")
    i_judul, kol = ketemu
    i_kode = kol["kode kbli"]
    i_jdl = kol.get("judul kbli")
    alt = [j for j, v in enumerate(b_hasil[i_judul]) if _norm(v).startswith("alternatif")]
    data = b_hasil[i_judul + 1:]

    def ambil(r, j):
        return r[j] if j is not None and j < len(r) else ""

    while data and not any(ambil(data[-1], j) for j in (i_kode, i_jdl, kol.get("catatan"))):
        data.pop()
    mulai = i_judul + 2                                   # nomor baris Excel data pertama

    peringatan: list[str] = []
    tujuan = {"kode": None, "judul": None, "berdampingan": True}
    lama: dict[int, tuple[str, str]] = {}
    isi_sumber: dict[int, bool] = {}
    if sumber:
        b_src = baca_baris(sumber, lembar)
        cari = _cari_judul(b_src, ["kode kbli", "judul kbli"])
        if not cari:
            peringatan.append("Kolom 'Kode KBLI' tidak ditemukan di sheet asal — tempel di kolom kode KBLI Anda "
                              f"mulai baris {mulai}.")
        else:
            j_src, k_src = cari
            if j_src != i_judul:
                peringatan.append(f"Baris judul sheet asal ({j_src + 1}) beda dgn berkas hasil ({i_judul + 1}) — "
                                  "periksa berkas hasilnya benar dari sheet ini.")
            ks, js = k_src["kode kbli"], k_src.get("judul kbli")
            tujuan["kode"] = f"{_huruf(ks + 1)}{mulai}"
            tujuan["judul"] = f"{_huruf(js + 1)}{mulai}" if js is not None else None
            tujuan["berdampingan"] = js == ks + 1
            for n, r in enumerate(b_src[i_judul + 1:], start=mulai):
                lama[n] = (ambil(r, ks), ambil(r, js))
                isi_sumber[n] = any(v for j, v in enumerate(r) if j not in (ks, js))

    baris_out: list[list[str]] = []
    ringkas = {"baris": len(data), "mesin": 0, "dipertahankan": 0, "kosong": 0, "periksa": 0}
    periksa, beda, hilang_nol, yatim = [], [], [], []
    for n, r in enumerate(data, start=mulai):
        kode, judul, catatan = ambil(r, i_kode), ambil(r, i_jdl), ambil(r, kol.get("catatan"))
        k_lama, j_lama = lama.get(n, ("", ""))
        if kode and sumber and n in isi_sumber and not isi_sumber[n]:
            yatim.append(n)
        if pertahankan and k_lama:
            if re.fullmatch(r"\d{4}", k_lama):
                hilang_nol.append(n)
            if kode and kode != k_lama:
                beda.append({"baris": n, "kode_sheet": k_lama, "kode_mesin": kode, "judul_mesin": judul})
            kode, judul = k_lama, j_lama
            ringkas["dipertahankan"] += 1
        elif kode:
            ringkas["mesin"] += 1
        else:
            ringkas["kosong"] += 1
        if catatan and not (pertahankan and k_lama):
            ringkas["periksa"] += 1
            periksa.append({"baris": n, "kode": kode, "judul": judul, "catatan": catatan,
                            "alternatif": [ambil(r, j) for j in alt if ambil(r, j)]})
        baris_out.append([kode, judul])

    if yatim:
        peringatan.append(f"{len(yatim)} baris punya kode di hasil tapi barisnya KOSONG di sheet asal "
                          f"(mis. baris {', '.join(map(str, yatim[:5]))}) — hasilnya mungkin dari sheet lain.")
    if hilang_nol:
        peringatan.append(f"{len(hilang_nol)} kode KBLI di sheet cuma 4 digit (nol di depan hilang?), mis. baris "
                          f"{', '.join(map(str, hilang_nol[:5]))} — dipertahankan apa adanya, periksa manual.")

    hasil_salin = {"gabung": {"teks": _tsv(baris_out), "html": _html(baris_out, {0})}}
    if not tujuan["berdampingan"]:
        hasil_salin = {"kode": {"teks": _tsv([[r[0]] for r in baris_out]), "html": _html([[r[0]] for r in baris_out], {0})},
                       "judul": {"teks": _tsv([[r[1]] for r in baris_out]), "html": _html([[r[1]] for r in baris_out], set())}}
    return {"salin": hasil_salin, "mulai": mulai, "tujuan": tujuan, "ringkas": ringkas, "periksa": periksa,
            "beda": beda, "peringatan": peringatan}


def salin_koordinat(hasil: str | Path, sumber: str | Path | None = None, lembar: str | None = None) -> dict:
    """Isi clipboard Latitude/Longitude dari hasil koordinat/koordinat_pengganti.py (A:B mulai baris 2)."""
    b = baca_baris(hasil)
    if not b or [_norm(v) for v in b[0][:2]] != ["latitude", "longitude"]:
        raise TabelSalah(f"{Path(hasil).name}: kolom A:B baris 1 bukan 'Latitude'/'Longitude' — bukan hasil "
                         "koordinat_pengganti?")
    data = [[(r[0] if len(r) > 0 else ""), (r[1] if len(r) > 1 else "")] for r in b[1:]]
    while data and not any(data[-1]):
        data.pop()
    judul = [_norm(v) for v in b[0]]
    j_ubah = judul.index("berubah") if "berubah" in judul else None
    berubah = (sum(1 for r in b[1:1 + len(data)] if j_ubah < len(r) and _norm(r[j_ubah]) not in ("", "0", "false"))
               if j_ubah is not None else None)
    tujuan, peringatan, dipertahankan = None, [], 0
    if sumber:
        b_src = baca_baris(sumber, lembar)
        cari = _cari_judul(b_src, ["latitude", "longitude"])
        if not cari or "longitude" not in cari[1]:
            peringatan.append("Kolom Latitude/Longitude tidak ditemukan di sheet asal — tempel di sel Latitude baris 2.")
        else:
            j, k = cari
            if j != 0:
                peringatan.append(f"Judul sheet asal di baris {j + 1}, bukan baris 1 — periksa kesejajaran baris.")
            tujuan = f"{_huruf(k['latitude'] + 1)}2"
            if k["longitude"] != k["latitude"] + 1:
                peringatan.append("Kolom Longitude tidak tepat di kanan Latitude di sheet asal — tempel per kolom.")
            # Baris yang tidak diproses alat (kosong di hasil) jangan sampai MENGOSONGKAN koordinat sheet.
            for n, r in enumerate(data):
                src = b_src[n + 1] if n + 1 < len(b_src) else []
                asli = [src[k["latitude"]] if k["latitude"] < len(src) else "",
                        src[k["longitude"]] if k["longitude"] < len(src) else ""]
                if not any(r) and any(asli):
                    data[n] = asli
                    dipertahankan += 1
    return {"salin": {"gabung": {"teks": _tsv(data), "html": _html(data, {0, 1})}}, "mulai": 2,
            "tujuan": {"kode": tujuan},
            "ringkas": {"baris": len(data), "berubah": berubah, "dipertahankan": dipertahankan},
            "peringatan": peringatan}
